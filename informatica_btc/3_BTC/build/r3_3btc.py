# -*- coding: utf-8 -*-
"""Informática 3.º BTC — copias R3 a partir de R2 (R2 queda intacta: se abre solo para lectura).

Cierre de la revisión: legibilidad de figuras (tomo a 15,5 cm; laboratorio con ajuste por práctica y piso de 7 pt),
figuras 2 y 19, y observaciones editoriales O-3-06 y O-3-07. Ver `CIERRE_R3_3BTC.md`.
Uso: python3 r3_3btc.py
"""
import copy, json, os, re, sys
import docx
from docx.oxml.ns import qn
AQUI = os.path.dirname(os.path.abspath(__file__)); BASE = os.path.dirname(AQUI)
sys.path.insert(0, os.path.join(os.path.dirname(BASE), 'build_comun'))
import docx_util as U

R2 = os.path.join(BASE, 'R2'); R3 = os.path.join(BASE, 'R3'); PDF = os.path.join(BASE, 'pdf_revision_R3')
FIG = os.path.join(BASE, 'figuras_R3')
SALIDA = ('Tomo', 'Laboratorio', 'Plan_anual', 'Solucionario', 'Planes_de_clase')
VERSION = ('Revisión R2, 22 de octubre de 2026.', 'Revisión R3, 22 de octubre de 2026.')
ANCHO = 15.5
LAB_FIGURAS = [1, 2, 3, 4, 5, 7, 8, 9, 10, 12, 13, 14, 15, 16, 17, 18, 19, 20]   # imagen i del cuaderno → figura
LAB_PRACTICAS = [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 13, 14, 15, 17, 18, 19, 20]  # práctica de la imagen i
# Cuerpo del texto más chico de cada figura R3 a 15,5 cm (LEGIBILIDAD_R3_3BTC.md). La 20 es vectorial.
MIN_PT = {1: 7.3, 2: 10.9, 3: 8.8, 4: 10.0, 5: 11.4, 6: 9.4, 7: 8.2, 8: 8.2, 9: 10.5, 10: 10.9, 11: 11.8, 12: 7.0,
          13: 9.7, 14: 10.9, 15: 12.2, 16: 8.2, 17: 10.0, 18: 7.0, 19: 8.8, 20: 8.5}
PISO_PT = 7.0
ANCHO_MIN = {n: min(ANCHO, round(ANCHO * PISO_PT / pt + 0.05, 1)) for n, pt in MIN_PT.items()}
NOTA_FIG2 = ' Esquema de laboratorio: no corresponde a un registro tributario real.'


def abrir(k):
    return docx.Document(os.path.join(R2, 'Informatica_3_BTC_%s_R2.docx' % k))


def ruta(k):
    return os.path.join(R3, 'Informatica_3_BTC_%s_R3.docx' % k)


def R(doc, k, viejo, nuevo, veces=1, motivo=''):
    U.reemplazar(doc, viejo, nuevo, veces, k, motivo)


def reg(doc, antes, despues, n, motivo):
    U.REGISTRO.append({'documento': doc, 'antes': antes, 'despues': despues, 'apariciones': n, 'motivo': motivo})


# ================================================================== TOMO
def tomo():
    d = abrir('Tomo'); k = 'Tomo'
    R(d, k, *VERSION, 1, 'Identificación de la versión')
    R(d, k, 'Pantalla de alta de la empresa con los datos de Despensa San Blas.',
      'Pantalla de alta de la empresa con los datos de Despensa San Blas.' + NOTA_FIG2, 1,
      'R3-3-02: la nota de la Figura 2 medía 6 pt; sale de la imagen y pasa al pie')
    R(d, k, 'Con esa configuración el sistema sabe qué reportes de IVA debe preparar cada mes.',
      'Con esa configuración el sistema sabe qué reportes de IVA debe preparar cada mes. Los registros que exige la ley dependen del régimen de cada contribuyente; en este curso San Blas lleva contabilidad completa en el sistema para practicar todos los libros y estados, y el profesional verifica en cada caso real qué registros corresponden.', 1,
      'R3-3-04 (O-3-06): nota de alcance sobre la contabilidad completa de un contribuyente en IRE SIMPLE')
    R(d, k, 'Un estudio contable que expone datos de terceros por descuido compromete su responsabilidad profesional.',
      'Un estudio contable que expone datos de terceros por descuido compromete su responsabilidad profesional. La Ley 7593/2025 de protección de datos personales ya fue publicada y amplía ese marco cuando entre en vigor (ver Clase 20).', 1,
      'R3-3-05 (O-3-07): la Clase 3 citaba solo la ley de datos crediticios')
    for i in range(1, 21):
        U.reemplazar_imagen(d, i, os.path.join(FIG, 'R3_Figura_%02d.png' % i), ANCHO)
    reg(k, '20 figuras a 14,2 cm (texto más chico 6,0-6,7 pt en las figuras 1, 2, 12, 18 y 19)',
        '20 figuras R3 a 15,5 cm: ningún texto por debajo de 7 pt', 20, 'R3-3-01: legibilidad')
    return d


# ================================================================== LABORATORIO
def lab(anchos=None, extra=None):
    anchos = anchos or {n: ANCHO for n in range(1, 21)}
    extra = extra or {}
    d = abrir('Laboratorio'); k = 'Laboratorio'
    R(d, k, *VERSION, 1, 'Identificación de la versión')
    R(d, k, 'Referencia gráfica de la práctica 2.', 'Referencia gráfica de la práctica 2.' + NOTA_FIG2, 1, 'R3-3-02')
    for i, f in enumerate(LAB_FIGURAS, 1):
        U.reemplazar_imagen(d, i, os.path.join(FIG, 'R3_Figura_%02d.png' % f), anchos[f])
    reg(k, 'Figuras a 9,6 cm (7,2 cm en la Práctica 2): textos de 4 a 6,6 pt',
        'Figuras entre %s y 15,5 cm; ningún texto de figura por debajo de 7 pt' % ('%.1f' % min(anchos[f] for f in LAB_FIGURAS)).replace('.', ','), 18,
        'R3-3-01: una práctica puede ocupar dos páginas antes que bajar la figura de 7 pt')
    n_extra = 0
    for p, cant in extra.items():
        ps = d.paragraphs
        i = next(j for j, q in enumerate(ps) if q.text.startswith('Qué hice:') and j > indice_practica(d, p))
        ult = ps[i + 2]._p                                       # «Qué aprendí»
        for _ in range(cant):
            nuevo = U.renglon_escritura(copy.deepcopy(ult)); ult.addnext(nuevo)
            U.poner_texto(docx.text.paragraph.Paragraph(nuevo, d._body), '_' * 74); n_extra += 1
    if extra:
        reg(k, 'Prácticas %s: segunda página con solo el registro del estudiante' % ', '.join(map(str, sorted(extra))),
            'Renglones adicionales del registro, de 8 mm, hasta el pie de esa página', n_extra,
            'R3-3-01: la segunda página se usa como espacio de trabajo')
    return d


def indice_practica(d, p):
    for j, q in enumerate(d.paragraphs):
        if q.text.startswith('Nombre y apellido') and j > 0:
            pass
    # el título de cada práctica está en una tabla; se ubica el párrafo siguiente a esa tabla
    body = list(d.element.body.iterchildren()); pars = d.paragraphs; idx = {id(q._p): j for j, q in enumerate(pars)}
    for n, el in enumerate(body):
        if el.tag == qn('w:tbl') and ''.join(x.text or '' for x in el.iter(qn('w:t'))).startswith('PRÁCTICA %d ·' % p):
            nxt = body[n + 1]
            return idx[id(nxt)] if id(nxt) in idx else idx.get(id(body[n + 2]), 0)
    raise SystemExit('práctica %d no encontrada' % p)


# ================================================================== PLAN, SOLUCIONARIO, PLANES
def solo_version(k):
    d = abrir(k); R(d, k, *VERSION, 1, 'Identificación de la versión'); return d


def solucionario():
    d = solo_version('Solucionario'); k = 'Solucionario'
    R(d, k, 'consultable en bacn.gov.py/leyes-paraguayas/2538/del-;',
      'consultable en la ficha de la Ley 4924/2013 de la Biblioteca y Archivo Central del Congreso (bacn.gov.py);', 1,
      'R3-3-06: el enlace venía truncado desde R1 y no llevaba a ninguna ficha')
    return d


# ================================================================== ajuste del laboratorio
def practicas_por_pagina(pdf):
    import pymupdf
    pags = [p.get_text() for p in pymupdf.open(pdf)]
    ini = {}
    for i, t in enumerate(pags):
        for l in t.split('\n'):
            m = re.match(r'PRÁCTICA (\d+) · ', l.strip())
            if m and int(m.group(1)) not in ini:
                ini[int(m.group(1))] = i
    orden = sorted(ini.items(), key=lambda x: x[1])
    return {n: list(range(p + 1, (orden[j + 1][1] if j + 1 < len(orden) else len(pags)) + 1)) for j, (n, p) in enumerate(orden)}


def probar(anchos, extra=None):
    n0 = len(U.REGISTRO); lab(anchos, extra).save(ruta('Laboratorio')); del U.REGISTRO[n0:]
    return practicas_por_pagina(U.a_pdf(ruta('Laboratorio'), PDF))


def ajustar_laboratorio():
    fig_de = dict(zip(LAB_PRACTICAS, LAB_FIGURAS))
    anchos = {n: ANCHO for n in range(1, 21)}; fijas = set()
    for vuelta in range(14):
        paginas = probar(anchos); cambio = False
        for p, pg in paginas.items():
            f = fig_de.get(p)
            if len(pg) > 1 and f and f not in fijas:
                nuevo = round(anchos[f] - 0.5, 1)
                if nuevo >= ANCHO_MIN[f]:
                    anchos[f] = nuevo
                else:
                    anchos[f] = ANCHO; fijas.add(f)
                cambio = True
        if not cambio:
            break
    paginas = probar(anchos)
    dobles = [p for p, pg in paginas.items() if len(pg) > 1]
    extra = {p: 0 for p in dobles}
    total = sum(map(len, paginas.values()))
    for paso in (32, 16, 8, 4, 2, 1):
        for p in dobles:
            prueba = dict(extra); prueba[p] += paso
            pg = probar(anchos, prueba)
            if len(pg[p]) == len(paginas[p]) and sum(map(len, pg.values())) == total:
                extra = prueba
    paginas = probar(anchos, extra)
    json.dump({'anchos_cm': anchos, 'ancho_minimo_cm': ANCHO_MIN, 'paginas_por_practica': paginas, 'renglones_extra': extra},
              open(os.path.join(AQUI, 'ajuste_laboratorio_r3_3btc.json'), 'w', encoding='utf8'), ensure_ascii=False, indent=1)
    return anchos, extra


# ================================================================== construcción
PATRON = r'^((?:Clase|EVALUACIÓN|Plan de Clase)\b.*?)\s+·\s+(\d+)$'
CLAVE = lambda t: t.rstrip('…').strip()

if __name__ == '__main__':
    os.makedirs(R3, exist_ok=True); os.makedirs(PDF, exist_ok=True)
    docs = {'Tomo': tomo(), 'Plan_anual': solo_version('Plan_anual'), 'Solucionario': solucionario(),
            'Planes_de_clase': solo_version('Planes_de_clase')}
    for k, d in docs.items():
        d.save(ruta(k))
    anchos, extra = ajustar_laboratorio()
    lab(anchos, extra).save(ruta('Laboratorio'))
    indices = {}
    for pasada in (1, 2):
        for k in ('Tomo', 'Planes_de_clase'):
            pdf = U.a_pdf(ruta(k), PDF)
            d = docx.Document(ruta(k))
            indices[k] = U.recalcular_indice(d, pdf, PATRON, clave=CLAVE)
            d.save(ruta(k))
    for k in SALIDA:
        U.a_pdf(ruta(k), PDF)
    json.dump({'cambios': U.REGISTRO, 'indices': indices}, open(os.path.join(AQUI, 'registro_r3_3btc.json'), 'w', encoding='utf8'),
              ensure_ascii=False, indent=1)
    print(len(U.REGISTRO), 'cambios registrados; anchos', anchos, 'extra', extra)
