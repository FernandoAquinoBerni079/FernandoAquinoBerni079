# -*- coding: utf-8 -*-
"""Informática 3.º BTC — copias R2 a partir de R1 (edición diferencial, sin reconstruir los libros).

Lee los cinco DOCX R1 (solo lectura), aplica las correcciones de la matriz de hallazgos H-3-xx, sincroniza los
documentos relacionados, sustituye las figuras por las versiones R2 (identificadas por contenido, no por nombre),
renderiza con LibreOffice, recalcula los índices estáticos (dos pasadas) y deja los DOCX y los PDF de revisión en
3_BTC/R2 y 3_BTC/pdf_revision. Copia también los CSV de laboratorio, sin cambios, a 3_BTC/recursos.
Uso: python3 r2_3btc.py <carpeta R1 de 3.º>
"""
import json, os, re, shutil, sys
import docx
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_LINE_SPACING
AQUI = os.path.dirname(os.path.abspath(__file__)); BASE = os.path.dirname(AQUI)
sys.path.insert(0, os.path.join(os.path.dirname(BASE), 'build_comun'))
import docx_util as U

R1 = sys.argv[1]
R2 = os.path.join(BASE, 'R2'); PDF = os.path.join(BASE, 'pdf_revision'); FIG = os.path.join(BASE, 'figuras')
NOMBRES = {'tomo': 'Tomo_corregido', 'lab': 'Laboratorio_corregido', 'plan': 'Plan_anual_corregido',
           'sol': 'Solucionario_corregido', 'planes': 'Planes_de_clase_corregidos'}
SALIDA = {'tomo': 'Tomo', 'lab': 'Laboratorio', 'plan': 'Plan_anual', 'sol': 'Solucionario', 'planes': 'Planes_de_clase'}
VERSION = ('Revisión de contenido octubre de 2026.', 'Revisión R2, 22 de octubre de 2026.')
# Orden de las imágenes del laboratorio → figura del tomo, cotejado por huella SHA-256 contra Antes_R1
# (las prácticas 12 y 16 no llevan imagen; la práctica 6 usa la figura 7, la 11 la 13, etc.)
LAB_FIGURAS = [1, 2, 3, 4, 5, 7, 8, 9, 10, 12, 13, 14, 15, 16, 17, 18, 19, 20]
CAPACIDADES = {
    1: 'Aplica los conocimientos teóricos administrativos contables procesando datos en un sistema.',
    4: 'Diseña imágenes rasterizadas logrando un diseño profesional. · Edita imágenes publicitarias acorde a la '
       'necesidad de la empresa y/o para utilizarlas en otros programas.',
    6: 'Identifica diferentes sitios WEB.'}
CAPACIDADES[2] = CAPACIDADES[3] = CAPACIDADES[1]; CAPACIDADES[5] = CAPACIDADES[4]
PROYECTO = ('cierre con la presentación del proyecto integrador del año',
            'preparación del guion del proyecto integrador del año (exposición de tres minutos por grupos o por muestreo)')


def abrir(k):
    return docx.Document(os.path.join(R1, 'Informatica_3_BTC_%s.docx' % NOMBRES[k]))


def ruta_r2(k):
    return os.path.join(R2, 'Informatica_3_BTC_%s_R2.docx' % SALIDA[k])


def R(doc, k, viejo, nuevo, veces=1, motivo=''):
    U.reemplazar(doc, viejo, nuevo, veces, SALIDA[k], motivo)


def renglones_3(doc, cantidad):
    """Renglones de respuesta para las preguntas abiertas del tomo de 3.º: ítems numerados del bloque
    «D. Respondé con tus palabras» y preguntas de la prueba diagnóstica. `cantidad(seccion)` según 'diag', 'clase',
    'unidad' o 'etapa'."""
    seccion, abierto, total = 'clase', False, 0
    for el in list(doc.element.body.iterchildren()):
        texto = ''.join(x.text or '' for x in el.iter(qn('w:t')))
        if el.tag == qn('w:tbl'):
            abierto = False
            if texto.startswith('Prueba diagnóstica'):
                seccion, abierto = 'diag', True
            elif texto.startswith('EVALUACIÓN DE LA UNIDAD'):
                seccion = 'unidad'
            elif texto.startswith('EVALUACIÓN DE LA ETAPA'):
                seccion = 'etapa'
            elif texto.startswith('Actividades de la clase'):
                seccion = 'clase'
            continue
        if el.tag != qn('w:p'):
            continue
        p = docx.text.paragraph.Paragraph(el, doc._body)
        if p.style.name.startswith('Heading'):
            abierto = False; continue
        if texto.startswith('D. Respondé'):
            abierto = True; continue
        if re.match(r'[A-C]\. ', texto):
            abierto = False; continue
        if not (abierto and re.match(r'\d+\. ', texto)):
            continue
        ancla = el
        for _ in range(cantidad(seccion)):
            nuevo = doc.add_paragraph()._p
            ancla.addnext(nuevo); ancla = nuevo
            q = docx.text.paragraph.Paragraph(nuevo, doc._body)
            q.style = doc.styles['Normal']
            pf = q.paragraph_format
            pf.space_before = Pt(0); pf.space_after = Pt(0)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY; pf.line_spacing = Pt(20)
            r = q.add_run('_' * 92); r.font.size = Pt(10); r.font.color.rgb = RGBColor.from_string('78909C')
            total += 1
        p.paragraph_format.keep_with_next = True
        e = el.getnext(); bloque = []
        while e is not None and e is not ancla.getnext():
            bloque.append(e); e = e.getnext()
        for b in bloque[:-1]:
            docx.text.paragraph.Paragraph(b, doc._body).paragraph_format.keep_with_next = True
    return total


# ================================================================== TOMO
def tomo():
    d = abrir('tomo'); k = 'tomo'
    R(d, k, *VERSION, 1, 'Identificación de la versión')
    m = 'H-3-02: la entrega depende del receptor (XML al electrónico; KuDE al no electrónico o consumidor final)'
    R(d, k, 'Lo que se imprime o se envía por celular es el KuDE, la representación gráfica del DTE, con un código QR que permite verificar su validez en el portal de la DNIT.',
      'La entrega depende del receptor: si también es facturador electrónico, recibe el archivo electrónico (XML); si no lo es, o si es consumidor final, recibe el KuDE en papel o en formato digital. El KuDE es la representación gráfica del DTE, con su código de control (CDC) y un código QR para verificar su validez en el portal de la DNIT; no reemplaza al documento electrónico.', 1, m)
    R(d, k, 'KuDE: representación gráfica del DTE, con código QR para su verificación.',
      'KuDE: representación gráfica del DTE, con CDC y código QR para su verificación; se entrega al receptor no electrónico o al consumidor final.', 1, m)
    R(d, k, 'Recorrido de un comprobante en el SIFEN: de DE a DTE y KuDE.',
      'Recorrido simplificado de un comprobante en el SIFEN: del DE al DTE y su entrega según el receptor.', 1, m)
    R(d, k, 'el artículo 76 exige autorización judicial para medios mecánicos u otros sistemas modernos, salvo disposición contraria de leyes especiales.',
      'el artículo 76, en el texto dado por la Ley 4924/2013, admite llevarlos por medios mecánicos u otros sistemas modernos mediante una comunicación del sistema, con firma certificada por escribano, ante la Dirección de Registro Público de Comercio, sin necesidad de orden judicial.', 1,
      'H-3-01: el art. 76 fue modificado por la Ley 4924/2013; la autorización judicial ya no rige')
    R(d, k, 'Casi siempre se elige rango de fechas, nivel de detalle y si se incluyen cuentas sin movimiento.',
      'Casi siempre se elige rango de fechas, nivel de detalle y si se incluyen cuentas sin movimiento. Además de libros y balances, el módulo emite otros registros: el registro de compras y de ventas que respalda la liquidación del IVA, los listados de comprobantes y el plan de cuentas.', 1,
      'H-3-08: el programa incluye «otros registros» y el tomo no los nombraba')
    R(d, k, 'en cuanto a las notas sobre utilidades, reservas y resultados distribuibles.',
      'en cuanto a las notas sobre utilidades, reservas y resultados distribuibles. En el caso didáctico San Blas liquida IRE SIMPLE: el profesional verifica si esa obligación le alcanza antes de preparar una presentación.', 1,
      'H-3-07: la obligación 948 no se presenta como aplicable automáticamente a una unipersonal en SIMPLE')
    R(d, k, 'hoja de 1080 × 1080 px, modo RGB y fondo transparente;', 'hoja de 1080 × 1080 px, modo RGB y fondo de color;', 1,
      'H-3-03: sincronizado con la tabla de la clase 12 y con la Práctica 12 (fondo de color)')
    R(d, k, 'Un filtro recalcula todos los píxeles: original, desenfoque, contraste e inversión.',
      'Cuatro filtros como operaciones sobre los valores de la imagen: desenfoque, contraste, escala de grises e inversión.', 1,
      'H-3-04: el pie no coincidía con los cuatro paneles de la figura 16')
    R(d, k, 'El estudio propone para la despensa el dominio sanblas.com.py y',
      'El estudio propone para la despensa el dominio sanblas.com.py (nombre de ejemplo: no se verificó su disponibilidad y no corresponde a un sitio publicado) y', 1,
      'H-3-06: el esquema web no se presenta como sitio publicado')
    nb = U.saltos_a_salto_antes(d)
    U.REGISTRO.append({'documento': 'Tomo', 'antes': 'Salto de página manual en párrafo propio', 'despues': 'Salto de página anterior en ese párrafo',
                       'apariciones': nb, 'motivo': 'H-3-12: evita páginas en blanco cuando la clase termina justo al pie'})
    n = renglones_3(d, lambda s: {'diag': 2, 'clase': 3, 'unidad': 2, 'etapa': 4}[s])
    U.REGISTRO.append({'documento': 'Tomo', 'antes': 'Preguntas abiertas sin espacio para responder',
                       'despues': '%d renglones (2 en la diagnóstica y en las evaluaciones de unidad, 3 en clases, 4 en etapas)' % n, 'apariciones': n,
                       'motivo': 'H-3-09: el estudiante trabaja sobre el tomo; las abiertas llevan renglones'})
    for i in range(1, 21):
        U.reemplazar_imagen(d, i, os.path.join(FIG, 'R2_Figura_%02d.png' % i))      # ancho R1 (14,2 cm) conservado
    U.REGISTRO.append({'documento': 'Tomo', 'antes': '20 figuras R1', 'despues': '20 figuras R2 (17 propuestas, 2 corregidas, 1 vectorial)',
                       'apariciones': 20, 'motivo': 'Decisiones de figuras; ancho de 14,2 cm conservado'})
    return d


# ================================================================== LABORATORIO
def lab():
    d = abrir('lab'); k = 'lab'
    R(d, k, *VERSION, 1, 'Identificación de la versión')
    R(d, k, 'Emitir cualquier reporte vacío y verificar que el encabezado', 'Emití cualquier reporte vacío y verificá que el encabezado', 1,
      'H-3-13: consigna en infinitivo; el resto del cuaderno usa voseo imperativo')
    R(d, k, 'redactá una conciliación bancaria suponiendo que el extracto muestra un cheque de ₲ 1.500.000 aún no cobrado.',
      'redactá una conciliación bancaria suponiendo que San Blas emitió y registró un cheque de ₲ 1.500.000 que el banco todavía no debitó, por lo que no figura en el extracto.', 1,
      'H-3-10: un cheque no cobrado no aparece en el extracto; se alinea con la orientación del solucionario')
    R(d, k, 'Armá la presentación final del proyecto integrador: contabilidad, pack gráfico y sitio en tres minutos.',
      'Prepará el guion de la presentación final del proyecto integrador (contabilidad, pack gráfico y sitio) para exponerlo en tres minutos; el docente organiza las exposiciones por grupos o por muestreo.', 1,
      'H-3-11: una exposición individual de tres minutos por estudiante no entra en el encuentro de cierre')
    # las propuestas son 3:2 y las R1 eran 1512 × 902: se conserva el ALTO R1 para que cada práctica siga en su página
    anchos = [int(e.get('cx')) / 360000 for e in d.element.body.iter(qn('wp:extent'))]
    for i, f in enumerate(LAB_FIGURAS, 1):
        U.reemplazar_imagen(d, i, os.path.join(FIG, 'R2_Figura_%02d.png' % f), round(anchos[i - 1] * (902 / 1512) * 1.5 * (0.85 if i == 2 else 1), 2))  # P2: 15 % menos para no partir el registro
    U.REGISTRO.append({'documento': 'Laboratorio', 'antes': '18 imágenes R1', 'despues': '18 figuras R2 según el cotejo por contenido',
                       'apariciones': 18, 'motivo': 'Imagen n.º i del cuaderno = figura %s del tomo; alto R1 conservado (ancho × 0,895)' % LAB_FIGURAS})
    return d


# ================================================================== PLAN ANUAL
def plan():
    d = abrir('plan'); k = 'plan'
    R(d, k, *VERSION, 1, 'Identificación de la versión')
    R(d, k, PROYECTO[0], PROYECTO[1], 1, 'H-3-11')
    R(d, k, 'Las capacidades subrayadas del programa oficial constituyen la base de la promoción en la disciplina.',
      'Las capacidades subrayadas del programa oficial constituyen la base de la promoción en la disciplina. Según el PDF consultado, son las capacidades «Aplica…», «Diseña…» y «Edita…»; la institución confirma esta lectura.', 1,
      'H-3-05: identificar las capacidades subrayadas sin atribuir al MEC una decisión editorial')
    for fila in d.tables[0].rows:
        cs = list(U.celdas_unicas(fila))
        if len(cs) == 1 and cs[0].text.startswith('UNIDAD '):
            n = int(cs[0].text.split()[1])
            r = cs[0].paragraphs[-1].add_run(' — Capacidad: ' + CAPACIDADES[n]); r.bold = False
            U.REGISTRO.append({'documento': 'Plan_anual', 'antes': cs[0].text.split(' — Capacidad')[0],
                               'despues': 'Banda de unidad con la capacidad del programa', 'apariciones': 1,
                               'motivo': 'H-3-05: el plan anual no mostraba las capacidades'})
    # las capacidades agregan líneas: se quitan dos de los tres párrafos vacíos previos a la firma para no dejarla sola
    vacios = [p for p in d.paragraphs[-4:-1] if not p.text.strip()]
    for p in vacios[:2]:
        p._p.getparent().remove(p._p)
    for p in d.paragraphs[-3:-1]:                   # última viñeta y párrafo vacío van con la firma
        p.paragraph_format.keep_with_next = True
    U.REGISTRO.append({'documento': 'Plan_anual', 'antes': 'Tres párrafos vacíos antes de la firma', 'despues': 'Uno',
                       'apariciones': len(vacios[:2]), 'motivo': 'H-3-05: la firma no queda sola en una página'})
    return d


# ================================================================== SOLUCIONARIO
def sol():
    d = abrir('sol'); k = 'sol'
    R(d, k, *VERSION, 1, 'Identificación de la versión')
    R(d, k, 'El caso asume IRE SIMPLE y la proyección verifica que no supera ₲ 2.000.000.000; el monto por sí solo no determina el régimen',
      'El caso asume IRE SIMPLE y la proyección solo verifica que no supera su tope; el monto por sí solo no determina el régimen', 1,
      'H-3-14: texto de la clave idéntico a la opción d) del tomo')
    R(d, k, 'conversión en DTE, entrega del KuDE al cliente e imputación',
      'conversión en DTE, entrega al receptor (XML si es facturador electrónico; KuDE si no lo es o si es consumidor final) e imputación', 1, 'H-3-02')
    R(d, k, '7. declarada', '7. cerrada (se aceptan «utilizada» o «declarada» con justificación coherente)', 1,
      'H-3-15: la clave usaba un término que el tomo no enseña')
    R(d, k, 'Ley 1034/83, artículos 74–76,', 'Ley 1034/83, artículos 74–76, con el art. 76 en el texto de la Ley 4924/2013 (comunicación al Registro Público de Comercio, sin orden judicial),', 1, 'H-3-01')
    R(d, k, 'P20: cinco pasos con responsables e indicadores medibles.',
      'P20: cinco pasos con responsables e indicadores medibles. La exposición de tres minutos se organiza por grupos o por muestreo.', 1, 'H-3-11')
    return d


# ================================================================== PLANES DE CLASE
def planes():
    d = abrir('planes'); k = 'planes'
    R(d, k, *VERSION, 1, 'Identificación de la versión')
    R(d, k, PROYECTO[0], PROYECTO[1], 2, 'H-3-11: sincronizado con el plan anual y la Práctica 20')
    n = U.escalar_letra(d, {7.5: 9.5})
    U.REGISTRO.append({'documento': 'Planes_de_clase', 'antes': 'Cuerpo de 7,5 pt en las tablas de los planes', 'despues': 'Cuerpo de 9,5 pt (igual que 2.º); cada plan sigue en su página',
                       'apariciones': n, 'motivo': 'H-3-16: legibilidad'})
    k2 = 0
    for p in U.parrafos(d):                       # el renglón del docente ya no entra en su celda con 9,5 pt
        if p.text == '_' * 30:
            U.poner_texto(p, '_' * 22); k2 += 1
    U.REGISTRO.append({'documento': 'Planes_de_clase', 'antes': 'Renglón de 30 guiones en la celda Docente', 'despues': '22 guiones',
                       'apariciones': k2, 'motivo': 'H-3-16: evitar que el renglón se parta en dos líneas'})
    return d


def bandas_huerfanas(ruta_docx, ruta_pdf):
    """Rótulos «Actividades de la clase N» que quedaron como última línea de una página: se les pone salto de página
    anterior (LibreOffice no respeta «conservar con el siguiente» en tablas)."""
    import pymupdf
    huerfanas = []
    for pg in pymupdf.open(ruta_pdf):
        L = [l.strip() for l in pg.get_text().split('\n') if l.strip()]
        if len(L) > 2 and L[-2].startswith('Actividades de la clase'):
            huerfanas.append(L[-2])
    d = docx.Document(ruta_docx)
    for t in d.element.body.iterchildren(qn('w:tbl')):
        texto = ''.join(x.text or '' for x in t.iter(qn('w:t')))
        if texto in huerfanas:
            docx.text.paragraph.Paragraph(next(t.iter(qn('w:p'))), d._body).paragraph_format.page_break_before = True
    d.save(ruta_docx)
    return huerfanas


# ================================================================== construcción
PATRON = r'^((?:Clase|EVALUACIÓN|Plan de Clase)\b.*?)\s+·\s+(\d+)$'
CLAVE = lambda t: t.rstrip('…').strip()

if __name__ == '__main__':
    os.makedirs(R2, exist_ok=True); os.makedirs(PDF, exist_ok=True)
    docs = {'tomo': tomo(), 'lab': lab(), 'plan': plan(), 'sol': sol(), 'planes': planes()}
    for k, d in docs.items():
        d.save(ruta_r2(k))
    h = bandas_huerfanas(ruta_r2('tomo'), U.a_pdf(ruta_r2('tomo'), PDF))
    U.REGISTRO.append({'documento': 'Tomo', 'antes': 'Rótulo de actividades al pie de página, separado de sus consignas',
                       'despues': 'Salto de página anterior en el rótulo', 'apariciones': len(h), 'motivo': 'H-3-12: ' + '; '.join(h)})
    indices = {}
    for pasada in (1, 2):
        for k in ('tomo', 'planes'):
            pdf = U.a_pdf(ruta_r2(k), PDF)
            d = docx.Document(ruta_r2(k))
            indices[k] = U.recalcular_indice(d, pdf, PATRON, clave=CLAVE)
            d.save(ruta_r2(k))
    for k in SALIDA:
        U.a_pdf(ruta_r2(k), PDF)
    rec = os.path.join(R1, 'Recursos_laboratorio'); dst = os.path.join(BASE, 'recursos')
    os.makedirs(dst, exist_ok=True)
    for f in sorted(os.listdir(rec)):
        shutil.copyfile(os.path.join(rec, f), os.path.join(dst, f))
    json.dump({'cambios': U.REGISTRO, 'indices': indices}, open(os.path.join(AQUI, 'registro_r2_3btc.json'), 'w', encoding='utf8'),
              ensure_ascii=False, indent=1)
    print(len(U.REGISTRO), 'cambios registrados')
    for k, v in indices.items():
        print(k, 'entradas', len(v), 'corregidas', sum(1 for _, a, b in v if a != b))
