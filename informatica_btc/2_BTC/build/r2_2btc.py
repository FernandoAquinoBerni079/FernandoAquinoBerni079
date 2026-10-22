# -*- coding: utf-8 -*-
"""Informática 2.º BTC — copias R2 a partir de R1 (edición diferencial, sin reconstruir los libros).

Lee los cinco DOCX R1 (solo lectura), aplica las correcciones de la matriz de hallazgos, sincroniza los documentos
relacionados, sustituye las figuras por las versiones R2, renderiza con LibreOffice, recalcula los índices estáticos
(dos pasadas) y deja los DOCX y los PDF de revisión en 2_BTC/R2 y 2_BTC/pdf_revision.
Uso: python3 r2_2btc.py <carpeta R1 de 2.º>
"""
import json, os, re, shutil, sys, copy
import docx
from docx.oxml.ns import qn
AQUI = os.path.dirname(os.path.abspath(__file__)); BASE = os.path.dirname(AQUI)
sys.path.insert(0, os.path.join(os.path.dirname(BASE), 'build_comun'))
import docx_util as U

R1 = sys.argv[1]
ANCHO_LAB = float(os.environ.get('ANCHO_LAB', '12.0'))
R2 = os.path.join(BASE, 'R2'); PDF = os.path.join(BASE, 'pdf_revision'); FIG = os.path.join(BASE, 'figuras')
NOMBRES = {'tomo': 'Tomo_corregido', 'lab': 'Laboratorio_corregido', 'plan': 'Plan_anual_corregido',
           'sol': 'Solucionario_corregido', 'planes': 'Planes_de_clase_corregidos'}
SALIDA = {'tomo': 'Tomo', 'lab': 'Laboratorio', 'plan': 'Plan_anual', 'sol': 'Solucionario', 'planes': 'Planes_de_clase'}


def abrir(k):
    return docx.Document(os.path.join(R1, 'Informatica_2_BTC_%s.docx' % NOMBRES[k]))


def ruta_r2(k):
    return os.path.join(R2, 'Informatica_2_BTC_%s_R2.docx' % SALIDA[k])


def R(doc, k, viejo, nuevo, veces=1, motivo=''):
    U.reemplazar(doc, viejo, nuevo, veces, SALIDA[k], motivo)


# ================================================================== TOMO
def tomo():
    d = abrir('tomo'); k = 'tomo'
    R(d, k, 'Revisión de contenido octubre de 2026.', 'Revisión R2, 22 de octubre de 2026.', 1, 'Identificación de la versión')
    m = 'H-2-01: la columna C guarda el honorario (A=Cliente, B=Régimen); B4 contenía una etiqueta de régimen'
    R(d, k, 'Si escribís B4*12 la planilla lo toma como texto; si escribís =B4*12 lo calcula.',
      'Si escribís C2*12 la planilla lo toma como texto; si escribís =C2*12 lo calcula.', 1, m)
    R(d, k, '=B4*12', '=C2*12', 2, m)            # tabla 4 y opción d) de la actividad 1.1
    R(d, k, '=B4/2', '=C2/2', 1, m)
    R(d, k, 'honorario anual', 'honorario anual del cliente de la fila 2', 1, m)
    R(d, k, 'honorario quincenal', 'mitad del honorario mensual', 1, m)
    R(d, k, 'a) B4*12', 'a) C2*12', 1, m); R(d, k, 'b) fórmula B4*12', 'b) fórmula C2*12', 1, m)
    R(d, k, 'c) 12*B4=', 'c) 12*C2=', 1, m)
    m2 = 'H-2-19: los distractores usaban la columna B; ahora solo cambia el separador, que es lo que se evalúa'
    R(d, k, 'b) B2-B17', 'b) C2-C17', 1, m2); R(d, k, 'c) B2..B17', 'c) C2..C17', 1, m2); R(d, k, 'd) B2;B17', 'd) C2;C17', 1, m2)
    R(d, k, '=C2+C3+…+B17', '=C2+C3+…+C17', 2, 'H-2-02: el rango de honorarios termina en C17')
    R(d, k, 'G. 3.044.000 es el salario mínimo vigente.',
      'G. 3.044.000 es el salario mínimo mensual vigente desde el 1 de julio de 2026 (Decreto 6225/2026).', 1,
      'H-2-12: fecha y norma del dato que más rápido cambia')
    R(d, k, 'Usa HOY() y cálculos entre fechas en documentos del estudio.',
      'Usa HOY(), AHORA() y cálculos entre fechas y horas en documentos del estudio.', 1,
      'H-2-13: indicador sincronizado con plan anual y planes de clase')
    # Clase 13: subtítulo vacío y párrafo largo con estilo de título
    ps = d.paragraphs
    i = next(j for j, p in enumerate(ps) if p.text.startswith('Descomponer un trazado:'))
    desc, sig = ps[i], ps[i + 1]
    assert ps[i - 1].text == 'Del rectángulo a la forma propia' and sig.text.startswith('Un rectángulo convertido a curvas')
    desc.style = d.styles['Normal']
    texto = desc.text
    U.poner_texto(desc, '')
    r1 = desc.runs[0]; r1.text = 'Descomponer un trazado: '; r1.bold = True
    r2 = desc.add_run(texto[len('Descomponer un trazado: '):]); r2.bold = False
    sig._p.addnext(desc._p)                       # va después de la explicación del rectángulo convertido a curvas
    U.REGISTRO.append({'documento': 'Tomo', 'antes': 'Párrafo «Descomponer un trazado…» con estilo Título 3, antes del texto',
                       'despues': 'Párrafo normal con entrada en negrita, después de «Un rectángulo convertido a curvas…»',
                       'apariciones': 1, 'motivo': 'H-2-14: un párrafo largo aparecía como título en la navegación'})
    m = 'H-2-15: redundancia y consigna demasiado absoluta (la conversión a curvas depende de la imprenta)'
    R(d, k, 'PDF para imprenta para la imprenta', 'PDF preparado para la imprenta', 1, m)
    R(d, k, 'PDF para imprenta (imprenta)', 'PDF para imprenta', 2, m)
    R(d, k, '2.2 Para la imprenta se exporta en PDF con las tipografías convertidas a __________.',
      '2.2 Si la imprenta lo solicita, las tipografías del PDF se convierten a __________.', 1, m)
    R(d, k, '2.1 Para la imprenta se exporta un PDF con las tipografías convertidas a __________.',
      '2.1 Si la imprenta lo solicita, las tipografías del PDF se convierten a __________.', 1, m)
    m = 'H-2-08: el MEF fue creado por la Ley 7158/2023, no por la ley de la DNIT'
    R(d, k, 'que unificó tributos internos y aduanas; el Ministerio de Hacienda pasó a llamarse Ministerio de Economía y Finanzas (MEF).',
      'que unificó tributos internos y aduanas. Por su parte, la Ley 7158/2023 creó el Ministerio de Economía y Finanzas (MEF), que absorbió al Ministerio de Hacienda.', 1, m)
    R(d, k, '2.3 El ex Ministerio de Hacienda hoy se llama __________.',
      '2.3 Las funciones del ex Ministerio de Hacienda las ejerce hoy el __________.', 1, m)
    m = 'H-2-04: la entrega depende del receptor (XML al electrónico; KuDE al no electrónico o consumidor final); envío diferido'
    R(d, k, 'recién ahí tiene efecto fiscal. El receptor recibe el KuDE, la representación gráfica con código QR que permite verificar el documento en el portal. El KuDE no es la factura: es su imagen verificable.',
      'recién ahí tiene efecto fiscal. En e-Kuatia el envío es diferido: el emisor firma el DE y debe transmitirlo al SIFEN dentro del plazo que fija la normativa técnica (hasta 72 horas desde la firma, según la DNIT). La entrega depende del receptor: si también es facturador electrónico, recibe el archivo electrónico (XML); si no lo es, o si es consumidor final, recibe el KuDE en papel o en formato digital. El KuDE es la representación gráfica del documento, con su código de control (CDC) y un código QR para verificarlo en el portal; no reemplaza al documento electrónico.', 1, m)
    R(d, k, 'útil para entregar y verificar, pero no reemplaza al documento electrónico.',
      'que se entrega al receptor no electrónico o al consumidor final y sirve para verificar, pero no reemplaza al documento electrónico.', 1, m)
    R(d, k, 'Figura 13. El recorrido de una factura electrónica: de DE a DTE y KuDE.',
      'Figura 13. Recorrido simplificado de una factura electrónica: del DE al DTE y su entrega según el receptor.', 1, m)
    R(d, k, 'Figura 5. Cuatro tipos de gráfico, cuatro historias con los mismos datos del estudio.',
      'Figura 5. Cuatro tipos de gráfico, cada uno para una pregunta distinta, con datos del estudio (millones de guaraníes).', 1,
      'H-2-06: figura reconstruida con los datos reales del caso')
    nb = U.saltos_a_salto_antes(d)
    U.REGISTRO.append({'documento': 'Tomo', 'antes': 'Salto de página manual en párrafo propio antes de cada evaluación',
                       'despues': 'Salto de página anterior en ese párrafo', 'apariciones': nb,
                       'motivo': 'H-2-21: evita páginas en blanco cuando la clase termina justo al pie'})
    n = U.renglones_respuesta(d, lambda sec: {'clase': 3, 'unidad': 3, 'etapa': 4}[sec])
    U.REGISTRO.append({'documento': 'Tomo', 'antes': 'Preguntas abiertas sin espacio para responder',
                       'despues': '%d renglones (3 por pregunta en clases y evaluaciones de unidad, 4 en las de etapa)' % n,
                       'apariciones': n, 'motivo': 'H-2-20: el estudiante trabaja sobre el tomo; las abiertas llevan renglones'})
    # figuras R2 (orden del cuerpo = Figura 1 … 13, verificado por contenido contra Antes_R1)
    for n in range(1, 14):
        U.reemplazar_imagen(d, n, os.path.join(FIG, 'R2_Figura_%02d.png' % n), 15.5)
    return d


# ================================================================== LABORATORIO
def lab():
    d = abrir('lab'); k = 'lab'
    R(d, k, 'Revisión de contenido octubre de 2026.', 'Revisión R2, 22 de octubre de 2026.', 1, 'Identificación de la versión')
    R(d, k, 'las Prácticas 1, 4, 9, 11, 15 y 20 se integran', 'las Prácticas 1, 4, 9, 11, 15 y 19 se integran', 1,
      'H-2-05: la Práctica 20 pasa a encuentro propio y la 19 se integra a su clase')
    R(d, k, '5. Activá el filtro automático sobre la fila 1.',
      '5. Seleccioná A1:C17, sin la fila 18 del total, y activá el filtro automático: así la fila del total no queda dentro del filtro.', 1,
      'H-2-09: con la fila 18 contigua, el autofiltro la incluye y la oculta al filtrar')
    R(d, k, '▸ Exportación: PDF para imprenta para imprenta · PNG para redes', '▸ Exportación: PDF para imprenta · PNG para redes', 1, 'H-2-15')
    m = 'H-2-04: entrega según el receptor'
    R(d, k, '▸ DE → validación del SIFEN → DTE → KuDE con QR',
      '▸ DE → validación del SIFEN → DTE → entrega: XML al receptor electrónico o KuDE con QR al no electrónico', 1, m)
    R(d, k, '5. Rotulá cada caja con su nombre técnico: DE, validación SIFEN, DTE y KuDE.',
      '5. Rotulá cada caja con su nombre técnico: DE, validación SIFEN, DTE y entrega (XML o KuDE, según el receptor).', 1, m)
    # 12 cm como en R1; en las Prácticas 10 y 14 (páginas más cargadas) la figura va a 8 y 9,5 cm y las anotaciones
    # pasan de tres a un renglón, para que cada práctica siga ocupando una sola página.
    for n in range(1, 14):
        U.reemplazar_imagen(d, n, os.path.join(FIG, 'R2_Figura_%02d.png' % n), {6: 8.0, 9: 9.5}.get(n, ANCHO_LAB))
    for prac in ('Práctica 10 ·', 'Práctica 14 ·'):
        ps = d.paragraphs
        i = next(j for j, p in enumerate(ps) if p.text.startswith(prac) and p.style.name.startswith('Heading'))
        j = next(j for j in range(i, len(ps)) if ps[j].text == 'Anotaciones')
        for q in (ps[j + 3], ps[j + 2]):
            q._p.getparent().remove(q._p)
    U.REGISTRO.append({'documento': 'Laboratorio', 'antes': 'Figuras a 12 cm y tres renglones de anotaciones en todas las prácticas',
                       'despues': 'Práctica 10: figura a 8 cm; Práctica 14: 9,5 cm; un renglón de anotaciones en ambas', 'apariciones': 2,
                       'motivo': 'Las figuras R2 (3:2) son más altas: se conserva una práctica por página'})
    U.REGISTRO.append({'documento': 'Laboratorio', 'antes': 'Imágenes 9 y 10 de una versión anterior (la 10 con =SI(B2>=…))',
                       'despues': 'Figuras R2 9 y 10, iguales a las del tomo', 'apariciones': 2,
                       'motivo': 'H-2-03: la figura de la Práctica 16 contradecía la estructura (C = honorario)'})
    return d


# ================================================================== PLAN ANUAL
CAPACIDADES = {
    1: 'Utiliza planillas electrónicas para el procesamiento de la información contable.',
    2: 'Utiliza planillas electrónicas para el procesamiento de la información contable.',
    3: 'Elabora informes financieros con datos estadísticos y gráficos.',
    4: 'Identifica el diseño gráfico a presentar. · Diseña materiales gráficos para el plan promocional de la empresa con creatividad.',
    5: 'Diseña planillas electrónicas para la inserción de datos administrativos – contables y la toma de decisiones de una empresa con un enfoque profesional.',
    6: 'Recaba en la página web de la SET informaciones útiles inherentes a su especialidad (programa de referencia; hoy, portal de la DNIT).'}


def plan():
    d = abrir('plan'); k = 'plan'
    R(d, k, 'Revisión de contenido octubre de 2026.', 'Revisión R2, 22 de octubre de 2026.', 1, 'Identificación de la versión')
    R(d, k, '(Prácticas 1, 4, 9, 11, 15 y 20)', '(Prácticas 1, 4, 9, 11, 15 y 19)', 1, 'H-2-05')
    R(d, k, 'Las tasas de Sueldos se guardan en G1 y G2.',
      'Las tasas de Sueldos se guardan en G1 y G2. El programa de referencia nombra la página web de la SET; desde 2024 esas funciones corresponden a la DNIT (Ley 7143/2023).', 1,
      'H-2-16: aclarar la actualización institucional sin atribuirla al MEC')
    t = d.tables[0]
    for fila in t.rows:
        cs = list(U.celdas_unicas(fila))
        if len(cs) == 1 and cs[0].text.startswith('UNIDAD '):
            n = int(cs[0].text.split()[1])
            p = cs[0].paragraphs[-1]
            r = p.add_run(' — Capacidad: ' + CAPACIDADES[n]); r.bold = False
            U.REGISTRO.append({'documento': 'Plan_anual', 'antes': cs[0].text.split(' — Capacidad')[0],
                               'despues': 'Banda de unidad con la capacidad del programa', 'apariciones': 1,
                               'motivo': 'H-2-07: el plan anual no mostraba las capacidades'})
        elif len(cs) == 7 and cs[0].text in ('19', '20'):
            if cs[0].text == '19':
                U.poner_celda(cs[3], ['2 h'])
                U.poner_celda(cs[4], [cs[4].text.rstrip('.') + ' · Práctica 19 del Cuaderno de Laboratorio integrada en este mismo encuentro.'])
            else:
                U.poner_celda(cs[3], ['4 h'])
                U.poner_celda(cs[4], [cs[4].text.replace(' · Práctica 20 del Cuaderno de Laboratorio integrada en este mismo encuentro.',
                                                         ' · Práctica 20 del Cuaderno de Laboratorio en encuentro propio.')])
            U.REGISTRO.append({'documento': 'Plan_anual', 'antes': 'Fila %s' % cs[0].text, 'despues': 'Tiempo y procedimiento redistribuidos',
                               'apariciones': 1, 'motivo': 'H-2-05: la Clase 20, la Práctica 20 y la evaluación de la Unidad 6 no caben en 80 min'})
    return d


# ================================================================== SOLUCIONARIO
def sol():
    d = abrir('sol'); k = 'sol'
    R(d, k, 'Revisión de contenido octubre de 2026.', 'Revisión R2, 22 de octubre de 2026.', 1, 'Identificación de la versión')
    R(d, k, '1.1 d) =B4*12', '1.1 d) =C2*12', 1, 'H-2-01')
    R(d, k, '1.1 b) una esquina con la proporción bloqueada', '1.1 b) una esquina', 1, 'Clave igual al texto de la opción impresa')
    m = 'H-2-04'
    R(d, k, '4) el receptor recibe el KuDE con QR para verificarlo.',
      '4) se entrega al receptor: el XML si es facturador electrónico, o el KuDE con QR si no lo es o si es consumidor final.', 1, m)
    R(d, k, '→ entrega del KuDE con QR al receptor.',
      '→ entrega al receptor (XML al receptor electrónico; KuDE con QR al no electrónico o al consumidor final).', 1, m)
    R(d, k, 'El volante exportado en PDF conserva su apariencia y el PNG del logo tiene fondo transparente. Los círculos soldados forman una sola figura; los agrupados siguen siendo dos.',
      'Hay tres PDF (volante A4, tarjeta 90 × 50 mm y banner 2000 × 800 mm), un PNG transparente y sus originales editables. La unión es una geometría y el grupo conserva dos objetos.', 1,
      'H-2-10: punto de control igual al del laboratorio')
    R(d, k, '3.1 V — Y por lo tanto no entra en las sumas.',
      '3.1 V — Suele indicarlo con la alineación predeterminada; se confirma con ESNUMERO y, si es texto, no entra en las sumas.', 1,
      'H-2-17: coherencia con la Clase 2 (la alineación no prueba el tipo de dato)')
    R(d, k, 'Doctrina legal aplicada (verificada en agosto de 2026)',
      'Doctrina legal aplicada (verificada en agosto de 2026; revisada el 22 de octubre de 2026)', 1, 'H-2-11')
    R(d, k, 'RG 41/2025 para contratos estatales desde 2/1/2026 y cronograma progresivo, ampliado por RG 52/2026.',
      'RG 41/2025 para contratos estatales desde 2/1/2026 y cronograma progresivo, ampliado por RG 52/2026. Entrega: XML al receptor electrónico; KuDE al no electrónico o al consumidor final. Pendiente: confirmar si la RG 54/2026 modifica la RG 41/2025.', 1, 'H-2-11')
    R(d, k, 'Salario mínimo vigente por Decreto 6225/2026;',
      'Salario mínimo mensual G. 3.044.000 desde el 1/7/2026 (Decreto 6225/2026);', 1, 'H-2-12')
    return d


# ================================================================== PLANES DE CLASE
def bloques_planes(d):
    """{n: [tabla institución, datos, capacidad, momentos, recursos, evaluación]} y el párrafo de título de cada plan."""
    body = list(d.element.body.iterchildren()); res = {}
    for i, el in enumerate(body):
        if el.tag == qn('w:p'):
            t = ''.join(x.text or '' for x in el.iter(qn('w:t')))
            m = re.match(r'Plan de Clase N\.º (\d+) — ', t)
            if m and body[i + 1].tag == qn('w:tbl'):
                tabs = [docx.table.Table(e, d._body) for e in body[i + 1:i + 9] if e.tag == qn('w:tbl')][:6]
                res[int(m.group(1))] = (docx.text.paragraph.Paragraph(el, d._body), tabs)
    return res


def fila(tabla, rotulo):
    for r in tabla.rows:
        cs = list(U.celdas_unicas(r))
        if cs[0].text.replace('\n', ' ').strip().startswith(rotulo):
            return cs
    raise SystemExit('fila %s no encontrada' % rotulo)


def poner_plan(bl, titulo, material, tema, inicio, desarrollo, cierre, recursos, proc, instr, ind=None, tiempos=None):
    tp, (inst, datos, cap, mom, rec, ev) = bl
    U.poner_texto(tp, titulo)
    U.poner_celda([c for c in U.celdas_unicas(datos.rows[3])][1], [material])
    U.poner_celda([c for c in U.celdas_unicas(datos.rows[4])][1], [tema])
    if ind:
        U.poner_celda(fila(cap, 'Indicadores')[1], [ind])
    for rot, lineas, t in (('Inicio', inicio, tiempos[0]), ('Desarrollo', desarrollo, tiempos[1]), ('Cierre', cierre, tiempos[2])):
        cs = fila(mom, rot); U.poner_celda(cs[1], lineas); U.poner_celda(cs[2], [t])
    U.poner_celda(fila(rec, 'Recursos')[1], [recursos])
    U.poner_celda(fila(ev, 'Procedimientos')[1], [proc]); U.poner_celda(fila(ev, 'Instrumentos')[1], [instr])
    if ind:
        U.poner_celda(fila(ev, 'Criterios')[1], [ind])


def planes():
    d = abrir('planes'); k = 'planes'
    R(d, k, 'Revisión de contenido octubre de 2026.', 'Revisión R2, 22 de octubre de 2026.', 1, 'Identificación de la versión')
    B = bloques_planes(d)
    assert sorted(B) == list(range(1, 37)), sorted(B)
    ind19 = fila(B[33][1][2], 'Indicadores')[1].text
    ind20 = fila(B[35][1][2], 'Indicadores')[1].text
    proc19 = 'Navegación guiada del portal de la DNIT con ficha de registro · distinción práctica entre servicios con y sin clave · Práctica 19 del Cuaderno de Laboratorio integrada en este mismo encuentro.'
    poner_plan(B[33], 'Plan de Clase N.º 33 — Portal de la DNIT: servicios al contribuyente y sistema Marangatú',
               'Tomo del estudiante · Clase 19 · Cuaderno de Laboratorio · Práctica 19 integrada',
               'Portal de la DNIT: servicios al contribuyente y sistema Marangatú',
               ['• Recuperación de saberes previos y presentación del propósito de la Clase 19.'],
               ['• ' + proc19, '• Práctica 19 integrada: Navegación guiada por el portal de la DNIT.'],
               ['• Síntesis, revisión de la ficha de navegación y retroalimentación.'],
               'Tomo del estudiante · Cuaderno de Laboratorio · Práctica 19 · Computadora del gabinete con conexión a internet',
               proc19, 'Ficha de navegación del portal con cinco datos, su sección de origen y su fecha', tiempos=('10 min', '60 min', '10 min'))
    proc20 = 'Reconstrucción gráfica del circuito del documento electrónico (DE, DTE, KuDE y entrega según el receptor) · consulta de RG DNIT 41/2025 y del cronograma progresivo vigente (incluida RG 52/2026)'
    poner_plan(B[34], 'Plan de Clase N.º 34 — Factura electrónica (e-Kuatia / e-Kuatia’i) y normativas legales vigentes',
               'Tomo del estudiante · Clase 20', 'Factura electrónica (e-Kuatia / e-Kuatia’i) y normativas legales vigentes',
               ['• Recuperación de saberes previos y presentación del propósito de la Clase 20 a partir de sus indicadores.'],
               ['• ' + proc20, '• Desarrollo conceptual y resolución guiada de las actividades del tomo correspondientes al tema.'],
               ['• Síntesis de los aprendizajes y preparación de la continuación práctica en laboratorio.'],
               'Tomo del estudiante · Clase 20 · Computadora del gabinete con conexión a internet · Pizarra',
               proc20, 'Producción escrita: respuesta del estudio al cliente sobre facturación electrónica', ind=ind20,
               tiempos=('10 min', '60 min', '10 min'))
    # el plan 34 heredaba la capacidad de la Unidad 6: es la misma; el título de la tabla de datos se actualiza abajo
    poner_plan(B[35], 'Plan de Clase N.º 35 — Continuación práctica de la Clase 20: El circuito de la factura electrónica',
               'Cuaderno de Laboratorio · Práctica 20 · Evaluación de la Unidad 6', 'Práctica 20 · El circuito de la factura electrónica',
               ['• Organización del puesto y recuperación de las consignas de la Clase 20.'],
               ['• Desarrollo de la Práctica 20: El circuito de la factura electrónica.',
                '• Búsqueda de la designación del contribuyente en el portal y redacción de la respuesta al cliente; comprobación del punto de control.'],
               ['• Aplicación de la evaluación integrada de la Unidad 6, entrega del producto y retroalimentación.'],
               'Cuaderno de Laboratorio · Práctica 20 · Computadora del gabinete con conexión a internet · Evaluación de la Unidad 6 impresa',
               'Práctica 20: El circuito de la factura electrónica · búsqueda de la designación del contribuyente y redacción de la respuesta al cliente',
               'Producción escrita: respuesta del estudio al cliente sobre facturación electrónica · Evaluación de la Unidad 6', ind=ind20,
               tiempos=('10 min', '50 min', '20 min'))
    U.REGISTRO.append({'documento': 'Planes_de_clase', 'antes': 'Plan 33 Clase 19 · Plan 34 Práctica 19 · Plan 35 Clase 20 + Práctica 20 + evaluación U6 en 80 min',
                       'despues': 'Plan 33 Clase 19 + Práctica 19 · Plan 34 Clase 20 · Plan 35 Práctica 20 + evaluación U6', 'apariciones': 3,
                       'motivo': 'H-2-05: redistribución sin cambiar los 36 encuentros ni las 72 HC'})
    n = U.escalar_letra(d, {7.5: 9.5, 8.0: 9.5})
    U.REGISTRO.append({'documento': 'Planes_de_clase', 'antes': 'Cuerpo de 7,5 pt en las tablas de los planes',
                       'despues': 'Cuerpo de 9,5 pt; cada plan sigue en su página', 'apariciones': n,
                       'motivo': 'H-2-22: legibilidad; cada plan ocupaba un tercio de la página de oficio'})
    k2 = 0
    for p in U.parrafos(d):                       # el renglón del docente ya no entra en su celda con 9,5 pt
        if p.text == '_' * 30:
            U.poner_texto(p, '_' * 22); k2 += 1
    U.REGISTRO.append({'documento': 'Planes_de_clase', 'antes': 'Renglón de 30 guiones en la celda Docente', 'despues': '22 guiones',
                       'apariciones': k2, 'motivo': 'H-2-22: evitar que el renglón se parta en dos líneas'})
    # índice de planes: títulos 34 y 35
    for p in U.parrafos(d):
        if p.text.startswith('Plan de Clase N.º 34 — ') and '·' in p.text[-6:]:
            U.poner_texto(p, 'Plan de Clase N.º 34 — Factura electrónica (e-Kuatia / e-Kuatia’i) y normativas legal…  ·  35')
        elif p.text.startswith('Plan de Clase N.º 35 — ') and '·' in p.text[-6:]:
            U.poner_texto(p, 'Plan de Clase N.º 35 — Continuación práctica de la Clase 20: El circuito de la factura…  ·  36')
    # celdas de Desarrollo con dos viñetas pegadas en un mismo párrafo
    arregladas = 0
    for n, (_, tabs) in B.items():
        cs = fila(tabs[3], 'Desarrollo')
        t = cs[1].text
        if len(cs[1].paragraphs) == 1 and re.search(r'\S• ', t):
            partes = [x.strip() for x in re.split(r'(?<=\S)(?=• )', t.strip()) if x.strip()]
            U.poner_celda(cs[1], partes); arregladas += 1
    U.REGISTRO.append({'documento': 'Planes_de_clase', 'antes': 'Viñetas pegadas («…contenido• Desarrollo…») y salto final',
                       'despues': 'Una viñeta por párrafo', 'apariciones': arregladas, 'motivo': 'H-2-18: maquetación de la celda Desarrollo'})
    return d


# ================================================================== construcción
CLAVES_INDICE = {k: (lambda t: t.rstrip('…').strip()) for k in ('tomo', 'lab', 'sol', 'planes')}

if __name__ == '__main__':
    os.makedirs(R2, exist_ok=True); os.makedirs(PDF, exist_ok=True)
    docs = {'tomo': tomo(), 'lab': lab(), 'plan': plan(), 'sol': sol(), 'planes': planes()}
    indices = {}
    for k, d in docs.items():
        d.save(ruta_r2(k))
    for pasada in (1, 2):
        for k in ('tomo', 'lab', 'sol', 'planes'):
            pdf = U.a_pdf(ruta_r2(k), PDF)
            d = docx.Document(ruta_r2(k))
            indices[k] = U.recalcular_indice(d, pdf, clave=CLAVES_INDICE[k])
            d.save(ruta_r2(k))
    U.a_pdf(ruta_r2('plan'), PDF)
    for k in ('tomo', 'lab', 'sol', 'planes'):
        U.a_pdf(ruta_r2(k), PDF)
    json.dump({'cambios': U.REGISTRO, 'indices': indices}, open(os.path.join(BASE, 'build', 'registro_r2_2btc.json'), 'w', encoding='utf8'),
              ensure_ascii=False, indent=1)
    print(len(U.REGISTRO), 'cambios registrados')
    for k, v in indices.items():
        print(k, 'entradas', len(v), 'corregidas', sum(1 for _, a, b in v if a != b))
