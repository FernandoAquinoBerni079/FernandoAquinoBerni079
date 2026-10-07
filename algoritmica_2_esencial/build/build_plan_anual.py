# -*- coding: utf-8 -*-
"""Plan Anual — Edición Esencial Comercial 2026. 36 filas (una por encuentro, misma numeración que los
Planes de Clase). Temas e indicadores leídos de las fichas del libro; procedimientos e instrumentos únicos."""
import re
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import maqueta as M
import planes_data
import ediciones
from build_libro import a_pdf

B = '/home/claude/alg2/build/'; OUT = '/home/claude/alg2/salida/'
NOMBRE = 'Algoritmica_2do_Curso_PLAN_ANUAL_ESENCIAL_COMERCIAL_2026'
ANCHOS = [1.0, 5.0, 6.8, 1.5, 5.3, 4.2, 2.8]
COLS = ['N.º', 'TEMAS', 'INDICADORES DE LOGRO', 'TIEMPO', 'PROCEDIMIENTOS', 'INSTRUMENTOS DE EVALUACIÓN', 'OBSERVACIÓN']
DARK, MED, LIGHT, AZUL, AZUL_CLARO = '1B5E20', '2E7D32', 'E8F5E9', '1B4F72', 'EAF0F6'
CAP_UNIDAD = {u: ediciones.caps(c) for u, c in ediciones.CAP_UNIDAD.items()}
NOTA_CAP = ('Las capacidades se reproducen textualmente del programa MEC vigente. Los indicadores de logro de este plan son operativizaciones '
            'observables elaboradas para organizar la enseñanza y la evaluación de cada encuentro.')
assert sum(ANCHOS) == 26.6


def celda(cell, texto, bold=False, blanco=False, size=8):
    par = cell.paragraphs[0]
    for i, linea in enumerate(texto if isinstance(texto, list) else [texto]):
        if i:
            par = cell.add_paragraph()
        par.paragraph_format.space_after = Pt(1)
        r = par.add_run(linea); r.bold = bold; r.font.size = Pt(size)
        if blanco:
            r.font.color.rgb = RGBColor.from_string('FFFFFF')


def filas_plan():
    PL = planes_data.planes()
    rows = []
    for p in PL:
        t = p['tipo']
        if t == 'C':
            tema = ['Clase %d — %s' % (p['n'], p['titulo'])] + ([p['tema']] if p['tema'] != p['titulo'] else [])
            obs = 'Clase y Práctica %d del libro.' % p['n']
            if p['n'] in planes_data.CONT:
                obs += ' Continúa en el encuentro siguiente.'
        elif t == 'P':
            tema = ['Continuidad práctica · Práctica %d' % p['n'], planes_data.practicas.PR[p['n']]['titulo'] + '.']
            obs = 'Actividades 2 y 3 de la Práctica %d.' % p['n']
        elif t == 'T':
            tema = [p['titulo']]
            obs = 'Proyecto Final Integrador · Feria de Informática.'
        else:
            tema = [p['titulo']]
            obs = 'Cierre de la %s etapa — %s.' % (('1.ª', 'junio') if p['n'] == 'E1' else ('2.ª', 'noviembre'))
        ev = [k for k, v in planes_data.EVAL_UNIDAD_EN.items() if k == (t, p['n'])]
        if ev:
            obs += ' Cierre de la Unidad %s: evaluación de unidad.' % planes_data.EVAL_UNIDAD_EN[ev[0]][1]
        rows.append(dict(n=p['idx'], tipo=t, unidad=p['unidad'], tema=tema, indicadores=p['indicadores'], tiempo='4 h',
                         proc=p['proc'], inst=p['inst'], obs=obs))
    # unicidad estricta y sin códigos internos
    assert len({r['proc'] for r in rows}) == 36 and len({r['inst'] for r in rows}) == 36
    for r in rows:
        for campo in ('proc', 'inst', 'obs'):
            assert not re.search(r'\b\d-\d\d\b|IMG_\d+|TXT_\w+', r[campo]), r
    return rows


def construir():
    d = M.nuevo_doc('Algorítmica · 2.º Curso · Plan anual', 'Material del docente: plan anual de 36 encuentros', 'Algorítmica; BTI; 2.º Curso; plan anual')
    s0 = d.sections[0]
    M.imagen_pagina_completa(d, B + 'portadas/portada_plan_anual.jpg')
    sec = d.add_section(WD_SECTION.NEW_PAGE)
    for hf in (s0.footer, s0.header):
        hf.is_linked_to_previous = False
        for p in hf.paragraphs:
            for el in list(p._p):
                if el.tag != qn('w:pPr'):
                    p._p.remove(el)
    M.formato_oficio(sec, horizontal=True)
    M.ANCHO = round(M.PAG_H - 2 * M.MARG, 2)   # 30,46 cm útiles en oficio horizontal
    M.par(d, 'PLAN ANUAL DETALLADO', 14, True, color=DARK, align='c', despues=2)
    M.par(d, 'Disciplina: Algorítmica  ·  Curso: 2.º Curso  ·  Bachillerato Técnico en Servicios · Especialidad Informática  ·  4 horas cátedra semanales', 9.5, align='c', despues=1)
    M.par(d, 'República del Paraguay · 2026', 9, italic=True, align='c', despues=1)
    M.par(d, 'Institución educativa: ______________________________     Docente: ______________________________', 9, align='c', despues=4)
    pp = M.par(d, '', 8.5, despues=4); M.run(pp, 'Nota sobre capacidades e indicadores. ', True, size=8.5, color=DARK); M.run(pp, NOTA_CAP, size=8.5)
    M.par(d, 'Distribución anual: 36 encuentros de 4 horas cátedra (144 HC), 18 en cada etapa. Cada fila es un encuentro y tiene su Plan de Clase con el mismo número. Las Clases 7, 10, 14 y 15 a 21 continúan en un encuentro propio (Actividades 2 y 3 de su práctica); el Proyecto Final Integrador ocupa tres talleres y cada evaluación integradora de etapa, un encuentro. Material de referencia: el libro del estudiante, que reúne cada Clase con su Práctica.', 8.5, despues=4)
    rows = filas_plan()
    t = d.add_table(rows=1, cols=7)
    M.bordes(t, '808080', '4'); M.margenes_celda(t, 30, 70); M.ancho_fijo(t, ANCHOS)
    hdr = t.rows[0]
    for j, h in enumerate(COLS):
        M.shd(hdr.cells[j], DARK); celda(hdr.cells[j], h, bold=True, blanco=True, size=8)

    def fila():
        r = t.add_row()
        for j, c in enumerate(r.cells):
            c.width = Cm(ANCHOS[j] * M.ANCHO / sum(ANCHOS))
        return r

    def banda(texto, fill, size=9):
        r = fila(); m = r.cells[0].merge(r.cells[6]); M.shd(m, fill); celda(m, texto, bold=True, blanco=True, size=size)

    u_prev = None
    for r_ in rows:
        if r_['n'] == 1:
            banda('1.ª ETAPA · febrero – junio · 18 encuentros de 4 HC', AZUL)
        if r_['n'] == 19:
            banda('2.ª ETAPA · julio – noviembre · 18 encuentros de 4 HC', AZUL)
        if r_['tipo'] in ('C', 'P') and r_['unidad'] != u_prev:
            u_prev = r_['unidad']
            un = int(re.search(r'UNIDAD (\d)', u_prev).group(1))
            banda('%s · Capacidades del programa MEC: %s' % (u_prev, ' · '.join(CAP_UNIDAD[un])), MED, 8.5)
        if r_['tipo'] == 'T' and u_prev != 'PFI':
            u_prev = 'PFI'
            banda('PROYECTO FINAL INTEGRADOR — El sistema de gestión de un emprendimiento · Feria de Informática · Capacidades del programa MEC: %s' % ' · '.join(ediciones.caps(ediciones.CAP_TALLER)), MED, 8.5)
        rr = fila()
        vals = [str(r_['n']), r_['tema'], ['• ' + i for i in r_['indicadores']], r_['tiempo'], r_['proc'], r_['inst'], r_['obs']]
        for j, v in enumerate(vals):
            celda(rr.cells[j], v, bold=(j == 0 or (j == 1 and r_['tipo'] == 'E')))
        if r_['tipo'] == 'E':
            for c in rr.cells:
                M.shd(c, AZUL_CLARO)
        else:
            M.shd(rr.cells[0], LIGHT)
    M.cant_split(t, header=True)
    M.par(d, '', despues=4)
    M.par(d, 'Criterios generales de planificación', 10, True, color=DARK, keep=True)
    for x in ['Enfoque constructivista «Aprender Haciendo»: cada clase parte de un problema del caso integrador del Copetín Karumbé; las prácticas usan otro contexto (la Librería escolar Arandu) para transferir lo aprendido.',
              'Cada clase se acompaña de su práctica, incluida en el libro del estudiante a continuación de la clase y con la misma numeración, en PSeInt, en el Explorador de archivos o en Microsoft Access.',
              'La evaluación es continua y formativa: los instrumentos por encuentro se complementan con las evaluaciones de unidad y las integradoras de etapa (junio y noviembre) y con la rúbrica del Proyecto Final Integrador.',
              'La numeración de encuentros es continua (1 a 36) y coincide con la de los Planes de Clase; la carga es de 4 horas cátedra semanales, con flexibilidad para reforzar según el ritmo del grupo.']:
        M.par(d, '• ' + x, 8.5, despues=1)
    M.par(d, 'Estrategias de inclusión', 10, True, color=DARK, keep=True, antes=4)
    for x in ['Consignas presentadas en forma oral y escrita, con ejemplos resueltos como modelo antes de la ejercitación individual.',
              'Trabajo en parejas o pequeños grupos heterogéneos en las prácticas, con roles rotativos y andamiaje entre pares.',
              'Tiempo adicional y ejercitación graduada para quienes lo requieran; desafíos finales para quienes terminan antes.',
              'Apoyos visuales (figuras, tablas y diagramas) y alternativa en papel cuando el laboratorio no está disponible.',
              'Adecuaciones razonables de acceso (ubicación, tamaño de fuente, apoyos técnicos) según las necesidades documentadas del estudiante.']:
        M.par(d, '• ' + x, 8.5, despues=1)
    M.par(d, '', despues=14)
    M.par(d, '__________________________________                                        __________________________________', 9.5, align='c', despues=0)
    M.par(d, 'Firma del docente                                                                         Dirección — Visto bueno', 9, align='c')
    M.pie_paginas(d, sec)
    M.no_actualizar_campos(d)
    M.purgar_relaciones(d)
    return d, rows


if __name__ == '__main__':
    d, rows = construir()
    out = OUT + NOMBRE + '.docx'; d.save(out); pdf = a_pdf(out, OUT)
    import pymupdf
    doc = pymupdf.open(pdf)
    print('Plan anual: páginas', doc.page_count, '· filas', len(rows))
