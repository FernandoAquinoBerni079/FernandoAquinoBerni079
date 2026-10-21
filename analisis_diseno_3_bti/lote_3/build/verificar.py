# -*- coding: utf-8 -*-
"""Controles automáticos del lote 3 (Unidad III). Uso: python3 lote_3/build/verificar.py   (sale con 1 si hay fallos)

Bloques del lote 2, extendidos a los Capítulos 7 a 9: A índice de Work · B capacidades · C trazabilidad de fuentes ·
E Algorítmica 3.º · F voseo y residuos · G maqueta (DOCX) · H maqueta (PDF) · I evaluaciones, prácticas e instrumentos ·
J numeración · L notaciones.
Bloques nuevos (plan del lote 3, §9): N obra, ejemplar e historial · D catálogo v1.1 y coherencia · M casos de uso,
mensajes y contratos · P arquitectura y decisión técnica · Q lotes 1 y 2 intactos."""
import os, re, sys, json, collections, subprocess
sys.dont_write_bytecode = True
AQUI = os.path.dirname(os.path.abspath(__file__))
L3 = os.path.dirname(AQUI)
ADS = os.path.dirname(L3)
sys.path.insert(0, AQUI)
sys.path.insert(1, os.path.join(L3, 'caso'))
sys.path.insert(2, os.path.join(ADS, 'lote_2', 'build'))
from docx import Document
from docx.oxml.ns import qn
import pymupdf

import oficio as O
K = O.K
import construir as B
import cap7, cap8, cap9, guia_u3 as GU
import diagramas_u3 as G
X, MB = B.X, B.MB
R, D, MO = X.R, X.D, X.MO
CAPS = (cap7, cap8, cap9)
EST, GUI, CUA = B.RUTA_EST[:-5], B.RUTA_GUIA[:-5], B.RUTA_CUAD[:-5]

ok = fallos = 0
RES = collections.OrderedDict()
DETALLE = []


def chk(bloque, desc, cond, detalle=''):
    global ok, fallos
    RES.setdefault(bloque, [0, 0])
    cond = bool(cond)
    if cond:
        ok += 1; RES[bloque][0] += 1
        print('CORRECTO [%s] %s' % (bloque, desc))
    else:
        fallos += 1; RES[bloque][1] += 1
        print('FALLA    [%s] %s %s' % (bloque, desc, detalle))
    DETALLE.append({'bloque': bloque, 'control': desc, 'resultado': 'correcto' if cond else 'FALLA'})


def texto_docx(ruta):
    d = Document(ruta)
    partes = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            vistos = set()
            for c in r.cells:
                if id(c._tc) in vistos:
                    continue
                vistos.add(id(c._tc))
                partes.append(c.text)
    return d, '\n'.join(partes)


dE, tE = texto_docx(EST + '.docx')
dG, tG = texto_docx(GUI + '.docx')
dC, tC = texto_docx(CUA + '.docx')
pE, pG, pC = pymupdf.open(EST + '.pdf'), pymupdf.open(GUI + '.pdf'), pymupdf.open(CUA + '.pdf')
res_build = json.load(open(os.path.join(L3, 'resultado_construccion.json'), encoding='utf8'))
todo = tE + '\n' + tG + '\n' + tC
j = lambda x: json.dumps(x, ensure_ascii=False)

# ---------------------------------------------------------------- A. índice de Work (§6, textual; §5.2)
WORK = {
    7: ('Datos que representan las reglas', (13, 14),
        ['7.1. Recuperación aplicada de modelos de Algorítmica.', '7.2. Obra/edición, ejemplar, lector y préstamo: identidad y responsabilidad.',
         '7.3. Relaciones, restricciones, historial y diccionario de datos.', '7.4. Revisión de normalización con ejemplos del proyecto.'],
        'justificar el modelo frente a pérdida, devolución y duplicación de solicitudes'),
    8: ('Diseñar para las tareas de las personas', (15, 16),
        ['8.1. Recorridos de usuario y organización de pantallas.', '8.2. Bocetos y prototipo de baja fidelidad.',
         '8.3. Estados vacío, carga, error y confirmación.', '8.4. Etiquetas, navegación por teclado y lectura comprensible.'],
        'observar el uso del boceto y modificarlo con evidencia'),
    9: ('Arquitectura y revisión del diseño', (17, 18),
        ['9.1. Presentación, reglas de negocio y persistencia.', '9.2. Responsabilidades, entradas, salidas y contratos de operaciones.',
         '9.3. Entorno, dependencias y decisiones técnicas justificadas.', '9.4. Coherencia entre requisitos, datos e interfaces.'], None),
}
sin_punto = lambda s: re.sub(r'^(\d\.\d)\. ', r'\1 ', s).rstrip('.')
heads = [p.text for p in dE.paragraphs if p.style.name.startswith('Heading')]
for C in CAPS:
    tit, enc, secs, prac = WORK[C.NUM]
    chk('A', 'Capítulo %d: título de Work «%s»' % (C.NUM, tit), C.TITULO == tit and '%s — %s' % (C.CAPITULO, tit) in heads)
    chk('A', 'Capítulo %d: encuentros %s' % (C.NUM, enc), tuple(C.ENCUENTROS) == enc)
    chk('A', 'Capítulo %d: secciones de Work §6, textuales y en orden' % C.NUM,
        C.SECCIONES == [sin_punto(s) for s in secs] and [h for h in heads if h in C.SECCIONES] == C.SECCIONES)
    if prac:
        chk('A', 'Capítulo %d: práctica de Work («%s»)' % (C.NUM, prac), C.PRACTICA_WORK == prac)
chk('A', 'Capítulo 9 cierra con la evaluación integradora de la primera etapa y el H3 (Work: «%s»)' % cap9.REVISION_WORK,
    cap9.H3_TITULO in heads and 'H3' in cap9.H3_TITULO and 'Evaluación integradora de la primera etapa' in cap9.H3_TITULO
    and heads.index(cap9.E18_TITULO) < heads.index(cap9.H3_TITULO))
orden = [B.UNIDAD_III] + ['%s — %s' % (C.CAPITULO, C.TITULO) for C in CAPS]
chk('A', 'Orden: apertura de la Unidad III y Capítulos 7, 8 y 9', all(t in heads for t in orden) and [heads.index(t) for t in orden] == sorted(heads.index(t) for t in orden))
chk('A', 'Apertura con el resultado de Work §5.1 («modelo de datos, interacción y arquitectura revisados»)',
    'modelo de datos, interacción y arquitectura revisados' in tE)
encs = [cap7.E13_TITULO, cap7.E14_TITULO, cap8.E15_TITULO, cap8.E16_TITULO, cap9.E17_TITULO, cap9.E18_TITULO]
chk('A', 'Encuentros 13 a 18 como títulos, en orden', all(e in heads for e in encs) and [heads.index(e) for e in encs] == sorted(heads.index(e) for e in encs)
    and [int(re.search(r'\d+', e).group()) for e in encs] == list(range(13, 19)))
chk('A', 'Fuera del lote: ningún título de la Unidad IV ni de los Capítulos 10 a 18', not [h for h in heads if re.match(r'(Unidad IV|Capítulo 1[0-8])\b', h)])

# ---------------------------------------------------------------- B. capacidades
CAPW = {'C5': 'Justifica el modelo de información y sus reglas de integridad.', 'C6': 'Diseña la interacción y la arquitectura del prototipo.',
        'C9': 'Verifica y mejora la solución a partir de evidencia.'}
ESPERADAS = {7: ('C5',), 8: ('C6', 'C9'), 9: ('C5', 'C6')}   # Work §5.2, encuentros 13 a 17
for C in CAPS:
    r = C.ROTULO_CAPACIDAD
    chk('B', 'Capítulo %d: rótulo «propuesta(s) para el programa institucional» y «no oficial del MEC»' % C.NUM,
        re.search(r'propuestas? para el programa institucional', r) and re.search(r'no (es una capacidad oficial|son capacidades oficiales) del MEC', r))
    chk('B', 'Capítulo %d: el rótulo figura en la ficha del libro y de la guía' % C.NUM, r in tE and r in tG)
    cods = ESPERADAS[C.NUM]
    chk('B', 'Capítulo %d: capacidades %s de Work §5.2, con su texto de §4.2' % (C.NUM, ', '.join(cods)),
        all(('%s. %s' % (c, CAPW[c])) in C.CAPACIDAD for c in cods) and all(c in r for c in cods))
chk('B', 'Capítulo 7: C9 solo como articulación, sin atribuirla a Work (no figura en la capacidad del capítulo)',
    'C9' not in cap7.CAPACIDAD and 'C9' not in cap7.ROTULO_CAPACIDAD and 'C9' in ' '.join(x[2] for x in GU.SEC_C7[1][1]))
chk('B', 'Capítulo 9: el H3 evalúa C1 a C7', 'C1 a C7' in cap9.CAPACIDAD and 'C1 a C7' in cap9.ROTULO_CAPACIDAD)
malas_of = [tx[max(0, m.start() - 40):m.end()] for tx in (tE, tG, tC) for m in re.finditer(r'(?i)capacidad(es)? oficial', tx)
            if not re.search(r'\b[Nn]o (es|son)( una)?\b', tx[max(0, m.start() - 40):m.start()])]
chk('B', 'Ninguna capacidad presentada como oficial del MEC', not malas_of, malas_of)

# ---------------------------------------------------------------- C. trazabilidad de fuentes
locales = dict(MB.FUENTES)
for C in CAPS:
    locales.update(C.EVAL_FUENTES)
locales.update(dict(X.ACCESIBILIDAD)); locales.update({v[0]: v[1] for v in X.BANCO_H3})
locales.update({m[0]: m[4] for m in X.MENSAJES})
# Comillas que rotulan un elemento de la pantalla justificado por una regla u observación, no citan una línea.
ROTULOS = {('Código del ejemplar', 'A-01'), ('Confirmar devolución', 'A-06'), ('Prestar', 'A-02'), ('Guardando…', 'OBS-4')}
def fuente(c):
    return R.TEXTO.get(c) or D.NECESIDAD.get(c) or X.REGLA.get(c) or locales.get(c) or (c if X.existe(c) else None)
PAT = (r'(?:E\d-L\d+|O1-\d+|D\d-A?\d+|N-\d\d|V-\d\d|M-[NF]\d|MN-\d\d|MP-\d|MS-\d|MC-\d|MR-\d\d|B-0\d|OBS-\d|VH-\d|[XZU]-\d|A-0\d'
       r'|RI-\d\d|PT-\d\d|ME-\d\d|OP-\d\d|DT-\d\d|CU-\d\d|P-\d\d|RF-\d\d|RNF-\d\d|RS-\d\d|RN-\d\d)')
for nombre, tx in (('libro', tE), ('guía', tG), ('cuaderno', tC)):
    usados = sorted(set(re.findall(r'(?<![\w-])' + PAT + r'(?![\w-])', tx)))
    falt = [c for c in usados if fuente(c) is None]
    chk('C', '%s: los %d códigos citados (dossier, requisitos, reglas, diseño, Mbokaja, evaluación) existen' % (nombre, len(usados)), not falt, falt)
    citas = re.findall(r'«([^»]{6,})»\s*\(((?:%s)(?:,\s*%s)*)\)' % (PAT, PAT), tx)
    malas = [(q[:50], cods) for q, cods in citas if (q, cods) not in ROTULOS
             if not any(fuente(c.strip()) and q.strip().rstrip('.') in fuente(c.strip()) for c in cods.split(','))]
    chk('C', '%s: %d citas textuales «…» (código) coinciden con la línea citada' % (nombre, len(citas)), not malas, malas)
chk('C', 'Diccionario: cada origen citado existe (requisito, regla, necesidad o línea del dossier)',
    all(fuente(o.strip()) for r in X.DICCIONARIO for o in r[5].split(',')), [o for r in X.DICCIONARIO for o in r[5].split(',') if not fuente(o.strip())])
chk('C', 'Diccionario: los 25 campos con significado, obligatoriedad, ejemplo y origen, y en el libro',
    len(X.DICCIONARIO) == 25 and all(all(r) for r in X.DICCIONARIO) and all(r[2] in tE for r in X.DICCIONARIO))
chk('C', 'Restricciones RI-01 a RI-08: cada una sale de una regla, un requisito, una necesidad o una línea que existe',
    [r[0] for r in X.RESTRICCIONES] == ['RI-%02d' % i for i in range(1, 9)] and all(r[2] and all(fuente(f) for f in r[2]) for r in X.RESTRICCIONES))
nuevos = set(c for c, _ in MB.DATOS + MB.SITUACIONES) | set(o[0] for o in MB.OBSERVACIONES)
chk('C', 'Mbokaja U3: los códigos nuevos (M-F8, M-F9, MS-3 a MS-5, OBS-1 a OBS-6) no repiten ninguno anterior',
    not (nuevos & set(MB.MB2.FUENTES)) and len(nuevos) == 11)
r_ = subprocess.run([sys.executable, '-B', os.path.join(ADS, 'lote_1', 'caso', 'verificar_caso.py')], capture_output=True, text=True)
chk('C', 'lote_1/caso/verificar_caso.py sigue sin fallos: %s' % r_.stdout.strip().split('\n')[-1], r_.returncode == 0)

# ---------------------------------------------------------------- N. obra, ejemplar e historial
dic = {(r[0], r[1]): r for r in X.DICCIONARIO}
chk('N', 'Modelo: Obra, Edición, Ejemplar, Lector, Préstamo, Parámetro y Cuenta', [e[0] for e in X.ENTIDADES] == ['Obra', 'Edición', 'Ejemplar', 'Lector', 'Préstamo', 'Parámetro', 'Cuenta'])
chk('N', 'Todo préstamo referencia a un ejemplar (no a la obra): relación Ejemplar 1:N Préstamo y campo Préstamo.ejemplar obligatorio',
    ('Ejemplar', 'se presta en', 'Préstamo', '1', 'N') == X.RELACIONES[2][:5] and dic[('Préstamo', 'ejemplar')][3] == 'Sí'
    and 'no la obra' in dic[('Préstamo', 'ejemplar')][2] and not any(r[0] == 'Préstamo' and r[1] == 'obra' for r in X.DICCIONARIO)
    and not any('Obra' in (r[0], r[2]) and 'Préstamo' in (r[0], r[2]) for r in X.RELACIONES))
chk('N', 'Obra → Edición → Ejemplar, con cardinalidad 1:N', X.RELACIONES[0][:5] == ('Obra', 'tiene', 'Edición', '1', 'N') and X.RELACIONES[1][:5] == ('Edición', 'tiene', 'Ejemplar', '1', 'N'))
chk('N', 'La disponibilidad es derivada: ningún campo «disponible» en el diccionario y figura en la Tabla 7.4',
    not any('disponib' in r[1].lower() for r in X.DICCIONARIO) and any(d[0].startswith('Ejemplar disponible') for d in X.DERIVADOS))
chk('N', 'Historial: la devolución no borra (RI-04, RN-05), lo que tiene historia se inactiva (RI-05, RN-06) y no hay campos para borrar',
    'no se borra' in X.RESTRICCION['RI-04'][0] and 'se inactiva' in X.RESTRICCION['RI-05'][0] and dic[('Préstamo', 'fecha de devolución')][3] == 'No')
chk('N', 'Ninguna operación borra: los resultados de OP-01 a OP-06 no borran nada',
    all(not re.search(r'\bborr', o[4]) or re.search(r'No se borra|no cambia', o[4]) for o in X.OPERACIONES) and 'eliminar' not in j(X.OPERACIONES).lower())
chk('N', 'Curso: obligatorio solo para estudiantes (RI-06), guardado al prestar (RI-08), y tipo de lector aparte (RF-11)',
    dic[('Lector', 'curso')][3] == 'Si es estudiante' and dic[('Préstamo', 'curso al prestar')][3] == 'Si es estudiante'
    and dic[('Lector', 'tipo')][3] == 'Sí' and 'RF-11' in dic[('Lector', 'tipo')][5])
inf = X.INFORME[1]
chk('N', 'Tabla 7.8: el total se recuenta y coincide con la suma de las filas, con la fila «Docentes» y sin «Sin curso»',
    int(inf[-1][2]) == sum(int(f[2]) for f in inf[:-1]) and any(f[0] == 'Docentes' for f in inf) and not any('Sin curso' in f[0] for f in inf))
chk('N', 'Tabla 7.8 no se presenta como el mes de marzo de D2 (cuyas cifras no se reconstruyen)', 'abril' in X.INFORME_ROTULO and 'marzo' not in X.INFORME_ROTULO)
chk('N', 'Vencimiento calculado al prestar y guardado (RI-07); plazo y máximos en Parámetro (7, 2 y 5)',
    'se guarda' in X.RESTRICCION['RI-07'][0] and [(p[0], p[1]) for p in X.PARAMETROS] == [('plazo_dias', '7'), ('maximo_estudiante', '2'), ('maximo_docente', '5')])
chk('N', 'Ejemplo resuelto 7.1 y Práctica 7 trabajan la distinción obra/ejemplar (E2-L2, M-F8)', 'E2-L2' in j(cap7.EJ71) and 'M-F8' in j(cap7.NIVELES))

# ---------------------------------------------------------------- D. catálogo v1.1 y coherencia
coh = {c[0]: c for c in X.COHERENCIA}
chk('D', 'Catálogo v1.1: 16 requisitos (los 15 de la v1.0 y RF-11 de SC-01)', len(X.CATALOGO_V11) == 16 and 'RF-11' in X.REQ and X.REQ['RF-11'] == R.RF11)
chk('D', 'RN-02 v1.1 (máximo de cinco para los docentes) reemplaza a la v1.0', X.REGLA['RN-02'] == R.REGLA_V11[1] and 'cinco' in X.REGLA['RN-02'])
chk('D', 'Matriz de coherencia: una fila por cada requisito del catálogo v1.1, sin repetir (RF-11 junto a los funcionales)',
    sorted(c[0] for c in X.COHERENCIA) == sorted(X.REQ) and len(X.COHERENCIA) == len(X.REQ)
    and [c[0] for c in X.COHERENCIA][:11] == ['RF-%02d' % i for i in range(1, 12)])
refs_coh = set(re.findall(r'(?:RI|PT|ME|OP|A|P)-\d\d|DT-01', j(X.COHERENCIA)))
chk('D', 'Matriz de coherencia: los %d códigos de datos, interfaz, operación y prueba existen' % len(refs_coh), all(fuente(c) for c in refs_coh), [c for c in refs_coh if not fuente(c)])
chk('D', 'Matriz de coherencia: la prueba de cada requisito es la de la matriz del Capítulo 6 (RF-11 usa P-06, desdoblada por SC-01)',
    all(coh[t[0]][4] == t[4] for t in R.TRAZABILIDAD) and coh['RF-11'][4] == 'P-06')
chk('D', 'Matriz de coherencia: cada «—» es de un requisito de calidad o una restricción del entorno',
    all(c[0].startswith(('RNF', 'RS')) for c in X.COHERENCIA if '—' in c[1:4]))
chk('D', 'RF-11 tiene respuesta en datos (Lector.tipo), interfaz (PT-08, ME-02) y operaciones (OP-01, OP-05)',
    'Lector.tipo' in coh['RF-11'][1] and 'PT-08' in coh['RF-11'][2] and 'OP-01' in coh['RF-11'][3] and 'RF-11' in X.OPERACION['OP-01'][6])
chk('D', 'Trazabilidad a SC-01: el Ejemplo resuelto 9.3 y el Capítulo 7 la citan', 'SC-01' in j(cap9.EJ93) and 'SC-01' in j(cap7.EJ72) and 'SC-01' in tE)
chk('D', 'Postergados fuera del núcleo: sin renovación ni reserva ni estado «en reparación» como decisiones del diseño',
    not re.search(r'(?i)renovaci|reserva del plan|en reparación', j([X.ENTIDADES, X.DICCIONARIO, X.OPERACIONES, X.PANTALLAS])))

# ---------------------------------------------------------------- M. casos de uso, mensajes y contratos
me_de = {m[3]: m for m in X.MENSAJES}
errores_op = ' '.join(e for o in X.OPERACIONES for e in o[5])
for cu in MO.PRIORITARIOS:
    p = MO.PLANTILLA[cu]
    for alt in p['alternativos']:
        clave = '%s, flujo %s' % (cu, alt[0])
        m = me_de.get(clave)
        chk('M', '%s tiene su estado y su mensaje en el Capítulo 8 (%s) y su error en un contrato del Capítulo 9' % (clave, m[0] if m else '—'),
            m and m[4] in tE and (m[0] in errores_op or m[2] == 'vacío'))
chk('M', 'Cada mensaje de error, salvo el de inicio de sesión (ME-12), es un error de algún contrato',
    all(m[0] in errores_op for m in X.MENSAJES if m[2] == 'error' and m[0] != X.AUTENTICACION))
chk('M', 'Cada error de contrato tiene su mensaje en la Tabla 8.6', all(re.match(r'ME-\d\d', e).group() in {m[0] for m in X.MENSAJES} for o in X.OPERACIONES for e in o[5]))
chk('M', 'Mensajes con la regla de 8.3 (qué pasó, por qué y qué hacer): los de error dicen qué hacer',
    all(re.search(r'(?i)(consultá|tiene que|elegí|revisá|apartalo)', m[4]) for m in X.MENSAJES if m[2] == 'error'))
chk('M', 'Recorridos: los pasos citados existen en las plantillas de CU-01 a CU-03 y las pantallas, en el mapa',
    all(n <= len(MO.PLANTILLA[cu]['escenario']) and pt in X.PANTALLA for cu, pasos in X.RECORRIDOS for n, pt, _, _ in pasos)
    and [cu for cu, _ in X.RECORRIDOS] == MO.PRIORITARIOS)
chk('M', 'Pantallas: los requisitos citados existen en el catálogo v1.1', all(r in X.REQ for p in X.PANTALLAS for r in p[4]))
chk('M', 'Consulta pública sin datos personales: PT-02, ME-09 y OP-03 no devuelven datos de lectores',
    'Ningún dato de lectores' in X.PANTALLA['PT-02'][3] and 'Nunca devuelve datos de lectores' in X.OPERACION['OP-03'][4] and 'RNF-03' in X.OPERACION['OP-03'][6])
chk('M', 'Práctica 8: el voluntario de 9.º es menor, supervisado y sin acceso supuesto a datos personales ni funciones administrativas',
    'menor de edad' in j(cap8.PRACTICA_INTRO) and 'no se puede suponer' in j(cap8.PRACTICA_INTRO) and 'MS-5' in j(cap8.PRACTICA_INTRO)
    and 'supervisión' in j(cap8.NIVELES) and 'no ve datos personales ni tiene funciones administrativas' in cap8.PRACTICA_CONTROL)

# ---------------------------------------------------------------- P. arquitectura y decisión técnica
dt = dict(X.DT01)
chk('P', 'Tres capas: presentación, reglas de negocio y persistencia', [c[0] for c in X.CAPAS] == ['Presentación', 'Reglas de negocio', 'Persistencia'])
chk('P', 'Contratos: OP-01 a OP-06 con entradas, precondiciones, resultado, errores y requisitos que existen',
    [o[0] for o in X.OPERACIONES] == ['OP-%02d' % i for i in range(1, 7)] and all(all(o[1:]) for o in X.OPERACIONES)
    and all(r in X.REQ for o in X.OPERACIONES for r in o[6]))
chk('P', 'Las reglas se comprueban en el servidor (Work §7.5) y RI-01 se protege también en la persistencia',
    'en el servidor' in j(cap9.S91) and 'protege' in X.CAPAS[2][1] and 'RI-01' in X.CAPAS[2][1])
chk('P', 'DT-01: contexto, condiciones excluyentes con RS-01 y RS-02, opciones, descarte, decisión, consecuencias y requisitos',
    list(dt) == ['Identificador', 'Contexto', 'Condiciones excluyentes', 'Opciones', 'Descarte', 'Decisión', 'Consecuencias', 'Requisitos que la condicionan']
    and 'RS-01' in dt['Condiciones excluyentes'] and 'RS-02' in dt['Condiciones excluyentes'])
chk('P', 'DT-01 sin puntajes ni pesos', not re.search(r'\d+\s*(puntos|%)|\bpeso|pondera', j(X.DT01)))
chk('P', 'HTML/CSS, PHP y SQLite como ejemplo provisional, no obligatorio, no aprobado para todos los colegios y sin probar en equipos',
    all(x in j(cap9.S93) for x in ('ejemplo provisional', 'no es obligatorio', 'no está aprobado para todos los colegios', 'no se probó'))
    and 'provisional' in dt['Decisión'] and 'no está verificada' in dt['Decisión'])
chk('P', 'Funcionamiento local y sin servicios pagos obligatorios', 'local' in dt['Opciones'] and 'servicios pagos' in dt['Condiciones excluyentes'])
chk('P', 'Sin la matriz de selección tecnológica de Algorítmica 3.º: la ARTICULACIÓN la remite y el capítulo no la repite',
    'matriz de decisión' in j(cap9.ARTICULACION) and 'No se ponen puntajes' in j(cap9.S93) and not re.search(r'(?i)criterios? cuantificad', j(cap9.S93) + j(X.DT01)))
chk('P', 'Práctica 9: complemento editorial, en el Encuentro 17 o como actividad complementaria, antes del Encuentro 18 y sin desplazar el H3',
    'complemento editorial' in cap9.PRACTICA_TITULO and 'Encuentro 17' in j(cap9.PRACTICA_INTRO) and 'no reemplaza ni desplaza el h3' in j(cap9.PRACTICA_INTRO).lower()
    and heads.index(cap9.PRACTICA_TITULO) < heads.index(cap9.E18_TITULO)
    and 'fuera de las 5 HC' in GU.SEC_C9[0][1][-1][0])
chk('P', 'Práctica 9 con decisiones nuevas: contrato de devolución, libro dañado, reposición y mantenimiento del entorno',
    all(x in j(cap9.NIVELES) for x in ('registrar devolución', 'dañado', 'reposición', 'mantener')) and 'M-F9' in j(cap9.NIVELES))

# ---------------------------------------------------------------- E. Algorítmica 3.º
chk('E', 'Sin personas, títulos, archivos ni tecnologías del caso de Algorítmica 3.º', not [p for p in D.PROHIBIDOS_ALGORITMICA if p in todo])
chk('E', 'Sin matriz ponderada ni pesos (ni «ponderad…», ni pesos 0,5 / 0,3 / 0,2)', not re.search(r'(?i)ponderad|\bpesos?\b', todo) and not re.search(r'0,5\D{0,30}0,3\D{0,30}0,2', todo))
chk('E', 'Sin las tablas de Algorítmica («Socios», Access): solo se nombran en la guía para decir que no se usan',
    'Socios' not in tE and 'Socios' not in tC and 'Access' not in tE and all('no tiene tablas' in s for s in re.findall(r'[^.]*Socios[^.]*\.', tG)))
chk('E', 'Sin sección ni ejercicios de formas normales (1FN a 3FN): solo revisión aplicada en 7.4',
    not re.search(r'\b[123]FN\b|(?i:primera forma normal|dependencia funcional)', todo) and not [h for h in heads if re.search(r'(?i)forma normal', h)])
casc = [s for s in re.findall(r'[^.]*cascada[^.]*\.[^.]*\.', tE + tG)]
chk('E', 'La «eliminación en cascada» aparece solo como contraste con RN-05 y RN-06', casc and all(re.search(r'violaría|Contradice|RN-05', s) for s in casc), casc)
chk('E', 'Notación entidad-relación clásica de Algorítmica, rotulada como representación simplificada', 'representación simplificada' in tE and 'notación de Algorítmica 3.º' in tE)
chk('E', 'Guía: acuerdo previo con Algorítmica antes del Encuentro 13, con las dos rutas de Work §3', 'Antes del Encuentro 13' in tG and 'Ruta A' in tG and 'Ruta B' in tG)

# ---------------------------------------------------------------- F. voseo y residuos
TUTEO = (r'\b(tienes|puedes|debes|sabes|quieres|eres|propón|tú)\b|(?:^\d+\.\s+|^\([a-h]\)\s+)'
         r'(Escribe|Lee|Completa|Indica|Explica|Clasifica|Responde|Anota|Marca|Elige|Justifica|Redacta|Compara|Describe|Identifica|Resume|Haz|Dibuja|Prioriza|Planifica|Señala|Decide|Diseña|Aplica|Revisa|Agrega)\b')
for nombre, tx in (('libro', tE), ('guía', tG), ('cuaderno', tC)):
    m = [x.group(0) for x in re.finditer(TUTEO, tx, re.M)]
    chk('F', '%s: sin formas de tuteo frecuentes' % nombre, not m, m[:5])
    chk('F', '%s: sin identificadores de modelos de IA' % nombre, not re.search(r'(?i)\b(claude|gpt|gemini|opus|sonnet|codex|chatgpt)\b', tx))
    chk('F', '%s: sin símbolos ✓ ✗' % nombre, not re.search('[✓✔✗✘✅]', tx))
    chk('F', '%s: sin marcadores sin reemplazar' % nombre, '§§' not in tx and 'TODO' not in tx and 'XXX' not in tx)
    chk('F', '%s: no se declara apto ni aprobado para la venta' % nombre,
        not re.search(r'(?i)\bapto\b|aprobado para (la )?(venta|comercializ)', tx.replace('No está aprobado para la venta', '')))

# ---------------------------------------------------------------- G. maqueta (DOCX)
tw = int(O.Cm(O.ANCHO).twips)
for nombre, d in (('libro', dE), ('guía', dG), ('cuaderno', dC)):
    s = d.sections[0]
    chk('G', '%s: oficio 21,59 × 33,02 cm, márgenes O3 y espejados' % nombre,
        abs(s.page_width.cm - 21.59) < 0.01 and abs(s.page_height.cm - 33.02) < 0.01 and abs(s.left_margin.cm - 2.5) < 0.01
        and abs(s.right_margin.cm - 1.8) < 0.01 and abs(s.top_margin.cm - 1.5) < 0.01 and d.settings.element.find(qn('w:mirrorMargins')) is not None)
    filas = [r for t in d.tables for r in t.rows]
    chk('G', '%s: %d filas de tabla, todas «no dividir»' % (nombre, len(filas)), all(r._tr.trPr is not None and r._tr.trPr.find(qn('w:cantSplit')) is not None for r in filas))
    multi = [t for t in d.tables if len(t.columns) > 1]
    malas = [i for i, t in enumerate(multi) if abs(sum(int(g.get(qn('w:w'))) for g in t._tbl.tblGrid) - tw) > 15]
    chk('G', '%s: %d tablas de varias columnas al ancho útil (%s cm)' % (nombre, len(multi), str(O.ANCHO).replace('.', ',')), not malas, malas)
    sin_enc = [t.rows[0].cells[0].text[:30] for t in multi if t.rows[0]._tr.trPr.find(qn('w:tblHeader')) is None]
    chk('G', '%s: toda tabla de datos repite su encabezado (las únicas sin encabezado son las fichas)' % nombre, all(x.startswith('Capacidad') for x in sin_enc), sin_enc)
chk('G', 'Se usan la maqueta del lote 1 y el perfil O3 del lote 2 (sin copias)',
    os.path.dirname(O.H.__file__).endswith(os.path.join('lote_1', 'build')) and os.path.dirname(O.__file__).endswith(os.path.join('lote_2', 'build')))

# ---------------------------------------------------------------- H. maqueta (PDF)
def lineas(pg):
    return [l.strip() for l in pg.get_text().split('\n') if l.strip() and not re.match(r'^Página \d+ de \d+$', l.strip())]
for nombre, pdf in (('libro', pE), ('guía', pG), ('cuaderno', pC)):
    vacias = [i + 1 for i, pg in enumerate(pdf) if not lineas(pg)]
    chk('H', '%s: ninguna página en blanco' % nombre, not vacias, vacias)
    sueltos = [i + 1 for i, pg in enumerate(pdf) if lineas(pg) and lineas(pg)[0].startswith('_' * 20)]
    chk('H', '%s: ninguna página empieza con renglones de respuesta sueltos' % nombre, not sueltos, sueltos)
for C in CAPS:
    pags = [i + 1 for i, pg in enumerate(pE) if C.EVAL_TITULO in lineas(pg)[:2]]
    sig = pags and pags[0] < len(pE) and lineas(pE[pags[0]])
    entra = pags and (pags[0] == len(pE) or (sig and (sig[0].startswith('Capítulo') or sig[0].startswith('Unidad'))))
    chk('H', '%s: en página propia y entera en esa página (pág. %s)' % (C.EVAL_TITULO, pags), len(pags) == 1 and entra)
    pp = [i + 1 for i, pg in enumerate(pE) if lineas(pg)[:1] and lineas(pg)[0].startswith(C.PRACTICA_TITULO[:25])]
    chk('H', '%s: empieza en página nueva' % C.PRACTICA_TITULO.split(' —')[0], len(pp) == 1)
    pc = [i + 1 for i, pg in enumerate(pC) if lineas(pg)[:1] and lineas(pg)[0].startswith(C.PRACTICA_TITULO[:25])]
    chk('H', 'Cuaderno: %s empieza en página nueva' % C.PRACTICA_TITULO.split(' —')[0], len(pc) == 1)
pH3 = [i + 1 for i, pg in enumerate(pE) if lineas(pg)[:1] == [cap9.E18_TITULO]]
chk('H', 'El Encuentro 18 (H3) empieza en página nueva (pág. %s)' % pH3, len(pH3) == 1)
paginas = res_build['paginas_libro']
chk('H', 'El índice ocupa una sola página y la siguiente es la apertura de la Unidad III', lineas(pE[2])[0] == B.UNIDAD_III)
recalc, _ = K.paginas_de_titulos(EST + '.pdf', list(paginas), desde=2)
chk('H', 'Índice estable: las %d entradas apuntan a la página donde está su título' % len(paginas), recalc == paginas and len(paginas) == len(B.entradas_indice()))
# Extensión de referencia del plan (§9): libro entre 38 y 46 págs., guía entre 16 y 22; el cuaderno, entre 6 y 12.
# El libro quedó por debajo del mínimo previsto: el control admite desde 32 y lo informa.
chk('H', 'Extensión del libro: %d págs. (plan: 38 a 46; se admite desde 32)' % len(pE), 32 <= len(pE) <= 46)
chk('H', 'Extensión de la guía: %d págs. (plan: 16 a 22)' % len(pG), 16 <= len(pG) <= 22)
chk('H', 'Extensión del cuaderno: %d págs. (referencia: 6 a 12)' % len(pC), 6 <= len(pC) <= 12)

# ---------------------------------------------------------------- I. evaluaciones, prácticas e instrumentos
claves = res_build['claves']
semillas_prev = ['ADS3-L1-C1', 'ADS3-L1-C2-i', 'ADS3-L1-C3', 'ADS3-C4v2-2', 'ADS3-L2-C5-c', 'ADS3-L2-C6']
for C in CAPS:
    ks = [c for c in claves if c[0] == C.EVAL_ID]
    mc = [c[2] for c in ks if c[2] in 'abcd' and len(c[2]) == 1]
    chk('I', '%s: %d ítems con clave; opción múltiple con claves %s' % (C.EVAL_TITULO, len(ks), ''.join(mc)), len(ks) == len(C.EVAL_ITEMS) and len(set(mc)) >= 3)
    pts = sum(1 for it in C.EVAL_ITEMS if it[0] in ('mc', 'vf')) + sum(it[3] for it in C.EVAL_ITEMS if it[0] == 'ab')
    chk('I', '%s: suma 15 puntos, como dice su remisión' % C.EVAL_TITULO, pts == 15 and 'Puntaje total: 15 puntos' in C.EVAL_PUNTAJE)
    chk('I', '%s: la guía trae su clave y los criterios de los ítems abiertos' % C.EVAL_TITULO, ('Clave de la %s' % C.EVAL_TITULO.lower()) in tG)
    chk('I', 'Capítulo %d: práctica en tres niveles con el caso Mbokaja' % C.NUM,
        [n.split(' —')[0] for n, _ in C.NIVELES] == ['Nivel 1', 'Nivel 2', 'Nivel 3'] and 'Mbokaja' in j(C.PRACTICA_INTRO))
semillas = [C.EVAL_ID for C in CAPS]
chk('I', 'Semillas de evaluación distintas entre sí y de las de los lotes 1 y 2 (%s)' % ', '.join(semillas), len(set(semillas + semillas_prev)) == 9)
ejemplos = j([getattr(C, n) for C in CAPS for n in dir(C) if n.startswith(('EJ', 'S7', 'S8', 'S9', 'TABLA'))])
chk('I', 'Los datos de las prácticas (M-F8, M-F9, MS-3 a MS-5, OBS) no aparecen en los ejemplos resueltos ni en las secciones',
    not re.search(r'\b(?:M-F[89]|MS-[345]|OBS-\d)\b', ejemplos))
chk('I', 'Rúbrica del H3: seis criterios que cubren C1 a C7, cuatro niveles, evidencia del equipo e individual',
    len(cap9.RUBRICA_H3) == 6 and all(len(r[2]) == 4 and r[3] and r[4] for r in cap9.RUBRICA_H3)
    and sorted(set(re.findall(r'C\d', ' '.join(r[1] for r in cap9.RUBRICA_H3)))) == ['C%d' % i for i in range(1, 8)]
    and all(('%s (%s)' % (r[0], r[1])) in tG for r in cap9.RUBRICA_H3))
chk('I', 'H3: 60/40 como propuesta editorial, ausencia de evidencia aparte y tiempo para corregir el dossier',
    '60 %' in tE and '40 %' in tE and 'propuesta editorial' in tE and 'ausencia de evidencia' in tG.lower() and 'Si falta una evidencia' in tE
    and 'no se confunde con un desempeño de nivel inicial' in tE and 'corrige el dossier' in tE and '1 HC de corrección' in tE)
chk('I', 'H3: variante individual no resuelta antes; banco de cinco variantes que no están en el dossier, el buzón, D3 ni las prácticas',
    'variante que no se resolvió antes' in tE and len(X.BANCO_H3) >= 5
    and not any(v[1][:30] in t for v in X.BANCO_H3 for t in list(R.TEXTO.values()) + list(MB.FUENTES.values()))
    and not any(v[1][:30] in j([C.NIVELES, C.PRACTICA_INTRO]) for v in X.BANCO_H3 for C in CAPS))
chk('I', 'Guía: planilla de registro del H3 con columna AE', 'Planilla de registro del H3' in tG and 'AE' in tG)
chk('I', 'Guía: secuencia en HC de los Encuentros 13 a 18', all(('Encuentro %d' % n) in tG for n in range(13, 19)) and tG.count('aprox.') >= 20)
chk('I', 'Guía: HC por categoría de Work §5.2 (13 D5; 14 D4 V1; 15 D5; 16 D3 V2; 17 D5; 18 D1 E4)',
    all(x in tG for x in ('(5 HC de diseño)', '(5 HC: 4 de diseño y 1 de verificación)', '(5 HC: 3 de diseño y 2 de verificación)', '(5 HC: 4 de evaluación y 1 de diseño)'))
    and tG.count('(5 HC de diseño)') == 3)
chk('I', 'Guía: claves de todas las actividades de los Encuentros 13 a 18',
    all(len(k) == sum(len(items) for _, items in acts) for k, acts in ((GU.CLAVES_E13, cap7.ACT_E13), (GU.CLAVES_E14, cap7.ACT_E14),
        (GU.CLAVES_E15, cap8.ACT_E15), (GU.CLAVES_E16, cap8.ACT_E16), (GU.CLAVES_E17, cap9.ACT_E17), (GU.CLAVES_E18, cap9.ACT_E18))))
chk('I', 'Guía: claves por nivel y lista de cotejo de las tres prácticas', all(('Lista de cotejo de la Práctica %d' % n) in tG for n in (7, 8, 9))
    and len(GU.CLAVES_P7) == len(GU.CLAVES_P8) == len(GU.CLAVES_P9) == 3)
chk('I', 'Guía: tarjetas de rol del Encuentro 16 sin datos de personas reales', 'Tarjetas de rol del Encuentro 16' in tG and len(GU.TARJETAS) >= 4
    and not [p for p in D.PROHIBIDOS_ALGORITMICA if p in j(GU.TARJETAS)])
chk('I', 'Calendario de 28 encuentros: recuadro en cada capítulo y tabla en la guía (se unen el 15 y el 16)',
    tE.count('Si el calendario se acorta') == 3 and 'Unidad III en 36 y en 28 encuentros' in tG and 'se unen en uno' in j(cap8.CALENDARIO))
chk('I', 'Cuaderno: las tres prácticas con el mismo texto que el libro y sin evaluaciones',
    all(t in tC for C in CAPS for t in [C.PRACTICA_TITULO, C.PRACTICA_CONTROL] + C.PRACTICA_INTRO + [i.split('\n')[0] for _, items in C.NIVELES for i in items])
    and not any(C.EVAL_TITULO in tC for C in CAPS))

# ---------------------------------------------------------------- J. numeración
for C in CAPS:
    n = C.NUM
    for tipo in ('Tabla', 'Figura', 'Ejemplo resuelto'):
        sec = [int(m.group(1)) for m in re.finditer(r'^%s %d\.(\d+)(?: —|:)' % (tipo, n), tE, re.M)]
        chk('J', 'Capítulo %d: %s %d.1 a %d.%d correlativas, en orden de aparición y sin repetir' % (n, tipo.lower(), n, n, len(sec)), sec == list(range(1, len(sec) + 1)) and len(sec) >= 1, sec)
    for tipo, pat in (('tabla', r'^Tabla %d\.(\d+) —'), ('figura', r'^Figura %d\.(\d+) —'), ('ejemplo resuelto', r'^Ejemplo resuelto %d\.(\d+):')):
        refs = set(int(x) for x in re.findall(r'%s %d\.(\d+)' % (tipo.capitalize(), n), tE + tG + tC))
        defs = set(int(x) for x in re.findall(pat % n, tE, re.M))
        chk('J', 'Capítulo %d: toda %s mencionada existe' % (n, tipo), refs <= defs, refs - defs)
_, t2 = texto_docx(os.path.join(ADS, 'lote_2', 'ADS_3BTI_LOTE2_ESTUDIANTE_OFICIO_O3.docx'))
_, t1 = texto_docx(os.path.join(ADS, 'lote_1_oficio_candidata', 'ADS_3BTI_LOTE1_ESTUDIANTE_OFICIO_O3_CANDIDATA.docx'))
otras = sorted(set(re.findall(r'\b(?:Tabla|Figura) [1-6]\.\d+', tE + tG + tC)))
chk('J', 'Las tablas y figuras de los lotes 1 y 2 que se citan existen allí (%s)' % (', '.join(otras) or 'ninguna'),
    all(re.search(r'^%s —' % re.escape(o), t1 + '\n' + t2, re.M) for o in otras))

# ---------------------------------------------------------------- L. notaciones
reg = json.load(open(os.path.join(AQUI, 'figs', 'simbolos.json'), encoding='utf8'))
figs_n = {'fig_7_1': 'der', 'fig_8_1': 'navegacion', 'fig_8_2': 'boceto', 'fig_8_3': 'boceto', 'fig_9_1': 'capas'}
for f, n in figs_n.items():
    usados = reg[f]['simbolos']
    chk('L', 'Sin mezcla: %s declara «%s» y usa solo sus símbolos (%s)' % (f, G.NOTACIONES[n], ', '.join(sorted(usados))),
        reg[f]['notacion'] == n and usados and all(G.SIMBOLOS[s][0] == n for s in usados))
for f in ('fig_7_2', 'fig_7_3'):
    chk('L', '%s (esquema didáctico) no usa símbolos de ninguna notación' % f, reg[f]['notacion'] == 'esquema' and not reg[f]['simbolos'])
s71 = reg['fig_7_1']['simbolos']
chk('L', 'Figura 7.1: siete entidades, cuatro relaciones y dos cardinalidades por relación, sin atributos dibujados',
    s71.get('der_entidad') == len(X.ENTIDADES) and s71.get('der_relacion') == len(X.RELACIONES) and s71.get('der_cardinalidad') == 2 * len(X.RELACIONES))
s81 = reg['fig_8_1']['simbolos']
chk('L', 'Figura 8.1: un inicio, nueve pantallas, tres zonas por rol y una transición por cada una de la fuente (más la de inicio)',
    s81.get('nav_inicio') == 1 and s81.get('nav_pantalla') == len(X.PANTALLAS) and s81.get('nav_zona') == len(X.ZONAS)
    and s81.get('nav_transicion') == len(X.TRANSICIONES) + 1)
chk('L', 'Mapa: transiciones rotuladas con la acción, con verbo (no sustantivos como en el contexto)',
    all(re.match(r'^\w+(ar|er|ir)\b', t[2]) for t in X.TRANSICIONES), [t[2] for t in X.TRANSICIONES if not re.match(r'^\w+(ar|er|ir)\b', t[2])])
zona_de = {p[0]: p[2] for p in X.PANTALLAS}
chk('L', 'Mapa: la zona pública no lleva a pantallas de la bibliotecaria sin pasar por el ingreso (PT-03)',
    all(not (zona_de[a] == 'Pública' and zona_de[b] != 'Pública') or a == 'PT-03' for a, b, _ in X.TRANSICIONES)
    and X.INICIO == 'PT-01' and zona_de['PT-01'] == 'Pública')
chk('L', 'Tabla 8.2: el mapa se distingue expresamente de los diagramas de actividades, contexto y casos de uso',
    cap8.TABLA_DIFERENCIAS_ROTULO in tE and len(cap8.TABLA_DIFERENCIAS[1]) == 4)
usados_todos = {s for f in figs_n for s in reg[f]['simbolos']}
chk('L', 'Cada símbolo usado en una figura figura en la leyenda del libro y de la guía (%d símbolos)' % len(usados_todos),
    all(G.SIMBOLOS[s][1] in tE and G.SIMBOLOS[s][3] in tG for s in usados_todos))
chk('L', 'Leyendas antes de la primera figura de cada notación',
    tE.index(cap7.LEYENDA_DER) < tE.index(cap7.EPI_FIG71) and tE.index(cap8.LEYENDA_NAV) < tE.index(cap8.EPI_FIG81)
    and tE.index(cap8.LEYENDA_BOC) < tE.index(cap8.EPI_FIG82) and tE.index(cap9.LEYENDA_CAPAS) < tE.index(cap9.EPI_FIG91))
chk('L', 'Guía: errores frecuentes de cada notación', all(e[0] in tG for e in GU.ERRORES_DER + GU.ERRORES_NAV + GU.ERRORES_CAPAS))
def _mezcla_detenida():
    L = G.Lienzo('prueba_mezcla', 'navegacion')
    try:
        L.usar('der_entidad')
    except AssertionError:
        return True
    finally:
        G.REGISTRO.pop('prueba_mezcla', None)
    return False
chk('L', 'Prueba del control: una figura que mezcla notaciones detiene la generación', _mezcla_detenida())

# ---------------------------------------------------------------- R. zonas, roles y mensajes de seguridad (v1.1)
zona_de = {p[0]: p[2] for p in X.PANTALLAS}
ROL_TXT = {'bibliotecaria': 'bibliotecaria', 'dirección': 'dirección'}
chk('R', 'Cada zona del mapa tiene sus roles definidos y cada pantalla está en una zona que existe',
    [z for z, _ in X.ZONAS] == list(X.ROLES_ZONA) and all(z in X.ROLES_ZONA for z in zona_de.values()))
chk('R', 'Cada operación tiene sus roles definidos y cada pantalla con operación la nombra',
    set(X.ROLES_OP) == set(X.OPERACION) and set(X.PANTALLA_OP.values()) == set(X.OPERACION) and all(p in X.PANTALLA for p in X.PANTALLA_OP))
contra = [(p, op, X.ROLES_ZONA[zona_de[p]], X.ROLES_OP[op]) for p, op in X.PANTALLA_OP.items() if sorted(X.ROLES_ZONA[zona_de[p]]) != sorted(X.ROLES_OP[op])]
chk('R', 'Sin contradicciones entre zonas y contratos: los roles de la zona de cada pantalla son los que admite su operación', not contra, contra)
mal_pre = [o for o, roles in X.ROLES_OP.items() if (not roles and 'pública' not in X.OPERACION[o][3]) or
           any(ROL_TXT[r] not in X.OPERACION[o][3] for r in roles) or
           any(ROL_TXT[r] in X.OPERACION[o][3] for r in ROL_TXT if r not in roles)]
chk('R', 'La precondición de cada contrato nombra exactamente los roles permitidos (o dice que es pública)', not mal_pre, mal_pre)
ingreso = [(b_, a_) for a_, b_, acc in X.TRANSICIONES if a_ == 'PT-03']
mal_ing = [(b_, acc) for a_, b_, acc in X.TRANSICIONES if a_ == 'PT-03' and not any(r in acc for r in X.ROLES_ZONA[zona_de[b_]])]
chk('R', 'Cada ingreso desde PT-03 lleva a una zona que admite el rol con que se ingresa', ingreso and not mal_ing, mal_ing)
alcanza = collections.defaultdict(set)
for a_, b_, acc in X.TRANSICIONES:
    for r in ROL_TXT:
        if r in X.ROLES_ZONA[zona_de[b_]] and (not X.ROLES_ZONA[zona_de[a_]] or r in X.ROLES_ZONA[zona_de[a_]]):
            alcanza[b_].add(r)
chk('R', 'Cada rol de una zona llega a sus pantallas por el mapa (PT-09: bibliotecaria por «Ver informe» y dirección al ingresar)',
    all(set(X.ROLES_ZONA[zona_de[p]]) <= alcanza[p] for p in X.PANTALLA if X.ROLES_ZONA[zona_de[p]]), dict(alcanza))
chk('R', 'PT-09 y OP-05: los consultan la bibliotecaria y la dirección, con el rol comprobado por el servidor',
    sorted(X.ROLES_ZONA[zona_de['PT-09']]) == ['bibliotecaria', 'dirección'] and 'servidor' in X.OPERACION['OP-05'][3]
    and 'Dirección' not in [z for z, _ in X.ZONAS] and zona_de['PT-09'] in tE)
me = X.MENSAJE
chk('R', 'ME-12 es solo de autenticación: se muestra en PT-03 por usuario o contraseña incorrectos',
    X.AUTENTICACION == 'ME-12' and me['ME-12'][1] == 'PT-03' and 'contraseña' in me['ME-12'][3] and 'permiso' not in me['ME-12'][4].lower())
chk('R', 'ME-13 es el de autorización: acceso denegado, distinto de ME-12 y fuera de PT-03',
    X.AUTORIZACION == 'ME-13' and 'permiso' in me['ME-13'][3] and 'PT-03' not in me['ME-13'][1] and me['ME-13'][4] != me['ME-12'][4])
reuso = [o[0] for o in X.OPERACIONES if any(e.startswith(X.AUTENTICACION) for e in o[5])]
chk('R', 'Ningún contrato usa el mensaje de autenticación (ME-12) para un error de autorización', not reuso, reuso)
falta13 = [o for o, roles in X.ROLES_OP.items() if roles and not any(e.startswith(X.AUTORIZACION) for e in X.OPERACION[o][5])]
sobra13 = [o for o, roles in X.ROLES_OP.items() if not roles and any(e.startswith(X.AUTORIZACION) for e in X.OPERACION[o][5])]
chk('R', 'Toda operación con roles tiene ME-13 entre sus errores, y la pública no', not falta13 and not sobra13, (falta13, sobra13))
mezcla = [m.group(0) for tx in (tE, tG, tC, j(X.OPERACIONES), j(X.COHERENCIA)) for m in re.finditer(r'ME-12[^.;\n]{0,40}(?:sin permiso|permiso\b(?! .{0,10}inicio))', tx)
          if 'no ME-12' not in tx[max(0, m.start() - 4):m.end()]]
chk('R', 'Ningún texto del libro, la guía, el cuaderno ni los contratos asocia ME-12 con falta de permiso', not mezcla, mezcla)
def _contradicciones(zonas, roles_op):
    return [p for p, op in X.PANTALLA_OP.items() if sorted(X.ROLES_ZONA[zonas[p]]) != sorted(roles_op[op])]
def _reuso(operaciones):
    return [o[0] for o in operaciones if any(e.startswith(X.AUTENTICACION) for e in o[5])]
prueba_zona = dict(zona_de, **{'PT-09': 'Bibliotecaria'})                       # la v1.0 corregida al revés
prueba_rol = dict(X.ROLES_OP, **{'OP-05': ['dirección']})
prueba_op = [o if o[0] != 'OP-04' else o[:5] + (['ME-12 (sin permiso)'],) + o[6:] for o in X.OPERACIONES]
chk('R', 'Prueba de los controles: detectan una zona exclusiva, un contrato con otros roles y ME-12 usado como «sin permiso»',
    _contradicciones(prueba_zona, X.ROLES_OP) == ['PT-09'] and _contradicciones(zona_de, prueba_rol) == ['PT-09']
    and _reuso(prueba_op) == ['OP-04'] and not _contradicciones(zona_de, X.ROLES_OP) and not _reuso(X.OPERACIONES))
chk('R', 'La matriz de coherencia, la Tabla 8.6 y el libro incluyen ME-13', 'ME-13' in X.COHERENCIA[[c[0] for c in X.COHERENCIA].index('RNF-02')][2]
    and me['ME-13'][4] in tE and 'ME-13' in tG)
chk('R', 'Leyenda del mapa: la zona agrupa las pantallas de un rol o de un grupo de roles (libro y guía)',
    'grupo de roles' in G.SIMBOLOS['nav_zona'][2] and G.SIMBOLOS['nav_zona'][2] in tE and G.SIMBOLOS['nav_zona'][3] in tG)

# ---------------------------------------------------------------- Q. lotes 1 y 2 intactos
st = subprocess.run(['git', 'status', '--porcelain', '--', 'lote_1', 'lote_2', 'lote_1_oficio_candidata'], cwd=ADS, capture_output=True, text=True).stdout.strip()
chk('Q', 'git: ningún archivo de lote_1, lote_1_oficio_candidata ni lote_2 cambió ni se agregó', not st, st[:300])
import hashlib
m10 = json.load(open(os.path.join(L3, 'version_1_0', 'MANIFIESTO_LOTE_3.json'), encoding='utf8'))
docs10 = [a for a in m10['archivos'] if a['ruta'].startswith('ADS_3BTI_LOTE3_')]
chk('Q', 'Versión 1.0 conservada: los seis DOCX y PDF de version_1_0/ coinciden con su manifiesto SHA-256',
    len(docs10) == 6 and all(hashlib.sha256(open(os.path.join(L3, 'version_1_0', a['ruta']), 'rb').read()).hexdigest() == a['sha256'] for a in docs10))

print('\nResultado por bloque:')
for b, (o, f) in RES.items():
    print('  %s: %d correctos, %d fallos' % (b, o, f))
print('Resultado: %d correctos, %d fallos' % (ok, fallos))
json.dump({'correctos': ok, 'fallos': fallos, 'bloques': RES, 'controles': DETALLE},
          open(os.path.join(L3, 'verificacion_resultado.json'), 'w', encoding='utf8'), ensure_ascii=False, indent=1)
sys.exit(1 if fallos else 0)
