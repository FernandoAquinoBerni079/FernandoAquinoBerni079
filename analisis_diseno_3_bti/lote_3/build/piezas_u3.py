# -*- coding: utf-8 -*-
"""Piezas de armado del lote 3 que no están en la maqueta ni en las piezas del lote 2: portada del lote, tablas de
leyenda con los símbolos de diagramas_u3 y recuadros de contrato de operación.
Usa sin modificarlas la maqueta O3 (lote_2/build/oficio.py) y las piezas del lote 2 (lote_2/build/piezas.py)."""
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

import oficio as O
H, M, K = O.H, O.M, O.K
import diseno as X     # agrega lote_2/caso al camino de importación
import piezas as Z2
import mbokaja_u3 as MB
import diagramas_u3 as G
R, D, MO = X.R, X.D, X.MO

VERSION = 'v1.0 (lote 3, para auditoría) · oficio O3'
LINEA = 'LOTE EDITORIAL 3 — UNIDAD III (CAPÍTULOS 7 A 9)'


def portada(doc, tipo_doc, estado):
    Z2.VERSION = VERSION
    Z2.portada(doc, tipo_doc, estado, linea=LINEA)


codigos = Z2.codigos


def tabla_leyenda(doc, notacion, rotulo, figs):
    """Tabla de símbolos de una notación: la primera columna lleva el símbolo dibujado con la misma función que las
    figuras (diagramas_u3.simbolo)."""
    sims = [s for s, v in G.SIMBOLOS.items() if v[0] == notacion]
    filas = [['', G.SIMBOLOS[s][1], G.SIMBOLOS[s][2], G.SIMBOLOS[s][3]] for s in sims]
    t = M.tabla(doc, (['Símbolo', 'Nombre', 'Significado', 'Regla de uso'], filas), [2.4, 3.0, 5.4, 5.8], rotulo=rotulo)
    for fila, s in zip(t.rows[1:], sims):
        p = fila.cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(figs['sim:' + s], width=Cm(1.75))
    return t


def contrato(doc, cod):
    """Recuadro con las cinco partes del contrato de una operación (caso/diseno.py, OPERACIONES)."""
    c, nombre, ent, pre, res, err, req = X.OPERACION[cod]
    H.caja(doc, 'Contrato %s — %s' % (c, nombre),
           ['Entradas: ' + ent, 'Precondiciones: ' + pre, 'Resultado: ' + res,
            'Errores:\n' + '\n'.join(err), 'Requisitos: ' + ', '.join(req)], tipo='neutra')


def encadenar_inicio(doc, filas=3):
    """Evita que una tabla larga quede cortada con el encabezado y una sola fila al pie de la página: los párrafos del
    encabezado y de las primeras filas se mantienen con el siguiente, así la tabla empieza con al menos `filas` filas
    o pasa entera a la página siguiente junto con su rótulo. Las tablas de hasta seis filas no se parten."""
    for t in doc.tables:
        if len(t.columns) < 2 or len(t.rows) < 3:
            continue
        n = len(t.rows) - 1 if len(t.rows) <= 6 else filas
        for r in t.rows[:n]:
            for c in r.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True
