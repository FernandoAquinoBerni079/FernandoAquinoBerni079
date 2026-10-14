# -*- coding: utf-8 -*-
"""Controles automáticos del Capítulo 4 v2. Uso: python3 verificar_v2.py [carpeta_pdf]"""
import os, re, sys, glob
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
import contenido as C
import guia_v2 as G
import helpers as H

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
ok = fallos = 0
def chk(desc, cond):
    global ok, fallos
    if cond: ok += 1
    else: fallos += 1
    print(('CORRECTO ' if cond else 'FALLA    ') + desc)

N_EST, N_GUI = 'ADS_3BTI_CAP4_v2_ESTUDIANTE.docx', 'ADS_3BTI_CAP4_v2_GUIA_DOCENTE.docx'
docs = {n: Document(os.path.join(RAIZ, n)) for n in (N_EST, N_GUI)}

def texto(doc):
    partes = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                partes.append(c.text)
    return '\n'.join(partes)

def todas_las_cadenas(mod):
    out = []
    def rec(x):
        if isinstance(x, str): out.append(x)
        elif isinstance(x, (list, tuple)):
            for y in x: rec(y)
        elif isinstance(x, dict):
            for k, v in x.items(): rec(k); rec(v)
    for k in dir(mod):
        if k.isupper(): rec(getattr(mod, k))
    return out

# 1. Alineación con el índice de Work (capítulo 4, secciones 4.1 a 4.4)
chk('Título del capítulo según Work: «Requisitos que se pueden comprobar»', C.TITULO == 'Requisitos que se pueden comprobar')
chk('Encuentros 7 y 8', C.ENCUENTROS == (7, 8))
sec = [C.S41[0], C.S42[0], C.S43[0], C.S44[0]]
work = ['Necesidad, requisito, regla de negocio y solución técnica', 'Requisitos funcionales, de calidad y restricciones',
        'Identificación, fuente, prioridad y criterios de aceptación', 'Ambigüedades frecuentes y cómo corregirlas']
chk('Secciones 4.1 a 4.4 con los títulos del índice de Work', all(s.startswith('4.%d ' % (i + 1)) and w in s for i, (s, w) in enumerate(zip(sec, work))))
chk('Capacidad C3 de Work, rotulada como no oficial del MEC', C.CAPACIDAD == 'Especifica y prioriza requisitos verificables.' and 'no es una capacidad oficial del MEC' in C.ROTULO_CAPACIDAD)

# 2. Trazabilidad de fuentes
cat_ids = [r[0] for r in C.CATALOGO]
chk('Catálogo: identificadores únicos (%d)' % len(cat_ids), len(cat_ids) == len(set(cat_ids)))
n_rf = sum(1 for i in cat_ids if i.startswith('RF-')); n_otros = sum(1 for i in cat_ids if i.startswith(('RNF-', 'RS-')))
chk('Catálogo: al menos 8 funcionales (%d) y 4 de calidad o restricciones (%d), como pide C3' % (n_rf, n_otros), n_rf >= 8 and n_otros >= 4)
citadas = [f for r in C.CATALOGO for f in r[3]] + [p[0] for p in C.POSTERGADOS]
chk('Catálogo: las %d fuentes citadas existen en el dossier o en las reglas' % len(citadas), all(f in C.FUENTES for f in citadas))
chk('Catálogo: ningún requisito sin fuente', all(r[3] for r in C.CATALOGO))
cad = todas_las_cadenas(C) + todas_las_cadenas(G)
patron = re.compile(r'\b(E\d-L\d+|RN-\d\d|B-0\d)\b')
mencionadas = set(m for s in cad for m in patron.findall(s))
previstas = {'RN-07'}   # regla nueva que el estudiante propone (práctica y evaluación)
faltan = sorted(m for m in mencionadas - previstas if m not in C.FUENTES)
chk('Todas las fuentes mencionadas en libro y guía existen (%d códigos; faltan: %s)' % (len(mencionadas), ', '.join(faltan) or 'ninguno'), not faltan)
# Citas textuales «…» (CÓDIGO): el texto citado está en ese fragmento
citas = re.findall(r'«([^»]+)»\s*\((E\d-L\d+|RN-\d\d)\)', '\n'.join(cad))
malas = [(c, t[:40]) for t, c in citas if t.rstrip('.') not in C.FUENTES[c]]
chk('Las %d citas textuales con código coinciden con su fragmento (%s)' % (len(citas), malas or 'sin diferencias'), not malas)
# Respaldo semántico revisado a mano: palabras clave que deben estar en la fuente de cada requisito
claves = {'RF-01': ('E1-L2', 'préstamos'), 'RF-03': ('E4-L2', 'devuelven'), 'RF-04': ('E4-L4', 'dos chicos'),
          'RF-06': ('E4-L6', 'no les presto'), 'RF-07': ('E3-L3', 'saber antes'), 'RF-08': ('E1-L4', 'quién lo llevó'),
          'RF-09': ('E5-L2', 'por curso'), 'RF-10': ('E4-L6', 'no los puedo borrar'), 'RNF-01': ('E1-L10', 'fácil'),
          'RNF-02': ('E4-L8', 'Solo yo'), 'RNF-03': ('E5-L6', 'a la vista'), 'RS-01': ('E4-L10', 'internet'), 'RS-02': ('E5-L4', 'presupuesto')}
fuentes_cat = {r[0]: r[3] for r in C.CATALOGO}
chk('Respaldo de %d requisitos: cada fuente contiene la expresión que justifica el requisito' % len(claves),
    all(f in fuentes_cat[i] and k in C.FUENTES[f] for i, (f, k) in claves.items()))
chk('El código de prueba P-07 (Work) no se confunde con el buzón (B-01 a B-08)', all(c.startswith('B-') for c, _, _ in C.BUZON))

# 3. Documentos
for nombre, doc in docs.items():
    tx = texto(doc)
    corto = 'ESTUDIANTE' if 'ESTUDIANTE' in nombre else 'GUÍA'
    chk('%s: sin símbolos prohibidos' % corto, not re.search('[✓✔✗✘✅]', tx))
    tuteo = re.findall(r'\b(tienes|puedes|debes|sabes|quieres|eres|haz|propón|tú|ti)\b', tx, re.I)
    chk('%s: sin formas de tuteo inequívocas (%s)' % (corto, ', '.join(sorted(set(tuteo))) or 'ninguna'), not tuteo)
    chk('%s: sin identificadores de modelos de IA' % corto, not re.search(r'(?i)\b(claude|gpt|gemini|opus|sonnet|codex)\b', tx))
    filas = [r for t in doc.tables for r in t.rows]
    chk('%s: %d filas de tabla, todas con «no dividir»' % (corto, len(filas)), all(r._tr.trPr is not None and r._tr.trPr.find(qn('w:cantSplit')) is not None for r in filas))
    multi = [t for t in doc.tables if len(t.columns) > 1]
    chk('%s: %d tablas de varias columnas con tblGrid y 16,6 cm' % (corto, len(multi)),
        all(t._tbl.tblPr.find(qn('w:tblLayout')) is not None and abs(sum(int(g.get(qn('w:w'))) for g in t._tbl.tblGrid) - 9411) < 15 for t in multi))

est, gui = texto(docs[N_EST]), texto(docs[N_GUI])
consignas = [it for _, its in C.ACT_E7 + C.ACT_E8 + [('n', i) for _, i in C.NIVELES] for it in its] + [it[1] for it in C.EVAL_ITEMS]
imp_tu = re.compile(r'^(Escribe|Lee|Clasifica|Marca|Completa|Elige|Redacta|Reescribe|Agrega|Propón|Justifica|Detecta|Diseña|Subraya|Explica|Arma|Encuentra|Anota|Nombra|Aplica|Asigna|Corrige|Señala|Decide)\b')
malas = [c[:40] for c in consignas if imp_tu.match(c)]
chk('Las %d consignas usan imperativo de voseo (%s)' % (len(consignas), '; '.join(malas) or 'sin tuteo'), not malas)
chk('Estudiante: sin claves ni criterios de corrección', 'Clave de la evaluación' not in est and 'Respuesta esperada' not in est)
letras = []
for it in C.EVAL_ITEMS:
    if it[0] == 'mc':
        baraja, letra = H.barajar_opciones(C.EVAL_ID, it[1], it[2])
        letras.append(letra)
        pos = est.find(it[1])
        chk('Ítem «%s…»: correcta en la letra %s del libro' % (it[1][:28], letra), '%s) %s' % (letra, it[2][0]) in est[pos:pos + 900])
chk('Claves de opción múltiple repartidas (%s)' % ', '.join(letras), len(set(letras)) == len(letras))
pts = sum(1 for it in C.EVAL_ITEMS if it[0] in ('mc', 'vf')) + sum(3 if i == 7 else 5 for i, it in enumerate(C.EVAL_ITEMS) if it[0] == 'ab')
chk('Puntaje total 15 = suma de los ítems (%d)' % pts, pts == 15 and '15 puntos' in est and '15 puntos' in gui)
chk('Guía: claves para las %d actividades' % (len(G.CLAVES_E7) + len(G.CLAVES_E8)),
    len(G.CLAVES_E7) == sum(len(i) for _, i in C.ACT_E7) and len(G.CLAVES_E8) == sum(len(i) for _, i in C.ACT_E8))
chk('RF-02: 2 de marzo + 7 días (RN-03) = 9 de marzo, como dice la Tabla 4.5', '9 de marzo' in C.TABLA_43[1][0][1])

for r in sorted(glob.glob(os.path.join(AQUI, 'figs', '*.png'))):
    im = Image.open(r).convert('L')
    px = list(im.get_flattened_data()) if hasattr(im, 'get_flattened_data') else list(im.getdata())
    tinta = sum(1 for v in px if v < 100) / len(px)
    chk('%s: imagen con contenido (tinta %.1f %%)' % (os.path.basename(r), 100 * tinta), tinta > 0.002)

alg = ['Lucía Ortiz', 'Marcos Rojas', 'Ana Paula Vera', 'Hijo de hombre', 'El trueno entre las hojas', 'La babosa', 'Yo el Supremo', 'Roa Bastos']
chk('Sin personas ni datos del caso de biblioteca de Algorítmica 3.º', not any(a in est + gui for a in alg))

carpeta = sys.argv[1] if len(sys.argv) > 1 else None
if carpeta:
    import pymupdf
    for n in (N_EST, N_GUI):
        d = pymupdf.open(os.path.join(carpeta, n.replace('.docx', '.pdf')))
        print('INFO     %s: %d páginas en la renderización de control' % (n, len(d)))
        if 'ESTUDIANTE' in n:
            pe = [i for i, p in enumerate(d) if p.get_text().strip().startswith('Evaluación del Capítulo 4')]
            chk('La evaluación empieza en página propia (pág. %s)' % (pe[0] + 1 if pe else '?'), bool(pe))
            sueltos = [i + 1 for i, p in enumerate(d) if p.get_text().strip().startswith('_' * 20)]
            chk('Ninguna página empieza con renglones sueltos', not sueltos)

print('\nResultado: %d correctos, %d fallos' % (ok, fallos))
sys.exit(1 if fallos else 0)
