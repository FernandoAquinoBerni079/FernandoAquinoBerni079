# -*- coding: utf-8 -*-
"""Construye el Capítulo 4 v2: libro del estudiante y guía docente (.docx).
Uso: python3 build_v2.py  →  ../ADS_3BTI_CAP4_v2_ESTUDIANTE.docx y ../ADS_3BTI_CAP4_v2_GUIA_DOCENTE.docx"""
import os
from docx.shared import Cm

import helpers as H
import maqueta as M
import contenido as C
import guia_v2
import diagramas

AQUI = os.path.dirname(os.path.abspath(__file__))
SALIDA = os.path.dirname(AQUI)

ESTADO = [
    'Piloto corregido (versión 2) para reauditoría. No está aprobado para la venta ni para uso en aula.',
    'Sigue el índice del informe de Work: Capítulo 4, secciones 4.1 a 4.4, encuentros 7 y 8.',
    'La capacidad de la ficha es una propuesta para el programa institucional (C3 del informe de Work). '
    'No es una capacidad oficial del MEC.',
]


def fuente_txt(cods):
    return ', '.join(cods)


def dossier(doc):
    for _, titulo, frs in C.DOSSIER:
        M.registro(doc, ('%s — %s' % (frs[0][0].split('-')[0], titulo), ['%s  %s' % (c, t) for c, t in frs]))
    M.registro(doc, (C.REGLAS_TITULO, ['%s  %s' % (c, t) for c, t in C.REGLAS] + [C.REGLAS_NOTA]))


def libro_estudiante(figs, ruta):
    doc = H.new_doc()
    M.portada(doc, C, 'Libro del estudiante — capítulo de muestra, versión 2', ESTADO)

    doc.add_heading('%s · %s' % (C.UNIDAD, C.CAPITULO), level=1)
    H.p_texto(doc, C.TITULO, bold=True, size=13, color=H.ID['MED'])
    M.ficha(doc, C)
    for t in C.APERTURA:
        H.p_texto(doc, t)
    H.caja(doc, 'Con qué se conecta este capítulo', C.ARTICULACION, tipo='neutra')

    H.salto_pagina(doc)
    doc.add_heading(C.DOSSIER_TITULO, level=2)
    H.p_texto(doc, 'Cada fragmento lleva un código: E1-L4 es la línea 4 de la entrevista E1. Las reglas RN '
                   'se acordaron en clase. Estos son los únicos datos que el capítulo usa como fuente.', size=10)
    dossier(doc)

    # Encuentro 7
    H.salto_pagina(doc)
    doc.add_heading(C.E7_TITULO, level=1)
    M.seccion(doc, C.S41)
    M.tabla(doc, C.TABLA_41, [2.8, 4.0, 6.0, 3.8], rotulo='Tabla 4.1 — Cuatro categorías que no hay que mezclar')
    H.caja(doc, *C.CONCEPTO_REQ, tipo='concepto')
    H.caja(doc, C.EJ1[0], C.EJ1[1], tipo='neutra')
    M.seccion(doc, C.S42)
    H.imagen(doc, figs[0], C.LEAD_FIG1, C.EPI_FIG1, ancho_cm=15.0)
    M.tabla(doc, C.TABLA_42, [3.0, 3.6, 4.4, 5.6], rotulo='Tabla 4.2 — Tipos de requisitos')
    H.caja(doc, C.EJ2[0], C.EJ2[1], tipo='neutra')
    filas = [[i, t, e, fuente_txt(f), p] for i, t, e, f, p in C.CATALOGO]
    M.tabla(doc, (['Id.', 'Tipo', 'Enunciado', 'Fuente', 'Prioridad'], filas), [1.5, 2.3, 7.8, 2.4, 2.6],
            rotulo='Tabla 4.3 — Catálogo inicial de requisitos del caso')
    filas = [[f, d, p, r] for f, d, p, r in C.POSTERGADOS]
    M.tabla(doc, (['Fuente', 'Pedido', 'Prioridad', 'Razón'], filas), [1.6, 4.4, 2.2, 8.4],
            rotulo='Tabla 4.4 — Pedidos registrados fuera del catálogo inicial')
    M.actividades(doc, 'Actividades del Encuentro 7', C.ACT_E7)

    # Encuentro 8
    H.salto_pagina(doc)
    doc.add_heading(C.E8_TITULO, level=1)
    M.seccion(doc, C.S43)
    H.imagen(doc, figs[1], C.LEAD_FIG2, C.EPI_FIG2, ancho_cm=15.5)
    H.caja(doc, *C.CONCEPTO_CA, tipo='concepto')
    H.caja(doc, C.EJ3[0], C.EJ3[1], tipo='neutra')
    M.tabla(doc, C.TABLA_43, [1.5, 10.2, 4.9], rotulo='Tabla 4.5 — Criterios de aceptación y prioridades de otros requisitos')
    M.seccion(doc, C.S44)
    M.tabla(doc, C.TABLA_44, [2.9, 3.9, 3.9, 5.9], rotulo='Tabla 4.6 — Ambigüedades frecuentes')
    H.caja(doc, C.EJ4[0], C.EJ4[1], tipo='neutra')
    H.caja(doc, *C.REVISION_PARES, tipo='neutra')
    M.actividades(doc, 'Actividades del Encuentro 8', C.ACT_E8)

    # Práctica
    H.salto_pagina(doc)
    doc.add_heading(C.PRACTICA_TITULO, level=1)
    for t in C.PRACTICA_INTRO:
        H.p_texto(doc, t)
    M.tabla(doc, (['Código', 'Quién lo pidió', 'Pedido'], [list(b) for b in C.BUZON]), [1.8, 4.0, 10.8],
            rotulo=C.BUZON_TITULO)
    for nivel, items in C.NIVELES:
        M.actividades(doc, nivel, [('Consignas', items)])
    H.caja(doc, 'Punto de control', [C.PRACTICA_CONTROL], tipo='neutra')
    H.caja(doc, *C.REPASO, tipo='neutra')
    H.caja(doc, *C.CALENDARIO, tipo='neutra')

    # Evaluación
    H.CLAVES.clear()
    H.evaluacion(doc, C.EVAL_ID, C.EVAL_TITULO, False, C.EVAL_TEXTO_BASE, C.EVAL_ITEMS,
                 remision='Evalúa los cuatro indicadores del capítulo. Puntaje total: 15 puntos (ítems 1 a 7: '
                          '1 punto cada uno; ítem 8: 3 puntos; ítem 9: 5 puntos).')
    M.pegar_renglones(doc)
    H.finalizar(doc, ruta)
    return list(H.CLAVES)


def main():
    figs = diagramas.generar()
    r1 = os.path.join(SALIDA, 'ADS_3BTI_CAP4_v2_ESTUDIANTE.docx')
    r2 = os.path.join(SALIDA, 'ADS_3BTI_CAP4_v2_GUIA_DOCENTE.docx')
    claves = libro_estudiante(figs, r1)
    guia_v2.construir(r2, claves, ESTADO)
    print(r1); print(r2)


if __name__ == '__main__':
    main()
