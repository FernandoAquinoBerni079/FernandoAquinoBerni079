# -*- coding: utf-8 -*-
"""Controles automáticos del piloto del Capítulo 4. Uso: python3 verificar.py [carpeta_pdf]"""
import os, re, sys, glob
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
import contenido as C
import diagramas as D
import helpers as H

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
ok = fallos = 0
def chk(desc, cond):
    global ok, fallos
    if cond: ok += 1
    else: fallos += 1
    print(('CORRECTO ' if cond else 'FALLA    ') + desc)

docs = {n: Document(os.path.join(RAIZ, n)) for n in ('ADS_3BTI_PILOTO_CAP4_ESTUDIANTE.docx', 'ADS_3BTI_PILOTO_CAP4_GUIA_DOCENTE.docx')}

def texto(doc):
    partes = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                partes.append(c.text)
    return '\n'.join(partes)

for nombre, doc in docs.items():
    tx = texto(doc)
    corto = nombre.split('_')[-1].replace('.docx', '')
    chk('%s: sin símbolos prohibidos (marcas de verificación)' % corto, not re.search('[✓✔✗✘✅]', tx))
    # Formas de tú que nunca coinciden con la 3.ª persona (los indicadores usan «Detecta», «Redacta»…).
    tuteo = re.findall(r'\b(tienes|puedes|debes|sabes|quieres|eres|haz|propón|tú|ti)\b', tx, re.I)
    chk('%s: sin formas de tuteo inequívocas (%s)' % (corto, ', '.join(sorted(set(tuteo))) or 'ninguna'), not tuteo)
    chk('%s: no presenta la capacidad como oficial del MEC' % corto,
        'no es una capacidad oficial del MEC' in tx and not re.search(r'capacidad oficial del MEC(?! de)', tx.replace('no es una capacidad oficial del MEC', '').replace('No se cita como capacidad oficial del MEC', '').replace('No es una capacidad oficial del MEC', '')))
    chk('%s: sin identificadores de modelos de IA' % corto, not re.search(r'(?i)claude|gpt|gemini|opus|sonnet', tx))
    # Filas no divisibles en todas las tablas
    filas = [r for t in doc.tables for r in t.rows]
    chk('%s: %d filas de tabla, todas con «no dividir»' % (corto, len(filas)), all(r._tr.trPr is not None and r._tr.trPr.find(qn('w:cantSplit')) is not None for r in filas))
    # Anchos fijos con tblGrid en las tablas de datos (más de una columna)
    multi = [t for t in doc.tables if len(t.columns) > 1]
    chk('%s: %d tablas de varias columnas con tblGrid y ancho de 16,6 cm' % (corto, len(multi)),
        all(t._tbl.tblPr.find(qn('w:tblLayout')) is not None and abs(sum(int(g.get(qn('w:w'))) for g in t._tbl.tblGrid) - 9411) < 15 for t in multi))

est = texto(docs['ADS_3BTI_PILOTO_CAP4_ESTUDIANTE.docx'])

# Consignas dirigidas al estudiante: imperativos de tú (ambiguos con la 3.ª persona) prohibidos al inicio.
consignas = [it for _, its in C.ACT_E7 + C.ACT_E8 + [('n', i) for _, i in C.NIVELES] for it in its] + [it[1] for it in C.EVAL_ITEMS]
imp_tu = re.compile(r'^(Escribe|Lee|Clasifica|Marca|Completa|Elige|Redacta|Reescribe|Agrega|Propón|Justifica|Detecta|Diseña|Subraya|Explica|Arma|Encuentra|Anota)\b')
malas = [c[:40] for c in consignas if imp_tu.match(c)]
chk('Las %d consignas del estudiante usan imperativo de voseo (%s)' % (len(consignas), '; '.join(malas) or 'sin tuteo'), not malas)
gui = texto(docs['ADS_3BTI_PILOTO_CAP4_GUIA_DOCENTE.docx'])

# Sin claves en el material del estudiante
chk('Estudiante: sin claves ni criterios de corrección', 'Clave de la evaluación' not in est and 'Respuesta esperada' not in est)

# Coherencia de la evaluación entre libro y guía
H.CLAVES.clear()
letras = []
for it in C.EVAL_ITEMS:
    if it[0] == 'mc':
        baraja, letra = H.barajar_opciones(C.EVAL_ID, it[1], it[2])
        letras.append(letra)
        pos = est.find(it[1])
        bloque = est[pos:pos + 900]
        chk('Ítem «%s…»: la opción correcta está en la letra %s del libro' % (it[1][:30], letra), '%s) %s' % (letra, it[2][0]) in bloque)
chk('Claves de opción múltiple repartidas (%s)' % ', '.join(letras), len(set(letras)) == len(letras))
chk('Guía: la tabla de claves tiene los %d ítems' % len(C.EVAL_ITEMS), all(re.search(r'\n%d\n' % n, gui) for n in range(1, len(C.EVAL_ITEMS) + 1)))
pts = sum(1 for it in C.EVAL_ITEMS if it[0] in ('mc', 'vf')) + 3 + 5
chk('Puntaje total declarado (15) coincide con los ítems (%d)' % pts, pts == 15 and '15 puntos' in est and '15 puntos' in gui)

# Figura 4.1 y Tabla 4.2 usan las mismas cifras
tabla = {f[0]: (float(f[1].replace(',', '.')), float(f[2].replace(',', '.'))) for f in C.TABLA_EJ1[1]}
fig = {n.replace('\n', ' '): (x, y) for n, x, y in D.INTERESADOS}
chk('Figura 4.1 y Tabla 4.2: mismos interesados y mismas cifras', tabla == fig)
cuad = {'Gestionar de cerca': lambda x, y: x >= 5 and y >= 5, 'Mantener satisfecha': lambda x, y: x < 5 and y >= 5,
        'Mantener satisfecho': lambda x, y: x < 5 and y >= 5, 'Mantener informados': lambda x, y: x >= 5 and y < 5,
        'Observar': lambda x, y: x < 5 and y < 5}
chk('Tabla 4.2: cada estrategia coincide con el cuadrante de la figura', all(cuad[f[4]](tabla[f[0]][0], tabla[f[0]][1]) for f in C.TABLA_EJ1[1]))

# Figuras: tinta suficiente
for r in sorted(glob.glob(os.path.join(AQUI, 'figs', '*.png'))):
    im = Image.open(r).convert('L')
    px = list(im.get_flattened_data()) if hasattr(im, 'get_flattened_data') else list(im.getdata())
    tinta = sum(1 for v in px if v < 100) / len(px)
    chk('%s: imagen con contenido (tinta %.1f %%)' % (os.path.basename(r), 100 * tinta), tinta > 0.002)

# Trazabilidad: toda fuente citada en hallazgos existe en el registro E1
lineas_e1 = {l.split()[0] for l in C.TRANSCRIPCION_E1[1]}
fuentes = re.findall(r'E1-(L\d+)', ' '.join(f[1] for f in C.TABLA_HALLAZGOS[1]))
chk('Hallazgos: las %d líneas citadas existen en el registro E1' % len(fuentes), all(f in lineas_e1 for f in fuentes))
hall = {f[0] for f in C.TABLA_HALLAZGOS[1]}
orig = re.findall(r'H\d+', ' '.join(f[2] for f in C.TABLA_RC[1]))
chk('Requisitos candidatos: los %d orígenes citados existen como hallazgos' % len(orig), all(o in hall for o in orig))

# No duplicar el caso de Algorítmica 3.º (observación 5)
alg = ['Lucía Ortiz', 'Marcos Rojas', 'Ana Paula Vera', 'Hijo de hombre', 'El trueno entre las hojas', 'La babosa', 'Yo el Supremo', 'Roa Bastos']
chk('Sin personas ni datos del caso de biblioteca de Algorítmica 3.º', not any(a in est + gui for a in alg))

# Páginas del PDF de control, si existe
carpeta = sys.argv[1] if len(sys.argv) > 1 else None
if carpeta:
    import pymupdf
    for n, esperado in (('ADS_3BTI_PILOTO_CAP4_ESTUDIANTE.pdf', 16), ('ADS_3BTI_PILOTO_CAP4_GUIA_DOCENTE.pdf', 6)):
        d = pymupdf.open(os.path.join(carpeta, n))
        chk('%s: %d páginas en la renderización de control' % (n, len(d)), len(d) == esperado)
        pag_eval = [i for i, p in enumerate(d) if 'Evaluación del Capítulo 4' in p.get_text()]
        if 'ESTUDIANTE' in n:
            chk('La evaluación empieza en página propia (pág. %s)' % (pag_eval[0] + 1 if pag_eval else '?'),
                bool(pag_eval) and d[pag_eval[0]].get_text().strip().startswith('Evaluación del Capítulo 4'))

print('\nResultado: %d correctos, %d fallos' % (ok, fallos))
sys.exit(1 if fallos else 0)
