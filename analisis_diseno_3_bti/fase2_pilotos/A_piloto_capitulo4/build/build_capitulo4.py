# -*- coding: utf-8 -*-
"""Construye el piloto del Capítulo 4: libro del estudiante y guía docente (.docx).
Uso: python3 build_capitulo4.py  →  ../ADS_3BTI_PILOTO_CAP4_ESTUDIANTE.docx y ../ADS_3BTI_PILOTO_CAP4_GUIA_DOCENTE.docx"""
import os
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import helpers as H
import contenido as C
import guia
import diagramas

AQUI = os.path.dirname(os.path.abspath(__file__))
SALIDA = os.path.dirname(AQUI)
ANCHO_UTIL = 16.6  # cm, A4 con márgenes de 2,2 cm

# Identidad PROVISIONAL del piloto (familia verde BTI del estándar gráfico).
# Falta el deconflicto contra el inventario del catálogo antes de fijarla.
H.set_identidad('VERDE BTI (provisional)', '1B5E20', '2E7D32', '43A047', 'E8F5E9')


def portada(doc, tipo_doc):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    H.sombrear_celda(c, H.ID['DARK'])
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('\nAnálisis y Diseño de Sistemas\n3.er Curso · Bachillerato Técnico en Informática\n')
    r.bold = True; r.font.size = Pt(22); r.font.color.rgb = RGBColor.from_string('FFFFFF')
    H.cant_split(t)
    doc.add_paragraph()
    H.p_texto(doc, 'PILOTO EDITORIAL — %s' % C.CAPITULO.upper(), bold=True, size=13, color=H.ID['DARK'], align='center')
    H.p_texto(doc, C.TITULO, bold=True, size=12, color=H.ID['MED'], align='center')
    H.p_texto(doc, tipo_doc, size=11.5, align='center')
    H.p_texto(doc, 'Plan Optativo · Biblioteca Comercial 2026', size=11, align='center')
    H.p_texto(doc, 'Proyecto conductor: Sistema de Gestión de Biblioteca Escolar', bold=True, size=11.5,
              color=H.ID['MED'], align='center')
    doc.add_paragraph()
    H.caja(doc, 'Estado de este documento', [
        'Material piloto para auditoría independiente. No está aprobado para la venta ni para uso en aula.',
        'La capacidad que figura en la ficha es institucional (propuesta editorial). No es una capacidad '
        'oficial del MEC: no se verificó un programa oficial de la asignatura.',
        'La ubicación del capítulo en la estructura de 18 capítulos es provisional hasta cotejarla con el '
        'documento base del proyecto.',
    ], tipo='neutra')
    H.p_texto(doc, 'República del Paraguay · 2026', bold=True, size=12, align='center', antes=12)
    H.salto_pagina(doc)


def ficha(doc):
    t = doc.add_table(rows=4, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    H._bordes(t, color=H.ID['SOFT'])
    t.autofit = False
    datos = ((C.ROTULO_CAPACIDAD, C.CAPACIDAD), ('Tema', C.TEMA),
             ('Indicadores de logro', '\n'.join('• ' + i for i in C.INDICADORES)), ('Tiempo', C.TIEMPO))
    for fila, (k, v) in zip(t.rows, datos):
        fila.cells[0].width, fila.cells[1].width = Cm(4.2), Cm(12.4)
        rk = fila.cells[0].paragraphs[0].add_run(k)
        rk.bold = True; rk.font.size = Pt(9.5); rk.font.color.rgb = RGBColor.from_string('FFFFFF')
        H.sombrear_celda(fila.cells[0], H.ID['MED'])
        rv = fila.cells[1].paragraphs[0].add_run(v); rv.font.size = Pt(9.5)
        H.sombrear_celda(fila.cells[1], H.ID['LIGHT'])
    fijar_anchos(t, [4.2, 12.4])
    H.cant_split(t); H.keep_next_tabla(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def seccion(doc, sec, nivel=2):
    titulo, parrafos = sec
    doc.add_heading(titulo, level=nivel)
    for t in parrafos:
        H.p_texto(doc, t)


def tabla(doc, spec, anchos, rotulo=None, size=9):
    enc, filas = spec
    assert abs(sum(anchos) - ANCHO_UTIL) < 0.05, (rotulo, sum(anchos))
    t = H.add_table_md(doc, enc, filas, anchos=anchos, rotulo=rotulo, size=size)
    fijar_anchos(t, anchos)
    # Fila de encabezado repetida si la tabla pasa de página.
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def fijar_anchos(t, anchos):
    """Ancho fijo real: tblW + tblLayout fixed + tblGrid columna por columna.
    Sin tblGrid, Word y LibreOffice ignoran los anchos de celda."""
    tblPr = t._tbl.tblPr
    tw = sum(int(Cm(a).twips) for a in anchos)
    for tag in ('w:tblW', 'w:tblLayout'):
        old = tblPr.find(qn(tag))
        if old is not None:
            tblPr.remove(old)
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), str(tw)); w.set(qn('w:type'), 'dxa'); tblPr.append(w)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc in list(grid):
        grid.remove(gc)
    for a in anchos:
        gc = OxmlElement('w:gridCol'); gc.set(qn('w:w'), str(int(Cm(a).twips))); grid.append(gc)
    for row in t.rows:
        for c, a in zip(row.cells, anchos):
            c.width = Cm(a)


def pegar_renglones(doc):
    """Cada consigna queda unida a sus renglones de respuesta, y los renglones entre sí."""
    ps = doc.paragraphs
    for i in range(len(ps) - 1):
        sig = ps[i + 1].text.startswith('_' * 20)
        if sig and ps[i].text.strip():
            ps[i].paragraph_format.keep_with_next = True
    # Enunciado y opciones de una opción múltiple, juntos.
    for i in range(len(ps) - 1):
        if ps[i + 1].text[:3] in ('a) ', 'b) ', 'c) ', 'd) '):
            ps[i].paragraph_format.keep_with_next = True


def registro(doc, spec):
    titulo, lineas = spec
    H.caja(doc, titulo, lineas, tipo='referencia')


def actividades(doc, titulo, grupos, inicio=1):
    doc.add_heading(titulo, level=3)
    n = inicio
    for rotulo, items in grupos:
        p = doc.add_paragraph()
        r = p.add_run(rotulo)
        r.bold = True; r.font.color.rgb = RGBColor.from_string(H.ID['MED']); r.font.size = Pt(10.5)
        H.keep_next(p)
        for it in items:
            primera, *resto = it.split('\n')
            pi = H.p_texto(doc, '%d. %s' % (n, primera), size=10)
            pi.paragraph_format.left_indent = Cm(0.5)
            H.keep_next(pi)
            for linea in resto:
                pl = H.p_texto(doc, linea, size=10, despues=1, align='left')
                pl.paragraph_format.left_indent = Cm(1.0)
                H.keep_next(pl)
            H.p_texto(doc, '_' * 88, size=9, despues=2)
            H.p_texto(doc, '_' * 88, size=9, despues=6)
            n += 1
    return n


def libro_estudiante(figs, ruta):
    doc = H.new_doc()
    portada(doc, 'Libro del estudiante — capítulo de muestra')

    doc.add_heading('%s · %s' % (C.UNIDAD, C.CAPITULO), level=1)
    H.p_texto(doc, C.TITULO, bold=True, size=13, color=H.ID['MED'])
    ficha(doc)
    for t in C.APERTURA:
        H.p_texto(doc, t)
    H.caja(doc, 'El caso del libro', C.CASO, tipo='neutra')
    H.caja(doc, 'Con qué se conecta este capítulo', C.ARTICULACION, tipo='neutra')

    # Encuentro 7
    H.salto_pagina(doc)
    doc.add_heading(C.E7_TITULO, level=1)
    seccion(doc, C.S41)
    tabla(doc, C.TABLA_NRS, [4.1, 3.6, 4.9, 4.0], rotulo='Tabla 4.1 — Necesidad, requisito y solución')
    H.caja(doc, *C.CONCEPTO_ELICITAR, tipo='concepto')
    seccion(doc, C.S42)
    H.imagen(doc, figs[0], C.LEAD_FIG1, C.EPI_FIG1, ancho_cm=12.5)
    H.caja(doc, C.EJ1[0], C.EJ1[1], tipo='neutra')
    tabla(doc, C.TABLA_EJ1, [3.2, 1.4, 1.7, 6.6, 3.7], rotulo='Tabla 4.2 — Puntajes y estrategias del Ejemplo resuelto 4.1')
    seccion(doc, C.S43)
    tabla(doc, C.TABLA_TECNICAS, [3.0, 4.6, 4.2, 4.8], rotulo='Tabla 4.3 — Técnicas de elicitación')
    H.caja(doc, *C.EN_PARAGUAY, tipo='paraguay')
    seccion(doc, C.S44)
    tabla(doc, C.TABLA_PREGUNTAS, [5.6, 4.2, 6.8], rotulo='Tabla 4.4 — Preguntas a evitar')
    H.caja(doc, C.EJ2[0], C.EJ2[1], tipo='neutra')
    H.imagen(doc, figs[1], C.LEAD_FIG2, C.EPI_FIG2, ancho_cm=15.5)
    actividades(doc, 'Actividades del Encuentro 7', C.ACT_E7)

    # Encuentro 8
    H.salto_pagina(doc)
    doc.add_heading(C.E8_TITULO, level=1)
    seccion(doc, C.S45)
    registro(doc, C.TRANSCRIPCION_E1)
    H.caja(doc, C.EJ3[0], C.EJ3[1], tipo='neutra')
    tabla(doc, C.TABLA_HALLAZGOS, [1.2, 2.2, 9.6, 3.6], rotulo='Tabla 4.5 — Hallazgos del registro E1')
    seccion(doc, C.S46)
    H.caja(doc, *C.CONCEPTO_RC, tipo='concepto')
    H.caja(doc, C.EJ4[0], C.EJ4[1], tipo='neutra')
    tabla(doc, C.TABLA_RC, [1.6, 9.4, 2.4, 3.2], rotulo='Tabla 4.6 — Requisitos candidatos del registro E1')
    seccion(doc, C.S47)
    H.caja(doc, C.EJ5[0], C.EJ5[1], tipo='neutra')
    seccion(doc, C.S48)
    actividades(doc, 'Actividades del Encuentro 8', C.ACT_E8)

    # Práctica diferenciada
    H.salto_pagina(doc)
    doc.add_heading(C.PRACTICA_TITULO, level=1)
    for t in C.PRACTICA_INTRO:
        H.p_texto(doc, t)
    registro(doc, C.TRANSCRIPCION_E2)
    registro(doc, C.OBSERVACION_O1)
    for nivel, items in C.NIVELES:
        actividades(doc, nivel, [('Consignas', items)])
    H.caja(doc, 'Punto de control', [C.PRACTICA_CONTROL], tipo='neutra')
    H.caja(doc, *C.REPASO, tipo='neutra')
    H.caja(doc, *C.CALENDARIO_REDUCIDO, tipo='neutra')

    # Evaluación
    H.CLAVES.clear()
    H.evaluacion(doc, C.EVAL_ID, C.EVAL_TITULO, False, C.EVAL_TEXTO_BASE,
                 [it if it[0] != 'ab' else ('ab', it[1], it[2], it[3]) for it in C.EVAL_ITEMS],
                 remision='Evalúa los cuatro indicadores de logro del capítulo. Puntaje total: 15 puntos (ítems 1 a 7: 1 punto cada uno; ítem 8: 3 puntos; ítem 9: 5 puntos).')
    pegar_renglones(doc)
    H.finalizar(doc, ruta)
    return list(H.CLAVES)


def main():
    figs = diagramas.generar()
    r1 = os.path.join(SALIDA, 'ADS_3BTI_PILOTO_CAP4_ESTUDIANTE.docx')
    r2 = os.path.join(SALIDA, 'ADS_3BTI_PILOTO_CAP4_GUIA_DOCENTE.docx')
    claves = libro_estudiante(figs, r1)
    guia.construir(r2, claves, portada, tabla, ficha)
    print(r1); print(r2)


if __name__ == '__main__':
    main()
