# -*- coding: utf-8 -*-
"""helpers.py — pipeline canónico cuadernillos-mec, reconstruido desde la doctrina
(las carpetas scripts/ y references/ de la skill no llegaron montadas).
Incorpora: cantSplit universal, keepNext en fichas/cajas, título fundido a la
primera fila de caja, páginas propias de evaluaciones, footer con campos reales,
updateFields ANTES de w:compat, barajado determinista de opciones."""
import hashlib
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ID = {}

def set_identidad(nombre, dark, med, soft, light):
    ID.update(nombre=nombre, DARK=dark, MED=med, SOFT=soft, LIGHT=light)

AMBAR = 'B8860B'
AZUL = '1B4F72'
REF = 'E0F2F1'      # texto de referencia propio
DP = 'F2FAFB'       # antología dominio público
OBS = 'FAF3F3'      # bloque «Observá en el texto»

# ---------- utilidades XML ----------

def _shd(el, fill):
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    el.append(shd)

def sombrear_celda(cell, fill):
    _shd(cell._tc.get_or_add_tcPr(), fill)

def cant_split(table):
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn('w:cantSplit')) is None:
            trPr.append(OxmlElement('w:cantSplit'))

def keep_next_tabla(table, salvo_ultima=True):
    rows = table.rows[:-1] if salvo_ultima else table.rows
    for row in rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True

def keep_next(p):
    p.paragraph_format.keep_with_next = True

def _bordes(table, color='BFBFBF', sz='4'):
    tbl = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), sz)
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        borders.append(e)
    tbl.append(borders)

def _sin_bordes_internos(table):
    tbl = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '6')
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), ID['MED'])
        borders.append(e)
    for edge in ('insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'none')
        borders.append(e)
    tbl.append(borders)

def salto_pagina(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    return p

# ---------- documento ----------

def new_doc():
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'; st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(4)
    for i, (sz, col) in enumerate((((16), ID['DARK']), ((13), ID['DARK']), ((11.5), ID['MED']))):
        h = doc.styles['Heading %d' % (i + 1)]
        h.font.name = 'Calibri'; h.font.size = Pt(sz); h.font.bold = True
        h.font.color.rgb = RGBColor.from_string(col)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.space_before = Pt(10); h.paragraph_format.space_after = Pt(4)
    sec = doc.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    for a in ('top_margin', 'bottom_margin'):
        setattr(sec, a, Cm(2.0))
    for a in ('left_margin', 'right_margin'):
        setattr(sec, a, Cm(2.2))
    return doc

def p_texto(doc, texto, bold=False, italic=False, size=10.5, color=None,
            align=None, antes=0, despues=4):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if align == 'center':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'right':
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(despues)
    return p

def portada(doc, materia, curso, subtitulo, proyecto, tipo_doc='Cuadernillo de trabajo del estudiante'):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    sombrear_celda(cell, ID['DARK'])
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('\n%s\n%s\n' % (materia, curso))
    r.bold = True; r.font.size = Pt(24); r.font.color.rgb = RGBColor.from_string('FFFFFF')
    cant_split(t)
    doc.add_paragraph()
    p_texto(doc, subtitulo, bold=True, size=13, color=ID['DARK'], align='center')
    p_texto(doc, 'Plan Común · Programa del Ministerio de Educación y Ciencias', size=11, align='center')
    p_texto(doc, tipo_doc, size=11.5, align='center')
    p_texto(doc, 'Enfoque constructivista «Aprender Haciendo»', italic=True, size=10.5, align='center')
    doc.add_paragraph()
    p_texto(doc, proyecto, bold=True, size=12, color=ID['MED'], align='center')
    p_texto(doc, 'Escuchar · Hablar · Leer · Escribir', italic=True, size=10.5, align='center')
    doc.add_paragraph()
    p_texto(doc, 'República del Paraguay · 2026', bold=True, size=12, align='center')
    salto_pagina(doc)

def presentacion(doc, parrafos):
    doc.add_heading('Presentación', level=1)
    for t in parrafos:
        p_texto(doc, t)
    salto_pagina(doc)

def indice_toc(doc):
    doc.add_heading('Índice', level=1)
    p_texto(doc, 'Índice — presioná Ctrl+E y luego F9 para actualizar.', italic=True, size=9.5, color='555555')
    p = doc.add_paragraph()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'TOC \\o "1-2" \\h \\z \\u')
    run = OxmlElement('w:r'); t = OxmlElement('w:t')
    t.text = 'Actualizá este índice: Ctrl+E → F9.'
    run.append(t); fld.append(run)
    p._p.append(fld)
    salto_pagina(doc)

def ficha(doc, capacidad, tema, indicadores):
    t = doc.add_table(rows=3, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _bordes(t, color=ID['SOFT'])
    t.autofit = False
    anchos = (Cm(3.2), Cm(13.4))
    datos = (('Capacidad', capacidad), ('Tema', tema),
             ('Indicadores de logro', '\n'.join('• ' + i for i in indicadores)))
    for fila, (k, v) in zip(t.rows, datos):
        fila.cells[0].width, fila.cells[1].width = anchos
        pk = fila.cells[0].paragraphs[0]
        rk = pk.add_run(k); rk.bold = True; rk.font.size = Pt(9.5)
        rk.font.color.rgb = RGBColor.from_string('FFFFFF')
        sombrear_celda(fila.cells[0], ID['MED'])
        pv = fila.cells[1].paragraphs[0]
        rv = pv.add_run(v); rv.font.size = Pt(9.5)
        sombrear_celda(fila.cells[1], ID['LIGHT'])
    cant_split(t); keep_next_tabla(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def caja(doc, titulo, bloques, tipo='concepto', fuente=None, nota=None, mono=False):
    """bloques: lista de str (párrafos). Título FUNDIDO a la primera fila.
    tipo: concepto|ejemplo|paraguay|neutra|taller|referencia|dp|observa"""
    fill = {'concepto': ID['LIGHT'], 'ejemplo': ID['LIGHT'], 'paraguay': ID['LIGHT'],
            'neutra': ID['LIGHT'], 'taller': ID['LIGHT'],
            'referencia': REF, 'dp': DP, 'observa': OBS}[tipo]
    pref = {'concepto': 'Concepto clave — ', 'ejemplo': 'Ejemplo — ',
            'paraguay': 'En Paraguay — ', 'taller': 'Taller — '}.get(tipo, '')
    n_extra = (1 if fuente else 0) + (1 if nota else 0)
    filas = max(1, len(bloques)) + n_extra
    t = doc.add_table(rows=filas, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _sin_bordes_internos(t)
    # fila 1: título + primer bloque (anti-huérfano)
    c0 = t.rows[0].cells[0]
    sombrear_celda(c0, fill)
    p0 = c0.paragraphs[0]
    r0 = p0.add_run(pref + titulo)
    r0.bold = True; r0.font.size = Pt(11)
    r0.font.color.rgb = RGBColor.from_string(ID['MED'])
    if bloques:
        p1 = c0.add_paragraph()
        _run_bloque(p1, bloques[0], mono)
    for i, b in enumerate(bloques[1:], start=1):
        c = t.rows[i].cells[0]
        sombrear_celda(c, fill)
        _run_bloque(c.paragraphs[0], b, mono)
    fila = len(bloques)
    if nota:
        c = t.rows[fila].cells[0]; fila += 1
        sombrear_celda(c, fill)
        rn = c.paragraphs[0].add_run(nota)
        rn.italic = True; rn.font.size = Pt(9.5)
    if fuente:
        c = t.rows[fila].cells[0]
        sombrear_celda(c, fill)
        rf = c.paragraphs[0].add_run(fuente)
        rf.italic = True; rf.font.size = Pt(9)
        rf.font.color.rgb = RGBColor.from_string('555555')
    cant_split(t); keep_next_tabla(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def _run_bloque(p, texto, mono=False):
    """Bloques con \n internos: saltos de línea sin justificar (versos)."""
    if '\n' in texto or mono:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        partes = texto.split('\n')
        for j, seg in enumerate(partes):
            if j:
                p.add_run().add_break()
            r = p.add_run(seg)
            r.font.size = Pt(10)
            if mono:
                r.font.name = 'Courier New'
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(texto)
        r.font.size = Pt(10)
    return p

def observa(doc, consignas):
    caja(doc, 'Observá en el texto',
         ['%d. %s' % (i + 1, c) for i, c in enumerate(consignas)], tipo='observa')

def actividades(doc, lee, analiza, produci, comparti):
    doc.add_heading('Actividades', level=3)
    for rotulo, items in (('Leé y comprendé', lee), ('Analizá', analiza),
                          ('Producí', produci), ('Compartí', comparti)):
        p = doc.add_paragraph()
        r = p.add_run(rotulo)
        r.bold = True; r.font.color.rgb = RGBColor.from_string(ID['MED']); r.font.size = Pt(10.5)
        keep_next(p)
        for i, it in enumerate(items, 1):
            pi = p_texto(doc, '%d. %s' % (i, it), size=10)
            pi.paragraph_format.left_indent = Cm(0.5)

def add_table_md(doc, encabezados, filas, anchos=None, rotulo=None, size=9.5):
    if rotulo:
        pr = p_texto(doc, rotulo, bold=True, size=10, color=ID['DARK'])
        keep_next(pr)
    t = doc.add_table(rows=1 + len(filas), cols=len(encabezados))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _bordes(t)
    if anchos:
        t.autofit = False
    for j, h in enumerate(encabezados):
        c = t.rows[0].cells[j]
        if anchos:
            c.width = Cm(anchos[j])
        sombrear_celda(c, ID['MED'])
        r = c.paragraphs[0].add_run(h)
        r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string('FFFFFF')
    for i, fila in enumerate(filas, 1):
        for j, v in enumerate(fila):
            c = t.rows[i].cells[j]
            if anchos:
                c.width = Cm(anchos[j])
            r = c.paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
            if i % 2 == 0:
                sombrear_celda(c, 'F5F5F5')
    cant_split(t)
    return t

# ---------- evaluaciones ----------

def barajar_opciones(eval_id, enunciado, opciones):
    """opciones[0] es la correcta. Devuelve (lista_barajada, letra_correcta)."""
    h = hashlib.sha256((eval_id + '|' + enunciado).encode('utf8')).hexdigest()
    orden = sorted(range(len(opciones)), key=lambda i: h[i * 4:(i + 1) * 4])
    baraja = [opciones[i] for i in orden]
    letra = 'abcd'[baraja.index(opciones[0])]
    return baraja, letra

def _banda(doc, texto, fill):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    sombrear_celda(c, fill)
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texto)
    r.bold = True; r.font.size = Pt(12.5); r.font.color.rgb = RGBColor.from_string('FFFFFF')
    cant_split(t)
    pid = doc.add_paragraph()
    ri = pid.add_run('Nombre y apellido: ______________________________  Curso: ______  Fecha: ____/____/____')
    ri.font.size = Pt(9.5)
    pid.paragraph_format.space_before = Pt(6)

def _texto_base(doc, titulo, parrafos):
    caja(doc, titulo, parrafos, tipo='neutra')

CLAVES = []   # (eval_id, item, letra/clave, detalle)

def _mc(doc, eval_id, n, enunciado, opciones, detalle=''):
    baraja, letra = barajar_opciones(eval_id, enunciado, opciones)
    p_texto(doc, '%d. %s' % (n, enunciado), size=10)
    for j, o in enumerate(baraja):
        po = p_texto(doc, '%s) %s' % ('abcd'[j], o), size=10)
        po.paragraph_format.left_indent = Cm(0.6)
    CLAVES.append((eval_id, n, letra, detalle or opciones[0]))

def _vf(doc, eval_id, n, enunciado, es_v, justif):
    p_texto(doc, '%d. (V / F) %s' % (n, enunciado), size=10)
    CLAVES.append((eval_id, n, 'V' if es_v else 'F', justif))

def _abierta(doc, eval_id, n, enunciado, orientacion, renglones=3):
    p_texto(doc, '%d. %s' % (n, enunciado), size=10)
    for _ in range(renglones):
        p_texto(doc, '_' * 92, size=9, despues=6)
    CLAVES.append((eval_id, n, 'abierta', orientacion))

def evaluacion(doc, eval_id, titulo, etapa, texto_base, items, remision=None):
    salto_pagina(doc)
    doc.add_heading(titulo, level=2)
    _banda(doc, titulo, AZUL if etapa else AMBAR)
    if remision:
        p_texto(doc, remision, italic=True, size=9.5)
    if texto_base:
        _texto_base(doc, texto_base[0], texto_base[1])
    n = 0
    for it in items:
        n += 1
        kind = it[0]
        if kind == 'mc':
            _mc(doc, eval_id, n, it[1], it[2], it[3] if len(it) > 3 else '')
        elif kind == 'vf':
            _vf(doc, eval_id, n, it[1], it[2], it[3])
        else:
            _abierta(doc, eval_id, n, it[1], it[2], it[3] if len(it) > 3 else 3)

def prueba_diagnostica(doc, intro, preguntas):
    doc.add_heading('Prueba diagnóstica — ¿Qué sabés ya?', level=1)
    p_texto(doc, intro, italic=True, size=10)
    for i, q in enumerate(preguntas, 1):
        p_texto(doc, '%d. %s' % (i, q), size=10)
        p_texto(doc, '_' * 92, size=9, despues=6)
    p_texto(doc, 'Esta prueba no lleva nota: sirve para que tu docente conozca tu punto de partida.', italic=True, size=9.5)

# ---------- imagen ----------

def imagen(doc, ruta, lead_in, epigrafe, ancho_cm=14.0):
    pl = p_texto(doc, lead_in)
    keep_next(pl)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(ruta, width=Cm(ancho_cm))
    keep_next(p)
    pe = doc.add_paragraph(); pe.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = pe.add_run(epigrafe)
    r.italic = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string('595959')

# ---------- cierre ----------

def _fld(run_parent, instr):
    for el, attrs in (('w:fldChar', {'w:fldCharType': 'begin'}),
                      ('w:instrText', None),
                      ('w:fldChar', {'w:fldCharType': 'end'})):
        r = OxmlElement('w:r')
        e = OxmlElement(el)
        if el == 'w:instrText':
            e.set(qn('xml:space'), 'preserve'); e.text = instr
        else:
            for k, v in attrs.items():
                e.set(qn(k), v)
        r.append(e)
        run_parent.append(r)

def footer_paginas(doc):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Página '); r.font.size = Pt(9)
    _fld(p._p, ' PAGE ')
    r2 = p.add_run(' de '); r2.font.size = Pt(9)
    _fld(p._p, ' NUMPAGES ')
    for rr in p.runs:
        rr.font.size = Pt(9)

def update_fields(doc):
    settings = doc.settings.element
    if settings.find(qn('w:updateFields')) is None:
        el = OxmlElement('w:updateFields')
        el.set(qn('w:val'), 'true')
        compat = settings.find(qn('w:compat'))
        if compat is not None:
            compat.addprevious(el)
        else:
            settings.append(el)

def finalizar(doc, ruta):
    footer_paginas(doc)
    update_fields(doc)
    doc.save(ruta)
    return ruta
