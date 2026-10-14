# -*- coding: utf-8 -*-
"""Controles de la guía docente v2.1 (correspondencia requisitos-criterios). Uso: python3 verificar_v2_1.py [carpeta_pdf]"""
import os, re, sys, hashlib
from docx import Document
from docx.oxml.ns import qn
import contenido as C
import guia_v2_1 as G
import correspondencia as K

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
ok = fallos = 0
def chk(desc, cond):
    global ok, fallos
    if cond: ok += 1
    else: fallos += 1
    print(('CORRECTO ' if cond else 'FALLA    ') + desc)

ids_cat = [r[0] for r in C.CATALOGO]
ids_k = [r[0] for r in K.CORRESPONDENCIA]
chk('Los 15 requisitos del catálogo, en el mismo orden y sin repetir', ids_k == ids_cat and len(ids_cat) == 15)
chk('Cada criterio tiene la forma «Dado… cuando… entonces…»', all(re.match(r'^Dad[oa]s? .+, cuando .+, entonces .+', c) for *_, c in K.CORRESPONDENCIA))
cuenta = {k: [r[0] for r in K.CORRESPONDENCIA if r[1] == k] for k in K.ETIQUETA}
print('INFO     Modelos: %s · Estudiante: %s · Docente: %s' % (', '.join(cuenta['MODELO']), ', '.join(cuenta['ESTUDIANTE']), ', '.join(cuenta['DOCENTE'])))

# Modelos: el criterio de la guía es idéntico al que figura en el libro, en el lugar indicado.
libro = {f[0]: f[1] for f in C.TABLA_43[1]}
ej3 = [b for b in C.EJ3[1] if b.startswith('Criterio de aceptación CA-04: ')][0].replace('Criterio de aceptación CA-04: ', '')
ej4 = [b for b in C.EJ4[1] if b.startswith('Paso 4. Criterio de aceptación: ')][0].replace('Paso 4. Criterio de aceptación: ', '')
libro['RF-04'] = ej3
libro['RNF-01'] = ej4
for i, o, donde, crit in K.CORRESPONDENCIA:
    if o == 'MODELO':
        chk('%s (modelo, %s): coincide palabra por palabra con el libro' % (i, donde), libro.get(i) == crit)

# Trabajo del estudiante: la actividad lo pide y la guía da la misma respuesta que en sus claves.
act_e8 = [it for _, its in C.ACT_E8 for it in its]
claves_e8 = {fila[0]: fila[1] for fila in G.CLAVES_E8}
for i, o, donde, crit in K.CORRESPONDENCIA:
    if o == 'ESTUDIANTE':
        n = int(re.search(r'Actividad (\d)', donde).group(1))
        chk('%s (estudiante, %s): la consigna lo nombra y la clave de esa actividad contiene el criterio' % (i, donde),
            i in act_e8[n - 1] and crit in claves_e8[str(n)])

# Referencias del docente: no tienen criterio en el libro.
texto_libro = ' '.join(str(x) for x in (C.TABLA_43, C.EJ3, C.EJ4, C.ACT_E7, C.ACT_E8, C.NIVELES))
for i in cuenta['DOCENTE']:
    chk('%s (referencia docente): el libro no trae su criterio' % i, i not in libro and not re.search(re.escape(i) + r'[^\n]{0,40}Dado', texto_libro))

# Documento
d = Document(os.path.join(RAIZ, 'ADS_3BTI_CAP4_v2_1_GUIA_DOCENTE.docx'))
tx = '\n'.join(p.text for p in d.paragraphs) + '\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
chk('La guía contiene la sección 3 y la tabla de correspondencia', '3. Requisitos del catálogo y sus criterios de aceptación' in tx and all(c in tx for *_, c in K.CORRESPONDENCIA))
chk('Secciones de la guía numeradas del 1 al 7 sin saltos', [int(m) for m in re.findall(r'^(\d)\. ', '\n'.join(p.text for p in d.paragraphs if p.style.name.startswith('Heading 2')), re.M)] == list(range(1, 8)))
filas = [r for t in d.tables for r in t.rows]
chk('%d filas de tabla, todas con «no dividir»' % len(filas), all(r._tr.trPr is not None and r._tr.trPr.find(qn('w:cantSplit')) is not None for r in filas))
multi = [t for t in d.tables if len(t.columns) > 1]
chk('%d tablas de varias columnas con tblGrid y 16,6 cm' % len(multi), all(abs(sum(int(g.get(qn('w:w'))) for g in t._tbl.tblGrid) - 9411) < 15 for t in multi))
chk('Sin formas de tuteo inequívocas', not re.search(r'\b(tienes|puedes|debes|sabes|quieres|eres|haz|propón|tú)\b', tx, re.I))
chk('Sin identificadores de modelos de IA', not re.search(r'(?i)\b(claude|gpt|gemini|opus|sonnet|codex)\b', tx))
chk('Sin símbolos prohibidos', not re.search('[✓✔✗✘✅]', tx))
chk('Capacidad rotulada como no oficial del MEC', 'No es una capacidad oficial del MEC' in tx)

# El libro del estudiante v2 no se tocó: el contenido es el mismo (comparación de los módulos con la v2).
v2 = os.path.join(RAIZ, '..', 'A2_capitulo4_v2', 'build')
for m in ('contenido.py', 'diagramas.py', 'maqueta.py', 'helpers.py', 'build_v2.py'):
    a = hashlib.sha256(open(os.path.join(AQUI, m), 'rb').read()).hexdigest()
    b = hashlib.sha256(open(os.path.join(v2, m), 'rb').read()).hexdigest()
    chk('%s idéntico al de la v2 (el libro del estudiante no cambia)' % m, a == b)

if len(sys.argv) > 1:
    import pymupdf
    p = pymupdf.open(os.path.join(sys.argv[1], 'ADS_3BTI_CAP4_v2_1_GUIA_DOCENTE.pdf'))
    print('INFO     Guía v2.1: %d páginas en la renderización de control' % len(p))
    pags = [i + 1 for i, pg in enumerate(p) if 'Correspondencia entre requisitos y criterios' in pg.get_text()]
    chk('La tabla de correspondencia aparece en el PDF (pág. %s)' % pags, bool(pags))

print('\nResultado: %d correctos, %d fallos' % (ok, fallos))
sys.exit(1 if fallos else 0)
