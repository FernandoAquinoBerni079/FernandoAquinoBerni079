# -*- coding: utf-8 -*-
"""Funciones de maqueta compartidas por el libro del estudiante y la guía docente (Capítulo 4 v2).
Tomadas del piloto anterior sin cambios de comportamiento: anchos fijos con tblGrid, encabezado
repetido, consignas unidas a sus renglones."""
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import helpers as H

ANCHO_UTIL = 16.6  # cm, A4 con márgenes de 2,2 cm

# Identidad PROVISIONAL (familia verde BTI). Pendiente de cotejo con el catálogo.
H.set_identidad('VERDE BTI (provisional)', '1B5E20', '2E7D32', '43A047', 'E8F5E9')


def portada(doc, C, tipo_doc, estado):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    H.sombrear_celda(c, H.ID['DARK'])
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('\nAnálisis y Diseño de Sistemas\n3.er Curso · Bachillerato Técnico en Informática\n')
    r.bold = True; r.font.size = Pt(22); r.font.color.rgb = RGBColor.from_string('FFFFFF')
    H.cant_split(t)
    doc.add_paragraph()
    H.p_texto(doc, 'PILOTO EDITORIAL — %s' % C.CAPITULO.upper(), bold=True, size=13, color=H.ID['DARK'], align='center')
    H.p_texto(doc, C.TITULO, bold=True, size=12, color=H.ID['MED'], align='center')
    H.p_texto(doc, tipo_doc, size=11.5, align='center')
    H.p_texto(doc, 'Plan Optativo · Biblioteca Comercial 2026', size=11, align='center')
    H.p_texto(doc, 'Proyecto conductor: Sistema de Gestión de Biblioteca Escolar', bold=True, size=11.5,
              color=H.ID['MED'], align='center')
    doc.add_paragraph()
    H.caja(doc, 'Estado de este documento', estado, tipo='neutra')
    H.p_texto(doc, 'República del Paraguay · 2026', bold=True, size=12, align='center', antes=12)
    H.salto_pagina(doc)


def ficha(doc, C):
    t = doc.add_table(rows=4, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    H._bordes(t, color=H.ID['SOFT'])
    t.autofit = False
    datos = ((C.ROTULO_CAPACIDAD, C.CAPACIDAD), ('Tema', C.TEMA),
             ('Indicadores de logro', '\n'.join('• ' + i for i in C.INDICADORES)), ('Tiempo', C.TIEMPO))
    for fila, (k, v) in zip(t.rows, datos):
        fila.cells[0].width, fila.cells[1].width = Cm(4.2), Cm(12.4)
        rk = fila.cells[0].paragraphs[0].add_run(k)
        rk.bold = True; rk.font.size = Pt(9.5); rk.font.color.rgb = RGBColor.from_string('FFFFFF')
        H.sombrear_celda(fila.cells[0], H.ID['MED'])
        rv = fila.cells[1].paragraphs[0].add_run(v); rv.font.size = Pt(9.5)
        H.sombrear_celda(fila.cells[1], H.ID['LIGHT'])
    fijar_anchos(t, [4.2, 12.4])
    H.cant_split(t); H.keep_next_tabla(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def seccion(doc, sec, nivel=2):
    titulo, parrafos = sec
    doc.add_heading(titulo, level=nivel)
    for t in parrafos:
        H.p_texto(doc, t)


def tabla(doc, spec, anchos, rotulo=None, size=9):
    enc, filas = spec
    assert abs(sum(anchos) - ANCHO_UTIL) < 0.05, (rotulo, sum(anchos))
    t = H.add_table_md(doc, enc, filas, anchos=anchos, rotulo=rotulo, size=size)
    fijar_anchos(t, anchos)
    # Fila de encabezado repetida si la tabla pasa de página.
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def fijar_anchos(t, anchos):
    """Ancho fijo real: tblW + tblLayout fixed + tblGrid columna por columna.
    Sin tblGrid, Word y LibreOffice ignoran los anchos de celda."""
    tblPr = t._tbl.tblPr
    tw = sum(int(Cm(a).twips) for a in anchos)
    for tag in ('w:tblW', 'w:tblLayout'):
        old = tblPr.find(qn(tag))
        if old is not None:
            tblPr.remove(old)
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), str(tw)); w.set(qn('w:type'), 'dxa'); tblPr.append(w)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc in list(grid):
        grid.remove(gc)
    for a in anchos:
        gc = OxmlElement('w:gridCol'); gc.set(qn('w:w'), str(int(Cm(a).twips))); grid.append(gc)
    for row in t.rows:
        for c, a in zip(row.cells, anchos):
            c.width = Cm(a)


def pegar_renglones(doc):
    """Cada consigna queda unida a sus renglones de respuesta, y los renglones entre sí."""
    ps = doc.paragraphs
    for i in range(len(ps) - 1):
        sig = ps[i + 1].text.startswith('_' * 20)
        if sig and ps[i].text.strip():
            ps[i].paragraph_format.keep_with_next = True
    # Enunciado y opciones de una opción múltiple, juntos.
    for i in range(len(ps) - 1):
        if ps[i + 1].text[:3] in ('a) ', 'b) ', 'c) ', 'd) '):
            ps[i].paragraph_format.keep_with_next = True


def registro(doc, spec):
    titulo, lineas = spec
    H.caja(doc, titulo, lineas, tipo='referencia')


def actividades(doc, titulo, grupos, inicio=1):
    doc.add_heading(titulo, level=3)
    n = inicio
    for rotulo, items in grupos:
        p = doc.add_paragraph()
        r = p.add_run(rotulo)
        r.bold = True; r.font.color.rgb = RGBColor.from_string(H.ID['MED']); r.font.size = Pt(10.5)
        H.keep_next(p)
        for it in items:
            primera, *resto = it.split('\n')
            pi = H.p_texto(doc, '%d. %s' % (n, primera), size=10)
            pi.paragraph_format.left_indent = Cm(0.5)
            H.keep_next(pi)
            for linea in resto:
                pl = H.p_texto(doc, linea, size=10, despues=1, align='left')
                pl.paragraph_format.left_indent = Cm(1.0)
                H.keep_next(pl)
            H.p_texto(doc, '_' * 88, size=9, despues=2)
            H.p_texto(doc, '_' * 88, size=9, despues=6)
            n += 1
    return n


