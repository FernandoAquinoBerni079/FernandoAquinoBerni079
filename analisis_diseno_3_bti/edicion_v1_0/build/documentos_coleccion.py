# -*- coding: utf-8 -*-
"""Edición comercial v1.0 del libro, de la guía docente y del cuaderno de prácticas, a partir de los DOCX aprobados de la
colección v0.2 (commit 88c5cb5), que solo se leen.

Cambios, y nada más:
1. Página 1 de los tres documentos: los textos de estado de la versión de auditoría («versión candidata», «no aprobado
   para la venta») se reemplazan por los de la edición comercial, con el aviso de propuesta institucional, el aviso de
   derechos y los datos de edición (titular, editor e ISBN) que se completan en datos_edicion.json.
2. Cuaderno de prácticas: los renglones de respuesta pasan de unos 5 mm a 8 mm entre renglones y de dos a tres por
   consigna, para escribir a mano; el índice se actualiza con las páginas nuevas.
Controles: el libro y la guía conservan la cantidad de páginas y el texto de cada página desde la 2; el cuaderno conserva
el texto de todas las consignas, en el mismo orden, y ninguna página empieza con un renglón suelto.
Uso: python3 -B analisis_diseno_3_bti/edicion_v1_0/build/documentos_coleccion.py
"""
import copy, json, os, re, subprocess, sys
sys.dont_write_bytecode = True
import docx, pymupdf
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import Table
from docx.text.paragraph import Paragraph

AQUI = os.path.dirname(os.path.abspath(__file__))
ED = os.path.dirname(AQUI)
ADS = os.path.dirname(ED)
COL = os.path.join(ADS, 'coleccion')
SAL = os.path.join(ED, 'salida'); os.makedirs(SAL, exist_ok=True)
DATOS = json.load(open(os.path.join(ED, 'datos_edicion.json'), encoding='utf8'))
norm = lambda s: re.sub(r'\s+', ' ', s or '').strip()

ORIGEN = {'libro': 'ADS_3BTI_LIBRO_ANUAL_ESTUDIANTE_OFICIO_O3', 'guia': 'ADS_3BTI_GUIA_DOCENTE_ANUAL_OFICIO_O3',
          'cuaderno': 'ADS_3BTI_CUADERNO_ANUAL_PRACTICAS_OFICIO_O3'}
DESTINO = {'libro': 'ADS_3BTI_LIBRO_DEL_ESTUDIANTE_v1_0', 'guia': 'ADS_3BTI_GUIA_DOCENTE_SOLUCIONARIO_v1_0',
           'cuaderno': 'ADS_3BTI_CUADERNO_DE_PRACTICAS_v1_0'}
PENDIENTE = DATOS['texto_si_falta']


def dato(clave, producto=None):
    v = DATOS[clave].get(producto) if isinstance(DATOS[clave], dict) else DATOS[clave]
    return v if v else PENDIENTE


PRIMERO = {
    'libro': 'Edición comercial v1.0 de la colección anual: 18 capítulos en seis unidades, 36 encuentros de 5 horas cátedra '
             '(180 HC), 16 prácticas en tres niveles y seis hitos integradores, con un proyecto que va del problema al '
             'prototipo web.'
             ' Cuando el texto dice «Work», se refiere al informe de diagnóstico y propuesta curricular de la colección (v0.1, 14 de octubre de 2026).',
    'guia': 'Edición comercial v1.0: guía del docente con la secuencia de los 36 encuentros (180 HC), las claves de todas las '
            'actividades, prácticas y evaluaciones, las rúbricas de los seis hitos y los instrumentos del año. En esta guía, '
            '«Work» designa el informe de diagnóstico y propuesta curricular de la colección (v0.1, 14 de octubre de 2026).',
    'cuaderno': 'Edición comercial v1.0: las 16 prácticas del libro del estudiante, en tres niveles, con renglones para '
                'responder a mano. Se usa junto con el libro, que tiene la teoría, los ejemplos y las evaluaciones.'
                ' Cuando el texto dice «Work», se refiere al informe de diagnóstico y propuesta curricular de la colección (v0.1, 14 de octubre de 2026).',
}
SEGUNDO = ('Propuesta editorial para el programa institucional del Plan Optativo, que cada institución puede elaborar según el '
           'Diseño Curricular del MEC (mayo de 2026, p. 29). No es un programa oficial ni está aprobado por el MEC: las '
           'capacidades C1 a C10 son propuestas para el programa institucional.')
TERCERO = ('El entorno técnico de referencia (HTML y CSS, PHP y SQLite, sin servicios pagos) se probó en Linux. Antes de usarlo '
           'en las computadoras de la institución se hace la prueba de campo de la guía docente (Parte IV, sección 0.2).')


def poner_texto(p, texto):
    runs = p.runs
    runs[0].text = texto
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def pagina_1(d, clave):
    body = list(d.element.body.iterchildren())
    p5 = Paragraph(body[5], d); assert 'Biblioteca Comercial 2026' in p5.text, p5.text
    poner_texto(p5, 'Plan Optativo · propuesta para el programa institucional · Edición 2026')
    t = Table(body[8], d)
    celdas = [c for r in t.rows for c in r.cells]
    assert celdas[0].paragraphs[0].text == 'Estado de este documento'
    poner_texto(celdas[0].paragraphs[0], 'Sobre esta edición')
    poner_texto(celdas[0].paragraphs[1], PRIMERO[clave])
    for c in celdas[1:-2]:   # la página 1 ya explica «Work»: la tabla de estado no lo repite
        for q in c.paragraphs:
            if 'al índice de Work' in q.text:
                poner_texto(q, q.text.replace('complementos editoriales al índice de Work', 'complementos editoriales'))
    # las filas intermedias del cuaderno (datos de práctica y complementos) se conservan; las dos últimas eran de estado
    poner_texto(celdas[-2].paragraphs[0], SEGUNDO)
    poner_texto(celdas[-1].paragraphs[0], TERCERO)
    for c, deja in zip([celdas[0], celdas[-2], celdas[-1]], (2, 1, 1)):   # sin restos del texto anterior
        for q in c.paragraphs[deja:]:
            q._p.getparent().remove(q._p)
    p10 = Paragraph(body[10], d); assert 'Versión v0.2' in p10.text
    poner_texto(p10, 'Edición comercial v1.0 · %s · formato oficio (21,59 × 33,02 cm)' % DATOS['anio'])
    # aviso de derechos y datos de edición, debajo de «República del Paraguay · 2026»
    p11 = Paragraph(body[11], d)
    aviso = [
        '© %s %s. Todos los derechos reservados. Ninguna parte de esta obra puede reproducirse, distribuirse ni comunicarse '
        'públicamente sin autorización escrita del titular, salvo las excepciones de la Ley N.º 1328/98 de Derecho de Autor '
        'y Derechos Conexos.' % (DATOS['anio'], dato('titular')),
        'Editor: %s · Domicilio: %s · ISBN: %s · Edición: %s · Tiraje: %s.' % (
            dato('editor'), dato('domicilio_editor'), dato('isbn', clave), dato('edicion'), dato('tiraje')),
        'Las capacidades y la planificación son propuestas para el programa institucional del Plan Optativo; no constituyen '
        'un programa oficial aprobado por el MEC. Los nombres de personas e instituciones de los casos son ficticios.',
    ]
    ancla = p11._p
    for texto in aviso:
        nuevo = copy.deepcopy(body[10])
        ancla.addnext(nuevo); ancla = nuevo
        q = Paragraph(nuevo, d); poner_texto(q, texto)
        q.runs[0].font.size = docx.shared.Pt(8.5)
        q.paragraph_format.space_before = docx.shared.Pt(3); q.paragraph_format.space_after = docx.shared.Pt(0)
    d.core_properties.title = DESTINO[clave].replace('_', ' ')
    d.core_properties.author = DATOS['titular'] or ''
    d.core_properties.comments = 'Edición comercial v1.0'


# ---------------------------------------------------------------- cuaderno: renglones de 8 mm, tres por consigna
RENGLON = re.compile(r'_{20,}$')


def renglones(d):
    ps = d.paragraphs
    grupos, actual = [], []
    for p in ps:
        if RENGLON.fullmatch(p.text.strip()):
            actual.append(p)
        else:
            if actual: grupos.append(actual)
            actual = []
    if actual: grupos.append(actual)
    for g in grupos:
        # un renglón más, igual al primero (keepNext), antes del último
        extra = copy.deepcopy(g[0]._p)
        g[-1]._p.addprevious(extra)
        for el in [x._p for x in g[:-1]] + [extra, g[-1]._p]:
            pPr = el.get_or_add_pPr()
            sp = pPr.find(qn('w:spacing'))
            if sp is None:
                sp = OxmlElement('w:spacing'); pPr.append(sp)
            sp.set(qn('w:line'), '454'); sp.set(qn('w:lineRule'), 'exact')      # 22,7 pt = 8 mm entre renglones
            sp.set(qn('w:before'), '0')
            sp.set(qn('w:after'), '120' if el is g[-1]._p else '0')
    return len(grupos), sum(len(g) for g in grupos)


def a_pdf(ruta):
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', os.path.dirname(ruta), ruta], check=True,
                   capture_output=True, env=dict(os.environ, HOME='/tmp/lohome_v1'))
    return ruta[:-5] + '.pdf'


def textos(pdf):
    return [norm(p.get_text()) for p in pymupdf.open(pdf)]


def entradas_indice(d):
    """Párrafos del índice del cuaderno (título, tabulación y página), antes de la primera práctica."""
    out = []
    for p in d.paragraphs:
        if p.style.name.startswith('Heading') and p.text.startswith('Práctica'):
            break
        if re.match(r'(Unidad|Práctica) .*\t\d+$', p.text):
            out.append(p)
    return out


def paginas_indice(d, pdf):
    """Página real de cada entrada: la de su título; las unidades llevan la de su primera práctica (como en la v0.2)."""
    tx = textos(pdf)
    ent = entradas_indice(d)
    pag = {}
    for k, p in enumerate(ent):
        titulo = p.text.rpartition('\t')[0]
        if titulo.startswith('Práctica'):
            pag[k] = next(i + 1 for i, t in enumerate(tx) if i > 1 and norm(titulo)[:50] in t)
    for k, p in enumerate(ent):
        if k not in pag:
            pag[k] = pag[min(j for j in pag if j > k)]
    return ent, pag


def indice_cuaderno(d, pdf):
    ent, pag = paginas_indice(d, pdf)
    cambios = 0
    for k, p in enumerate(ent):
        if p.text.rpartition('\t')[2].strip() != str(pag[k]):
            p.runs[-1].text = re.sub(r'\d+$', str(pag[k]), p.runs[-1].text); cambios += 1
    return cambios


def construir():
    res = {}
    for clave in ('libro', 'guia', 'cuaderno'):
        d = docx.Document(os.path.join(COL, ORIGEN[clave] + '.docx'))
        pagina_1(d, clave)
        extra = renglones(d) if clave == 'cuaderno' else None
        ruta = os.path.join(SAL, DESTINO[clave] + '.docx')
        d.save(ruta); pdf = a_pdf(ruta)
        if clave == 'cuaderno':
            for _ in range(3):   # el índice estático se actualiza hasta que las páginas no cambien
                d = docx.Document(ruta)
                if not indice_cuaderno(d, pdf):
                    break
                d.save(ruta); pdf = a_pdf(ruta)
        res[clave] = {'docx': ruta, 'pdf': pdf, 'renglones': extra}
    return res


def verificar(res):
    out = []
    for clave in ('libro', 'guia'):
        a = textos(os.path.join(COL, ORIGEN[clave] + '.pdf')); b = textos(res[clave]['pdf'])
        iguales = len(a) == len(b) and a[1:] == b[1:]
        out.append((clave, 'mismas %d páginas y mismo texto desde la página 2' % len(a), iguales))
    a = textos(os.path.join(COL, ORIGEN['cuaderno'] + '.pdf')); b = textos(res['cuaderno']['pdf'])
    limpio = lambda ts: re.sub(r'_{5,}|Página \d+ de \d+|\s+', ' ', ' '.join(ts[2:]))
    out.append(('cuaderno', 'mismo texto de consignas, en el mismo orden (sin renglones ni números de página)',
                norm(limpio(a)) == norm(limpio(b))))
    pdf = pymupdf.open(res['cuaderno']['pdf'])
    sueltos = []
    for i, p in enumerate(pdf):
        lineas = [l.strip() for l in p.get_text().split('\n') if l.strip()]
        if lineas and RENGLON.fullmatch(lineas[0]):
            sueltos.append(i + 1)
    out.append(('cuaderno', 'ninguna página empieza con un renglón de respuesta separado de su consigna', not sueltos))
    n_doc = sum(1 for p in docx.Document(res['cuaderno']['docx']).paragraphs if RENGLON.fullmatch(p.text.strip()))
    n_pdf = sum(1 for p in pdf for l in p.get_text().split('\n') if RENGLON.fullmatch(l.strip()))
    out.append(('cuaderno', '%d renglones en el DOCX y %d en el PDF (ninguno perdido)' % (n_doc, n_pdf), n_doc == n_pdf))
    # espacio entre renglones medido en el PDF
    ys = []
    for p in pdf:
        prev = None
        for b in p.get_text('dict')['blocks']:
            for l in b.get('lines', []):
                t = ''.join(s['text'] for s in l['spans']).strip()
                if RENGLON.fullmatch(t):
                    if prev is not None and 0 < l['bbox'][3] - prev < 40: ys.append(l['bbox'][3] - prev)
                    prev = l['bbox'][3]
                else:
                    prev = None
    mm = sorted(ys)[len(ys) // 2] * 25.4 / 72 if ys else 0
    out.append(('cuaderno', 'distancia entre renglones medida en el PDF: %.1f mm (mediana)' % mm, 7.5 <= mm <= 8.5))
    dc = docx.Document(res['cuaderno']['docx'])
    ent, pag = paginas_indice(dc, res['cuaderno']['pdf'])
    malos = [p.text for k, p in enumerate(ent) if p.text.rpartition('\t')[2].strip() != str(pag[k])]
    out.append(('cuaderno', 'índice: las %d entradas (16 prácticas y 6 unidades) en la página que indican' % len(ent),
                not malos and len(ent) == 22))
    for clave in res:
        t1 = textos(res[clave]['pdf'])[0]
        out.append((clave, 'página 1 sin textos de la versión de auditoría y con el aviso de propuesta institucional',
                    'candidata' not in t1 and 'no está aprobad' not in t1 and 'no constituyen un programa oficial' in t1
                    and 'corrieron solo en el entorno' not in t1 and 'pendiente de revisión jurídica' not in t1))
    return out


if __name__ == '__main__':
    res = construir()
    ver = verificar(res)
    for c, d_, ok in ver:
        print(('ok    ' if ok else 'FALLA ') + '[%s] %s' % (c, d_))
    info = {c: {'paginas': len(pymupdf.open(r['pdf'])), 'renglones_agregados': r['renglones']} for c, r in res.items()}
    json.dump({'documentos': info, 'controles': [{'documento': c, 'detalle': d_, 'resultado': 'correcto' if ok else 'FALLA'}
                                                 for c, d_, ok in ver]},
              open(os.path.join(SAL, 'resultado_documentos_coleccion.json'), 'w', encoding='utf8'), ensure_ascii=False, indent=1)
    print(info)
    sys.exit(0 if all(ok for *_, ok in ver) else 1)
