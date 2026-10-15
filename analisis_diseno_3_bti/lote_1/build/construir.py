# -*- coding: utf-8 -*-
"""Construye el lote 1: libro del estudiante (preliminares y Capítulos 1 a 3) y guía docente, en DOCX, con PDF
de control. El índice del libro es estático: se calcula sobre el PDF y se repite hasta que la paginación queda
estable. Uso: python3 build/construir.py   →  ../ADS_3BTI_LOTE1_ESTUDIANTE.docx/.pdf y ../ADS_3BTI_LOTE1_GUIA_DOCENTE.docx/.pdf"""
import os, sys, json
AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
from docx.shared import Pt, Cm

import helpers as H
import maqueta as M
import comun as K
import diagramas
import prelim as P
import cap1, cap2, cap3
import guia
D = K.D
MB = cap1.MB

RAIZ = K.RAIZ
RUTA_EST = os.path.join(RAIZ, 'ADS_3BTI_LOTE1_ESTUDIANTE.docx')
RUTA_GUIA = os.path.join(RAIZ, 'ADS_3BTI_LOTE1_GUIA_DOCENTE.docx')
CAPS = (cap1, cap2, cap3)

ESTADO = [
    'Lote editorial 1 para auditoría independiente. No está aprobado para la venta ni para uso en aula; la '
    'aprobación final corresponde a Fernando.',
    'Sigue el índice del informe de Work v0.1: preliminares y Unidad I (Capítulos 1 a 3, Encuentros 1 a 6).',
    'Las capacidades C1 a C10 son propuestas para el programa institucional. No son capacidades oficiales del MEC.',
]


# ---------------------------------------------------------------- preliminares
def preliminares(doc):
    doc.add_heading(P.TITULOS[0], level=1)
    for t in P.PRESENTACION:
        H.p_texto(doc, t)
    M.tabla(doc, P.UNIDADES, [5.6, 2.0, 1.2, 7.8], rotulo='Las seis unidades del libro')

    doc.add_heading(P.TITULOS[1], level=1)
    for t in P.COMO_USAR:
        H.p_texto(doc, t)

    K.titulo_en_pagina_nueva(doc, P.TITULOS[2])
    M.tabla(doc, P.CAPACIDADES, [1.4, 8.8, 6.4], rotulo='Capacidades y evidencias del año')
    H.p_texto(doc, P.CAPACIDADES_NOTA, italic=True, size=9.5)
    M.tabla(doc, cap3.TABLA_38, [1.6, 2.4, 12.6], rotulo='Hitos del proyecto')

    doc.add_heading(P.TITULOS[3], level=1)
    H.caja(doc, 'Acuerdo de trabajo del equipo', P.ACUERDO, tipo='neutra')

    K.titulo_en_pagina_nueva(doc, P.TITULOS[4])
    H.p_texto(doc, P.DIAGNOSTICO_INTRO, italic=True, size=10)
    n = 1
    for area, items in P.DIAGNOSTICO:
        n = M.actividades(doc, area, [('Consignas', items)], inicio=n)
    M.tabla(doc, P.AUTOEVALUACION, [5.2, 3.8, 3.8, 3.8], rotulo='Autoevaluación: marcá una opción por área')

    K.titulo_en_pagina_nueva(doc, P.TITULOS[5])
    H.caja(doc, 'El caso del libro', P.CASO_INTRO, tipo='neutra')
    M.tabla(doc, P.PERSONAS, [5.0, 11.6], rotulo='Personas del caso (ficticias)')
    K.registro(doc, 'D0', lead=P.CASO_D0_LEAD)

    doc.add_heading(P.TITULOS[6], level=1)
    for t in P.USO_RESPONSABLE:
        H.p_texto(doc, t)

    doc.add_heading(P.TITULOS[7], level=1)
    M.tabla(doc, P.CONSULTA, [5.6, 4.4, 6.6], rotulo='Dónde repasar lo que el proyecto necesita')
    H.p_texto(doc, P.CONSULTA_NOTA, italic=True, size=9.5)


def apertura_unidad(doc):
    K.titulo_en_pagina_nueva(doc, 'Unidad I — Comprender el problema y su contexto')
    H.p_texto(doc, 'Resultado de la unidad: un problema bien delimitado, un relevamiento con sus fuentes y un alcance '
                   'que el equipo puede cumplir. Tres capítulos, seis encuentros y 30 HC. La unidad cierra con el hito '
                   'H1: la defensa del problema y del alcance.')
    M.tabla(doc, (['Capítulo', 'Encuentros', 'Evidencias'],
                  [['1. Un sistema comienza con una necesidad', '1 y 2', 'Resolución individual breve (diagnóstico); ficha de problema; mapa de actores'],
                   ['2. Investigar antes de diseñar', '3 y 4', 'Guion y criterios de resguardo; notas, fuentes y registro de contradicciones'],
                   ['3. Elegir un alcance posible', '5 y 6', 'Factibilidad y riesgos; H1: problema y alcance defendidos, con correcciones']]),
            [5.4, 2.2, 9.0], rotulo='La Unidad I en un vistazo')


def cierre_capitulo(doc, C, figs, practica_cuerpo):
    K.titulo_en_pagina_nueva(doc, C.PRACTICA_TITULO)
    for t in C.PRACTICA_INTRO:
        H.p_texto(doc, t)
    practica_cuerpo()
    for nivel, items in C.NIVELES:
        M.actividades(doc, nivel, [('Consignas', items)])
    H.caja(doc, 'Punto de control', [C.PRACTICA_CONTROL], tipo='neutra')
    H.evaluacion(doc, C.EVAL_ID, C.EVAL_TITULO, False, C.EVAL_TEXTO_BASE, C.EVAL_ITEMS,
                 remision='Evaluación opcional. Evalúa los indicadores del capítulo. ' + C.EVAL_PUNTAJE)


# ---------------------------------------------------------------- capítulo 1
def capitulo1(doc, figs):
    C = cap1
    K.encabezado_capitulo(doc, C)
    doc.add_heading(C.E1_TITULO, level=1)
    M.seccion(doc, C.S11)
    M.tabla(doc, C.TABLA_11, [3.0, 4.4, 4.8, 4.4], rotulo='Tabla 1.1 — Síntoma, problema y solución anticipada')
    H.caja(doc, *C.CONCEPTO_PROBLEMA, tipo='concepto')
    H.caja(doc, C.EJ11[0], C.EJ11[1], tipo='neutra')
    M.tabla(doc, C.TABLA_12, [4.4, 8.2, 4.0], rotulo='Tabla 1.2 — Ficha de problema del caso (primera versión)')
    M.seccion(doc, C.S12)
    H.imagen(doc, figs['fig_1_1'], C.LEAD_FIG11, C.EPI_FIG11, ancho_cm=15.0)
    M.tabla(doc, C.TABLA_13, [2.4, 5.4, 8.8], rotulo='Tabla 1.3 — Los cinco elementos de un sistema')
    H.caja(doc, *C.CONCEPTO_SISTEMA, tipo='concepto')
    H.caja(doc, C.EJ12[0], C.EJ12[1], tipo='neutra')
    M.actividades(doc, 'Actividades del Encuentro 1', C.ACT_E1)

    doc.add_heading(C.E2_TITULO, level=1)
    M.seccion(doc, C.S13)
    H.imagen(doc, figs['fig_1_2'], C.LEAD_FIG12, C.EPI_FIG12, ancho_cm=11.5)
    H.caja(doc, C.EJ13[0], C.EJ13[1], tipo='neutra')
    M.tabla(doc, C.TABLA_14, [3.6, 1.7, 2.0, 6.0, 3.3], rotulo='Tabla 1.4 — Interesados de la biblioteca: puntajes y estrategia')
    M.tabla(doc, C.TABLA_15, [4.6, 4.6, 7.4], rotulo='Tabla 1.5 — Actores e interesados del caso')
    M.seccion(doc, C.S14)
    H.caja(doc, C.EJ14[0], C.EJ14[1], tipo='neutra')
    M.tabla(doc, C.TABLA_16, [5.2, 4.4, 7.0], rotulo='Tabla 1.6 — Resultados del diagnóstico de Noelia')
    M.tabla(doc, C.TABLA_17, [4.0, 3.6, 2.9, 2.2, 3.9], rotulo='Tabla 1.7 — Plan de apoyo de Noelia')
    M.actividades(doc, 'Actividades del Encuentro 2', C.ACT_E2)

    H.caja(doc, *C.REPASO, tipo='neutra')
    H.caja(doc, *C.CALENDARIO, tipo='neutra')
    cierre_capitulo(doc, C, figs, lambda: K.tabla_codigos(
        doc, MB.VISITA_TITULO, MB.VISITA, anchos=(1.6, 11.6, 3.4), extra_col=('Clasificación',)))


# ---------------------------------------------------------------- capítulo 2
def capitulo2(doc, figs):
    C = cap2
    K.encabezado_capitulo(doc, C)
    doc.add_heading(C.E3_TITULO, level=1)
    M.seccion(doc, C.S21)
    M.tabla(doc, C.TABLA_21, [5.4, 3.6, 3.4, 4.2], rotulo='Tabla 2.1 — Necesidad y solución en lo que dice un interesado')
    H.caja(doc, *C.CONCEPTO_RELEVAR, tipo='concepto')
    doc.add_heading(C.S44_TITULO, level=3)
    for t in C.S44:
        H.p_texto(doc, t)
    M.tabla(doc, C.TABLA_PREGUNTAS, [5.2, 4.4, 7.0], rotulo='Tabla 2.2 — Preguntas que conviene evitar')
    H.caja(doc, C.EJ21[0], C.EJ21[1], tipo='neutra')
    K.registro(doc, 'E1', lead=C.LEAD_E1)
    H.caja(doc, *C.EN_PARAGUAY, tipo='paraguay')
    H.caja(doc, *C.RESGUARDO_PREVIO, tipo='neutra')
    M.seccion(doc, C.S22)
    M.tabla(doc, C.TABLA_TECNICAS, [3.0, 4.6, 4.2, 4.8], rotulo='Tabla 2.3 — Técnicas de relevamiento')
    H.imagen(doc, figs['fig_2_1'], C.LEAD_FIG21, C.EPI_FIG21, ancho_cm=15.5)
    K.registro(doc, 'O1', lead='Este es el registro de la observación del primer recreo. Fijate que solo tiene conteos y hechos, sin nombres.')
    K.registro(doc, 'D1')
    K.registro(doc, 'D2')
    H.caja(doc, C.EJ22[0], C.EJ22[1], tipo='neutra')
    M.tabla(doc, C.TABLA_24, [1.6, 4.2, 5.6, 5.2], rotulo='Tabla 2.4 — Análisis del reglamento y del informe mensual')
    M.actividades(doc, 'Actividades del Encuentro 3', C.ACT_E3)

    doc.add_heading(C.E4_TITULO, level=1)
    H.p_texto(doc, 'Con el guion del Encuentro 3, el equipo modelo completó el relevamiento. Estos son los demás registros '
                   'del dossier. La segunda entrevista a la bibliotecaria (E4) se preparó con los documentos D1 y D2.')
    for cod in ('E2', 'E3', 'E4', 'E5'):
        K.registro(doc, cod)
    M.seccion(doc, C.S23)
    H.imagen(doc, figs['fig_2_2'], C.LEAD_FIG22, C.EPI_FIG22, ancho_cm=15.0)
    H.caja(doc, C.EJ23[0], C.EJ23[1], tipo='neutra')
    H.caja(doc, C.EJ24[0], C.EJ24[1], tipo='neutra')
    M.seccion(doc, C.S24)
    H.caja(doc, C.EJ25[0], C.EJ25[1], tipo='neutra')
    M.tabla(doc, C.TABLA_HALLAZGOS, [1.1, 2.4, 6.3, 4.0, 2.8], rotulo='Tabla 2.5 — Hallazgos del registro E1')
    doc.add_heading(C.RESGUARDO_TITULO, level=3)
    for t in C.RESGUARDO:
        H.p_texto(doc, t)
    M.tabla(doc, C.TABLA_26, [2.8, 6.6, 7.2], rotulo='Tabla 2.6 — Cinco principios para cuidar a las personas y sus datos')
    H.caja(doc, *C.CONFIRMAR, tipo='concepto')
    H.p_texto(doc, C.NECESIDADES_INTRO)
    M.tabla(doc, (['Cód.', 'Necesidad', 'Fuentes', 'Estado'],
                  [[c, t, ', '.join(f), e] for c, t, f, e in D.NECESIDADES]),
            [1.3, 8.0, 3.5, 3.8], rotulo='Tabla 2.7 — Necesidades del dossier del caso')
    M.actividades(doc, 'Actividades del Encuentro 4', C.ACT_E4)

    H.caja(doc, *C.REPASO, tipo='neutra')
    H.caja(doc, *C.CALENDARIO, tipo='neutra')
    cierre_capitulo(doc, C, figs, lambda: K.tabla_codigos(doc, MB.NORMAS_TITULO, MB.NORMAS, anchos=(1.8, 14.8)))


# ---------------------------------------------------------------- capítulo 3
def capitulo3(doc, figs):
    C = cap3
    K.encabezado_capitulo(doc, C)
    doc.add_heading(C.E5_TITULO, level=1)
    M.seccion(doc, C.S31)
    H.caja(doc, C.EJ31[0], C.EJ31[1], tipo='neutra')
    M.tabla(doc, C.TABLA_31, [6.6, 2.4, 7.6], rotulo='Tabla 3.1 — Objetivos específicos y resultados esperados')
    M.seccion(doc, C.S32)
    M.tabla(doc, C.TABLA_32, [3.4, 6.6, 6.6], rotulo='Tabla 3.2 — Tres tipos de solución para la biblioteca')
    M.tabla(doc, C.TABLA_33, [2.6, 4.6, 4.7, 4.7], rotulo='Tabla 3.3 — Factibilidad cualitativa de las alternativas')
    H.caja(doc, *C.CAJA_SIN_PUNTAJES, tipo='concepto')
    H.caja(doc, C.EJ32[0], C.EJ32[1], tipo='neutra')
    M.seccion(doc, C.S33)
    M.tabla(doc, C.TABLA_34, [1.2, 3.6, 1.4, 1.6, 4.4, 4.4], rotulo='Tabla 3.4 — Riesgos del proyecto', size=8.5)
    H.imagen(doc, figs['fig_3_1'], C.LEAD_FIG31, C.EPI_FIG31, ancho_cm=11.0)
    H.caja(doc, C.EJ33[0], C.EJ33[1], tipo='neutra')
    M.tabla(doc, (['Función del núcleo', 'Resultado que se tiene que poder comprobar'], [list(x) for x in D.ALCANCE_NUCLEO]),
            [4.2, 12.4], rotulo='Tabla 3.5 — Alcance mínimo del proyecto (núcleo)')
    H.caja(doc, 'Reglas iniciales para negociar en clase', [C.REGLAS_INTRO] + ['• ' + t for _, t in D.REGLAS], tipo='neutra')
    M.tabla(doc, C.TABLA_36, [5.0, 3.2, 8.4], rotulo='Tabla 3.6 — Exclusiones y pedidos postergados')
    M.actividades(doc, 'Actividades del Encuentro 5', C.ACT_E5)

    doc.add_heading(C.E6_TITULO, level=1)
    M.seccion(doc, C.S34)
    H.caja(doc, C.EJ34[0], C.EJ34[1], tipo='neutra')
    M.tabla(doc, C.TABLA_37, [4.6, 12.0], rotulo='Tabla 3.7 — Acuerdo de proyecto del equipo modelo (v1.0)')
    M.tabla(doc, C.TABLA_38, [1.6, 2.4, 12.6], rotulo='Tabla 3.8 — Hitos del proyecto')
    doc.add_heading(C.H1_TITULO, level=2)
    for t in C.H1_TEXTO:
        H.p_texto(doc, t)
    H.caja(doc, *C.H1_LISTA, tipo='neutra')
    H.caja(doc, *C.H1_EVALUACION, tipo='concepto')
    M.tabla(doc, (['Criterio del H1', 'Capacidad'], [[r[0], r[1]] for r in C.RUBRICA_H1]), [12.6, 4.0],
            rotulo='Criterios del H1 (la rúbrica completa está en la guía docente)')
    M.tabla(doc, (['Nivel de desempeño', 'Descriptor común'], [list(x) for x in C.NIVELES_DESEMPENO]), [3.6, 13.0],
            rotulo='Niveles de desempeño (propuesta editorial)')
    M.actividades(doc, 'Actividades del Encuentro 6', C.ACT_E6)

    def cuerpo():
        K.tabla_codigos(doc, MB.DATOS_TITULO, MB.DATOS, anchos=(1.8, 14.8))
        M.tabla(doc, (['Cód.', 'Necesidad', 'De dónde sale'], [list(x) for x in MB.NECESIDADES]), [1.6, 9.6, 5.4],
                rotulo=MB.NECESIDADES_TITULO)
    H.caja(doc, *C.REPASO, tipo='neutra')
    H.caja(doc, *C.CALENDARIO, tipo='neutra')
    cierre_capitulo(doc, C, figs, cuerpo)


# ---------------------------------------------------------------- libro completo
def entradas_indice():
    e = [(1, t) for t in P.TITULOS]
    e.append((1, 'Unidad I — Comprender el problema y su contexto'))
    for C in CAPS:
        e.append((1, '%s — %s' % (C.CAPITULO, C.TITULO)))
        encs = [getattr(C, n) for n in ('E1_TITULO', 'E2_TITULO', 'E3_TITULO', 'E4_TITULO', 'E5_TITULO', 'E6_TITULO') if hasattr(C, n)]
        secs = list(C.SECCIONES)
        mitad = len(secs) // 2
        e.append((2, encs[0])); e += [(3, s) for s in secs[:mitad]]
        e.append((2, encs[1])); e += [(3, s) for s in secs[mitad:]]
        if C is cap3:
            e.append((3, cap3.H1_TITULO))
        e.append((2, C.PRACTICA_TITULO)); e.append((2, C.EVAL_TITULO))
    return e


def libro_estudiante(figs, paginas=None):
    doc = H.new_doc()
    K.portada(doc, 'Libro del estudiante', ESTADO)
    K.indice_marcador(doc)
    preliminares(doc)
    apertura_unidad(doc)
    H.CLAVES.clear()
    capitulo1(doc, figs)
    capitulo2(doc, figs)
    capitulo3(doc, figs)
    M.pegar_renglones(doc)
    K.saltos_a_titulos(doc)
    K.rellenar_indice(doc, entradas_indice(), paginas or {})
    H.finalizar(doc, RUTA_EST)
    return list(H.CLAVES)


def construir_estudiante(figs, pdf_dir):
    titulos = [t for _, t in entradas_indice()]
    paginas = {}
    for vuelta in range(5):
        claves = libro_estudiante(figs, paginas)
        pdf = K.a_pdf(RUTA_EST, pdf_dir)
        nuevas, total = K.paginas_de_titulos(pdf, titulos, desde=2)
        faltan = [t for t in titulos if t not in nuevas]
        assert not faltan, 'títulos del índice no encontrados en el PDF: %s' % faltan
        if nuevas == paginas:
            break
        paginas = nuevas
    else:
        raise AssertionError('la paginación del índice no se estabilizó')
    return claves, paginas, total


if __name__ == '__main__':
    riesgos = [(r[0], r[2], r[3]) for r in cap3.TABLA_34[1]]
    figs = diagramas.generar(riesgos)
    salida_pdf = RAIZ
    claves, paginas, total = construir_estudiante(figs, salida_pdf)
    print('Libro del estudiante: %d páginas' % total)
    guia.construir(RUTA_GUIA, claves, paginas, ESTADO)
    pdf_g = K.a_pdf(RUTA_GUIA, salida_pdf)
    import pymupdf
    print('Guía docente: %d páginas' % len(pymupdf.open(pdf_g)))
    json.dump({'paginas_libro': paginas, 'claves': claves}, open(os.path.join(AQUI, 'resultado_construccion.json'), 'w',
              encoding='utf8'), ensure_ascii=False, indent=1)
