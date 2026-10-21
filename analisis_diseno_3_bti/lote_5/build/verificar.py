# -*- coding: utf-8 -*-
"""Controles automáticos del lote 5 (Unidad V). Uso: python3 lote_5/build/verificar.py   (sale con 1 si hay fallos)
Con la variable ADS_DR=<ruta privada de dr.json>, comprueba además que ningún código del conjunto reservado esté en el
repositorio ni en los documentos (sin imprimirlos).

Bloques de los lotes 2 a 4, extendidos a los Capítulos 14 a 16: A índice de Work · B capacidades · C trazabilidad de
fuentes · E no duplicar Laboratorio ni Algorítmica · F voseo y residuos · G maqueta (DOCX) · H maqueta (PDF) ·
I evaluaciones, prácticas e instrumentos · J numeración · L notaciones · T fragmentos de código · V pruebas y entorno.
Bloques nuevos (decisiones de Fernando del 21-oct-2026): M matriz de pruebas (doce casos, seis límites o rechazos,
requisito, esperado, obtenido, entorno y condición) · N condiciones de los resultados · D datos de prueba DP
independientes y cargador · DR conjunto reservado fuera del repositorio · P integración, permisos, persistencia,
restauración y regresión · S IN-02 y SC-02 · U sesiones de uso · W H5, recuperación y calendario de 28 encuentros ·
Q lotes 1 a 4 intactos."""
import os, re, sys, json, collections, subprocess, hashlib
sys.dont_write_bytecode = True
AQUI = os.path.dirname(os.path.abspath(__file__))
L5 = os.path.dirname(AQUI)
ADS = os.path.dirname(L5)
REPO = os.path.dirname(ADS)
sys.path.insert(0, AQUI)
from docx import Document
from docx.oxml.ns import qn
import pymupdf

import construir as B
import oficio as O
K = O.K
import cap14, cap15, cap16, guia_u5 as GU
import fragmentos_u5 as F
import diagramas_u5 as G
import diferencias as DF
V, DP, MB = cap14.V, cap14.DP, cap14.MB
C4, X = V.C4, V.X
R, D = X.R, X.D
CAPS = (cap14, cap15, cap16)
EST, GUI, CUA = B.RUTA_EST[:-5], B.RUTA_GUIA[:-5], B.RUTA_CUAD[:-5]
PROY = os.path.join(L5, 'proyecto')

ok = fallos = 0
RES = collections.OrderedDict()
DETALLE = []


def chk(bloque, desc, cond, detalle=''):
    global ok, fallos
    RES.setdefault(bloque, [0, 0])
    cond = bool(cond)
    if cond:
        ok += 1; RES[bloque][0] += 1
        print('CORRECTO [%s] %s' % (bloque, desc))
    else:
        fallos += 1; RES[bloque][1] += 1
        print('FALLA    [%s] %s %s' % (bloque, desc, detalle))
    DETALLE.append({'bloque': bloque, 'control': desc, 'resultado': 'correcto' if cond else 'FALLA'})


def texto_docx(ruta):
    d = Document(ruta)
    partes = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            vistos = set()
            for c in r.cells:
                if id(c._tc) in vistos:
                    continue
                vistos.add(id(c._tc))
                partes.append(c.text)
    return d, '\n'.join(partes)


def leer(*p):
    return open(os.path.join(*p), encoding='utf8').read()


def git(*args):
    return subprocess.run(['git'] + list(args), cwd=ADS, capture_output=True, text=True).stdout.strip()


dE, tE = texto_docx(EST + '.docx')
dG, tG = texto_docx(GUI + '.docx')
dC, tC = texto_docx(CUA + '.docx')
pE, pG, pC = pymupdf.open(EST + '.pdf'), pymupdf.open(GUI + '.pdf'), pymupdf.open(CUA + '.pdf')
res_build = json.load(open(os.path.join(L5, 'resultado_construccion.json'), encoding='utf8'))
todo = tE + '\n' + tG + '\n' + tC
j = lambda x: json.dumps(x, ensure_ascii=False)
heads = [p.text for d in (dE, dG, dC) for p in d.paragraphs if p.style.name.startswith('Heading')]
r5 = V.resumen()
norm = lambda s: re.sub(r'\s+', ' ', s)

# ---------------------------------------------------------------- A. índice de Work (§6, textual; §5.1 y §5.2)
WORK = {
    14: ('Probar con criterios y datos independientes', (27, 28),
         ['14.1. Verificación, validación y aceptación.', '14.2. Casos normales, límites, excepciones y rechazos.',
          '14.3. Resultado esperado, resultado obtenido y evidencia.', '14.4. Cobertura de requisitos y reporte de defectos.'],
         'elaborar y ejecutar una batería nueva de pruebas'),
    15: ('Probar el sistema completo con usuarios', (29, 30),
         ['15.1. Pruebas de integración y persistencia.', '15.2. Permisos, recuperación de datos y manejo de errores.',
          '15.3. Sesiones de uso y observación de dificultades.', '15.4. Aceptación por criterios acordados y prioridad de correcciones.'],
         'revisión cruzada entre equipos con datos reservados para evaluación'),
    16: ('Corregir y volver a comprobar', (31, 32),
         ['16.1. Diagnóstico de incidencias e impacto de cambios.', '16.2. Corrección, revisión entre pares y regresión.',
          '16.3. Actualización de modelos, pruebas y documentación.', '16.4. Recuperación técnica y versión candidata.'],
         None),
}
work_txt = leer('/tmp/claude-0/fuentes/work.txt') if os.path.exists('/tmp/claude-0/fuentes/work.txt') else ''
w = norm(work_txt)
for Cm in CAPS:
    t, enc, secs, prac = WORK[Cm.NUM]
    chk('A', 'Capítulo %d: título y encuentros %s de Work §6' % (Cm.NUM, enc), Cm.TITULO == t and Cm.ENCUENTROS == enc
        and ('%s — %s' % (Cm.CAPITULO, t)) in tE)
    chk('A', 'Capítulo %d: las cuatro secciones con el texto de Work §6, como títulos del libro' % Cm.NUM,
        [s.replace('. ', ' ', 1).rstrip('.') for s in secs] == Cm.SECCIONES and all(s in heads for s in Cm.SECCIONES))
    if prac:
        chk('A', 'Capítulo %d: la práctica es la de Work §6 («%s»)' % (Cm.NUM, prac), Cm.PRACTICA_WORK == prac)
chk('A', 'Capítulo 16: cierre de Work «Revisión de unidad V: evidencia de calidad y límites conocidos; H5», sin práctica',
    cap16.REVISION_WORK == 'Revisión de unidad V: evidencia de calidad y límites conocidos; H5' and cap16.PRACTICA_WORK is None
    and 'evidencia de calidad y límites conocidos' in cap16.H5_TITULO)
if work_txt:
    chk('A', 'Las 12 secciones, las dos prácticas y el cierre están textuales en el informe de Work',
        all(s in w for v in WORK.values() for s in v[2]) and all(('Práctica: %s' % v[3]) in w for v in WORK.values() if v[3])
        and 'Revisión de unidad V: evidencia de calidad y límites conocidos; H5' in w,
        [s for v in WORK.values() for s in v[2] if s not in w])
    chk('A', 'Las evidencias de Work §5.2 de los Encuentros 27 a 32 que cita la apertura están en el informe',
        all(x in w for x in ('Matriz con resultados esperados y', 'Resultados obtenidos e incidencias', 'Evidencia del circuito completo y',
                             'Registro de regresión; trazabilidad', 'H5: versión candidata con límites')))   # tabla de Work en columnas
chk('A', 'Resultado de la unidad de Work §5.1 en la apertura («evidencia de pruebas, correcciones y aceptación»)',
    'evidencia de pruebas, correcciones y aceptación' in tE)
chk('A', 'Apertura: las evidencias de Work §5.2 de los seis encuentros',
    all(x in tE for x in ('Matriz con resultados esperados y requisitos cubiertos', 'resultados obtenidos e incidencias reproducibles',
                          'Evidencia del circuito completo y recuperación', 'observaciones de dos usuarios o roles y prioridades',
                          'Registro de regresión y trazabilidad actualizada', 'H5: versión candidata con límites conocidos')))
chk('A', 'Unidad V: seis encuentros y 30 HC (27-29 V5; 30 A1 V4; 31 P2 V3; 32 P2 V1 E2)',
    'seis encuentros y 30 HC' in tE and all(('Encuentro %d' % n) in tG for n in range(27, 33)))

# ---------------------------------------------------------------- B. capacidades
CAPW = {'C3': 'Especifica y prioriza requisitos verificables.', 'C6': 'Diseña la interacción y la arquitectura del prototipo.',
        'C7': 'Organiza y registra el trabajo colaborativo.', 'C8': 'Integra un prototipo funcional con persistencia.',
        'C9': 'Verifica y mejora la solución a partir de evidencia.', 'C10': 'Documenta, entrega y defiende el prototipo.'}
ESPERADAS = {14: ('C8', 'C9'), 15: ('C3', 'C6', 'C8', 'C9', 'C10'), 16: ('C7', 'C8', 'C9')}   # Work §5.2, enc. 27 a 32
for Cm in CAPS:
    r = Cm.ROTULO_CAPACIDAD
    chk('B', 'Capítulo %d: rótulo «propuestas para el programa institucional» y «no son capacidades oficiales del MEC»' % Cm.NUM,
        'propuestas para el programa institucional' in r and 'no son capacidades oficiales del MEC' in r and r in tE and r in tG)
    chk('B', 'Capítulo %d: capacidades %s de Work §5.2, con su texto de §4.2' % (Cm.NUM, ', '.join(ESPERADAS[Cm.NUM])),
        all(('%s. %s' % (c, CAPW[c])) in Cm.CAPACIDAD and c in r for c in ESPERADAS[Cm.NUM]))
    if work_txt:
        chk('B', 'Capítulo %d: el texto de cada capacidad está en el informe de Work' % Cm.NUM,
            all(('%s. %s' % (c, CAPW[c].split()[0])) in w for c in ESPERADAS[Cm.NUM]))
chk('B', 'Encuentros 27 (C9), 28 (C8, C9), 29 (C8 a C10), 30 (C3, C6, C9), 31 (C7 a C9) y 32 (C8, C9) en los rótulos',
    'C9 en el Encuentro 27; C8 y C9 en el Encuentro 28' in cap14.ROTULO_CAPACIDAD
    and 'C8 a C10 en el Encuentro 29; C3, C6 y C9 en el Encuentro 30' in cap15.ROTULO_CAPACIDAD
    and 'C7 a C9 en el Encuentro 31; en el H5, C8 y C9' in cap16.ROTULO_CAPACIDAD)
chk('B', 'Indicadores de C9 de Work §4.2: doce casos, esperado/obtenido, seis rechazos o límites, corrige y vuelve a probar, dos sesiones',
    'doce casos' in j(cap14.INDICADORES) and 'seis de ellos de límite o de rechazo' in j(cap14.INDICADORES)
    and 'esperado, obtenido' in j(cap14.INDICADORES) and 'regresión' in j(cap16.INDICADORES) and 'sesión de uso' in j(cap15.INDICADORES)
    and 'dos sesiones' in cap15.S153[1][0])
malas_of = [tx[max(0, m.start() - 40):m.end()] for tx in (tE, tG, tC) for m in re.finditer(r'(?i)capacidad(es)? oficial', tx)
            if not re.search(r'\b[Nn]o (es|son)( una)?\b', tx[max(0, m.start() - 40):m.start()])]
chk('B', 'Ninguna capacidad presentada como oficial del MEC', not malas_of, malas_of)

# ---------------------------------------------------------------- C. trazabilidad de fuentes
mens_sc02 = {m: m for m in re.findall(r"'(ME-\d\d)' =>", leer(PROY, 'referencia_docente_v2_sc02', 'src', 'mensajes.php'))}
locales = dict(MB.FUENTES)
for Cm in CAPS:
    locales.update(Cm.EVAL_FUENTES)
locales.update({m[0]: m[4] for m in X.MENSAJES})
locales.update(C4.MENSAJES_PHP)
locales.update(mens_sc02)
locales.update({i[0]: i[2] for i in C4.INCIDENCIAS})
locales.update({e[0]: e[1] for e in C4.FICHA_ENTORNO})
locales.update({c[0]: c[4] for c in V.CASOS})
locales.update({c[0]: c[6] for c in MB.CASOS_P14})
locales.update({v[0]: v[1] for v in V.BANCO_H5})
locales.update({c: c for c in re.findall(r"\('(R-\d\d)'", leer(L5, 'caso', 'dr_plantilla.py'))})
locales.update({'SC-01': dict(R.SC01)['Identificador'], 'SC-02': 'SC-02', 'IN-01': 'IN-01'})


def fuente(c):
    return (R.TEXTO.get(c) or D.NECESIDAD.get(c) or X.REGLA.get(c) or locales.get(c) or (c if C4.existe(c) else None)
            or (c if c in ('E-6', 'P-06a', 'P-06b') else None))


PAT = (r'(?:E\d-L\d+|D\d-A?\d+|N-\d\d|M-[NF]\d\d?|MN-\d\d|MP-\d|MS-\d|MR-\d\d|OBS-\d|MD-\d\d|ML-\d|MI-\d\d|IN-\d\d'
       r'|V5-\d|C-\d\d|M-0\d|R-0\d|TA-\d|EE-\d|SO-\d|SC-\d\d|E-\d+[ab]?|OA-\d\d|PA-\d\d|MA-\d\d|MV-\d\d'
       r'|RI-\d\d|PT-\d\d|ME-\d\d|OP-\d\d|P-\d\d[ab]?|RF-\d\d|RNF-\d\d|RS-\d\d|RN-\d\d)')
for nombre, tx in (('libro', tE), ('guía', tG), ('cuaderno', tC)):
    usados = sorted(set(re.findall(r'(?<![\w-])' + PAT + r'(?![\w-])', tx)))
    falt = [c for c in usados if fuente(c) is None]
    chk('C', '%s: los %d códigos citados existen en las fuentes' % (nombre, len(usados)), not falt, falt)
nuevos = set(MB.FUENTES) - set(MB.MB4.FUENTES)
chk('C', 'Mbokaja U5: %d códigos nuevos (M-F12, MS-7, MS-8) que no repiten ninguno anterior' % len(nuevos),
    nuevos == {'M-F12', 'MS-7', 'MS-8'})
chk('C', 'ME-14 es un mensaje nuevo solo de la versión con SC-02 (no está en la v2 ni en el lote 4)',
    'ME-14' in mens_sc02 and "'ME-14'" not in leer(PROY, 'referencia_docente_v2', 'src', 'mensajes.php')
    and "'ME-14'" not in leer(ADS, 'lote_4', 'proyecto', 'referencia_docente', 'src', 'mensajes.php'))

# ---------------------------------------------------------------- M. matriz de pruebas
casos = V.CASOS
lim_rech = [c for c in casos if c[3] in ('Límite', 'Rechazo')]
chk('M', 'Matriz del equipo modelo: %d casos identificados (Work: al menos doce)' % len(casos),
    len(casos) >= 12 and len(set(c[0] for c in casos)) == len(casos))
chk('M', 'Matriz: %d casos de límite o de rechazo (Work: al menos seis)' % len(lim_rech), len(lim_rech) >= 6)
chk('M', 'Cada caso tiene requisito existente, prueba prevista, clase, entradas y resultado esperado escrito',
    all(fuente(c[1]) and c[2] and c[3] in [k for k, _ in V.CLASES] and c[7] and c[10] for c in casos))
_, f146 = cap14.tabla_146()
m1 = {c['id']: c for c in r5['matriz']['v1_lote4']['casos']}
chk('M', 'Tabla 14.6: cada fila con esperado, obtenido y condición, y el obtenido sale de la ejecución real',
    len(f146) == len(casos) and all(f[3] and f[4] and f[5] for f in f146)
    and all(('Comprobado' in f[5]) == bool(m1[f[0]]['correcto']) for f in f146 if f[0] != 'C-20'))
chk('M', 'Tabla 14.6 y su nota de entorno impresas en el libro', all(f[0] in tE and f[4] in tE for f in f146)
    and (cap14.NOTA_146 % V.entorno_texto()) in tE)
chk('M', 'Matriz del lote 4: 19 de 20; solo C-20 (IN-02) no coincide y se marca reproducido',
    r5['matriz']['v1_lote4']['correctos'] == 19 and [c for c in m1 if not m1[c]['correcto']] == ['C-20']
    and [f for f in f146 if f[0] == 'C-20'][0][5].startswith('Reproducido'))
chk('M', 'La consigna del libro pide doce casos y seis límites o rechazos al equipo, con requisito, esperado, obtenido, entorno y evidencia',
    'al menos doce casos' in j(cap14.ACT_E27) and 'Al menos seis' in j(cap14.ACT_E27)
    and 'requisito, esperado, obtenido, entorno y evidencia' in j(cap14.INDICADORES))
chk('M', 'Cobertura (Tabla 14.7) calculada de la matriz: señala RF-01, RF-02 y RNF-03 con solo casos normales',
    [f[0] for f in cap14.tabla_147()[1] if f[4].startswith('Solo casos normales')] == ['RF-01', 'RF-02', 'RNF-03'])
chk('M', 'Reporte de defecto de IN-02 con pasos, esperado, obtenido, gravedad, entorno y condición',
    [k for k, _ in V.REPORTE_IN02] == ['Código', 'Pantalla y operación', 'Requisito', 'Pasos para reproducirlo', 'Esperado', 'Obtenido',
                                       'Gravedad', 'Entorno', 'Condición'])

# ---------------------------------------------------------------- N. condiciones de los resultados
chk('N', 'Las cuatro condiciones (comprobado, reproducido, simulado, pendiente) definidas en la Tabla 14.5',
    [c for c, _ in V.CONDICIONES] == ['Comprobado', 'Reproducido', 'Simulado', 'Pendiente'] and all(c in tE for c, _ in V.CONDICIONES))
cond = {i[0]: i[3] for i in V.INCIDENCIAS_PARTIDA}
chk('N', 'Incidencias de partida con su condición real: IN-01 simulado, IN-01 bis e IN-02 reproducidos, IN-03 reproducido como mejora, campo pendiente',
    cond['IN-01'] == 'Simulado' and cond['IN-01 bis'] == 'Reproducido' and cond['IN-02'] == 'Reproducido'
    and cond['IN-03'] == 'Reproducido' and 'mejora' in j(V.INCIDENCIAS_PARTIDA) and cond['E-1 a E-12'] == 'Pendiente')
rep = leer(ADS, 'lote_5_plan', 'reproduccion_incidencias_u4.txt')
chk('N', 'La condición de IN-01, IN-01 bis, IN-02 e IN-03 coincide con la reproducción registrada del plan',
    'IN-02' in rep and 'IN-03' in rep and 'La fecha  no es válida' in rep)
chk('N', 'Sesiones del equipo modelo rotuladas como simulación; las del equipo, reales', 'simulación' in cap15.SESIONES_ROTULO
    and V.SESIONES_NOTA in tE and 'simulación rotulada' in tG)
chk('N', 'Lo que no se probó en la computadora del colegio queda pendiente (Tablas 15.1, 15.4, 16.7 y 16.8)',
    'Pendiente' in j(cap15.tabla_151()) and 'Pendiente' in j(cap15.tabla_154()) and 'pendiente en Windows' in j(cap16.tabla_167())
    and 'Pendiente' in j(cap16.TABLA_168))
chk('N', 'Datos de Mbokaja y de las evaluaciones rotulados como ficticios', 'caso de práctica, ficticio' in tE and tE.count('ficticia)') >= 3)

# ---------------------------------------------------------------- D. datos de prueba DP independientes y cargador
dp = DP.datos()
codigos_dp = {x[0] for k in ('obras', 'ediciones', 'ejemplares', 'lectores') for x in dp[k]}
chk('D', 'DP: %d códigos, todos en el rango 1000 a 1999' % len(codigos_dp),
    codigos_dp and all(1000 <= int(c.split('-')[1]) <= 1999 for c in codigos_dp))
previos = ''
for lote in ('lote_1', 'lote_2', 'lote_3', 'lote_4'):
    for d_, _, fs in os.walk(os.path.join(ADS, lote)):
        for f_ in fs:
            if f_.endswith(('.py', '.php', '.json', '.md', '.txt', '.sql', '.sh')):
                try:
                    previos += open(os.path.join(d_, f_), encoding='utf8').read()
                except UnicodeDecodeError:
                    pass
for ruta in ('lote_4/ADS_3BTI_LOTE4_ESTUDIANTE_OFICIO_O3.docx', 'lote_4/ADS_3BTI_LOTE4_GUIA_DOCENTE_OFICIO_O3.docx',
             'lote_4/ADS_3BTI_LOTE4_CUADERNO_PRACTICAS_OFICIO_O3.docx'):
    previos += texto_docx(os.path.join(ADS, ruta))[1]
cruce = sorted(c for c in codigos_dp if re.search(r'(?<![\w-])%s(?![\w-])' % c, previos))
chk('D', 'DP es independiente: ningún código ni título de DP aparece en los lotes 1 a 4 (fuentes, proyecto y documentos)',
    not cruce and not [o[1] for o in dp['obras'] if o[1] in previos], cruce)
ejemplos = j([getattr(Cm, n) for Cm in CAPS for n in dir(Cm) if n.startswith(('EJ', 'S14', 'S15', 'S16'))])
chk('D', 'Los ejemplos resueltos usan DP, y los datos de la Práctica 14 (Mbokaja) no aparecen en ellos',
    'LE-1001' in ejemplos and not re.search(r'\b(?:M-F12|MS-[78]|EJ-03\d\d|LE-032\d)\b', ejemplos))
c_dp = r5['cargador_dp']
chk('D', 'Cargador: estado esperado correcto sobre la v2 y sobre la base del estudiante, con la misma huella',
    c_dp['estado_esperado_v2'] and c_dp['estado_esperado_base_estudiante'] and c_dp['misma_huella'] and c_dp['huella'] in tG)
chk('D', 'Cargador: se niega a sobrescribir un destino existente y no crea una base en el origen',
    c_dp['se_niega_a_sobrescribir'] and c_dp['origen_sin_base_creada'])
src_c = leer(PROY, 'cargador_dp', 'cargar_dp.php')
chk('D', 'Cargador: valida los datos y rechaza el destino ANTES de copiar, y marca la copia como descartable',
    src_c.index('file_exists($destino)') < src_c.index('json_decode') < src_c.index('copiar($origen, $destino)'))
chk('D', 'cargar_dp.php marca la copia (COPIA_DESCARTABLE_DP.txt) y ejecutar_matriz.php se niega sin esa marca',
    'COPIA_DESCARTABLE_DP.txt' in src_c and 'COPIA_DESCARTABLE_DP.txt' in leer(PROY, 'matriz', 'ejecutar_matriz.php'))
chk('D', 'dp.json y mbokaja_p14.json salen de las fuentes del caso (datos_prueba_u5.py y mbokaja_u5.py)',
    json.load(open(os.path.join(PROY, 'cargador_dp', 'dp.json'), encoding='utf8')) == json.loads(j(DP.datos()))
    and json.load(open(os.path.join(PROY, 'cargador_dp', 'mbokaja_p14.json'), encoding='utf8')) == json.loads(j(MB.p14_json())))
chk('D', 'Parte manual de DP (Tabla 14.4) con las operaciones del equipo: alta de obra, dos ejemplares y un lector',
    [m[0] for m in DP.DP_MANUAL] == ['OA-02', 'OA-01', 'OA-01', 'OA-03'] and cap14.TABLA_144_ROTULO in tE)
chk('D', 'Las cuentas de la copia reciben clave al azar o la dada; ninguna contraseña real en los documentos',
    'random_bytes' in src_c and not re.search(r'(?i)contraseña:\s*\S{4,}', todo))

# ---------------------------------------------------------------- DR. conjunto reservado fuera del repositorio
en_repo = [os.path.relpath(os.path.join(d_, f_), REPO) for d_, _, fs in os.walk(REPO) if '.git' not in d_.split(os.sep)
           for f_ in fs if f_ in ('dr.json', 'comprobacion_dr.txt')]
chk('DR', 'Ningún dr.json ni comprobacion_dr.txt en el repositorio', not en_repo, en_repo)
chk('DR', 'Ningún archivo versionado con «reservado del docente» como conjunto de datos, salvo la plantilla',
    [l for l in git('grep', '--untracked', '-l', 'reservado del docente, Unidad V; ' + 'no publicar').split('\n') if l] == ['lote_5/caso/dr_plantilla.py'])
pl = leer(L5, 'caso', 'dr_plantilla.py')
chk('DR', 'La plantilla no trae datos ni semilla: el generador pide la semilla como argumento y no la guarda',
    "op.get('semilla')" in pl and 'semilla' not in j(re.findall(r"open\([^)]*\)\.write\(([^)]*)\)", pl)) and 'random.Random(semilla)' in pl)
q = subprocess.run([sys.executable, '-B', os.path.join(L5, 'caso', 'dr_plantilla.py'), '--semilla=prueba-del-verificador',
                    '--salida=' + os.path.join(L5, 'salida_dr_prohibida')], capture_output=True, text=True)
chk('DR', 'El generador se niega a escribir dentro del repositorio', q.returncode != 0 and 'FUERA del repositorio' in q.stderr
    and not os.path.exists(os.path.join(L5, 'salida_dr_prohibida')))
chk('DR', 'Libro y cuaderno: la Práctica 15 no trae datos reservados; la guía trae el procedimiento y la plantilla sin datos',
    'Plantilla de DR (estructura sin datos)' in tG and 'no están en este cuaderno' in tC
    and not re.search(r'(?<![\w-])(?:OB|ED|EJ|LE)-[2-8]\d{3}(?![\w-])', todo))
dr_ruta = os.environ.get('ADS_DR')
if dr_ruta and os.path.exists(dr_ruta):
    dr = json.load(open(dr_ruta, encoding='utf8'))
    cods = {x[0] for k in ('obras', 'ediciones', 'ejemplares', 'lectores') for x in dr[k]}
    archivos = [f_ for f_ in git('ls-files').split('\n') + git('ls-files', '--others', '--exclude-standard').split('\n')
                if f_ and os.path.isfile(os.path.join(ADS, f_))]
    versionado = ''.join(open(os.path.join(ADS, f_), encoding='utf8', errors='ignore').read() for f_ in archivos)
    hallados = [c for c in cods if c in versionado or c in todo]
    chk('DR', 'ADS_DR: ninguno de los %d códigos del conjunto reservado está en el repositorio ni en los documentos' % len(cods),
        not hallados, '%d hallados' % len(hallados))
    chk('DR', 'ADS_DR: la carpeta privada está fuera del repositorio',
        not (os.path.realpath(dr_ruta) + os.sep).startswith(os.path.realpath(REPO) + os.sep))

# ---------------------------------------------------------------- P. integración, permisos, persistencia, restauración, regresión
for v in ('v1_lote4', 'v2'):
    i = r5['integracion'][v]
    chk('P', 'Integración %s: permisos (%s), persistencia entre procesos y restauración con copia previa' % (v, i['permisos']),
        i['permisos_ok'] and i['persistencia'] and i['restauracion'] and i['copia_previa_a_restaurar'])
chk('P', 'Permisos: diez operaciones con las tres cuentas, cada rechazo ME-13 sin cambiar datos (Tabla 15.2 del resultado real)',
    len(cap15.tabla_152()[1]) == 10 and all('ME-13 y ningún dato cambia' in d_ or 'permitida' in d_ or 'pública' in d_.lower()
                                            or 'ME-12' in d_ for d_ in r5['integracion']['v1_lote4']['permisos_detalle']))
chk('P', 'Regresión de la v2: todas las pruebas del lote 4 y prueba_in02.sh, sin fallos',
    all(x['codigo_salida'] == 0 for x in r5['regresion']['v2'].values()) and 'prueba_in02.sh' in r5['regresion']['v2']
    and r5['matriz']['v2']['correctos'] == r5['matriz']['v2']['total'] == 20)
chk('P', 'Regresión de la v2 con SC-02: sin fallos; contratos con las tres comprobaciones nuevas; matriz con C-19 = ME-14',
    all(x['codigo_salida'] == 0 for x in r5['regresion']['v2_sc02'].values())
    and r5['regresion']['v2_sc02']['prueba_contratos.php']['resultado'] == 'Resultado: 93 de 93 correctas'
    and r5['matriz']['v2_sc02']['correctos'] == 20 and V.ESPERADO_SC02 == {'C-19': 'ME-14'})
chk('P', 'IN-02: la prueba nueva falla en el lote 4 (se reproduce) y pasa en la v2', r5['in02']['lote4_falla'] and r5['in02']['v2_correcta'])
chk('P', 'Tablas 16.3 y 16.5 del libro con los resultados reales de la regresión',
    all(f[1] in tE for f in cap16.tabla_163()[1] + cap16.tabla_165()[1]))
chk('P', 'Práctica 14: carga de Mbokaja y casos M-01 a M-06 comprobados sobre la v2',
    r5['practica_14']['carga'] and all(c['correcto'] for c in r5['practica_14']['casos']) and len(r5['practica_14']['casos']) == 6)
chk('P', 'Acta de aceptación (Tabla 15.4) derivada de los resultados: RNF-01 no aceptado por C-20; RS-01 y RS-02 pendientes',
    [a[3] for a in V.acta('v1_lote4') if a[0] == 'RNF-01'] == ['No aceptado (C-20)']
    and all(a[3] == 'Pendiente' for a in V.acta('v1_lote4') if a[0] in ('RS-01', 'RS-02'))
    and all(a[3] in ('Aceptado', 'Pendiente') for a in V.acta('v2')))

# ---------------------------------------------------------------- S. IN-02 (ejemplo central) y SC-02 (complementario)
chk('S', 'IN-02 es el ejemplo central del Capítulo 16: reproducción, diagnóstico, corrección, revisión entre pares y regresión',
    'IN-02' in cap16.S161[1][0] and 'Tabla 16.1' in cap16.S161[1][2] and cap16.TABLA_162_ROTULO in tE and cap16.TABLA_163_ROTULO in tE
    and 'Revisar' in leer(AQUI, 'diagramas_u5.py'))
frases = [s for s in re.split(r'(?<=[.:;])\s+', todo) if 'incumplimiento' in s.lower()]
chk('S', 'SC-02 e IN-03 nunca se presentan como incumplimiento del contrato anterior (%d menciones, todas negadas)' % len(frases),
    frases and all(re.search(r'(?i)\bno (es|era)\b[^.]*incumplimiento|\bno\b[^.]{0,30}incumplimiento', s) for s in frases),
    [s for s in frases if not re.search(r'(?i)\bno\b', s)])
sc = dict(V.SC02)
chk('S', 'SC-02 actualiza requisito, contrato, pantalla, mensaje y pruebas, con su versión, y conserva la v1.1',
    all(k in sc for k in ('Requisitos afectados', 'Contrato y pantalla', 'Mensaje nuevo', 'Pruebas afectadas', 'Nueva versión'))
    and 'se conservan' in sc['Nueva versión'] and 'ME-14' in sc['Mensaje nuevo'] and 'P-09' in sc['Pruebas afectadas'])
chk('S', 'SC-02 se rotula como mejora y solicitud de cambio, con el mismo formato de SC-01',
    'solicitud de mejora' in tE.lower() and 'mismo formato que SC-01' in tE and 'como mejora' in j(cap16.S163))
chk('S', 'Referencia del lote 4 intacta y dos copias nuevas (v2 y v2 con SC-02), sin bases de datos dentro',
    not git('status', '--porcelain', '--', 'lote_4/proyecto') and not git('diff', '45bc5e9', '--stat', '--', 'lote_4')
    and not [f_ for v_ in ('referencia_docente_v2', 'referencia_docente_v2_sc02')
             for d_, _, fs in os.walk(os.path.join(PROY, v_)) for f_ in fs if f_.endswith(('.sqlite', '.db'))])
dif = {n: [r_ for r_, _, _ in DF.diferencias(a, os.path.join(PROY, n))] for n, a, _ in DF.PARES}
chk('S', 'Diferencias de la v2: solo LEEME, la prueba nueva, _plantilla.php y lectores.php',
    dif['referencia_docente_v2'] == ['LEEME.txt', 'pruebas/prueba_in02.sh', 'public/_plantilla.php', 'public/lectores.php'], dif)
chk('S', 'Diferencias de SC-02 sobre la v2: LEEME, contratos, informe.php, mensajes.php y reglas.php',
    dif['referencia_docente_v2_sc02'] == ['LEEME.txt', 'pruebas/prueba_contratos.php', 'public/informe.php', 'src/mensajes.php',
                                          'src/reglas.php'], dif)
md = leer(PROY, 'DIFERENCIAS_v2.md')
chk('S', 'DIFERENCIAS_v2.md está al día con los archivos', all(('`%s`' % r_) in md for v_ in dif.values() for r_ in v_)
    and all(l in md for n, a, _ in DF.PARES for _, _, d_ in DF.diferencias(a, os.path.join(PROY, n)) for l in d_[2:]))
chk('S', 'La corrección guarda lo escrito sin la ficha CSRF ni la contraseña, y lo muestra escapado',
    "unset($entradas['csrf'], $entradas['clave'])" in leer(PROY, 'referencia_docente_v2', 'public', '_plantilla.php')
    and 'e((string) $e[$campo])' in leer(PROY, 'referencia_docente_v2', 'public', '_plantilla.php'))
chk('S', 'Libro: solo los fragmentos necesarios (16.1 a 16.3); la solución completa (R-4 a R-7) en la guía',
    sorted(F.FRAGMENTOS) == ['Código 16.1', 'Código 16.2', 'Código 16.3'] and all(('%s —' % r_) in tG for r_ in F.REFERENCIA))

# ---------------------------------------------------------------- U. sesiones de uso
ses = V.SESIONES
chk('U', 'Dos sesiones de uso (S1 y S2) con roles distintos y tareas nuevas, con observación y requisito',
    sorted(set(s[0] for s in ses)) == ['S1', 'S2'] and len(set(s[1] for s in ses)) == 2 and all(s[2] and s[3] and s[5] for s in ses))
chk('U', 'Sesiones con compañeros en roles; persona de la biblioteca real solo con autorización institucional, datos ficticios y privacidad',
    all('compañer' in s[1] for s in ses) and 'autorización de la institución' in j(cap15.PERSONA_REAL)
    and 'datos ficticios' in j(cap15.PERSONA_REAL) and 'No se graban ni se fotografían' in j(cap15.PERSONA_REAL))
chk('U', 'Las sesiones no se confunden con la revisión de bocetos del Capítulo 8',
    'No es la sesión del Capítulo 8' in j(cap15.S153) and 'no son la revisión de bocetos del Capítulo 8' in tG)
chk('U', 'Observación y opinión separadas; el observador no ayuda', 'No ayuda ni explica' in j(cap15.S153) and 'Las opiniones' in j(cap15.S153))

# ---------------------------------------------------------------- W. H5, recuperación y calendario de 28 encuentros
rh = cap16.RUBRICA_H5
chk('W', 'Capítulo 16 sin práctica independiente: el H5 es la actividad integradora',
    cap16.PRACTICA_TITULO is None and not cap16.NIVELES and 'Práctica 16' not in todo and 'actividad integradora' in j(cap16.H5_TEXTO))
chk('W', 'H5: seis criterios que cubren C8 y C9, cuatro niveles, evidencia del equipo e individual',
    len(rh) == 6 and all(len(x[2]) == 4 and x[3] and x[4] for x in rh) and sorted(set(x[1] for x in rh)) == ['C8', 'C9']
    and all(('%s (%s)' % (x[0], x[1])) in tG for x in rh))
chk('W', 'H5: 2 HC de evaluación; 60/40 rotulado como propuesta editorial; AE aparte; pendientes del entorno separados',
    '2 de evaluación (H5)' in cap16.TIEMPO and '60 %' in tE and '40 %' in tE and 'propuesta editorial' in j(cap16.H5_EVALUACION)
    and 'no se confunde con un desempeño de nivel inicial' in j(cap16.H5_EVALUACION) and 'pendiente del entorno' in j(cap16.H5_EVALUACION).lower())
chk('W', 'Condiciones mínimas de Work §4.3 (seis) revisadas una por una en la Tabla 16.7',
    [f[0] for f in cap16.tabla_167()[1]] == ['Circuito de préstamo y devolución ejecutable', 'Persistencia comprobada',
                                              'Ninguna pérdida de historial', 'Rechazo de préstamo simultáneo del mismo ejemplar',
                                              'Evidencia de pruebas', 'Instalación y restauración reproducibles'])
chk('W', 'Si falta una condición mínima: no se declara aprobado, se registra la recuperación y se ofrece el procedimiento',
    re.search(r'no se (lo )?declara aprobado', j(cap16.RECUPERACION)) and 'Planilla de recuperación' in tG
    and 'Procedimiento de recuperación técnica' in tG and re.search(r'no se (lo )?declara aprobado', j(cap16.CALENDARIO))
    and 'se ofrece el procedimiento' in j(cap16.CALENDARIO))
extra = [s for s in re.split(r'(?<=[.])\s+', todo) if 'horas extraordinarias' in s]
chk('W', 'No se suponen horas extraordinarias que la institución no disponga (%d menciones, todas negadas)' % len(extra),
    extra and all(re.search(r'No se suponen', s) for s in extra) and 'fuera del horario' not in todo)
chk('W', 'Calendario de 28 encuentros: 20 a 23 sin cambios, 24 = 31+32 «Sin encuentro de reserva para recuperación técnica.»',
    [f[1] for f in GU.CALENDARIO_28[1]] == ['20', '21', '22', '23', '24'] and 'Sin encuentro de reserva para recuperación técnica.' in tE
    and 'Sin encuentro de reserva para recuperación técnica.' in tG and tE.count('Si el calendario se acorta') == 3)
chk('W', 'Adaptación de 28: prioriza condiciones mínimas, pruebas esenciales, regresión y defensa; reduce complementarias; documenta el alcance',
    all(x in j(GU.CALENDARIO_TEXTO) for x in ('condiciones mínimas', 'regresión', 'defensa del H5', 'pasan a lectura', 'bitácora del curso')))
textos_previos = list(R.TEXTO.values()) + list(MB.FUENTES.values()) + [v[1] for v in C4.BANCO_H4]
chk('W', 'Banco V5: cinco reportes que no están en el dossier, en Mbokaja, en el banco del H4 ni en las prácticas',
    len(V.BANCO_H5) >= 5 and not any(v[1][:30] in t for v in V.BANCO_H5 for t in textos_previos)
    and not any(v[1][10:40] in tE + tC for v in V.BANCO_H5))

# ---------------------------------------------------------------- T. fragmentos de código
for fid in F.TODOS:
    ls = F.extraer(fid)
    archivo = os.path.join(F.PROY, F.TODOS[fid][0])
    chk('T', '%s: leído tal cual de %s (%d líneas), ninguna de más de %d caracteres' % (fid, F.TODOS[fid][0], len(ls), F.ANCHO),
        '\n'.join(ls) in leer(archivo).replace('\r', '') and max(len(l) for l in ls) <= F.ANCHO and len(ls) <= 40)
    dest = tE if fid in F.FRAGMENTOS else tG
    chk('T', '%s: con su ficha (archivo, propósito, requisito o contrato, entradas, resultado y prueba) en %s'
        % (fid, 'el libro' if fid in F.FRAGMENTOS else 'la guía'), all(v in dest for v in F.ficha(fid).values())
        and all(l.strip() in dest for l in ls if l.strip()))
for fid in F.REFERENCIA:
    del_libro = {l.strip() for f_ in F.FRAGMENTOS for l in F.extraer(f_)}
    propias = [l.strip() for l in F.extraer(fid) if len(l.strip()) > 25 and l.strip() not in del_libro]
    chk('T', '%s no se publica en el libro ni en el cuaderno (solución completa del docente)' % fid,
        not [l for l in propias if l in tE or l in tC])
_citas = re.sub(r'^Código \d+\.\d+ —', '', tE, flags=re.M)
chk('T', 'Cada fragmento del libro se cita en el texto y tiene explicación',
    all(re.search(r'Códigos? (?:\d+\.\d+(?:,| y) )*%s\b' % re.escape(f_.split()[1]), _citas) for f_ in F.FRAGMENTOS)
    and all(f_ in cap16.EXPLICA for f_ in F.FRAGMENTOS))
orden = [m.group(1) for m in re.finditer(r'^Código (\d+\.\d+) —', tE, re.M)]
chk('T', 'Códigos 16.1 a 16.3 correlativos y en orden de aparición', orden == ['16.1', '16.2', '16.3'], orden)
txtpdf = '\n'.join(pg.get_text() for pg in pE) + '\n'.join(pg.get_text() for pg in pG)
faltan_cod = [l for f_ in F.TODOS for l in F.extraer(f_) if l.strip() and re.sub(r'\s+', ' ', l.strip()) not in re.sub(r'[ \t]+', ' ', txtpdf)]
chk('T', 'Cada línea de código del libro y de la guía está impresa entera en el PDF', not faltan_cod, faltan_cod[:3])

# ---------------------------------------------------------------- V. pruebas y entorno
chk('V', 'Última ejecución de build/reproducir.py sin fallos y con el entorno registrado',
    r5 and r5['fallos'] == 0 and r5['entorno']['php'] and 'NO es una computadora escolar' in r5['entorno']['tipo'])
chk('V', 'Cada resultado de prueba guardado tiene su línea de entorno',
    all('Entorno' in leer(L5, 'resultados_pruebas', f_) or f_ == 'resumen.json' for f_ in os.listdir(os.path.join(L5, 'resultados_pruebas'))))
chk('V', 'Cada capítulo y la apertura advierten que el entorno no se probó en la institución', tE.count(C4.AVISO_ENTORNO[0]) >= 4)
chk('V', 'Las tablas de resultados del libro y de la guía llevan el entorno donde corrieron',
    tE.count(V.entorno_texto()) >= 3 and V.entorno_texto() in tG)
malas = re.findall(r'(?i)(?:entorno|base|sistema|prototipo|computadora|material|libro)[^.\n]{0,40}\bvalidad[oa]s?\b(?! en aulas)'
                   r'|probad[oa]s? en (?:los )?colegios|\bapto para todos', todo)
chk('V', 'Nada se declara validado ni probado en colegios', not malas, malas[:5])
chk('V', 'La guía dice que el material no está validado en aulas', 'no está validado en aulas' in tG)

# ---------------------------------------------------------------- E. no duplicar Laboratorio ni Algorítmica
chk('E', 'Sin unidad de herramientas de pruebas ni manual de SQL (ningún título lo es)',
    not [h for h in heads if re.search(r'(?i)sintaxis|manual de sql|phpunit|junit|selenium|introducci[oó]n a php', h)])
chk('E', 'Sin Access ni tablas «Socios»', 'Access' not in todo and 'Socios' not in todo)
chk('E', 'Sin personas, títulos ni archivos del caso de Algorítmica 3.º', not [p for p in D.PROHIBIDOS_ALGORITMICA if p in todo])
chk('E', 'Articulación con Laboratorio: acá no se enseña una herramienta de pruebas', 'no se enseña una herramienta de pruebas' in tE)

# ---------------------------------------------------------------- F. voseo y residuos
TUTEO = (r'\b(tienes|puedes|debes|sabes|quieres|eres|propón|tú)\b|(?:^\d+\.\s+|^\([a-h]\)\s+)'
         r'(Escribe|Lee|Completa|Indica|Explica|Clasifica|Responde|Anota|Marca|Elige|Justifica|Redacta|Compara|Describe|Identifica|Resume|Haz|Dibuja|Construye|Ejecuta|Registra|Señala|Decide|Diseña|Aplica|Revisa|Agrega|Calcula|Demuestra|Prepara|Arma|Ordena)\b')
sin_codigo = lambda tx: '\n'.join(l for l in tx.split('\n') if l.rstrip() not in {x.rstrip() for f_ in F.TODOS for x in F.extraer(f_)})
for nombre, tx in (('libro', tE), ('guía', tG), ('cuaderno', tC)):
    m = [x.group(0) for x in re.finditer(TUTEO, tx, re.M)]
    chk('F', '%s: sin formas de tuteo frecuentes' % nombre, not m, m[:5])
    chk('F', '%s: sin identificadores de modelos de IA' % nombre, not re.search(r'(?i)\b(claude|gpt|gemini|opus|sonnet|codex|chatgpt)\b', tx))
    chk('F', '%s: sin símbolos ✓ ✗' % nombre, not re.search('[✓✔✗✘✅]', tx))
    chk('F', '%s: sin marcadores sin reemplazar (fuera del código impreso)' % nombre, '§§' not in tx and 'TODO' not in tx
        and 'XXX' not in tx and '%s' not in sin_codigo(tx) and '%d' not in sin_codigo(tx))
    chk('F', '%s: no se declara apto ni aprobado para la venta' % nombre,
        not re.search(r'(?i)\bapto\b|aprobado para (la )?(venta|comercializ)', tx.replace('No está aprobado para la venta', '')))

# ---------------------------------------------------------------- G. maqueta (DOCX)
tw = int(O.Cm(O.ANCHO).twips)
for nombre, d in (('libro', dE), ('guía', dG), ('cuaderno', dC)):
    s = d.sections[0]
    chk('G', '%s: oficio 21,59 × 33,02 cm, márgenes O3 y espejados' % nombre,
        abs(s.page_width.cm - 21.59) < 0.01 and abs(s.page_height.cm - 33.02) < 0.01 and abs(s.left_margin.cm - 2.5) < 0.01
        and abs(s.right_margin.cm - 1.8) < 0.01 and abs(s.top_margin.cm - 1.5) < 0.01 and d.settings.element.find(qn('w:mirrorMargins')) is not None)
    filas = [r_ for t in d.tables for r_ in t.rows]
    chk('G', '%s: %d filas de tabla, todas «no dividir»' % (nombre, len(filas)),
        all(r_._tr.trPr is not None and r_._tr.trPr.find(qn('w:cantSplit')) is not None for r_ in filas))
    tablas = [t for t in d.tables if len(t.columns) > 1 or t.rows[0].cells[0].paragraphs[0].runs and
              t.rows[0].cells[0].paragraphs[0].runs[0].font.name == 'Courier New']
    malas_t = [i for i, t in enumerate(tablas) if abs(sum(int(g.get(qn('w:w'))) for g in t._tbl.tblGrid) - tw) > 15]
    chk('G', '%s: %d tablas de varias columnas y bloques de código al ancho útil' % (nombre, len(tablas)), not malas_t, malas_t)
bloques = [t for t in dE.tables if len(t.columns) == 1 and t.rows[0].cells[0].paragraphs[0].runs
           and t.rows[0].cells[0].paragraphs[0].runs[0].font.name == 'Courier New']
chk('G', 'Libro: %d bloques de código, uno por fragmento, en letra monoespaciada de 7,5 pt' % len(bloques),
    len(bloques) == len(F.FRAGMENTOS) and all(r_.font.size.pt == 7.5 for t in bloques for p in t.rows[0].cells[0].paragraphs for r_ in p.runs))
chk('G', 'Se usan la maqueta del lote 1, el perfil O3 del lote 2 y las piezas del lote 4 (sin copias)',
    os.path.dirname(O.H.__file__).endswith(os.path.join('lote_1', 'build')) and os.path.dirname(O.__file__).endswith(os.path.join('lote_2', 'build'))
    and os.path.dirname(B.Z.Z4.__file__).endswith(os.path.join('lote_4', 'build')))
chk('G', 'Portada: versión del lote 5 en los tres documentos', all('v1.0 (lote 5, para auditoría)' in t for t in (tE, tG, tC)))


# ---------------------------------------------------------------- H. maqueta (PDF)
def lineas(pg):
    return [l.strip() for l in pg.get_text().split('\n') if l.strip() and not re.match(r'^Página \d+ de \d+$', l.strip())]


for nombre, pdf in (('libro', pE), ('guía', pG), ('cuaderno', pC)):
    vacias = [i + 1 for i, pg in enumerate(pdf) if not lineas(pg)]
    chk('H', '%s: ninguna página en blanco' % nombre, not vacias, vacias)
    sueltos = [i + 1 for i, pg in enumerate(pdf) if lineas(pg) and lineas(pg)[0].startswith('_' * 20)]
    chk('H', '%s: ninguna página empieza con renglones de respuesta sueltos' % nombre, not sueltos, sueltos)
for Cm in CAPS:
    pags = [i + 1 for i, pg in enumerate(pE) if Cm.EVAL_TITULO in lineas(pg)[:2]]
    sig = pags and pags[0] < len(pE) and lineas(pE[pags[0]])
    entra = pags and (pags[0] == len(pE) or (sig and (sig[0].startswith('Capítulo') or sig[0].startswith('Unidad'))))
    chk('H', '%s: en página propia y entera en esa página (pág. %s)' % (Cm.EVAL_TITULO, pags), len(pags) == 1 and entra)
    if Cm.PRACTICA_TITULO:
        pp = [i + 1 for i, pg in enumerate(pE) if lineas(pg)[:1] and lineas(pg)[0].startswith(Cm.PRACTICA_TITULO[:25])]
        chk('H', '%s: empieza en página nueva' % Cm.PRACTICA_TITULO.split(' —')[0], len(pp) == 1)
        pc = [i + 1 for i, pg in enumerate(pC) if lineas(pg)[:1] and lineas(pg)[0].startswith(Cm.PRACTICA_TITULO[:25])]
        chk('H', 'Cuaderno: %s empieza en página nueva' % Cm.PRACTICA_TITULO.split(' —')[0], len(pc) == 1)
p32 = [i + 1 for i, pg in enumerate(pE) if lineas(pg)[:1] == [cap16.E32_TITULO]]
chk('H', 'El Encuentro 32 (recuperación y H5) empieza en página nueva (pág. %s)' % p32, len(p32) == 1)
paginas = res_build['paginas_libro']
chk('H', 'El índice ocupa una sola página y la siguiente es la apertura de la Unidad V', lineas(pE[2])[0] == B.UNIDAD_V)
recalc, _ = K.paginas_de_titulos(EST + '.pdf', list(paginas), desde=2)
chk('H', 'Índice estable: las %d entradas apuntan a la página donde está su título' % len(paginas),
    recalc == paginas and len(paginas) == len(B.entradas_indice()))
chk('H', 'Extensión del libro: %d págs. (estimación del plan: 32 a 40; no se rellena para alcanzarla)' % len(pE), 26 <= len(pE) <= 40)
chk('H', 'Extensión de la guía: %d págs. (estimación del plan: 16 a 22)' % len(pG), 16 <= len(pG) <= 22)
chk('H', 'Extensión del cuaderno: %d págs. (dos prácticas; el Capítulo 16 no tiene)' % len(pC), 3 <= len(pC) <= 10)
chk('H', 'Control de filas: ninguna tabla pierde filas en el PDF', not res_build['tablas_en_hoja_nueva']
    or all(O.incompletas(dx + '.docx', dx + '.pdf') == [] for dx in (EST, GUI, CUA)))

# ---------------------------------------------------------------- I. evaluaciones, prácticas e instrumentos
claves = res_build['claves']
semillas_prev = ['ADS3-L1-C1', 'ADS3-L1-C2-i', 'ADS3-L1-C3', 'ADS3-C4v2-2', 'ADS3-L2-C5-c', 'ADS3-L2-C6', 'ADS3-L3-C7', 'ADS3-L3-C8',
                 'ADS3-L3-C9', 'ADS3-L4-C10', 'ADS3-L4-C11', 'ADS3-L4-C12', 'ADS3-L4-C13']
for Cm in CAPS:
    ks = [c for c in claves if c[0] == Cm.EVAL_ID]
    mc = [c[2] for c in ks if c[2] in 'abcd' and len(c[2]) == 1]
    chk('I', '%s: %d ítems con clave; opción múltiple con claves %s' % (Cm.EVAL_TITULO, len(ks), ''.join(mc)),
        len(ks) == len(Cm.EVAL_ITEMS) and len(set(mc)) >= 3)
    pts = sum(1 for it in Cm.EVAL_ITEMS if it[0] in ('mc', 'vf')) + sum(it[3] for it in Cm.EVAL_ITEMS if it[0] == 'ab')
    chk('I', '%s: suma 15 puntos, como dice su remisión' % Cm.EVAL_TITULO, pts == 15 and 'Puntaje total: 15 puntos' in Cm.EVAL_PUNTAJE)
    chk('I', '%s: la guía trae su clave y los criterios de los ítems abiertos' % Cm.EVAL_TITULO, ('Clave de la %s' % Cm.EVAL_TITULO.lower()) in tG)
for Cm in (cap14, cap15):
    chk('I', 'Capítulo %d: práctica en tres niveles, con claves y lista de cotejo en la guía' % Cm.NUM,
        [n.split(' —')[0] for n, _ in Cm.NIVELES] == ['Nivel 1', 'Nivel 2', 'Nivel 3'] and ('Lista de cotejo de la Práctica %d' % Cm.NUM) in tG)
chk('I', 'Semillas de evaluación distintas entre sí y de las de los lotes 1 a 4', len(set([Cm.EVAL_ID for Cm in CAPS] + semillas_prev)) == 16)
chk('I', 'Práctica 14: datos nuevos de Mbokaja en una copia cargada con el cargador; nivel 3 en papel',
    'mbokaja_p14.json' in j(cap14.PRACTICA_INTRO) and 'en papel' in j(cap14.PRACTICA_INTRO) and 'M-F12' in j(cap14.NIVELES))
chk('I', 'Guía: secuencia en HC de los Encuentros 27 a 32 con las categorías de la v2.1 (27-29 V5; 30 A1 V4; 31 P2 V3; 32 P2 V1 E2)',
    tG.count('(5 HC de verificación)') == 3 and '(5 HC: 1 de análisis y 4 de verificación)' in tG
    and '(5 HC: 2 de práctica y 3 de verificación)' in tG and '(5 HC: 2 de práctica, 1 de verificación y 2 de evaluación)' in tG)
chk('I', 'Guía: claves de todas las actividades de los Encuentros 27 a 32',
    all(len(k) == sum(len(items) for _, items in acts) for k, acts in (
        (GU.CLAVES_E27, cap14.ACT_E27), (GU.CLAVES_E28, cap14.ACT_E28), (GU.CLAVES_E29, cap15.ACT_E29), (GU.CLAVES_E30, cap15.ACT_E30),
        (GU.CLAVES_E31, cap16.ACT_E31), (GU.CLAVES_E32, cap16.ACT_E32))))
chk('I', 'Cuaderno: las dos prácticas con el mismo texto que el libro y sin evaluaciones',
    all(t in tC for Cm in (cap14, cap15) for t in [Cm.PRACTICA_TITULO, Cm.PRACTICA_CONTROL] + Cm.PRACTICA_INTRO
        + [i.split('\n')[0] for _, items in Cm.NIVELES for i in items]) and not any(Cm.EVAL_TITULO in tC for Cm in CAPS))

# ---------------------------------------------------------------- J. numeración
for Cm in CAPS:
    n = Cm.NUM
    for tipo in ('Tabla', 'Figura', 'Ejemplo resuelto'):
        sec = [int(m.group(1)) for m in re.finditer(r'^%s %d\.(\d+)(?: —|:)' % (tipo, n), tE, re.M)]
        chk('J', 'Capítulo %d: %s %d.1 a %d.%d correlativas, en orden de aparición' % (n, tipo.lower(), n, n, len(sec)),
            sec == list(range(1, len(sec) + 1)) and (len(sec) >= 1 or tipo == 'Figura'), sec)
    for tipo, pat in (('tabla', r'^Tabla %d\.(\d+) —'), ('figura', r'^Figura %d\.(\d+) —'), ('ejemplo resuelto', r'^Ejemplo resuelto %d\.(\d+):'),
                      ('código', r'^Código %d\.(\d+) —')):
        refs = set(int(x) for x in re.findall(r'%s %d\.(\d+)' % (tipo.capitalize(), n), tE + tG + tC))
        defs = set(int(x) for x in re.findall(pat % n, tE, re.M))
        chk('J', 'Capítulo %d: toda %s mencionada existe' % (n, tipo), refs <= defs, refs - defs)
lotes = {}
for ruta in ('lote_1_oficio_candidata/ADS_3BTI_LOTE1_ESTUDIANTE_OFICIO_O3_CANDIDATA.docx', 'lote_2/ADS_3BTI_LOTE2_ESTUDIANTE_OFICIO_O3.docx',
             'lote_3/ADS_3BTI_LOTE3_ESTUDIANTE_OFICIO_O3.docx', 'lote_4/ADS_3BTI_LOTE4_ESTUDIANTE_OFICIO_O3.docx'):
    lotes[ruta] = texto_docx(os.path.join(ADS, ruta))[1]
otras = sorted(set(re.findall(r'\b(?:Tabla|Figura) (?:[1-9]|1[0-3])\.\d+', tE + tG + tC)))
chk('J', 'Las tablas y figuras de los lotes 1 a 4 que se citan existen allí (%s)' % ', '.join(otras),
    all(re.search(r'^%s —' % re.escape(o), '\n'.join(lotes.values()), re.M) for o in otras),
    [o for o in otras if not re.search(r'^%s —' % re.escape(o), '\n'.join(lotes.values()), re.M)])

# ---------------------------------------------------------------- L. notaciones
reg = json.load(open(os.path.join(AQUI, 'figs', 'simbolos.json'), encoding='utf8'))
for f_ in G.FIGURAS:
    chk('L', '%s (esquema didáctico) no usa símbolos de ninguna notación' % f_, reg[f_]['notacion'] == 'esquema' and not reg[f_]['simbolos'])

# ---------------------------------------------------------------- Q. lotes 1 a 4 intactos
st = git('status', '--porcelain', '--', 'lote_1', 'lote_2', 'lote_3', 'lote_4', 'lote_1_oficio_candidata', 'lote_5_plan')
chk('Q', 'git: ningún archivo de los lotes 1 a 4 ni del plan del lote 5 cambió ni se agregó', not st, st[:300])
chk('Q', 'git: el lote 4 es el de la v1.1 (45bc5e9), sin cambios', not git('diff', '--stat', '45bc5e9', '--', 'lote_4'))
for lote, n in (('lote_3', 3), ('lote_4', 4)):
    m_ = json.load(open(os.path.join(ADS, lote, 'MANIFIESTO_LOTE_%d.json' % n), encoding='utf8'))
    docs = [a for a in m_['archivos'] if a['ruta'].startswith('ADS_3BTI_LOTE%d_' % n)]
    chk('Q', 'Lote %d: sus seis DOCX y PDF coinciden con su manifiesto SHA-256' % n,
        len(docs) == 6 and all(hashlib.sha256(open(os.path.join(ADS, lote, a['ruta']), 'rb').read()).hexdigest() == a['sha256'] for a in docs))
chk('Q', 'Sin capítulos 17 ni 18 ni Unidad VI construidos en este lote',
    not re.search(r'^Capítulo 1[78] —', tE, re.M) and 'Unidad VI —' not in tE and not os.path.exists(os.path.join(L5, 'build', 'cap17.py')))

print('\nResultado por bloque:')
for b, (o_, f_) in RES.items():
    print('  %s: %d correctos, %d fallos' % (b, o_, f_))
print('Resultado: %d correctos, %d fallos' % (ok, fallos))
json.dump({'correctos': ok, 'fallos': fallos, 'bloques': RES, 'controles': DETALLE},
          open(os.path.join(L5, 'verificacion_resultado.json'), 'w', encoding='utf8'), ensure_ascii=False, indent=1)
sys.exit(1 if fallos else 0)
