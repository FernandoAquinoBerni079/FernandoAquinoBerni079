# -*- coding: utf-8 -*-
"""Guía docente con solucionario del lote 2 (Unidad II, Capítulos 4 a 6): secuencia de cada encuentro en HC, claves de
las actividades, prácticas y evaluaciones, listas de cotejo, leyenda de notaciones con errores frecuentes, rúbrica del
H2 con planilla y banco de necesidades nuevas, calendario de 28 encuentros y procedencia de los recursos.

Capítulo 4: las claves, la secuencia, la lista de cotejo y la correspondencia de la guía v2.1 se leen del piloto
(`fase2_pilotos/A2_capitulo4_v2_1/build/guia_v2_1.py`) y se les aplican solo los cambios de la especificación C3 (§3)
y los de integración, declarados en CAMBIOS_GUIA."""
import os, sys, importlib.util
from docx.shared import Cm
import oficio as O
H, M, K = O.H, O.M, O.K
import piezas as Z
import cap4, cap5, cap6
R, MO, MB, G = Z.R, Z.MO, Z.MB, Z.G
D = R.D

AQUI = os.path.dirname(os.path.abspath(__file__))
PIL_DIR = os.path.join(os.path.dirname(os.path.dirname(AQUI)), 'fase2_pilotos', 'A2_capitulo4_v2_1', 'build')
sys.path.append(PIL_DIR)
_spec = importlib.util.spec_from_file_location('guia_v2_1_piloto', os.path.join(PIL_DIR, 'guia_v2_1.py'))
GP = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(GP)
sys.path.remove(PIL_DIR)

ANCHO_SEC = [3.2, 2.2, 11.2]
CAMBIOS_GUIA = []


def _cambio(qué, motivo):
    CAMBIOS_GUIA.append((qué, motivo))


# ---------------------------------------------------------------- capítulo 4 (piloto v2.1 + especificación C3)
SEC_C4 = [GP.SECUENCIA[0], (GP.SECUENCIA[1][0], [
    GP.SECUENCIA[1][1][0],
    ['Práctica', 'aprox. 3 HC', 'Sección 4.4 con la Tabla 4.6 y el Ejemplo resuelto 4.4. Actividades 1 (a) a (c) y 2 del Encuentro 8. Práctica 4: consigna común para todos (aprox. 0,25 HC) y luego los niveles (aprox. 1,25 HC).'],
    ['Revisión y cierre', 'aprox. 1 HC', 'Revisión entre pares con la lista del libro (Actividad 3). Evidencia del encuentro (Work, encuentro 8): requisitos comprobables y revisión entre pares. La Actividad 1 (d) y las Actividades 4 y 5 quedan como tarea.'],
])]
_cambio('Secuencia del Encuentro 8', 'Especificación C3: consigna común de 0,25 HC y Actividad 1 (d) como tarea.')

CLAVES_P4 = [['Consigna común', 'RNF-04. La consulta de disponibilidad de un título debe mostrar el resultado en menos de 3 '
              'segundos en la computadora de la biblioteca (cifra a confirmar con quien lo pidió). Fuente: B-03. Prioridad: '
              'importante; la consulta funciona igual sin esa meta de tiempo, pero es lo que pidió el lector. Criterio: '
              + R.CORRESPONDENCIA_C3[3] + ' Se acepta otra cifra o situación si es comprobable. No se acepta «rápido», '
              '«enseguida» ni ninguna otra palabra vaga en el criterio.']]
CLAVES_P4 += [GP.CLAVES_PRACTICA[0],
              ['Nivel 2', GP.CLAVES_PRACTICA[1][1]
               .replace('RNF-04 tiempo de respuesta de la consulta (B-03, a confirmar); RNF-05 solo los usuarios autorizados ven datos personales (B-08 y E5-L6)',
                        'RNF-05 solo los usuarios autorizados ven datos personales (B-08 y E5-L6); RNF-06, por ejemplo, que la docente vea solo los préstamos de su curso (B-01, seguridad, a confirmar)')
               .replace('Ambigüedades: B-03, preguntar «¿cuánto esperarías como máximo?»; B-08, preguntar «¿quién necesita ver qué datos?».',
                        'Ambigüedad de B-08: preguntar «¿quién necesita ver qué datos?». La de B-03 ya la resolvió la consigna común.')],
              GP.CLAVES_PRACTICA[2]]
assert 'RNF-04' not in CLAVES_P4[2][1] and 'RNF-06' in CLAVES_P4[2][1] and 'B-03 ya la resolvió' in CLAVES_P4[2][1], 'reemplazo de la clave del Nivel 2'
_cambio('Claves de la Práctica 4', 'Especificación C3, §3: clave de la consigna común; el Nivel 2 numera desde RNF-05 y corrige solo B-08.')
COTEJO_P4 = GP.LISTA_COTEJO + [['Redacta al menos un criterio de aceptación comprobable para un requisito de calidad (consigna común).', 'Indicador 3']]
_cambio('Lista de cotejo de la Práctica 4', 'Especificación C3, §3: fila del criterio de calidad, vinculada al indicador 3.')
CORRESPONDENCIA = [[i, R.ETIQUETA[o], d, c] for i, o, d, c in R.CORRESPONDENCIA] + \
                  [[R.CORRESPONDENCIA_C3[0], R.ETIQUETA[R.CORRESPONDENCIA_C3[1]], R.CORRESPONDENCIA_C3[2], R.CORRESPONDENCIA_C3[3]]]
NOTA_C3 = ('Evidencia mínima de C3: cada estudiante escribe un criterio de aceptación funcional (Encuentro 8, Actividad 2) '
           'y uno de calidad (Práctica 4, consigna común).')
_cambio('Correspondencia de requisitos y criterios', 'Especificación C3, §3: fila RNF-04 (Trabajo del estudiante) y nota de evidencia mínima.')
COBERTURA_C3 = [GP.C3[0], GP.C3[1], GP.C3[2],
                ['Registra una modificación', 'Sí en el lote: Capítulo 6, sección 6.4 (solicitud de cambio SC-01 y catálogo v1.1) y H2. El Desafío del Encuentro 7 la anticipa.']]
_cambio('Cobertura de los indicadores de C3', 'La modificación registrada ya está en el lote (Capítulo 6).')
CRIT_C4 = [['8', '3: identificador y enunciado (1), fecha de referencia explícita (1), criterio con resultado esperado observable (1).', 'Cada parte por separado. 0 en el criterio si no dice qué se observa.'],
           ['9', '5: RN-07 con fuente R-07 (1), requisito sin tecnología (1), criterio comprobable (1), prioridad justificada (1), identificadores y fuentes correctos (1).', 'Cada parte por separado. Se valora la justificación de la prioridad, no la etiqueta.']]

# ---------------------------------------------------------------- capítulo 5
SEC_C5 = [
    ('Encuentro 9 — Proceso actual y proceso propuesto (5 HC: 3 de análisis y 2 de diseño)', [
        ['Problema y explicación', 'aprox. 1 HC', 'Sección 5.1 con la Tabla 5.1 (símbolos). Figura 5.1 y Tabla 5.2: el docente reconstruye en la pizarra el proceso actual con el Ejemplo resuelto 5.1.'],
        ['Práctica', 'aprox. 3 HC', 'Figura 5.2, Tabla 5.3 y Ejemplo resuelto 5.2 (proceso propuesto, la parte de diseño). Sección 5.2 con la Tabla 5.4, la Figura 5.3, la Tabla 5.5 y el Ejemplo resuelto 5.3. Actividades 1 a 3, en parejas.'],
        ['Revisión y cierre', 'aprox. 1 HC', 'Puesta en común de la Actividad 3 (proceso propuesto de devolución). Evidencia del encuentro (Work, encuentro 9): flujos comentados y decisiones justificadas. Las Actividades 4 y 5 quedan como tarea.'],
    ]),
    ('Encuentro 10 — Casos de uso y excepciones (5 HC de análisis)', [
        ['Problema y explicación', 'aprox. 1 HC', 'Sección 5.3 con la Tabla 5.6, la Figura 5.4 y el Ejemplo resuelto 5.4.'],
        ['Práctica', 'aprox. 3 HC', 'Sección 5.4, Ejemplo resuelto 5.5 y plantillas de CU-01 a CU-03. Actividades 1 a 3. Práctica 5 por niveles (aprox. 1,5 HC).'],
        ['Revisión y cierre', 'aprox. 1 HC', 'Revisión cruzada de la Práctica 5 con el punto de control. Evidencia del encuentro (Work, encuentro 10): préstamo, devolución y consulta con alternativas. Las Actividades 4 y 5 quedan como tarea.'],
    ]),
]
CLAVES_E9 = [
    ['1', '(a) Acción, en el carril del lector: «Entregar el libro». (b) Decisión con dos guardas: [lector activo] y [lector inactivo]. (c) Nodo inicial. (d) Partición (carril) «Sistema».', 'Dibujar la condición dentro de una acción.'],
    ['2', '(a) «pedido de informe» o «mes del informe»: sin verbo. (b) Demasiado general: «datos del préstamo» o «código del ejemplar devuelto». (c) Describe una acción: «datos del préstamo». (d) Junta tres flujos: se separan en flechas distintas.', 'Dejar el verbo en infinitivo.'],
    ['3', 'Carriles lector, bibliotecaria y sistema. Lector: entregar el ejemplar. Bibliotecaria: indicar el código del ejemplar. Sistema: buscar el préstamo activo del ejemplar (RF-08); decisión [tiene préstamo activo] / [sin préstamo activo]; con préstamo, cerrar el préstamo y conservarlo en el historial (RF-03) y la bibliotecaria confirma; sin préstamo, informar y la bibliotecaria aparta el ejemplar. Dos nodos finales posibles.', 'Borrar el préstamo al devolver (contradice RN-05).'],
    ['4', 'Sí, si el sistema le da o recibe información: por ejemplo, los flujos «pedido de ejemplares para el plan lector» y «cantidad disponible de un título», con fuentes E2-L2 y E2-L4. Se acepta que no vaya si se justifica que la reserva para el plan lector está postergada (Capítulo 6, Tabla 6.1) y que la docente hoy consulta como cualquier lector.', 'Incluirla sin flujos.'],
    ['5', 'Volvería el problema (2) del proceso actual: la regla RN-01 dependería del apuro de la persona, como en E4-L4. En el carril del sistema, la comprobación se hace siempre, también con fila.', 'Responder que «da lo mismo».'],
    ['6', 'Mezcla notaciones: la figura de persona es del diagrama de casos de uso y la elipse, del de casos de uso; en el de contexto, el lector es un rectángulo (entidad externa) y el sistema, un círculo. Basta con cambiar esos dos símbolos.', 'Proponer rehacer todo con otra notación.'],
]
CLAVES_E10 = [
    ['1', '(a) Persona: el actor es «Docente» (si usa el sistema). (b) Caso de uso. (c) Paso de un escenario, no un caso de uso. (d) Actor. (e) Nada de eso: es una solución técnica; no va en el diagrama. (f) Caso de uso (CU-05).', 'Aceptar «Base de datos» como actor.'],
    ['2', 'Actor: Bibliotecaria. Precondición: sesión iniciada (RNF-02). Escenario: 1. indica el lector o el ejemplar; 2. el sistema muestra sus datos e historial; 3. la bibliotecaria confirma la inactivación; 4. el sistema lo marca inactivo sin borrar su historial (RF-10, RN-06). Requisitos: RF-10, RNF-02.', 'Usar «borrar» en el escenario.'],
    ['3', 'Flujo 3a: el lector tiene un préstamo activo. Decisión posible: el sistema avisa y no inactiva hasta la devolución, o inactiva y conserva el préstamo activo. El equipo propone opciones; decide la institución con la bibliotecaria, porque afecta una regla (RN-04 y RN-06) y el control de libros que faltan (E1-L4).', 'Elegir una opción sin consultar.'],
    ['4', 'Puede hacer falta: «RF-11 El sistema debe registrar una devolución sin préstamo activo identificado, con la fecha, para revisarla con el historial» (a confirmar). Se acepta que no haga falta si se justifica que se resuelve con un registro manual aparte y se dice quién lo revisa. En el catálogo del caso, RF-11 se usa en el Capítulo 6 para otra cosa (SC-01): el número es el próximo libre en la actividad, no en el caso.', 'Inventar una regla sin consultarla.'],
    ['5', 'El primer rombo tiene dos flujos (4a y 4b) porque [no habilitado] reúne dos condiciones; el segundo rombo, el flujo 5a. Diferencia posible: la Figura 5.2 no distingue lector inactivo de máximo alcanzado; la plantilla sí. Es una simplificación aceptable si el motivo del rechazo lo dice.', 'Contar solo los rombos.'],
]
CLAVES_P5 = [
    ['Nivel 1', 'Carriles Lector y Encargado. Acciones: pedir un libro (lector); buscarlo en el estante (encargado); decisión [está] / [no está]; [no está]: decirle que vuelva otro día (MP-2) y final; [está]: buscar la ficha en la caja (MP-3); anotar nombre, grado y fecha y entregar (MP-4); devolver (lector); buscar la ficha y tachar la línea (MP-5). (1) en «buscar la ficha» al devolver: si no la encuentra, el libro vuelve al estante sin anotar (MP-6, V-11). (2) en «anotar en la ficha»: los préstamos a docentes no se anotan en la ficha (MN-01; registro paralelo de la Práctica 2). Plantilla: «con el sistema, el paso MP-3 lo hace el sistema y se apoya en MR-01» (se acepta MR-03).'],
    ['Nivel 2', 'Proceso propuesto de devolución: lector entrega; encargado indica el libro; sistema busca el préstamo (MR-03) y decide [con préstamo] / [sin préstamo]; con préstamo, cierra el préstamo (MR-02). Caso de uso «Registrar devolución»: actor Encargado; precondición: sesión del encargado; escenario de 3 a 5 pasos; flujo alternativo para MS-1 numerado desde el paso en que el encargado revisa el libro (por ejemplo, 3a): el libro vuelve dañado, se cierra el préstamo y se registra el daño y quién debe reponerlo (M-N5). Requisito nuevo: «MR-09 El sistema debe registrar que un libro devuelto está dañado y no prestarlo hasta que se reponga» (a confirmar).'],
    ['Nivel 3', 'Caso de uso «Registrar préstamo» con flujo alternativo para MS-2: el sistema no lo resuelve; el flujo termina con «no se registra el préstamo» o con un pedido pendiente. La decisión es de la dirección: abrir un recreo de la mañana con los voluntarios (M-F5), un buzón de pedidos a la mañana con entrega a la tarde, o préstamos por medio del docente guía. Se presentan opciones con ventajas y desventajas, sin elegir. Contexto: sistema al centro; entidades Encargado (datos del préstamo, código del libro devuelto / préstamos activos), Lector (título buscado / disponibilidad, sin datos de estudiantes, MR-07) y Dirección (período del informe / libros más prestados, MR-06). Cada flujo con fuente V, M o MN.'],
]
COTEJO_P5 = [
    ['Cada diagrama usa una sola notación y sus símbolos con su significado.', 'Indicador 5'],
    ['Toda acción está en un carril; toda salida de una decisión tiene guarda.', 'Indicador 1'],
    ['El proceso propuesto cita un requisito MR en cada acción del sistema.', 'Indicadores 1 y 5'],
    ['Los actores son roles y los casos de uso, objetivos con verbo y objeto.', 'Indicador 3'],
    ['Cada flujo alternativo se numera desde su paso, dice qué hace el sistema y cómo termina.', 'Indicador 4'],
    ['Si un flujo necesita un requisito nuevo, lo redacta con identificador y «a confirmar».', 'Indicadores 4 y 5'],
]
CRIT_C5 = [['8', '4: carriles Docente y Secretaria (1), acciones con verbo y en orden (1), decisión de Q-3 con sus dos guardas (1), nodos inicial y final (1).', 'Cada parte por separado. 0 en la decisión si no tiene guardas.'],
           ['9', '4: actor como rol (1), escenario de tres pasos o más, numerado (1), flujo alternativo numerado desde su paso (1), resultado del flujo y aviso a la dirección (1).', 'Cada parte por separado.']]
ERRORES_NOTACION = [
    ['Mezclar símbolos de dos notaciones (una elipse en un diagrama de actividades, una figura de persona en el de contexto)', 'Una notación por figura; la leyenda de la notación al lado'],
    ['Poner verbos en los flujos del diagrama de contexto («pedir el informe»)', 'Rótulo con sustantivo: «pedido de informe»'],
    ['Nombrar personas en lugar de roles («Mirta» como actor)', 'El rol: «Bibliotecaria»'],
    ['Dibujar acciones del sistema en el carril de la bibliotecaria', 'Las comprobaciones de reglas van en el carril «Sistema», con su requisito'],
    ['Casos de uso que son pasos («Escribir el código») en lugar de objetivos', 'El objetivo del actor: «Registrar préstamo»; el paso va en la plantilla'],
    ['Salidas de una decisión sin guarda', 'Cada salida con su condición entre corchetes'],
]

# ---------------------------------------------------------------- capítulo 6
SEC_C6 = [
    ('Encuentro 11 — Prioridades, trazabilidad e incrementos (5 HC: 2 de análisis y 3 de gestión)', [
        ['Problema y explicación', 'aprox. 1 HC', 'Sección 6.1 con la Tabla 6.1 y el Ejemplo resuelto 6.1. Sección 6.2 con la Figura 6.1.'],
        ['Práctica', 'aprox. 3 HC', 'Tabla 6.2 y Ejemplo resuelto 6.2: cada equipo arma su matriz. Sección 6.3 con las Tablas 6.3 y 6.4 y el Ejemplo resuelto 6.3: cada equipo planifica un incremento con tareas y responsables. Actividades 1 y 2. Práctica 6, Niveles 1 o 2 (aprox. 1 HC).'],
        ['Revisión y cierre', 'aprox. 1 HC', 'Actividad 3 en plenario. Evidencia del encuentro (Work, encuentro 11): requisito-modelo-prueba; tareas y responsables. La Actividad 4 queda como tarea.'],
    ]),
    ('Encuentro 12 — Solicitud de cambio y revisión de la Unidad II (5 HC: 3 de análisis y 2 de evaluación)', [
        ['Solicitud de cambio', 'aprox. 1,5 HC', 'Sección 6.4 con el registro D3, la Tabla 6.5, la Figura 6.2, el Ejemplo resuelto 6.4 y la Tabla 6.6. Actividades 1 y 2.'],
        ['Revisión del H2', 'aprox. 2 HC', 'El docente entrega a cada equipo una necesidad nueva del banco (sección 4 de esta guía). El equipo la registra, la tramita como solicitud de cambio y presenta catálogo, modelos y matriz en la versión nueva (unos 10 minutos), con preguntas individuales.'],
        ['Correcciones y cierre', 'aprox. 1,5 HC', 'Cada equipo registra las observaciones en la bitácora y planifica las correcciones (evidencia: requisitos y modelos funcionales versionados). Práctica 6, Nivel 3, como ampliación; Actividades 3 y 4 como tarea.'],
    ]),
]
CLAVES_E11 = [
    ['1', '(a) Imprescindible: atiende N-03 y está en el alcance mínimo (Consultas). (b) Deseable: no cambia el circuito ni una regla; depende de tener imágenes de los libros. (c) Imprescindible: sin él se viola N-08 y el permiso mínimo del alcance. (d) Importante o deseable: valor para la dirección y las compras; depende de RF-09 o de un requisito nuevo.', 'Dar la prioridad sin razón.'],
    ['2', 'RF-03: el criterio de referencia de la guía (Correspondencia, RF-03). RF-08: el criterio de referencia (RF-08). Se acepta otra redacción con situación de partida, acción y resultado observable.', 'Repetir el enunciado del requisito.'],
    ['3', 'Tareas del Incremento 2, por ejemplo: cargar ejemplares de prueba (Desarrollo A, 1,5 HC, depende del Incremento 1); cargar lectores de prueba con su estado (Desarrollo B, 1,5 HC, Incremento 1); inactivar lector o ejemplar sin borrar el historial (Desarrollo A, 2 HC, depende de las dos anteriores); comprobar el permiso mínimo (Verificación, 1 HC, depende de las anteriores); bitácora y versión (Versiones, 0,5 HC).', 'Tareas sin responsable o «todos».'],
    ['4', 'RF-09 depende de que haya préstamos registrados (Incremento 3) y es importante, no imprescindible. Empezar por él retrasa el circuito y no se podría probar con datos reales del prototipo. Se le responde con las cuatro preguntas de 6.1 y la columna «depende de» de la Tabla 6.3.', 'Aceptar porque «lo pide la dirección».'],
    ['5', 'RNF-04 | Necesidad: N-03 (saber antes de ir si hay un ejemplar) | Regla: — | Caso de uso: CU-03, paso 4 | Criterio: Práctica 4, consigna común | Prueba: P-16, medir el tiempo de cinco consultas seguidas.', 'Usar un código de prueba ya usado.'],
]
CLAVES_E12 = [
    ['1', '(a) RN-03, RF-02, el criterio de RF-02 (Tabla 4.5: vence el 9 de marzo, ahora el 12), P-02 y CU-01, paso 7. (b) RF-09 y su criterio, P-09, CU-04; depende de RF-11 (tipo de lector) si se aceptó SC-01. (c) RS-01, su criterio y P-13; hay que revisar si alguna decisión del diseño dependía de no tener internet.', 'Revisar solo el requisito.'],
    ['2', 'Solicitud con los campos de la Tabla 6.5: pide la dirección; cambia RN-03 a diez días; RF-02 no cambia de texto («plazo vigente»); cambian el criterio de RF-02 y P-02; impacto bajo; decide la dirección; versión v1.2 del catálogo.', 'Cambiar el texto de RF-02.'],
    ['3', 'Porque cambia una restricción del catálogo y el diseño puede depender de ella: hay que analizar qué se revisa y registrar la versión. Eliminar RS-01 sería riesgoso: si internet falla, el sistema dejaría de funcionar en la única computadora (E4-L10); conviene mantenerla o convertirla en un requisito de calidad de tolerancia a cortes.', 'Eliminar RS-01 porque «ya hay internet».'],
    ['4', 'Respuesta abierta. Se valora que siga la fila de la matriz: regla, requisito, caso de uso, criterio y prueba, y que diga qué se revisa aunque no cambie.', 'Respuesta sin elementos concretos.'],
]
CLAVES_P6 = [
    ['Nivel 1', 'MR-01 | MN-01 o V-03 | Registrar préstamo, paso de registro | Dado un libro disponible y un lector, cuando el encargado registra el préstamo, entonces figura con libro, lector y fecha | PM-01. MR-02 | V-11 | Registrar devolución | dado un préstamo activo, cuando se registra la devolución, entonces el préstamo se cierra | PM-02. MR-04 | MN-04 | Registrar préstamo, flujo alternativo del libro ya prestado | rechazo con motivo | PM-03. MR-05 | MN-02 | Consultar disponibilidad (si se escribió) | cantidad disponible sin datos de estudiantes | PM-04. Omisión típica: MR-06 (libros más prestados) sin caso de uso; es del modelo, que no lo desarrolló.'],
    ['Nivel 2', 'Imprescindibles: MR-01, MR-02, MR-03, MR-04 y MR-08 (sin ellos no hay circuito ni funciona en la biblioteca). Importantes: MR-05 (turno mañana), MR-07 (privacidad de menores; se acepta imprescindible si se justifica con MN-05). Deseable o importante: MR-06. Incremento 1: libros, lectores y préstamo; Incremento 2: devolución y consultas. Con M-F4 (menos horas) conviene que los voluntarios de los miércoles (M-F5) carguen los libros, bajo supervisión. Cada tarea con responsable, HC, dependencia y estado.'],
    ['Nivel 3', 'Solicitud: pide la directora (MC-1, MC-2); cambia la norma M-N1 para docentes (un mes) y los docentes pasan a registrarse como lectores; afecta MR-01 y MR-03 (que ya nombra a los docentes), el cálculo del vencimiento si hay un requisito de plazo, el caso de uso «Registrar préstamo» (tipo de lector) y sus pruebas. Decide la directora, que aprueba las normas. Opciones: aceptar, aceptar con un máximo de libros por docente o postergar al segundo incremento. Nueva versión del catálogo con el requisito del tipo de lector y el plazo por tipo.'],
]
COTEJO_P6 = [
    ['Cada fila de la matriz cita una fuente que existe y una prueba prevista con código.', 'Indicador 2'],
    ['Detecta al menos una omisión y dice si es del modelo o del catálogo.', 'Indicador 2'],
    ['Cada prioridad tiene una razón con las cuatro preguntas; no se suman puntajes.', 'Indicador 1'],
    ['Cada tarea tiene responsable, estimación, dependencia y estado (Nivel 2 o 3).', 'Indicador 3'],
    ['La solicitud de cambio dice qué cambia, qué se revisa, quién decide y la nueva versión (Nivel 3).', 'Indicador 4'],
]
CRIT_C6 = [['8', '4: dos filas correctas (1), otras dos (1), criterios observables (1), una omisión justificada (1).', 'Cada parte por separado.'],
           ['9', '4: qué cambia (1), requisitos y pruebas afectados (1), quién decide (1), nueva versión del catálogo (1).', 'Cada parte por separado.']]

BANCO_H2 = [
    ['Préstamo en sala', 'La bibliotecaria quiere prestar libros para leer en la biblioteca, que vuelven el mismo día y no cuentan para el máximo.', 'Regla nueva (relación con RN-02), un tipo de préstamo, el plazo (RF-02), CU-01 y su flujo 4b, pruebas P-02 y P-06.'],
    ['Vacaciones de invierno', 'La dirección propone préstamos de dos semanas durante las vacaciones de julio.', 'RN-03 por período, RF-02 y su criterio, P-02; quién decide (E5-L8) y desde cuándo rige.'],
    ['Aviso de disponibilidad', 'Un grupo de lectores pide que les avisen cuando un título vuelva a estar disponible.', 'Notificaciones externas fuera del núcleo y RS-01 sin internet: se espera una decisión razonada de postergar o rechazar, registrada.'],
    ['Corrección de un préstamo', 'La bibliotecaria pide poder corregir un préstamo cargado al lector equivocado.', 'RN-05 (el historial no se borra), RNF-02 (quién corrige), un caso de uso nuevo y sus pruebas.'],
]
RUBRICA_EXPLICACION = [
    'El H2 se califica con la misma lógica que el H1: para cada criterio se registran el nivel del producto del equipo y '
    'el nivel individual de cada integrante, y se combinan en la proporción 60 % (equipo) y 40 % (individual). Por '
    'ejemplo, con 4, 3, 2 y 1 puntos para Sólido, Logrado, En desarrollo e Inicial, un producto Logrado (3) con una '
    'evidencia individual Sólida (4) da 3 × 0,6 + 4 × 0,4 = 3,4. La conversión a calificación la decide la institución.',
    'La evidencia individual es la explicación del impacto de un cambio acotado que propone el docente durante la '
    'revisión, el aporte registrado en la bitácora y la observación del docente. La necesidad nueva de la revisión sale '
    'del banco de esta guía y no está en el dossier ni en el buzón.',
    'La ausencia de evidencia se registra aparte («AE»), con la fecha acordada para entregarla; no se confunde con un '
    'desempeño de nivel Inicial. La proporción 60/40 y los niveles son una propuesta editorial tomada del informe de Work '
    '(§4.3); no son una escala oficial del MEC.',
]
CALENDARIO_28 = (['Encuentro (36)', 'Encuentro (28)', 'Qué se hace'],
                 [['7', '6', 'Sin cambios: requisitos funcionales y reglas.'],
                  ['8', '7', 'Sin cambios: calidad, restricciones y criterios; Práctica 4 con su consigna común.'],
                  ['9 y 10', '8', 'Se unen: proceso actual simplificado (Figura 5.1, aprox. 1 HC), proceso propuesto y los tres casos de uso con un flujo alternativo cada uno. El contexto pasa a lectura.'],
                  ['11 y 12', '9', 'Se unen: matriz de trazabilidad, plan de incrementos como lista breve (Tabla 6.3) y solicitud de cambio resuelta dentro de la revisión del H2.']])
PROCEDENCIA = (['Recurso', 'Procedencia', 'Derechos'],
               [['Texto del Capítulo 4', 'Capítulo 4 v2 y guía v2.1 (Fase 2), integrados sin reescribir; cambios declarados en cap4.py y guia.py (especificación C3 e integración).', 'Propios del proyecto.'],
                ['Texto de los Capítulos 5 y 6 y de la apertura de la unidad', 'Redacción original para este libro.', 'Propios del proyecto.'],
                ['Figuras 4.1 y 4.2', 'Mismo código del Capítulo 4 v2 (matplotlib), regenerado.', 'Propias; sin recursos de terceros.'],
                ['Figuras 5.1 a 5.4 y leyendas de símbolos', 'Generadas con build/diagramas.py desde caso/modelos.py, una función por símbolo y por notación.', 'Propias; sin recursos de terceros.'],
                ['Figuras 6.1 y 6.2', 'Esquemas didácticos generados con build/diagramas.py.', 'Propias.'],
                ['Notaciones UML (actividades con particiones y casos de uso) y diagrama de contexto', 'Conocimiento general de la disciplina, con la selección de símbolos de NOTACIONES_CAPITULO_5.md, aprobada por Fernando.', 'No se copió material de terceros.'],
                ['Caso Arapoty: catálogo, buzón, D3, trazabilidad y pruebas previstas', 'Inventados. Fuente única: lote_1/caso/dossier.py y lote_2/caso/requisitos.py y modelos.py.', 'Propios del proyecto.'],
                ['Caso de práctica Mbokaja (Unidad II)', 'Inventado para este lote: lote_2/caso/mbokaja_u2.py, que amplía lote_1/caso/mbokaja.py.', 'Propio.'],
                ['Capacidades, niveles de desempeño, proporción 60/40, cadena N-02 a P-07 e incrementos', 'Informe de Work v0.1 (§4.2, §4.3, §5.2 y §7.3), como propuesta editorial.', 'Insumo interno del proyecto.'],
                ['Relación con Algorítmica 3.º', 'No se reutilizaron ejercicios, personas ni títulos de ese libro; no hay diagramas de datos ni tablas.', 'No aplica.']])


def _sec(doc, secuencia):
    for titulo, filas in secuencia:
        M.tabla(doc, (['Momento', 'Tiempo', 'Qué se hace'], filas), ANCHO_SEC, rotulo=titulo)


def _claves(doc, rotulo, filas):
    M.tabla(doc, (['Act.', 'Respuesta esperada', 'Error frecuente'], filas), [1.2, 11.0, 4.4], rotulo=rotulo)


def _eval(doc, C, claves, crit):
    filas = [[str(n), clave, detalle] for eid, n, clave, detalle in claves if eid == C.EVAL_ID]
    M.tabla(doc, (['Ítem', 'Clave', 'Respuesta o criterio'], filas), [1.2, 1.8, 13.6], rotulo='Clave de la %s' % C.EVAL_TITULO.lower())
    M.tabla(doc, (['Ítem', 'Puntaje completo', 'Puntaje parcial'], crit), [1.2, 8.4, 7.0], rotulo='Criterios de corrección de los ítems abiertos')


def _niveles(doc, rotulo, filas, ancho=(2.2, 14.4)):
    M.tabla(doc, (['Nivel', 'Respuesta esperada'], filas), list(ancho), rotulo=rotulo)


def construir(ruta, claves, paginas, estado, figs):
    import cap3   # niveles de desempeño del lote 1
    doc = H.new_doc()
    Z.portada(doc, 'Guía docente con solucionario', estado)
    doc.add_heading('Guía docente — Lote 2', level=1)
    H.caja(doc, 'Sobre las capacidades', [
        'Las capacidades de las fichas (C3, C4, C7 y C9 en esta unidad) son propuestas para el programa institucional del '
        'Plan Optativo, tomadas del informe de Work. No son capacidades oficiales del MEC. Según el informe de Work, el '
        'diseño curricular consultado (mayo de 2026, p. 29) permite que la institución elabore el programa del Plan '
        'Optativo; esa lectura no se verificó por separado en este lote.'], tipo='neutra')
    H.p_texto(doc, 'Cada encuentro sigue la estructura de taller de Work: 1 HC de problema y explicación, 3 HC de práctica y '
                   '1 HC de revisión, con la distribución por categorías de la versión 2.1 (análisis, diseño, verificación, '
                   'gestión y evaluación). Los tiempos son aproximados y la duración de la hora cátedra la fija cada '
                   'institución. Las páginas indicadas son las del libro del estudiante.', italic=True, size=9.5)

    # 1. capítulo 4
    doc.add_heading('1. Capítulo 4 — %s' % cap4.TITULO, level=2)
    M.ficha(doc, cap4)
    M.tabla(doc, (['Indicador de C3 (Work)', 'Cobertura en el lote'], COBERTURA_C3), [6.6, 10.0], rotulo='Cobertura de los indicadores de C3')
    _sec(doc, SEC_C4)
    _claves(doc, 'Claves de las actividades del Encuentro 7', GP.CLAVES_E7)
    _claves(doc, 'Claves de las actividades del Encuentro 8', GP.CLAVES_E8)
    n = {k: sum(1 for r in R.CORRESPONDENCIA if r[1] == k) for k in R.ETIQUETA}
    doc.add_heading('Requisitos del catálogo y sus criterios de aceptación', level=3)
    H.p_texto(doc, 'Los 15 requisitos de la Tabla 4.3 y el RNF-04 de la consigna común tienen un criterio «Dado… cuando… '
                   'entonces…». Modelo resuelto en el libro (%d): el estudiante lo encuentra hecho. Trabajo del estudiante '
                   '(%d más RNF-04): lo escribe en la actividad indicada; acá va la respuesta esperada. Referencia para el '
                   'docente (%d): no figura en el libro y sirve para corregir el catálogo de cada equipo; en la matriz del '
                   'Capítulo 6 aparece como «a escribir por el equipo».' % (n['MODELO'], n['ESTUDIANTE'], n['DOCENTE']), size=10)
    M.tabla(doc, (['Id.', 'Origen', 'Dónde', 'Criterio de aceptación'], CORRESPONDENCIA), [1.9, 2.6, 3.0, 9.1],
            rotulo='Correspondencia entre requisitos y criterios de aceptación')
    H.p_texto(doc, NOTA_C3, italic=True, size=9.5)
    doc.add_heading('Práctica 4: niveles, claves y lista de cotejo', level=3)
    H.p_texto(doc, 'La consigna común la resuelven todos antes de su nivel: asegura que cada estudiante escriba un criterio '
                   'sobre un requisito de calidad. Los buzones B-01 a B-08 son datos nuevos del mismo caso, que no aparecen '
                   'en los ejemplos resueltos (decisión editorial: buzón de Arapoty). El Nivel 1 es el piso; el Nivel 3, '
                   'ampliación.', size=10)
    _niveles(doc, 'Práctica 4: claves de la consigna común y de los niveles', CLAVES_P4)
    M.tabla(doc, (['Criterio', 'Indicador'], COTEJO_P4), [13.6, 3.0], rotulo='Lista de cotejo de la Práctica 4')
    _eval(doc, cap4, claves, CRIT_C4)

    # 2. capítulo 5
    K.titulo_en_pagina_nueva(doc, '2. Capítulo 5 — %s' % cap5.TITULO, nivel=2)
    M.ficha(doc, cap5)
    _sec(doc, SEC_C5)
    doc.add_heading('Leyenda de notaciones y errores frecuentes', level=3)
    H.p_texto(doc, 'Una notación por modelo (NOTACIONES_CAPITULO_5.md). Las tres tablas son las mismas del libro; las '
                   'figuras se dibujan con estas funciones, y el control automático comprueba que ninguna mezcle símbolos.', size=10)
    Z.tabla_leyenda(doc, 'actividad', 'Leyenda: diagrama de actividades con particiones', figs)
    Z.tabla_leyenda(doc, 'contexto', 'Leyenda: diagrama de contexto', figs)
    Z.tabla_leyenda(doc, 'casos_de_uso', 'Leyenda: diagrama de casos de uso', figs)
    M.tabla(doc, (['Error frecuente', 'Cómo se corrige'], ERRORES_NOTACION), [9.6, 7.0], rotulo='Errores frecuentes con las notaciones')
    _claves(doc, 'Claves de las actividades del Encuentro 9', CLAVES_E9)
    _claves(doc, 'Claves de las actividades del Encuentro 10', CLAVES_E10)
    _niveles(doc, 'Práctica 5: claves por nivel', CLAVES_P5)
    M.tabla(doc, (['Criterio', 'Indicador'], COTEJO_P5), [12.8, 3.8], rotulo='Lista de cotejo de la Práctica 5')
    _eval(doc, cap5, claves, CRIT_C5)

    # 3. capítulo 6
    K.titulo_en_pagina_nueva(doc, '3. Capítulo 6 — %s' % cap6.TITULO, nivel=2)
    M.ficha(doc, cap6)
    _sec(doc, SEC_C6)
    _claves(doc, 'Claves de las actividades del Encuentro 11', CLAVES_E11)
    _claves(doc, 'Claves de las actividades del Encuentro 12', CLAVES_E12)
    M.tabla(doc, (['Prueba', 'Requisito', 'Qué se comprueba (se diseña en los Capítulos 14 a 16)'], [list(p) for p in R.PRUEBAS]),
            [1.6, 2.0, 13.0], rotulo='Pruebas previstas P-01 a P-15 del caso')
    _niveles(doc, 'Práctica 6: claves por nivel', CLAVES_P6)
    M.tabla(doc, (['Criterio', 'Indicador'], COTEJO_P6), [13.0, 3.6], rotulo='Lista de cotejo de la Práctica 6')
    _eval(doc, cap6, claves, CRIT_C6)

    # 4. rúbrica H2
    K.titulo_en_pagina_nueva(doc, '4. Rúbrica del hito H2: requisitos y modelos funcionales versionados', nivel=2)
    for t in RUBRICA_EXPLICACION:
        H.p_texto(doc, t)
    M.tabla(doc, (['Criterio', '4. Sólido', '3. Logrado', '2. En desarrollo', '1. Inicial'],
                  [['%s (%s)' % (r[0], r[1])] + r[2] for r in cap6.RUBRICA_H2]),
            [2.8, 3.6, 3.4, 3.4, 3.4], rotulo='Rúbrica del H2: descriptores por nivel', size=8.5)
    M.tabla(doc, (['Criterio', 'Evidencia del equipo (60 %)', 'Evidencia individual (40 %)'],
                  [[r[0], r[3], r[4]] for r in cap6.RUBRICA_H2]), [4.4, 5.8, 6.4], rotulo='Evidencias de cada criterio del H2')
    M.tabla(doc, (['Nivel de desempeño', 'Descriptor común (Work §4.3)'], [list(x) for x in cap3.NIVELES_DESEMPENO]),
            [3.6, 13.0], rotulo='Niveles de desempeño (propuesta editorial)')
    M.tabla(doc, (['Estudiante (código)', 'Criterio', 'Nivel del equipo', 'Nivel individual', 'Resultado', 'AE'],
                  [['', r[0], '', '', '', ''] for r in cap6.RUBRICA_H2]),
            [3.0, 5.0, 2.2, 2.2, 2.2, 2.0], rotulo='Planilla de registro del H2 (una por estudiante)')
    M.tabla(doc, (['Necesidad nueva', 'Qué pide el usuario', 'Qué tiene que encontrar el análisis de impacto'], BANCO_H2),
            [3.4, 6.4, 6.8], rotulo='Banco de necesidades nuevas para la revisión del H2')
    H.p_texto(doc, 'Ninguna de estas necesidades está en el dossier, en el buzón ni en la nota D3. Si el curso trabaja con su '
                   'propia biblioteca, el docente formula una necesidad equivalente a partir del relevamiento autorizado.', italic=True, size=9.5)

    # 5. calendario y 6. procedencia
    doc.add_heading('5. Calendario con menos encuentros', level=2)
    H.p_texto(doc, 'En la adaptación de 28 encuentros (distribución v2.1), la Unidad II pasa de seis a cuatro encuentros. '
                   'Cada capítulo del libro tiene su recuadro «Si el calendario se acorta».')
    M.tabla(doc, CALENDARIO_28, [2.6, 2.6, 11.4], rotulo='Unidad II en 36 y en 28 encuentros')
    doc.add_heading('6. Procedencia y derechos de los recursos', level=2)
    M.tabla(doc, PROCEDENCIA, [4.4, 7.8, 4.4], rotulo='Inventario de recursos del lote')
    M.pegar_renglones(doc)
    K.saltos_a_titulos(doc)
    H.finalizar(doc, ruta)
