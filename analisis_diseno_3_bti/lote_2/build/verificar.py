# -*- coding: utf-8 -*-
"""Controles automáticos del lote 2 (Unidad II). Uso: python3 lote_2/build/verificar.py   (sale con 1 si hay fallos)

Bloques del lote 1, extendidos a los Capítulos 4 a 6: A índice de Work · B capacidades · C trazabilidad de fuentes ·
D integración del Capítulo 4 (piloto) · E Algorítmica 3.º y Matemática Aplicada · F voseo y residuos · G maqueta (DOCX) ·
H maqueta (PDF) · I evaluaciones, prácticas e instrumentos · J numeración.
Bloques nuevos: K especificación de evidencia C3 (§4, cuatro comprobaciones y sus ajustes) · L notaciones del Capítulo 5
(NOTACIONES_CAPITULO_5.md, §4, siete controles) · M trazabilidad entre capítulos y solicitud de cambio."""
import os, re, sys, json, hashlib, collections, subprocess, importlib.util
sys.dont_write_bytecode = True
AQUI = os.path.dirname(os.path.abspath(__file__))
L2 = os.path.dirname(AQUI)
ADS = os.path.dirname(L2)
sys.path.insert(0, AQUI)
sys.path.insert(1, os.path.join(L2, 'caso'))
from docx import Document
from docx.oxml.ns import qn
import pymupdf

import oficio as O
K = O.K
import cap4, cap5, cap6, guia, construir as B
import diagramas as G
import requisitos as R, modelos as MO, mbokaja_u2 as MB
D = R.D
CAPS = (cap4, cap5, cap6)
EST, GUI, CUA = B.RUTA_EST[:-5], B.RUTA_GUIA[:-5], B.RUTA_CUAD[:-5]

ok = fallos = 0
RES = collections.OrderedDict()
DETALLE = []


def chk(bloque, desc, cond, detalle=''):
    global ok, fallos
    RES.setdefault(bloque, [0, 0])
    if cond:
        ok += 1; RES[bloque][0] += 1
        print('CORRECTO [%s] %s' % (bloque, desc))
    else:
        fallos += 1; RES[bloque][1] += 1
        print('FALLA    [%s] %s %s' % (bloque, desc, detalle))
    DETALLE.append({'bloque': bloque, 'control': desc, 'resultado': 'correcto' if cond else 'FALLA'})


def texto_docx(ruta):
    d = Document(ruta)
    partes = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            vistos = set()
            for c in r.cells:
                if id(c._tc) in vistos:
                    continue
                vistos.add(id(c._tc))
                partes.append(c.text)
    return d, '\n'.join(partes)


dE, tE = texto_docx(EST + '.docx')
dG, tG = texto_docx(GUI + '.docx')
dC, tC = texto_docx(CUA + '.docx')
pE, pG, pC = pymupdf.open(EST + '.pdf'), pymupdf.open(GUI + '.pdf'), pymupdf.open(CUA + '.pdf')
res_build = json.load(open(os.path.join(L2, 'resultado_construccion.json'), encoding='utf8'))


def cargar(ruta, nombre):
    spec = importlib.util.spec_from_file_location(nombre, ruta)
    m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
    return m


PIL = cargar(os.path.join(ADS, 'fase2_pilotos', 'A2_capitulo4_v2', 'build', 'contenido.py'), 'piloto_v2')

# ---------------------------------------------------------------- A. índice de Work (§6, textual; §5.2)
WORK = {
    4: ('Requisitos que se pueden comprobar', (7, 8),
        ['4.1. Necesidad, requisito, regla de negocio y solución técnica.', '4.2. Requisitos funcionales, calidad y restricciones.',
         '4.3. Identificación, fuente, prioridad y criterios de aceptación.', '4.4. Ambigüedades frecuentes y cómo corregirlas.'],
        'convertir pedidos de usuarios en requisitos verificables'),
    5: ('Procesos, contexto y casos de uso', (9, 10),
        ['5.1. Proceso actual y proceso propuesto.', '5.2. Diagrama de contexto y límites de intercambio de información.',
         '5.3. Actores, casos de uso y escenarios principales.', '5.4. Excepciones, decisiones y flujos alternativos.'],
        'modelar préstamo y devolución con situaciones problemáticas nuevas'),
    6: ('Priorizar y mantener la coherencia', (11, 12),
        ['6.1. Alcance mínimo y funcionalidades postergables.', '6.2. Trazabilidad entre necesidad, requisito, modelo y prueba.',
         '6.3. Iteraciones, tareas, dependencias y responsabilidades.', '6.4. Solicitud de cambio y análisis de impacto.'], None),
}
sin_punto = lambda s: re.sub(r'^(\d\.\d)\. ', r'\1 ', s).rstrip('.')
heads = [p.text for p in dE.paragraphs if p.style.name.startswith('Heading')]
for C in CAPS:
    tit, enc, secs, prac = WORK[C.NUM]
    chk('A', 'Capítulo %d: título de Work «%s»' % (C.NUM, tit), C.TITULO == tit and '%s — %s' % (C.CAPITULO, tit) in heads)
    chk('A', 'Capítulo %d: encuentros %s' % (C.NUM, enc), tuple(C.ENCUENTROS) == enc)
    chk('A', 'Capítulo %d: secciones de Work §6, textuales y en orden' % C.NUM,
        C.SECCIONES == [sin_punto(s) for s in secs] and [h for h in heads if h in C.SECCIONES] == C.SECCIONES)
    if prac:
        chk('A', 'Capítulo %d: práctica de Work («%s»)' % (C.NUM, prac), C.PRACTICA_WORK == prac)
chk('A', 'Capítulo 6 cierra con la revisión de la Unidad II y el hito H2 (Work: «%s»)' % cap6.REVISION_WORK,
    cap6.H2_TITULO in heads and 'H2' in cap6.H2_TITULO and 'Revisión de la Unidad II' in cap6.H2_TITULO)
orden = [B.UNIDAD_II] + ['%s — %s' % (C.CAPITULO, C.TITULO) for C in CAPS]
chk('A', 'Orden: apertura de la Unidad II y Capítulos 4, 5 y 6', all(t in heads for t in orden) and [heads.index(t) for t in orden] == sorted(heads.index(t) for t in orden))
chk('A', 'Apertura de la Unidad II con el resultado de Work §5.1 («Requisitos, procesos y trazabilidad inicial»)',
    'requisitos, procesos y trazabilidad inicial' in tE)

# ---------------------------------------------------------------- B. capacidades
CAPW = {'C2': 'Obtiene y contrasta información sobre necesidades de usuarios.', 'C3': 'Especifica y prioriza requisitos verificables.',
        'C4': 'Modela procesos e interacciones del sistema.', 'C7': 'Organiza y registra el trabajo colaborativo.',
        'C9': 'Verifica y mejora la solución a partir de evidencia.'}
ESPERADAS = {4: ('C3', 'C9'), 5: ('C4', 'C3'), 6: ('C3', 'C4', 'C7')}   # Work §5.2, encuentros 7 a 12
for C in CAPS:
    r = C.ROTULO_CAPACIDAD
    chk('B', 'Capítulo %d: rótulo «propuestas para el programa institucional» y «no son capacidades oficiales del MEC»' % C.NUM,
        'propuestas para el programa institucional' in r and 'no son capacidades oficiales del MEC' in r)
    chk('B', 'Capítulo %d: el rótulo figura en la ficha del libro y de la guía' % C.NUM, r in tE and r in tG)
    cods = ESPERADAS[C.NUM]
    chk('B', 'Capítulo %d: capacidades %s de Work §5.2, con su texto de §4.2' % (C.NUM, ', '.join(cods)),
        all(('%s. %s' % (c, CAPW[c])) in C.CAPACIDAD for c in cods) and all(c in r for c in cods))
malas_of = [tx[max(0, m.start() - 40):m.end()] for tx in (tE, tG, tC) for m in re.finditer(r'(?i)capacidad(es)? oficial', tx)
            if not re.search(r'\b[Nn]o (es|son)\b', tx[max(0, m.start() - 40):m.start()])]
chk('B', 'Ninguna capacidad presentada como oficial del MEC', not malas_of, malas_of)
chk('B', 'Guía: aclaración sobre las capacidades propuestas', 'no son capacidades oficiales del MEC' in tG.replace('\n', ' '))

# ---------------------------------------------------------------- C. trazabilidad de fuentes
locales = dict(MB.FUENTES)
for C in CAPS:
    locales.update(C.EVAL_FUENTES)
def fuente(c):
    return R.TEXTO.get(c) or D.NECESIDAD.get(c) or dict(D.REGLAS).get(c) or locales.get(c)
PAT = (r'(?:E\d-L\d+|O1-\d+|D\d-A?\d+|N-\d\d|V-\d\d|M-[NF]\d|MN-\d\d|MP-\d|MS-\d|MC-\d|B-0\d|Q-\d|W-\d|R-07)')
# Comillas que rotulan un estado o un nombre, no citan una línea (clave del Nivel 2 de la guía v2.1 del piloto).
ROTULOS = {('en reparación', 'B-02')}
for nombre, tx in (('libro', tE), ('guía', tG), ('cuaderno', tC)):
    usados = sorted(set(re.findall(r'(?<![\w-])' + PAT + r'(?![\w-])', tx)))
    falt = [c for c in usados if fuente(c) is None]
    chk('C', '%s: los %d códigos de fuente citados existen en el caso, en Mbokaja o en la situación de evaluación' % (nombre, len(usados)), not falt, falt)
    citas = re.findall(r'«([^»]{6,})»\s*\(((?:%s)(?:,\s*%s)*)\)' % (PAT, PAT), tx)
    malas = [(q[:50], cods) for q, cods in citas if (q, cods) not in ROTULOS
             and not any(fuente(c.strip()) and q.strip().rstrip('.') in fuente(c.strip()) for c in cods.split(','))]
    chk('C', '%s: %d citas textuales «…» (código) coinciden con la línea citada' % (nombre, len(citas)), not malas, malas)
REQS = set(R.REQ) | {R.RF11[0]}
chk('C', 'Catálogo v1.0: 15 requisitos (10 RF, 3 RNF, 2 RS), identificadores únicos y fuentes existentes',
    len(R.CATALOGO) == 15 and len(R.REQ) == 15 and all(fuente(f) for r in R.CATALOGO for f in r[3]))
chk('C', 'Catálogo: al menos ocho funcionales y cuatro de calidad o restricciones (Work §4.2, C3)',
    sum(i.startswith('RF-') for i in R.REQ) >= 8 and sum(i.startswith(('RNF-', 'RS-')) for i in R.REQ) >= 4)
chk('C', 'El libro reproduce el catálogo, el buzón y los postergados de la fuente única',
    all(r[2] in tE for r in R.CATALOGO) and all(t in tE for _, _, t in R.BUZON) and all(p[1] in tE for p in R.POSTERGADOS))
chk('C', 'Extracto del dossier: cada línea con el texto completo de lote_1/caso/dossier.py',
    all(('%s  %s' % (c, D.TEXTO[c])) in tE for _, _, frs in cap4.DOSSIER for c, _ in frs))
chk('C', 'Registro D3 completo en el libro, con sus tres líneas', all(('%s  %s' % (c, t)) in tE for c, _, t in R.REGISTROS_U2[0]['lineas']))
chk('C', 'Mbokaja: los requisitos MR citan fuentes que existen (V, M-F, MP, MN)',
    all(all(fuente(f.strip()) for f in r[3].split(',')) for r in MB.CATALOGO))
chk('C', 'Mbokaja: los datos nuevos no repiten ningún código del caso Mbokaja del lote 1',
    not (set(c for c, _ in [MB.DATO_NUEVO] + MB.PROCESO + MB.SITUACIONES + MB.CAMBIO) & set(MB.MB.FUENTES)))
r = subprocess.run([sys.executable, '-B', os.path.join(ADS, 'lote_1', 'caso', 'verificar_caso.py')], capture_output=True, text=True)
chk('C', 'lote_1/caso/verificar_caso.py sigue sin fallos: %s' % r.stdout.strip().split('\n')[-1], r.returncode == 0)

# ---------------------------------------------------------------- D. integración del Capítulo 4
cambiados = {n for n, _ in cap4.CAMBIOS}
iguales = [n for n in dir(PIL) if n.isupper() and n not in cambiados]
dif = [n for n in iguales if getattr(cap4, n) != getattr(PIL, n)]
chk('D', 'Capítulo 4: %d piezas del piloto v2 sin cambios; solo cambian las %d declaradas' % (len(iguales), len(cambiados)), not dif, dif)
h = {c: hashlib.sha256(open(os.path.join(ADS, 'fase2_pilotos', c, 'build', 'contenido.py'), 'rb').read()).hexdigest() for c in ('A2_capitulo4_v2', 'A2_capitulo4_v2_1')}
chk('D', 'El piloto no se modificó (contenido.py v2 = v2.1, SHA-256 %s…)' % h['A2_capitulo4_v2'][:12], len(set(h.values())) == 1)
chk('D', 'Traslado a la fuente única: catálogo, postergados y buzón idénticos al piloto',
    R.CATALOGO == PIL.CATALOGO and R.POSTERGADOS == PIL.POSTERGADOS and R.BUZON == PIL.BUZON and cap4.BUZON == PIL.BUZON)
chk('D', 'Reglas RN-01 a RN-06 idénticas al piloto y al dossier', cap4.REGLAS == PIL.REGLAS == list(D.REGLAS))
chk('D', 'Extracto del dossier: mismos registros y códigos que el piloto, y cada fragmento del piloto está en su línea completa',
    [(a, [c for c, _ in f]) for a, _, f in cap4.DOSSIER] == [(a, [c for c, _ in f]) for a, _, f in PIL.DOSSIER]
    and all(t in D.TEXTO[c] for _, _, f in PIL.DOSSIER for c, t in f))
chk('D', 'S42: párrafos iguales al piloto; solo cambia el título (textual de Work)', cap4.S42[1] == PIL.S42[1] and cap4.S42[0] != PIL.S42[0])
chk('D', 'Niveles 1 y 3 de la Práctica 4 iguales al piloto', cap4.NIVELES[0] == PIL.NIVELES[0] and cap4.NIVELES[2] == PIL.NIVELES[2])
chk('D', 'Actividades del Encuentro 8: solo cambia la Actividad 1 (d)', cap4.ACT_E8[1:] == PIL.ACT_E8[1:]
    and cap4.ACT_E8[0][1][0].replace('(d) Para resolver en casa: ', '(d) ') == PIL.ACT_E8[0][1][0])
chk('D', 'Indicadores: los cuatro del piloto más el de C9', cap4.INDICADORES[:4] == PIL.INDICADORES and len(cap4.INDICADORES) == 5)
chk('D', 'Sin portada ni recuadro de piloto: el libro no dice «piloto» ni «Piloto C»', not re.search(r'(?i)\bpiloto\b', tE))
chk('D', 'Guía: claves de los Encuentros 7 y 8 de la guía v2.1, sin cambios', all(row[1] in tG for row in guia.GP.CLAVES_E7 + guia.GP.CLAVES_E8))
chk('D', 'Guía: correspondencia de los 15 requisitos igual a la guía v2.1', [list(x) for x in R.CORRESPONDENCIA] == [list(x) for x in guia.GP.K.CORRESPONDENCIA])

# ---------------------------------------------------------------- K. especificación de evidencia C3
cc = ' '.join(cap4.CONSIGNA_COMUN)
chk('K', '1. Consigna fuera de NIVELES rotulada «Consigna común», con «requisito de calidad», «RNF-04» y «Dado… cuando… entonces…»',
    'Consigna común' in cap4.CONSIGNA_COMUN_TITULO and all(x in cc for x in ('requisito de calidad', 'RNF-04', 'Dado… cuando… entonces…', 'B-03'))
    and cap4.CONSIGNA_COMUN_TITULO in tE and cap4.CONSIGNA_COMUN[0].split('\n')[0] in tE)
chk('K', '2. Lista de cotejo con la fila del criterio de calidad, vinculada al indicador 3',
    any('requisito de calidad' in f[0] and f[1] == 'Indicador 3' for f in guia.COTEJO_P4) and 'consigna común' in tG)
chk('K', '3. Correspondencia con RNF-04 y origen «Trabajo del estudiante»',
    R.CORRESPONDENCIA_C3[0].startswith('RNF-04') and R.ETIQUETA[R.CORRESPONDENCIA_C3[1]] == 'Trabajo del estudiante' and R.CORRESPONDENCIA_C3[3] in tG)
chk('K', '4. Ninguna consigna de nivel reutiliza RNF-04', 'RNF-04' not in repr(cap4.NIVELES))
chk('K', 'Ajustes: Nivel 2 numera desde RNF-05 y corrige solo B-08; Actividad 1 (d) del Encuentro 8 como tarea',
    'RNF-05' in cap4.NIVELES[1][1][0] and 'B-08' in cap4.NIVELES[1][1][1] and 'B-03' not in cap4.NIVELES[1][1][1]
    and '(d) Para resolver en casa' in tE)
chk('K', 'Guía: clave de la consigna común y nota de evidencia mínima de C3', 'Consigna común' in tG and guia.NOTA_C3 in tG)
chk('K', 'Todo camino por la Práctica 4 incluye un criterio sobre un requisito de calidad (la consigna común está antes de los tres niveles)',
    tE.index(cap4.CONSIGNA_COMUN_TITULO) < tE.index('Nivel 1 — Guiado', tE.index(cap4.PRACTICA_TITULO)))

# ---------------------------------------------------------------- L. notaciones del Capítulo 5
reg = json.load(open(os.path.join(AQUI, 'figs', 'simbolos.json'), encoding='utf8'))
figs5 = {'fig_5_1': 'actividad', 'fig_5_2': 'actividad', 'fig_5_3': 'contexto', 'fig_5_4': 'casos_de_uso'}
for f, n in figs5.items():
    usados = reg[f]['simbolos']
    chk('L', '1. Sin mezcla: %s declara «%s» y usa solo sus símbolos (%s)' % (f, G.NOTACIONES[n], ', '.join(sorted(usados))),
        reg[f]['notacion'] == n and usados and all(G.SIMBOLOS[s][0] == n for s in usados))
for f in ('fig_6_1', 'fig_6_2'):
    chk('L', '1. %s (esquema del Capítulo 6) no usa símbolos UML ni del diagrama de contexto' % f, reg[f]['notacion'] == 'esquema' and not reg[f]['simbolos'])
for f, m in (('fig_5_1', MO.PROCESO_ACTUAL), ('fig_5_2', MO.PROCESO_PROPUESTO)):
    tipos = collections.Counter(n[1] for n in m['nodos'])
    salidas = collections.defaultdict(list)
    for a, b, g in m['flujos']:
        salidas[a].append(g)
    dec = [n[0] for n in m['nodos'] if n[1] == 'decision']
    guardas_ok = all((len(salidas[d]) == 1 and salidas[d][0] is None) or (len(salidas[d]) >= 2 and all(g and g.startswith('[') and g.endswith(']') for g in salidas[d])) for d in dec)
    sin_guarda_fuera = all(g is None for a, b, g in m['flujos'] if a not in dec)
    chk('L', '2. %s: un nodo inicial, %d nodos finales, guardas en todas las salidas de cada decisión y ninguna fuera de ellas' % (f, tipos['final']),
        tipos['inicial'] == 1 and tipos['final'] >= 1 and guardas_ok and sin_guarda_fuera)
    chk('L', '2. %s: acciones con verbo en infinitivo' % f, all(re.match(r'^\w+(ar|er|ir)\b', n[5]) for n in m['nodos'] if n[1] == 'accion'),
        [n[5] for n in m['nodos'] if n[1] == 'accion' and not re.match(r'^\w+(ar|er|ir)\b', n[5])])
    chk('L', '3. %s: toda acción tiene carril' % f, all(n[2] in m['carriles'] for n in m['nodos']))
sis = [n for n in MO.PROCESO_PROPUESTO['nodos'] if n[2] == 'Sistema' and n[1] == 'accion']
citas_sis = [re.findall(r'R[FN]?F?-\d\d|RS-\d\d', n[5].replace('\n', ' ')) for n in sis]
chk('L', '3. Proceso propuesto: las %d acciones del carril «Sistema» citan requisitos del catálogo' % len(sis),
    sis and all(c for c in citas_sis) and all(x in REQS for c in citas_sis for x in c))
chk('L', '2. Contexto: un solo sistema, entidades externas y flujos rotulados con sustantivos',
    reg['fig_5_3']['simbolos'].get('contexto_sistema') == 1 and reg['fig_5_3']['simbolos'].get('contexto_entidad') == len(MO.CONTEXTO['entidades'])
    and all(not re.match(r'^\w+(ar|er|ir)$', f[2].split()[0]) for f in MO.CONTEXTO['flujos']))
chk('L', '2. Casos de uso: actores fuera del límite (dibujados afuera), roles y no personas',
    reg['fig_5_4']['simbolos'].get('cu_limite') == 1 and not set(MO.CASOS_USO['actores']) & {c for c, _ in MO.CASOS_USO['casos']}
    and not any(a in ' '.join(n for _, n, _ in D.PERSONAS) for a in MO.CASOS_USO['actores']))
citados_modelos = set()
for txt in [n[5] or '' for m in (MO.PROCESO_ACTUAL, MO.PROCESO_PROPUESTO) for n in m['nodos']] + \
           [json.dumps(p, ensure_ascii=False) for p in MO.PLANTILLAS] + [json.dumps(MO.CAMBIOS_PROCESO, ensure_ascii=False)]:
    citados_modelos |= set(re.findall(r'\b(?:RF|RNF|RS)-\d\d\b', txt))
chk('L', '4. Los %d requisitos citados en diagramas, casos de uso y flujos alternativos existen en la fuente única' % len(citados_modelos),
    citados_modelos and citados_modelos <= REQS, citados_modelos - REQS)
cubre = set().union(*(MO.PLANTILLA[c]['requisitos'] for c in MO.PRIORITARIOS))
chk('L', '4. «Registrar préstamo», «Registrar devolución» y «Consultar disponibilidad» cubren entre los tres RF-01 a RF-08',
    {'RF-%02d' % i for i in range(1, 9)} <= cubre and [MO.PLANTILLA[c]['nombre'] for c in MO.PRIORITARIOS] == ['Registrar préstamo', 'Registrar devolución', 'Consultar disponibilidad'])
chk('L', '4. Los requisitos de cada plantilla figuran en su escenario o en sus flujos',
    all(all(r in json.dumps(p['escenario'] + [a[1] for a in p['alternativos']] + [p['precondicion']], ensure_ascii=False) for r in p['requisitos']) for p in MO.PLANTILLAS))
carriles = set(MO.PROCESO_PROPUESTO['carriles'])
chk('L', '5. Cada actor de los casos de uso es una entidad externa del contexto o un carril del proceso propuesto',
    all(a in MO.CONTEXTO['entidades'] or a in carriles for a in MO.CASOS_USO['actores']))
chk('L', '5. Cada flujo de información del contexto tiene una fuente del dossier que existe',
    all(f[3] and all(c in D.TEXTO for c in f[3]) for f in MO.CONTEXTO['flujos']))
chk('L', '6. Cada caso de uso prioritario tiene al menos un flujo alternativo numerado desde un paso existente',
    all(p['alternativos'] and all(int(re.match(r'(\d+)[a-z]$', a[0]).group(1)) <= len(p['escenario']) for a in p['alternativos']) for p in MO.PLANTILLAS))
alt01 = ' '.join(a[1] for a in MO.PLANTILLA['CU-01']['alternativos'])
chk('L', '6. Los flujos de «Registrar préstamo» cubren RF-04 (ya prestado), RF-05 (máximo) y RF-06 (lector inactivo)',
    all(x in alt01 for x in ('RF-04', 'RF-05', 'RF-06')))
usados_todos = {s for f in figs5 for s in reg[f]['simbolos']}
chk('L', '7. Cada símbolo usado en una figura figura en la leyenda de la guía (%d símbolos)' % len(usados_todos),
    all(G.SIMBOLOS[s][1] in tG and G.SIMBOLOS[s][3] in tG for s in usados_todos))
chk('L', 'Leyenda en el libro: un recuadro de símbolos en 5.1, 5.2 y 5.3, antes de la primera figura de cada notación',
    all(r in tE for r in (cap5.LEYENDA_ACTIVIDAD, cap5.LEYENDA_CONTEXTO, cap5.LEYENDA_CU))
    and tE.index(cap5.LEYENDA_ACTIVIDAD) < tE.index(cap5.EPI_FIG51) and tE.index(cap5.LEYENDA_CONTEXTO) < tE.index(cap5.EPI_FIG53)
    and tE.index(cap5.LEYENDA_CU) < tE.index(cap5.EPI_FIG54))
chk('L', 'Guía: errores frecuentes de la NOTACIONES_CAPITULO_5.md §3 (mezclar, verbos, personas, carril, pasos)',
    len(guia.ERRORES_NOTACION) >= 5 and all(e[0] in tG for e in guia.ERRORES_NOTACION))
def _mezcla_detenida():
    """Prueba del propio control: una figura de contexto que intenta dibujar un actor tiene que detenerse."""
    L = G.Lienzo('prueba_mezcla', 'contexto')
    try:
        L.usar('cu_actor')
    except AssertionError:
        return True
    finally:
        G.REGISTRO.pop('prueba_mezcla', None)
    return False
chk('L', 'Prueba del control: una figura que mezcla notaciones detiene la generación', _mezcla_detenida())
chk('L', 'La Figura 2.1 del lote 1 se declara esquema didáctico, no UML', 'no usaba UML' in tE)

# ---------------------------------------------------------------- M. trazabilidad entre capítulos y cambio
tr_req = [t[0] for t in R.TRAZABILIDAD]
chk('M', 'Matriz: una fila por cada requisito del catálogo v1.0, en orden', tr_req == [r[0] for r in R.CATALOGO])
chk('M', 'Matriz: necesidades y reglas citadas existen', all(n in D.NECESIDAD for t in R.TRAZABILIDAD for n in t[1]) and all(rn in dict(D.REGLAS) for t in R.TRAZABILIDAD for rn in t[2]))
chk('M', 'Matriz: cada requisito tiene necesidad o regla (ninguno sin origen)', all(t[1] or t[2] for t in R.TRAZABILIDAD))
chk('M', 'Matriz: las reglas coinciden con las fuentes del catálogo', all(set(t[2]) == {f for f in R.REQ[t[0]][3] if f.startswith('RN-')} for t in R.TRAZABILIDAD))
chk('M', 'Pruebas previstas: P-01 a P-15, una por requisito y la misma en la matriz',
    [p for p, _, _ in R.PRUEBAS] == ['P-%02d' % i for i in range(1, 16)] and all(R.PRUEBA[t[4]][0] == t[0] for t in R.TRAZABILIDAD)
    and sorted(r for _, r, _ in R.PRUEBAS) == sorted(R.REQ))
cus = {c for c, _ in MO.CASOS_USO['casos']}
def ref_ok(txt):
    for cu in re.findall(r'CU-\d\d', txt):
        if cu not in cus:
            return False
    for cu, paso in re.findall(r'(CU-\d\d), pasos? (\d+)', txt):
        if cu not in MO.PLANTILLA or int(paso) > len(MO.PLANTILLA[cu]['escenario']):
            return False
    for cu, fl in re.findall(r'(CU-\d\d), flujo (\d+[a-z])', txt):
        if cu not in MO.PLANTILLA or fl not in [a[0] for a in MO.PLANTILLA[cu]['alternativos']]:
            return False
    return True
chk('M', 'Matriz: cada caso de uso, paso y flujo citado existe en el Capítulo 5', all(ref_ok(t[3]) for t in R.TRAZABILIDAD))
chk('M', 'Cadena de Work §7.3 (N-02, RN-01, RF-04, P-07) igual en la Figura 4.2, la Figura 6.1 y la matriz',
    R.TRAZABILIDAD[3][:3] == ('RF-04', ['N-02'], ['RN-01']) and R.TRAZABILIDAD[3][4] == 'P-07' and 'flujo 5a' in R.TRAZABILIDAD[3][3]
    and [x[0].split()[-1] for x in G.CADENA_6] == ['N-02', 'RN-01', 'RF-04', 'CU-01', 'CA-04', 'P-07']
    and G.CADENA[0][0] == 'Necesidad\nN-02' and 'P-07' in G.CADENA[-1][0])
chk('M', 'P-07 dice lo mismo que la prueba de Work §7.3 (segunda solicitud, rechazo y sin segundo préstamo activo)',
    'segunda solicitud' in R.PRUEBA['P-07'][1] and 'rechazo' in R.PRUEBA['P-07'][1] and 'segundo préstamo activo' in R.PRUEBA['P-07'][1])
chk('M', 'Matriz del libro: criterio «a escribir por el equipo» exactamente en los requisitos con criterio de referencia docente',
    all((row[4] == 'a escribir por el equipo') == (R.DONDE_CA[row[0]][0] == 'DOCENTE') for row in cap6.TABLA_62[1]))
chk('M', 'SC-01: la regla RN-02 v1.1 sale de D3, RF-11 es nuevo y RF-05 conserva el texto («máximo vigente»)',
    R.REGLA_V11[0] == 'RN-02' and 'cinco' in R.REGLA_V11[1] and R.RF11[0] not in R.REQ and all(f in R.TEXTO for f in R.RF11[3])
    and 'máximo vigente' in R.REQ['RF-05'][2])
chk('M', 'Tabla 6.6: los elementos que cambian existen (reglas, requisitos, casos de uso, pruebas)',
    all(e in dict(D.REGLAS) or e in REQS or e in cus or e in R.PRUEBA for e, _, _ in R.CAMBIOS_V11))
chk('M', 'La solicitud SC-01 tiene todos los campos de 6.4 (quién, qué, por qué, afectados, impacto, decisión, versión)',
    [c for c, _ in R.SC01] == ['Identificador', 'Quién la pide y cuándo', 'Qué se pide', 'Por qué', 'Reglas afectadas', 'Requisitos afectados',
                               'Casos de uso afectados', 'Criterios y pruebas afectados', 'Impacto en el plan', 'Decisión', 'Nueva versión'])
chk('M', 'Incrementos de la Tabla 6.3: los requisitos citados existen y cada uno depende del anterior',
    all(set(re.findall(r'\b(?:RF|RNF|RS)-\d\d\b', f[3])) <= REQS for f in cap6.TABLA_63[1])
    and [f[4] for f in cap6.TABLA_63[1]] == ['—', 'Incremento 1', 'Incremento 2', 'Incremento 3'])
est = sum(float(f[2].replace(',', '.')) for f in cap6.TABLA_64[1])
chk('M', 'Tabla 6.4: tareas con responsable, estimación y dependencia; suman %.1f HC, como dice el Ejemplo resuelto 6.3' % est,
    all(f[1] and f[2] and f[3] for f in cap6.TABLA_64[1]) and ('%s HC' % ('%.1f' % est).replace('.', ',')) in ' '.join(cap6.EJ63[1]))
chk('M', 'Ejemplo resuelto 6.2: los requisitos con criterio «a escribir» que nombra son los de la matriz',
    all(r in ' '.join(cap6.EJ62[1]) for r in [x[0] for x in R.CORRESPONDENCIA if x[1] == 'DOCENTE']))

# ---------------------------------------------------------------- E. Algorítmica 3.º y Matemática Aplicada
todo = tE + '\n' + tG + '\n' + tC
chk('E', 'Sin personas ni títulos del caso de biblioteca de Algorítmica 3.º', not [p for p in D.PROHIBIDOS_ALGORITMICA if p in todo])
chk('E', 'Sin matriz ponderada ni pesos (ni «ponderad…», ni pesos 0,5 / 0,3 / 0,2)', not re.search(r'(?i)ponderad|\bpesos?\b', todo) and not re.search(r'0,5\D{0,30}0,3\D{0,30}0,2', todo))
chk('E', 'Sin diagramas de datos ni tablas de base de datos en el Capítulo 5 (solo se remite al Capítulo 7)',
    not re.search(r'(?i)entidad-relaci|diagrama de clases|clave primaria|clave foránea', todo))
chk('E', 'Sin PERT ni teoría de grafos en 6.3 (Work §3, Matemática Aplicada)', not re.search(r'(?i)\bPERT\b|\bgrafos?\b|camino crítico|ruta crítica', todo))

# ---------------------------------------------------------------- F. voseo y residuos
TUTEO = (r'\b(tienes|puedes|debes|sabes|quieres|eres|propón|tú)\b|(?:^\d+\.\s+|^\([a-h]\)\s+)'
         r'(Escribe|Lee|Completa|Indica|Explica|Clasifica|Responde|Anota|Marca|Elige|Justifica|Redacta|Compara|Describe|Identifica|Resume|Haz|Dibuja|Prioriza|Planifica|Señala)\b')
for nombre, tx in (('libro', tE), ('guía', tG), ('cuaderno', tC)):
    m = [x.group(0) for x in re.finditer(TUTEO, tx, re.M)]
    chk('F', '%s: sin formas de tuteo frecuentes' % nombre, not m, m[:5])
    chk('F', '%s: sin identificadores de modelos de IA' % nombre, not re.search(r'(?i)\b(claude|gpt|gemini|opus|sonnet|codex|chatgpt)\b', tx))
    chk('F', '%s: sin símbolos ✓ ✗' % nombre, not re.search('[✓✔✗✘✅]', tx))
    chk('F', '%s: sin marcadores sin reemplazar' % nombre, '§§' not in tx and 'TODO' not in tx and 'XXX' not in tx)
    chk('F', '%s: no se declara apto ni aprobado para la venta' % nombre,
        not re.search(r'(?i)\bapto\b|aprobado para (la )?(venta|comercializ)', tx.replace('No está aprobado para la venta', '')))

# ---------------------------------------------------------------- G. maqueta (DOCX)
for p in ('helpers.py', 'maqueta.py'):
    a = hashlib.sha256(open(os.path.join(ADS, 'lote_1', 'build', p), 'rb').read()).hexdigest()
    b = hashlib.sha256(open(os.path.join(ADS, 'fase2_pilotos', 'A2_capitulo4_v2_1', 'build', p), 'rb').read()).hexdigest()
    chk('G', '%s del lote 1 (el que se importa) idéntico al del Capítulo 4 v2.1' % p, a == b and os.path.dirname({'helpers.py': O.H, 'maqueta.py': O.M}[p].__file__).endswith(os.path.join('lote_1', 'build')))
tw = int(O.Cm(O.ANCHO).twips)
for nombre, d in (('libro', dE), ('guía', dG), ('cuaderno', dC)):
    s = d.sections[0]
    chk('G', '%s: oficio 21,59 × 33,02 cm, márgenes O3 y espejados' % nombre,
        abs(s.page_width.cm - 21.59) < 0.01 and abs(s.page_height.cm - 33.02) < 0.01 and abs(s.left_margin.cm - 2.5) < 0.01
        and abs(s.right_margin.cm - 1.8) < 0.01 and abs(s.top_margin.cm - 1.5) < 0.01 and d.settings.element.find(qn('w:mirrorMargins')) is not None)
    filas = [r for t in d.tables for r in t.rows]
    chk('G', '%s: %d filas de tabla, todas «no dividir»' % (nombre, len(filas)), all(r._tr.trPr is not None and r._tr.trPr.find(qn('w:cantSplit')) is not None for r in filas))
    multi = [t for t in d.tables if len(t.columns) > 1]
    malas = [i for i, t in enumerate(multi) if abs(sum(int(g.get(qn('w:w'))) for g in t._tbl.tblGrid) - tw) > 15]
    chk('G', '%s: %d tablas de varias columnas al ancho útil (%s cm)' % (nombre, len(multi), str(O.ANCHO).replace('.', ',')), not malas, malas)
    sin_enc = [t.rows[0].cells[0].text[:30] for t in multi if t.rows[0]._tr.trPr.find(qn('w:tblHeader')) is None]
    chk('G', '%s: toda tabla de datos repite su encabezado (las únicas sin encabezado son las fichas)' % nombre, all(x.startswith('Capacidad') for x in sin_enc), sin_enc)

# ---------------------------------------------------------------- H. maqueta (PDF)
def lineas(pg):
    return [l.strip() for l in pg.get_text().split('\n') if l.strip() and not re.match(r'^Página \d+ de \d+$', l.strip())]
for nombre, pdf in (('libro', pE), ('guía', pG), ('cuaderno', pC)):
    vacias = [i + 1 for i, pg in enumerate(pdf) if not lineas(pg)]
    chk('H', '%s: ninguna página en blanco' % nombre, not vacias, vacias)
    sueltos = [i + 1 for i, pg in enumerate(pdf) if lineas(pg) and lineas(pg)[0].startswith('_' * 20)]
    chk('H', '%s: ninguna página empieza con renglones de respuesta sueltos' % nombre, not sueltos, sueltos)
for C in CAPS:
    pags = [i + 1 for i, pg in enumerate(pE) if C.EVAL_TITULO in lineas(pg)[:2]]
    sig = pags and pags[0] < len(pE) and lineas(pE[pags[0]])
    entra = pags and (pags[0] == len(pE) or (sig and (sig[0].startswith('Capítulo') or sig[0].startswith('Unidad'))))
    chk('H', '%s: en página propia y entera en esa página (pág. %s)' % (C.EVAL_TITULO, pags), len(pags) == 1 and entra)
    pp = [i + 1 for i, pg in enumerate(pE) if lineas(pg)[:1] and lineas(pg)[0].startswith(C.PRACTICA_TITULO[:25])]
    chk('H', '%s: empieza en página nueva' % C.PRACTICA_TITULO.split(' —')[0], len(pp) == 1)
    pc = [i + 1 for i, pg in enumerate(pC) if lineas(pg)[:1] and lineas(pg)[0].startswith(C.PRACTICA_TITULO[:25])]
    chk('H', 'Cuaderno: %s empieza en página nueva' % C.PRACTICA_TITULO.split(' —')[0], len(pc) == 1)
paginas = res_build['paginas_libro']
chk('H', 'El índice ocupa una sola página y la siguiente es la apertura de la Unidad II', lineas(pE[2])[0] == B.UNIDAD_II)
recalc, _ = K.paginas_de_titulos(EST + '.pdf', list(paginas), desde=2)
chk('H', 'Índice estable: las %d entradas apuntan a la página donde está su título' % len(paginas), recalc == paginas and len(paginas) == len(B.entradas_indice()))
# Extensión de referencia: el lote 1 en O3 tiene 49 págs. para preliminares y tres capítulos; la Unidad II, sin
# preliminares, se espera entre 36 y 52; la guía, entre 14 y 22 (lote 1 en O3: 14); el cuaderno, entre 6 y 12.
chk('H', 'Extensión del libro: %d págs. (referencia: 36 a 52)' % len(pE), 36 <= len(pE) <= 52)
chk('H', 'Extensión de la guía: %d págs. (referencia: 14 a 22)' % len(pG), 14 <= len(pG) <= 22)
chk('H', 'Extensión del cuaderno: %d págs. (referencia: 6 a 12)' % len(pC), 6 <= len(pC) <= 12)

# ---------------------------------------------------------------- I. evaluaciones, prácticas e instrumentos
claves = res_build['claves']
semillas_l1 = ['ADS3-L1-C1', 'ADS3-L1-C2-i', 'ADS3-L1-C3']
for C in CAPS:
    ks = [c for c in claves if c[0] == C.EVAL_ID]
    mc = [c[2] for c in ks if c[2] in 'abcd' and len(c[2]) == 1]
    chk('I', '%s: %d ítems con clave; opción múltiple con claves %s' % (C.EVAL_TITULO, len(ks), ''.join(mc)), len(ks) == len(C.EVAL_ITEMS) and len(set(mc)) >= 3)
    pts = sum(1 for it in C.EVAL_ITEMS if it[0] in ('mc', 'vf')) + sum(it[3] for it in C.EVAL_ITEMS if it[0] == 'ab')
    chk('I', '%s: suma 15 puntos, como dice su remisión' % C.EVAL_TITULO, pts == 15 and 'Puntaje total: 15 puntos' in C.EVAL_PUNTAJE)
    chk('I', '%s: la guía trae su clave y los criterios de los ítems abiertos' % C.EVAL_TITULO, ('Clave de la %s' % C.EVAL_TITULO.lower()) in tG)
    chk('I', 'Capítulo %d: práctica en tres niveles' % C.NUM, [n.split(' —')[0] for n, _ in C.NIVELES] == ['Nivel 1', 'Nivel 2', 'Nivel 3'])
semillas = [C.EVAL_ID for C in CAPS]
chk('I', 'Semillas de evaluación distintas entre sí y de las del lote 1 (%s)' % ', '.join(semillas), len(set(semillas + semillas_l1)) == 6)
chk('I', 'Capítulo 4 conserva la semilla del piloto (ADS3-C4v2-2) y sus claves', cap4.EVAL_ID == 'ADS3-C4v2-2' == PIL.EVAL_ID and cap4.EVAL_ITEMS == PIL.EVAL_ITEMS)
chk('I', 'Prácticas 5 y 6 con el caso Mbokaja y datos que no están en los ejemplos (MP, MS, MR, MC)',
    all('Mbokaja' in ' '.join(C.PRACTICA_INTRO) for C in (cap5, cap6))
    and not re.search(r'\bM[PSRC]-\d', ' '.join(json.dumps([getattr(C, n) for n in dir(C) if n.startswith(('EJ', 'S5', 'S6', 'TABLA'))], ensure_ascii=False) for C in (cap5, cap6))))
chk('I', 'Situaciones nuevas de la Práctica 5 (libro dañado, turno mañana) ausentes de los ejemplos resueltos',
    not re.search(r'(?i)hojas arrancadas|turno mañana', json.dumps([getattr(cap5, n) for n in dir(cap5) if n.startswith('EJ')], ensure_ascii=False)))
chk('I', 'Rúbrica del H2: cuatro niveles de Work, cinco criterios con descriptores y evidencia del equipo e individual',
    len(cap6.RUBRICA_H2) == 5 and all(len(r[2]) == 4 and r[3] and r[4] for r in cap6.RUBRICA_H2)
    and all(('%s (%s)' % (r[0], r[1])) in tG for r in cap6.RUBRICA_H2))
chk('I', 'El H2 evalúa requisitos (C3), modelos funcionales (C4), solicitud de cambio (C3, C7) y dossier (C2, C7)',
    [r[1] for r in cap6.RUBRICA_H2] == ['C3', 'C4', 'C4, C3', 'C3, C7', 'C2, C7'])
chk('I', 'Proporción 60/40 rotulada como propuesta editorial y ausencia de evidencia registrada aparte',
    '60 %' in tE and '40 %' in tE and 'propuesta editorial' in tE and 'ausencia de evidencia' in tG.lower() and 'Si falta una evidencia' in tE)
chk('I', 'Guía: planilla de registro del H2 y banco de necesidades nuevas que no están en el dossier, el buzón ni D3',
    'Planilla de registro del H2' in tG and len(guia.BANCO_H2) >= 3
    and not any(b[1][:30] in t for b in guia.BANCO_H2 for t in R.TEXTO.values()))
chk('I', 'Guía: secuencia en HC de los Encuentros 7 a 12', all(('Encuentro %d' % n) in tG for n in range(7, 13)) and tG.count('aprox.') >= 18)
chk('I', 'Guía: claves de todas las actividades de los Encuentros 7 a 12',
    all(len(k) == sum(len(items) for _, items in acts) for k, acts in ((guia.GP.CLAVES_E7, cap4.ACT_E7), (guia.GP.CLAVES_E8, cap4.ACT_E8),
        (guia.CLAVES_E9, cap5.ACT_E9), (guia.CLAVES_E10, cap5.ACT_E10), (guia.CLAVES_E11, cap6.ACT_E11), (guia.CLAVES_E12, cap6.ACT_E12))))
chk('I', 'Guía: claves y lista de cotejo de las tres prácticas', all(('Lista de cotejo de la Práctica %d' % n) in tG for n in (4, 5, 6))
    and len(guia.CLAVES_P4) == 4 and len(guia.CLAVES_P5) == 3 and len(guia.CLAVES_P6) == 3)
chk('I', 'Calendario de 28 encuentros: recuadro en cada capítulo y tabla en la guía', tE.count('Si el calendario se acorta') == 3 and 'Unidad II en 36 y en 28 encuentros' in tG)
chk('I', 'Capítulo 5 en 28 encuentros: se conserva el proceso actual simplificado de 5 a 7 pasos (adenda v2.1)',
    'de 5 a 7 pasos' in ' '.join(cap5.CALENDARIO[1]) and 5 <= sum(1 for n in MO.PROCESO_ACTUAL['nodos'] if n[1] == 'accion') <= 7)
chk('I', 'Cuaderno: las tres prácticas con el mismo texto que el libro',
    all(t in tC for C in CAPS for t in [C.PRACTICA_TITULO, C.PRACTICA_CONTROL] + C.PRACTICA_INTRO + [i.split('\n')[0] for _, items in C.NIVELES for i in items])
    and not any(C.EVAL_TITULO in tC for C in CAPS))

# ---------------------------------------------------------------- J. numeración
EXTERNAS = {'Tabla 3.5', 'Tabla 3.6', 'Figura 1.1', 'Figura 2.1'}   # del lote 1, citadas a propósito
for C in CAPS:
    n = C.NUM
    for tipo in ('Tabla', 'Figura', 'Ejemplo resuelto'):
        sec = [int(m.group(1)) for m in re.finditer(r'^%s %d\.(\d+)(?: —|:)' % (tipo, n), tE, re.M)]
        chk('J', 'Capítulo %d: %s %d.1 a %d.%d correlativas, sin repetir' % (n, tipo.lower(), n, n, len(sec)), sec == list(range(1, len(sec) + 1)) and len(sec) >= 2, sec)
    for tipo, pat in (('tabla', r'^Tabla %d\.(\d+) —'), ('figura', r'^Figura %d\.(\d+) —'), ('ejemplo resuelto', r'^Ejemplo resuelto %d\.(\d+):')):
        refs = set(int(x) for x in re.findall(r'%s %d\.(\d+)' % (tipo.capitalize(), n), tE + tG))
        defs = set(int(x) for x in re.findall(pat % n, tE, re.M))
        chk('J', 'Capítulo %d: toda %s mencionada existe' % (n, tipo), refs <= defs, refs - defs)
otras = set(re.findall(r'\b(?:Tabla|Figura) [1-3]\.\d+', tE + tG))
chk('J', 'Las referencias a tablas y figuras del lote 1 son solo las previstas (%s)' % ', '.join(sorted(otras)), otras <= EXTERNAS, otras - EXTERNAS)

print('\nResultado por bloque:')
for b, (o, f) in RES.items():
    print('  %s: %d correctos, %d fallos' % (b, o, f))
print('Resultado: %d correctos, %d fallos' % (ok, fallos))
json.dump({'correctos': ok, 'fallos': fallos, 'bloques': RES, 'controles': DETALLE},
          open(os.path.join(L2, 'verificacion_resultado.json'), 'w', encoding='utf8'), ensure_ascii=False, indent=1)
sys.exit(1 if fallos else 0)
