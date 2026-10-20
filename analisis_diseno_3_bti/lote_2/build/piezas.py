# -*- coding: utf-8 -*-
"""Piezas de armado del lote 2 que no están en la maqueta del lote 1: portada del lote, registros de la Unidad II,
extracto del dossier, tablas de leyenda con los símbolos dibujados y plantillas de casos de uso.
Cada capítulo se arma con funciones `capituloN(doc, figs)` y `practicaN(doc)`, para que el tomo anual pueda llamarlas
igual que las del lote 1 (lote_1/build/construir.py)."""
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

import oficio as O
H, M, K = O.H, O.M, O.K
import requisitos as R
import modelos as MO
import mbokaja_u2 as MB
import diagramas as G
D = R.D

VERSION = 'v1.0 (lote 2, para auditoría) · oficio O3'


def portada(doc, tipo_doc, estado, linea='LOTE EDITORIAL 2 — UNIDAD II (CAPÍTULOS 4 A 6)'):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    H.sombrear_celda(c, H.ID['DARK'])
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('\nAnálisis y Diseño de Sistemas\n3.er Curso · Bachillerato Técnico en Informática\n')
    r.bold = True; r.font.size = Pt(22); r.font.color.rgb = RGBColor.from_string('FFFFFF')
    H.cant_split(t)
    doc.add_paragraph()
    H.p_texto(doc, 'Del problema al prototipo web', bold=True, size=14, color=H.ID['DARK'], align='center')
    H.p_texto(doc, linea, bold=True, size=12, color=H.ID['MED'], align='center')
    H.p_texto(doc, tipo_doc, size=11.5, align='center')
    H.p_texto(doc, 'Plan Optativo · propuesta para el programa institucional · Biblioteca Comercial 2026', size=11, align='center')
    H.p_texto(doc, 'Proyecto conductor: Sistema de Gestión de Biblioteca Escolar', bold=True, size=11.5,
              color=H.ID['MED'], align='center')
    doc.add_paragraph()
    H.caja(doc, 'Estado de este documento', estado, tipo='neutra')
    H.p_texto(doc, 'Versión %s' % VERSION, size=10, align='center', antes=6)
    H.p_texto(doc, 'República del Paraguay · 2026', bold=True, size=12, align='center', antes=6)
    H.salto_pagina(doc)


def registro(doc, cod, lead=None):
    """Caja de un registro del dossier (lote 1 o Unidad II) con todas sus líneas, con el formato de comun.registro."""
    r = R.REGISTRO[cod]
    if lead:
        H.keep_next(H.p_texto(doc, lead))
    lineas = ['%s  %s%s' % (c, (hab + ': ') if hab else '', txt) for c, hab, txt in r['lineas']]
    H.caja(doc, 'Registro %s — %s (caso ficticio)' % (cod, r['titulo']), lineas, tipo='referencia', nota=r['contexto'])


def extracto(doc, C):
    doc.add_heading(C.DOSSIER_TITULO, level=2)
    H.p_texto(doc, 'Cada fragmento lleva un código: E1-L4 es la línea 4 de la entrevista E1. Las reglas RN se acordaron '
                   'en clase. Son las líneas del dossier del Capítulo 2 que usa este capítulo.', size=10)
    for reg, titulo, frs in C.DOSSIER:
        H.caja(doc, '%s — %s' % (reg, titulo), ['%s  %s' % (c, t) for c, t in frs], tipo='referencia')
    H.caja(doc, C.REGLAS_TITULO, ['%s  %s' % (c, t) for c, t in C.REGLAS] + [C.REGLAS_NOTA], tipo='referencia')


def codigos(doc, rotulo, pares, anchos=(1.8, 14.8), enc=('Código', 'Texto')):
    return M.tabla(doc, (list(enc), [[c, t] for c, t in pares]), list(anchos), rotulo=rotulo)


def tabla_leyenda(doc, notacion, rotulo, figs):
    """Tabla de símbolos de una notación: la primera columna lleva el símbolo dibujado con la misma función que las
    figuras (diagramas.simbolo)."""
    sims = [s for s, v in G.SIMBOLOS.items() if v[0] == notacion]
    filas = [['', G.SIMBOLOS[s][1], G.SIMBOLOS[s][2], G.SIMBOLOS[s][3]] for s in sims]
    t = M.tabla(doc, (['Símbolo', 'Nombre', 'Significado', 'Regla de uso'], filas), [2.4, 3.0, 5.4, 5.8], rotulo=rotulo)
    for fila, s in zip(t.rows[1:], sims):
        p = fila.cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(figs['sim:' + s], width=Cm(1.75))
    return t


def plantilla(doc, cu):
    p = MO.PLANTILLA[cu]
    bloques = ['Actor: ' + p['actor'], 'Precondición: ' + p['precondicion'],
               'Escenario principal:\n' + '\n'.join('%d. %s' % (i, s) for i, s in enumerate(p['escenario'], 1)),
               'Flujos alternativos:\n' + '\n'.join('%s. %s' % a for a in p['alternativos']),
               'Requisitos que cubre: ' + ', '.join(p['requisitos'])]
    H.caja(doc, 'Caso de uso %s — %s' % (p['id'], p['nombre']), bloques, tipo='neutra')


def catalogo_mbokaja(doc):
    M.tabla(doc, (['Id.', 'Tipo', 'Enunciado', 'Fuentes'], [list(x) for x in MB.CATALOGO]), [1.6, 2.8, 9.4, 2.8],
            rotulo=MB.CATALOGO_TITULO)
