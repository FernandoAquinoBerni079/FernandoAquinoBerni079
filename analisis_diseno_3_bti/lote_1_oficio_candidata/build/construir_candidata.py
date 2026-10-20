# -*- coding: utf-8 -*-
"""Versión CANDIDATA EDITORIAL del lote 1 en oficio, perfil O3 (decisión de Fernando del 20-oct-2026).
No es un libro aprobado para comercialización.

Uso: python3 lote_1_oficio_candidata/build/construir_candidata.py
Salida: lote_1_oficio_candidata/ADS_3BTI_LOTE1_ESTUDIANTE_OFICIO_O3_CANDIDATA.docx/.pdf y …_GUIA_DOCENTE_…

Parte de `lote_1_oficio/build/construir_oficio.py` (versión de prueba, que no se modifica) y agrega:
1. ajuste de anchos por palabra: ninguna palabra de una celda queda más ancha que su columna (se mide con la métrica
   de Carlito, igual a Calibri, al tamaño y peso de la celda); el ancho que falta se toma de la columna con más
   holgura, sin cambiar el ancho total de la tabla;
2. control de filas en el ciclo de reconstrucción: cada fila de cada tabla de 2 o más columnas del DOCX tiene que
   aparecer, en orden, en las tablas que el PDF muestra (detección de tablas de PyMuPDF); si falta alguna, el rótulo
   de esa tabla abre hoja nueva y se reconstruye.
Importa las fuentes de `lote_1/build/` sin modificarlas. No escribe nada dentro de `lote_1/` ni de `lote_1_oficio/`.
"""
import os, re, sys, json, importlib
sys.dont_write_bytecode = True   # no deja cachés dentro de lote_1/

AQUI = os.path.dirname(os.path.abspath(__file__))
OFI = os.path.dirname(AQUI)
ADS = os.path.dirname(OFI)
FUENTE = os.path.join(ADS, 'lote_1', 'build')
sys.path.insert(0, AQUI)
sys.path.insert(1, FUENTE)

import perfiles as PF
import controles_impresion as CI
from docx.shared import Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

clave = 'O3'   # perfil adoptado
P = PF.PERFILES[clave]
ANCHO = PF.ancho_util(P)
F = ANCHO / PF.ANCHO_A4
SALIDA = OFI
os.makedirs(SALIDA, exist_ok=True)

import helpers as H
import maqueta as M
import comun as K
import diagramas
import construir as C
import guia
import cap3

assert os.path.dirname(H.__file__) == FUENTE and os.path.dirname(C.__file__) == FUENTE, 'se importó otra copia de las fuentes'

# ---------------------------------------------------------------- geometría
_new_doc = H.new_doc


def new_doc():
    doc = _new_doc()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(P['papel'][0]), Cm(P['papel'][1])
    s.left_margin, s.right_margin = Cm(P['interior']), Cm(P['exterior'])
    s.top_margin, s.bottom_margin = Cm(P['superior']), Cm(P['inferior'])
    s.header_distance = s.footer_distance = Cm(PF.ENCABEZADO_PIE)
    if P['espejo']:   # márgenes simétricos: en Word, «izquierdo» pasa a ser el interior
        st = doc.settings.element
        if st.find(qn('w:mirrorMargins')) is None:
            st.insert(0, OxmlElement('w:mirrorMargins'))
    return doc


H.new_doc = new_doc

_fijar = M.fijar_anchos


AJUSTES = []   # (primera celda, columna, ancho previo, ancho nuevo) para el informe
_FUENTES = {}
CARLITO = {False: '/usr/local/share/fonts/carlito/Carlito-Regular.ttf', True: '/usr/local/share/fonts/carlito/Carlito-Bold.ttf'}
MARGEN_CELDA = 2 * 0.19 + 0.06   # márgenes izquierdo y derecho de la celda (Word y LibreOffice: 0,19 cm) y holgura


def _ancho_cm(palabra, pt, negrita):
    from PIL import ImageFont
    clave_f = (pt, negrita)
    if clave_f not in _FUENTES:
        _FUENTES[clave_f] = ImageFont.truetype(CARLITO[negrita], size=int(round(pt * 20)))
    return _FUENTES[clave_f].getlength(palabra) / 20 / 72 * 2.54


def _minimo_columna(t, j):
    m = 0.0
    for fila in t.rows:
        if j >= len(fila.cells):
            continue
        for p in fila.cells[j].paragraphs:
            for r in p.runs:
                pt = r.font.size.pt if r.font.size else 10.5
                for w in re.findall(r'\S+', r.text):
                    m = max(m, _ancho_cm(w, pt, bool(r.bold)))
    return m + MARGEN_CELDA


def fijar_anchos(t, anchos):
    a = [round(x * F, 3) for x in anchos]
    minimos = [_minimo_columna(t, j) for j in range(len(a))]
    for j in range(len(a)):
        falta = minimos[j] - a[j]
        if falta > 0:
            k = max(range(len(a)), key=lambda c: a[c] - minimos[c] if c != j else -1e9)
            assert a[k] - minimos[k] >= falta, 'no hay holgura para la columna %d de «%s»' % (j, t.rows[0].cells[0].text[:30])
            AJUSTES.append((t.rows[0].cells[0].text[:40], j + 1, round(a[j], 2), round(minimos[j], 2)))
            a[k] = round(a[k] - falta, 3); a[j] = round(minimos[j], 3)
    _fijar(t, a)


M.fijar_anchos = fijar_anchos

_rellenar = K.rellenar_indice
TAB_A4 = str(int(Cm(PF.ANCHO_A4).twips))


def rellenar_indice(doc, entradas, paginas):
    _rellenar(doc, entradas, paginas)
    for tab in doc.element.body.iter(qn('w:tab')):
        if tab.get(qn('w:pos')) == TAB_A4:
            tab.set(qn('w:pos'), str(int(Cm(ANCHO).twips)))


K.rellenar_indice = rellenar_indice

_finalizar = H.finalizar


def finalizar(doc, ruta):
    """Quita el párrafo vacío de separación que deja `maqueta.tabla` cuando lo sigue un título con salto de página
    o el final del documento: ahí no separa nada y, si la tabla termina justo al pie, pasa solo a la hoja
    siguiente y deja una hoja en blanco (se vio en la guía con los perfiles O2 y O3). No cambia ningún texto."""
    body = doc.element.body
    for p in list(body.iterchildren(qn('w:p'))):
        vacio = (not ''.join(t.text or '' for t in p.iter(qn('w:t'))).strip() and p.find('.//' + qn('w:drawing')) is None
                 and p.find('.//' + qn('w:br')) is None and p.find('.//' + qn('w:fldChar')) is None)
        sig = p.getnext()
        if vacio and sig is not None and (sig.tag == qn('w:sectPr') or sig.find('.//' + qn('w:pageBreakBefore')) is not None):
            body.remove(p)
    # Tablas que LibreOffice dejó con celdas sin imprimir en la vuelta anterior: su rótulo abre hoja nueva.
    for t in body.iterchildren(qn('w:tbl')):
        prev = t.getprevious()
        rot = ''.join(x.text or '' for x in prev.iter(qn('w:t'))) if prev is not None and prev.tag == qn('w:p') else ''
        if rot and rot in HOJA_NUEVA:
            ppr = prev.find(qn('w:pPr'))
            if ppr is None:
                ppr = OxmlElement('w:pPr'); prev.insert(0, ppr)
            if ppr.find(qn('w:pageBreakBefore')) is None:
                ppr.insert(0, OxmlElement('w:pageBreakBefore'))
    return _finalizar(doc, ruta)


HOJA_NUEVA = set()   # rótulos de las tablas que deben empezar en hoja nueva


def tablas_con_perdida(ruta_docx, ruta_pdf, minimo=4):
    """Rótulos de las tablas con encabezado repetido a las que les faltan palabras en el PDF. Es la combinación
    «encabezado repetido» + «no dividir fila»: si la tabla empieza cerca del pie, LibreOffice imprime el encabezado
    y la primera fila y omite las demás (se vio en la Tabla 3.5 del PDF A4 original y en la guía con el perfil O2).
    Se cuenta el déficit de palabras del documento (DOCX menos PDF) que cae en las celdas de cada tabla; con menos de
    `minimo` palabras no se marca: son palabras partidas por el ancho de una celda angosta, no filas perdidas."""
    import re, collections, docx as _docx, pymupdf
    W = lambda t: re.findall(r'[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]{2,}', t)
    pdf = collections.Counter()
    for pg in pymupdf.open(ruta_pdf):
        pdf.update(W(pg.get_text()))
    d = _docx.Document(ruta_docx)
    doc = collections.Counter()
    for p in d.element.body.iter(qn('w:p')):
        doc.update(W(''.join(x.text or '' for x in p.iter(qn('w:t')))))
    deficit = doc - pdf
    puntajes = []
    for t in d.tables:
        trp = t.rows[0]._tr.trPr
        if trp is None or trp.find(qn('w:tblHeader')) is None:
            continue
        tw = collections.Counter()
        for tc in t._tbl.iter(qn('w:tc')):
            tw.update(W(''.join(x.text or '' for x in tc.iter(qn('w:t')))))
        puntaje = sum(min(c, deficit[w]) for w, c in tw.items())
        if puntaje >= minimo:
            prev = t._tbl.getprevious()
            puntajes.append((puntaje, ''.join(x.text or '' for x in prev.iter(qn('w:t')))))
    # una sola tabla por vuelta, la de mayor déficit: las palabras que pierde una tabla también figuran en otras
    return [max(puntajes)[1]] if puntajes else []


H.finalizar = finalizar

K.VERSION = K.VERSION + ' · candidata editorial en oficio O3; no aprobada para comercialización'
diagramas.AQUI = AQUI   # figuras regeneradas fuera de lote_1/
C.RUTA_EST = os.path.join(SALIDA, 'ADS_3BTI_LOTE1_ESTUDIANTE_OFICIO_O3_CANDIDATA.docx')
C.RUTA_GUIA = os.path.join(SALIDA, 'ADS_3BTI_LOTE1_GUIA_DOCENTE_OFICIO_O3_CANDIDATA.docx')

if __name__ == '__main__':
    riesgos = [(r[0], r[2], r[3]) for r in cap3.TABLA_34[1]]
    figs = diagramas.generar(riesgos)
    assert all(os.path.dirname(f).startswith(AQUI) for f in figs.values()), figs
    for vuelta in range(8):   # construir, detectar tablas con celdas no impresas, corregir y volver a construir
        claves, paginas, total = C.construir_estudiante(figs, SALIDA)
        guia.construir(C.RUTA_GUIA, claves, paginas, C.ESTADO)
        pdf_g = K.a_pdf(C.RUTA_GUIA, SALIDA)
        nuevas = set(tablas_con_perdida(C.RUTA_EST, C.RUTA_EST[:-5] + '.pdf')) | set(tablas_con_perdida(C.RUTA_GUIA, pdf_g))   # a lo sumo una por documento
        for dx, pdf in ((C.RUTA_EST, C.RUTA_EST[:-5] + '.pdf'), (C.RUTA_GUIA, pdf_g)):   # control de filas
            incompletas = [x['rotulo'] for x in CI.filas_de_tablas(dx, pdf) if x['filas_impresas'] < x['filas_docx'] and x['rotulo']]
            if incompletas:
                nuevas.add(incompletas[0])
        if not nuevas - HOJA_NUEVA:
            assert not nuevas, 'tablas que siguen perdiendo celdas: %s' % nuevas
            break
        HOJA_NUEVA.update(nuevas)
    else:
        raise AssertionError('no se pudo evitar la pérdida de celdas: %s' % HOJA_NUEVA)
    import pymupdf
    print('%s · %s · ancho útil %.2f cm · libro %d págs. · guía %d págs.' % (clave, P['nombre'], ANCHO, total, len(pymupdf.open(pdf_g))))
    print('tablas que empiezan en hoja nueva para no perder celdas:', sorted(HOJA_NUEVA))
    print('columnas ensanchadas para no partir palabras:', len(dict.fromkeys(AJUSTES)))
    json.dump({'perfil': clave, 'ancho_util': ANCHO, 'factor_tablas': F, 'tablas_en_hoja_nueva': sorted(HOJA_NUEVA),
               'columnas_ensanchadas': [{'tabla': a, 'columna': b, 'antes_cm': c, 'despues_cm': d} for a, b, c, d in dict.fromkeys(AJUSTES)],
               'paginas_libro': paginas, 'claves': claves},
              open(os.path.join(SALIDA, 'resultado_construccion.json'), 'w', encoding='utf8'), ensure_ascii=False, indent=1)
