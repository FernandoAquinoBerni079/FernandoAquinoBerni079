# -*- coding: utf-8 -*-
"""Controles de impresión: lo que está en el DOCX tiene que estar en el PDF, en su lugar y entero.

1. filas_de_tablas: cada fila de cada tabla de 2 o más columnas del DOCX se busca, en orden, entre las filas de las
   tablas que PyMuPDF detecta en el PDF (página por página, de arriba abajo). Se cuentan las filas impresas, las
   páginas que ocupa la tabla y los encabezados repetidos. Una celda coincide si su texto sin espacios ni guiones es
   igual (o tiene las mismas letras: la extracción de ligaduras puede cambiar el orden de dos letras).
2. parrafos_no_impresos: todo párrafo de 12 caracteres o más tiene que aparecer en el texto del PDF (en trozos de 12
   caracteres sin espacios; falta si falta más de la mitad).
3. palabras_partidas: toda palabra de 5 letras o más de una celda de tabla tiene que aparecer entera en el PDF
   tantas veces como en el DOCX; si no, la columna la está cortando.
"""
import re, collections
import docx, pymupdf
from docx.oxml.ns import qn

LETRAS = r'[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]'


def _norm(s):
    return re.sub(r'[\s\-­]+', '', s or '')


def _igual(a, b):
    a, b = _norm(a), _norm(b)
    return a == b or (len(a) > 3 and sorted(a) == sorted(b))


def _texto(el):
    return ''.join(x.text or '' for x in el.iter(qn('w:t')))


def _tablas_docx(dx):
    out = []
    for t in docx.Document(dx).element.body.iterchildren(qn('w:tbl')):
        filas = [[_texto(tc) for tc in tr.findall(qn('w:tc'))] for tr in t.findall(qn('w:tr'))]
        if max(len(f) for f in filas) >= 2:
            prev = t.getprevious()
            rot = _texto(prev) if prev is not None and prev.tag == qn('w:p') else ''
            out.append((rot, filas))
    return out


def filas_de_tablas(dx, pdf):
    pf = []
    for i, pg in enumerate(pymupdf.open(pdf)):
        for tb in sorted(pg.find_tables().tables, key=lambda x: x.bbox[1]):
            if tb.col_count >= 2:
                for r in tb.extract():
                    pf.append((i + 1, [c or '' for c in r]))
    k = 0; res = []
    for n, (rot, filas) in enumerate(_tablas_docx(dx)):
        enc = filas[0]; pags = set(); impresas = 0; repet = 0; faltan = []
        for fila in filas:
            j = k
            while j < len(pf) and not (len(pf[j][1]) == len(fila) and all(_igual(a, b) for a, b in zip(fila, pf[j][1]))):
                j += 1
            if j < len(pf):
                repet += sum(1 for p in pf[k:j] if len(p[1]) == len(enc) and all(_igual(a, b) for a, b in zip(enc, p[1])))
                impresas += 1; pags.add(pf[j][0]); k = j + 1
            else:
                faltan.append(fila[0][:40])
        res.append({'n': n + 1, 'rotulo': rot[:80], 'primera_celda': enc[0][:40], 'filas_docx': len(filas),
                    'filas_impresas': impresas, 'paginas': sorted(pags), 'encabezado_repetido': repet, 'filas_faltantes': faltan})
    return res


def parrafos_no_impresos(dx, pdf):
    P = re.sub(r'\s+', '', ' '.join(p.get_text() for p in pymupdf.open(pdf)))
    malas = []
    for el in docx.Document(dx).element.body.iter(qn('w:p')):
        txt = _texto(el); s = re.sub(r'\s+', '', txt)
        if len(s) >= 12:
            trozos = [s[i:i + 12] for i in range(0, len(s) - 11, 12)]
            if sum(x not in P for x in trozos) > len(trozos) * 0.5:
                malas.append(txt[:120])
    return malas


def palabras_partidas(dx, pdf):
    pal = collections.Counter()
    for pg in pymupdf.open(pdf):
        pal.update(re.findall(LETRAS + '+', pg.get_text()))
    d = docx.Document(dx); celdas = collections.Counter()
    for tc in d.element.body.iter(qn('w:tc')):
        celdas.update(w for w in re.findall(LETRAS + '+', _texto(tc)) if len(w) >= 5)
    # una palabra larga de celda que el PDF muestra menos veces que el DOCX está partida (o falta)
    todas = collections.Counter()
    for p in d.element.body.iter(qn('w:p')):
        todas.update(re.findall(LETRAS + '+', _texto(p)))
    return {w: todas[w] - pal[w] for w in celdas if pal[w] < todas[w]}
