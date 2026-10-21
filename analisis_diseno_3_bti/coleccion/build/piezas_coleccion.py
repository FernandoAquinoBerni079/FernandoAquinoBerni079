# -*- coding: utf-8 -*-
"""Piezas propias de la colección: lo único que el armado agrega a las fuentes de los lotes.

1. Portada y estado de cada documento anual (con la portada del lote 2, que usan los lotes 2 a 6, y otra línea).
2. Índice con páginas buscadas en orden: el índice anual ocupa varias páginas, así que cada título se busca a partir
   de la página del título anterior y después del propio índice (comun.paginas_de_titulos busca desde una página fija).
3. REEMPLAZOS: la lista cerrada de cambios de texto del libro del estudiante (remisiones a «lote» que pasan a nombrar
   la unidad, y la frase de alcance de la presentación). Cada uno con el texto exacto, el reemplazo y la cantidad de
   apariciones esperada; si la cantidad no coincide, el armado se detiene.
4. Títulos de parte de la guía anual y notas de integración (texto nuevo, que la lista de diferencias declara).
No modifica ninguna fuente: los cambios de texto se aplican sobre el documento armado."""
import re
import fuentes as FU

O, LOTES, GUIA, DIAG = FU.cargar()
H, M, K = O.H, O.M, O.K
Z2 = LOTES[2].Z   # lote_2/build/piezas.py: la portada de los lotes 2 a 6

VERSION = 'v0.1 de la colección (integración editorial, para auditoría) · oficio O3'
TITULOS_UNIDAD = {1: 'Unidad I — Comprender el problema y su contexto',
                  2: LOTES[2].UNIDAD_II, 3: LOTES[3].UNIDAD_III, 4: LOTES[4].UNIDAD_IV, 5: LOTES[5].UNIDAD_V,
                  6: LOTES[6].UNIDAD_VI}
ENCUENTROS = {1: '1 a 6', 2: '7 a 12', 3: '13 a 18', 4: '19 a 26', 5: '27 a 32', 6: '33 a 36'}

ESTADO_COMUN = [
    'Las capacidades C1 a C10 son propuestas para el programa institucional del Plan Optativo, tomadas del informe de '
    'Work. No son capacidades oficiales del MEC.',
    'Las pruebas del prototipo corrieron solo en el entorno de construcción (Linux). Windows y las computadoras escolares '
    'siguen pendientes de validación. El borrador de licencia de la Unidad VI está pendiente de revisión jurídica.',
]
ESTADO = {
    'libro': ['Tomo anual integrado desde las fuentes aprobadas de los seis lotes: 18 capítulos, seis unidades, 36 '
              'encuentros y 180 HC de la propuesta editorial. Versión para la auditoría de integración: no está aprobado '
              'para la venta ni para uso en aula; la aprobación final corresponde a Fernando.'] + ESTADO_COMUN,
    'guia': ['Guía docente anual con solucionario, integrada desde las guías aprobadas de los seis lotes. Versión para la '
             'auditoría de integración: no está aprobada para la venta ni para uso en aula; la aprobación final '
             'corresponde a Fernando.'] + ESTADO_COMUN,
    'cuaderno': ['Cuaderno anual de prácticas: las prácticas de los Capítulos 1 a 15 y 17, extraídas de las mismas fuentes '
                 'que el libro del estudiante, sin teoría ni evaluaciones. Versión para la auditoría de integración: no '
                 'está aprobado para la venta ni para uso en aula; la aprobación final corresponde a Fernando.',
                 'Cada práctica trabaja con datos que no aparecen en los ejemplos resueltos del libro. Tiene tres niveles que '
                 'usan los mismos datos: tu docente te indica en cuál trabajar. Si el trabajo es un diagrama o un boceto, '
                 'hacelo en una hoja aparte y adjuntala. Las tablas, figuras y códigos que se citan son los del libro del '
                 'estudiante.',
                 'Las Prácticas 9 y 13 son complementos editoriales al índice de Work: no reemplazan los hitos H3 y H4. Los '
                 'Capítulos 16 y 18 no tienen práctica: sus instrumentos están en el libro.'] + ESTADO_COMUN,
}
LINEA = 'TOMO ANUAL — UNIDADES I A VI (CAPÍTULOS 1 A 18)'
TIPO = {'libro': 'Libro del estudiante', 'guia': 'Guía docente anual con solucionario', 'cuaderno': 'Cuaderno anual de prácticas'}


def portada(doc, clave):
    anterior = Z2.VERSION
    Z2.VERSION = VERSION
    try:
        Z2.portada(doc, TIPO[clave], ESTADO[clave], linea=LINEA)
    finally:
        Z2.VERSION = anterior


# ---------------------------------------------------------------- índice con búsqueda en orden
def _lineas(pg):
    ls = [l.strip() for l in pg.get_text().split('\n')]
    return ls + [ls[k] + ' ' + ls[k + 1] for k in range(len(ls) - 1)]


LINEA_INDICE = re.compile(r'\.{6,}\s*(\d+|—)\s*$', re.M)


def fin_del_indice(pdf):
    """Última página (0-based) del índice: la última de las primeras páginas con líneas de puntos y número."""
    pags = [i for i, pg in enumerate(pdf) if i < 12 and LINEA_INDICE.search(pg.get_text())]
    return max(pags) if pags else 1


def ordenar_por_documento(doc, entradas):
    """Las entradas en el orden en que sus títulos aparecen en el documento. Un título de nivel 1 se busca después del
    título de nivel 1 anterior; uno de nivel 2 o 3, después del título de nivel 1 que lo contiene. Las entradas que no
    son títulos del documento (los grupos por unidad del cuaderno) quedan delante de la entrada siguiente. Devuelve
    (entradas ordenadas, [entradas que cambiaron de lugar])."""
    titulos = [p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
    pos, cap, nivel1 = [], 0, 0
    for nivel, t in entradas:
        desde = nivel1 if nivel == 1 else cap
        i = next((k for k in range(desde, len(titulos)) if titulos[k] == t), None)
        if i is not None and nivel == 1:
            cap = nivel1 = i
        pos.append(i)
    for k in range(len(pos) - 1, -1, -1):
        if pos[k] is None:
            pos[k] = (pos[k + 1] - 0.5) if k + 1 < len(pos) and pos[k + 1] is not None else float('inf')
    orden = sorted(range(len(entradas)), key=lambda k: (pos[k], k))
    movidas = [entradas[k] for j, k in enumerate(orden) if k != j]
    return [entradas[k] for k in orden], movidas


def paginas_en_orden(ruta_pdf, entradas):
    """Página (1-based) de cada entrada [(nivel, título)], buscada como línea propia y en orden: después del índice y
    desde la página de la entrada anterior, así un título repetido (por ejemplo, «Calendario con menos encuentros» en
    cada parte de la guía) toma la página de su aparición. Las entradas con página None son las que no se encontraron.
    Devuelve (lista de páginas alineada con las entradas, total de páginas)."""
    import pymupdf
    pdf = pymupdf.open(ruta_pdf)
    lineas = [_lineas(pg) for pg in pdf]
    desde = fin_del_indice(pdf) + 1
    res = []
    for _, t in entradas:
        pag = None
        for i in range(desde, len(pdf)):
            if t in lineas[i]:
                pag = i + 1; desde = i
                break
        res.append(pag)
    return res, len(pdf)


def rellenar_indice(doc, entradas, paginas):
    """Como comun.rellenar_indice, con la tabulación al ancho útil del perfil O3 (como oficio.rellenar_indice), pero con
    la página de cada entrada en una lista alineada: los títulos repetidos no comparten página."""
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
    from docx.shared import Pt, Cm
    marca = [p for p in doc.paragraphs if p.text == K.MARCA_INDICE][0]
    ancla = marca
    for (nivel, texto), pag in zip(entradas, paginas or [None] * len(entradas)):
        np_ = OxmlElement('w:p')
        ancla._p.addnext(np_)
        p = Paragraph(np_, ancla._parent)
        pf = p.paragraph_format
        pf.space_after = Pt(1.5); pf.space_before = Pt(3 if nivel == 1 else 0)
        pf.left_indent = Cm({1: 0, 2: 0.5, 3: 1.0}[nivel])
        pf.tab_stops.add_tab_stop(Cm(O.ANCHO), alignment=2, leader=1)
        r = p.add_run(texto); r.font.size = Pt(9.8 if nivel == 1 else 9.3); r.bold = (nivel == 1)
        r2 = p.add_run('\t%s' % (pag or '—')); r2.font.size = Pt(9.5)
        ancla = p
    marca._p.getparent().remove(marca._p)


# ---------------------------------------------------------------- lista cerrada de reemplazos (libro del estudiante)
# (id, capítulo o pieza, texto exacto, reemplazo, apariciones, motivo)
REEMPLAZOS = [
    ('R-01', 'Preliminares (Presentación y alcance)',
     'Esta edición de control contiene los preliminares y la Unidad I (Capítulos 1 a 3, Encuentros 1 a 6).',
     'Este tomo anual contiene los preliminares y las seis unidades (Capítulos 1 a 18, Encuentros 1 a 36).', 1,
     'Alcance del documento: el libro del lote 1 contenía solo la Unidad I.'),
    ('R-02', 'Capítulo 11 (ficha de un bloque de código)', '(lote 3 v1.1)', '(diseño de la Unidad III, v1.1)', 1,
     'Referencia editorial: la versión del diseño se nombra por su unidad.'),
    ('R-03', 'Capítulo 14 (rótulo de la Tabla 14.6) y Capítulo 16 (dos párrafos)', 'prototipo del lote 4', 'prototipo de la Unidad IV', 3,
     'Referencia editorial: el prototipo se nombra por la unidad que lo construye.'),
    ('R-04', 'Capítulo 16', 'la del lote 4 queda como estaba', 'la de la Unidad IV queda como estaba', 1,
     'Referencia editorial.'),
    ('R-05', 'Apertura de la Unidad V (caja «Las copias descartables»)', 'la del lote anterior no cambian', 'la de la Unidad IV no cambian', 1,
     'Referencia editorial: «lote anterior» no tiene sentido en el tomo anual.'),
    ('R-06', 'Capítulo 14 (tabla de resultados)', 'Referencia docente del lote 4;', 'Referencia docente de la Unidad IV;', 1,
     'Referencia editorial.'),
    ('R-07', 'Capítulo 14 (tabla de resultados)', 'la referencia del lote 4 (falla)', 'la referencia de la Unidad IV (falla)', 1,
     'Referencia editorial.'),
]


def _bloques_texto(doc):
    """Párrafos del cuerpo y de las celdas, en orden de documento."""
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn
    for p in doc.element.body.iter(qn('w:p')):
        yield Paragraph(p, doc)


def aplicar_reemplazos(doc, lista=REEMPLAZOS):
    """Aplica cada reemplazo dentro de un tramo de texto (run), para no cambiar ningún formato. Devuelve
    [(id, apariciones)] y se detiene si alguna cantidad no coincide con la esperada."""
    hechos = []
    for rid, _, viejo, nuevo, esperadas, _ in lista:
        n = 0
        for p in _bloques_texto(doc):
            if viejo not in p.text:
                continue
            tramos = [r for r in p.runs if viejo in r.text]
            assert tramos, '%s: «%s» está repartido en varios tramos de texto' % (rid, viejo)
            for r in tramos:
                n += r.text.count(viejo)
                r.text = r.text.replace(viejo, nuevo)
        assert n == esperadas, '%s: se esperaban %d apariciones de «%s» y hay %d' % (rid, esperadas, viejo, n)
        hechos.append((rid, n))
    return hechos


def reemplazar_texto(texto, lista=REEMPLAZOS):
    """Los mismos reemplazos, sobre un texto (lo usa la equivalencia para transformar la versión aprobada)."""
    for _, _, viejo, nuevo, _, _ in lista:
        texto = texto.replace(viejo, nuevo)
    return texto


# ---------------------------------------------------------------- guía anual: partes y notas de integración
def titulo_parte(n):
    caps = FU.CAPITULOS[n]
    return 'Parte %s — %s (Capítulos %d %s %d, Encuentros %s)' % (
        FU.UNIDADES[n], TITULOS_UNIDAD[n].split(' — ', 1)[1], caps[0], 'y' if len(caps) == 2 else 'a', caps[-1], ENCUENTROS[n])


ENCABEZADO_GUIA_LOTE = 'Guía docente — Lote %d'   # el título de nivel 1 con que empieza la guía de cada lote

NOTA_UNIDAD_V = ('Nota de integración: preparación del Encuentro 30 y conjunto reservado', [
    'Observaciones de la auditoría del lote 5, que se agregan acá porque el lote 5 está congelado:',
    '• El docente prepara de antemano las copias de evaluación con DR, las cuentas, las tareas impresas y las planillas.',
    '• Las 5 HC del Encuentro 30 se reparten entre las dos sesiones de uso, la revisión cruzada y el acta. Si falta '
    'equipamiento, se documenta una rotación, sin omitir evidencias esenciales.',
    '• El conjunto DR se guarda fuera de GitHub y de las carpetas del alumnado. Se recomienda agregar casos propios. Sus '
    'resultados privados no cuentan como auditados de forma independiente hasta que se examine ese paquete.'])

NOTA_UNIDAD_VI = ('Nota de integración: pruebas de código recibido de otros', [
    'Un paquete recibido de otro equipo, o de alguien ajeno a la clase, es código desconocido. Las pruebas funcionales de '
    'ese código requieren una máquina virtual o un equipo verdaderamente aislado: sin datos institucionales ni de '
    'personas, sin red y con una cuenta sin permisos de administrador.',
    'La lista cerrada de las herramientas del docente limita qué pruebas se ejecutan, pero no vuelve confiable su código: '
    'una prueba con el nombre permitido puede contener cualquier cosa. Por eso la lista cerrada no reemplaza el '
    'aislamiento. Primero se hace la inspección estática, que no ejecuta nada.',
    'Las herramientas se probaron solo en el entorno de construcción (Linux). En Windows y en los equipos escolares siguen '
    'pendientes de validación.'])

NOTAS_PARTE = {5: NOTA_UNIDAD_V, 6: NOTA_UNIDAD_VI}
