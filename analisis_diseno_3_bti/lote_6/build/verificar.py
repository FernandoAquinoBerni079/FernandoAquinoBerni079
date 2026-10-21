# -*- coding: utf-8 -*-
"""Controles automáticos del lote 6 (Unidad VI). Uso: python3 lote_6/build/verificar.py   (sale con 1 si hay fallos)
Con ADS_V6=<ruta privada de variantes_v6.json> y ADS_DR=<ruta privada de dr.json>, comprueba además que ninguna variante
reservada ni ningún código de DR estén en el repositorio ni en los documentos (sin imprimirlos).

Bloques de los lotes anteriores, extendidos a los Capítulos 17 y 18: A índice de Work · B capacidades · C trazabilidad de
fuentes · F voseo y residuos · G maqueta (DOCX) · H maqueta (PDF) · I evaluaciones e instrumentos · J numeración ·
L notaciones · T fragmentos de código · V pruebas y entorno · Q lotes 1 a 5 intactos.
Bloques nuevos (plan del lote 6 y decisiones de Fernando del 21-oct-2026): P paquete modelo (D-1) · R reproducibilidad
(guía como prueba, herramientas del docente, pruebas negativas, Mbokaja) · LI licencias y procedencia (D-2) · H6 hito · D18 Capítulo 18 sin práctica ni
prueba escrita (D-3, D-4) · TI planilla de tiempos (D-5) · V6 variantes reservadas · W recuperación y calendario."""
import os, re, sys, json, collections, subprocess, hashlib
sys.dont_write_bytecode = True
AQUI = os.path.dirname(os.path.abspath(__file__))
L6 = os.path.dirname(AQUI)
ADS = os.path.dirname(L6)
REPO = os.path.dirname(ADS)
sys.path.insert(0, AQUI)
from docx import Document
from docx.oxml.ns import qn
import pymupdf

import construir as B
import oficio as O
K = O.K
import cap17, cap18, guia_u6 as GU
import fragmentos_u6 as F
import diagramas_u6 as G
import planilla_tiempos as PT
import armar_paquetes as A
E, MB, I = cap17.E, cap17.MB, cap17.I
import v6_plantilla as V6
import construccion as C4
X = C4.X
R, D = X.R, X.D
CAPS = (cap17, cap18)
EST, GUI, CUA = B.RUTA_EST[:-5], B.RUTA_GUIA[:-5], B.RUTA_CUAD[:-5]

ok = fallos = 0
RES = collections.OrderedDict()
DETALLE = []


def chk(bloque, desc, cond, detalle=''):
    global ok, fallos
    RES.setdefault(bloque, [0, 0])
    cond = bool(cond)
    if cond:
        ok += 1; RES[bloque][0] += 1
        print('CORRECTO [%s] %s' % (bloque, desc))
    else:
        fallos += 1; RES[bloque][1] += 1
        print('FALLA    [%s] %s %s' % (bloque, desc, detalle))
    DETALLE.append({'bloque': bloque, 'control': desc, 'resultado': 'correcto' if cond else 'FALLA'})


def texto_docx(ruta):
    d = Document(ruta)
    partes = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            vistos = set()
            for c in r.cells:
                if id(c._tc) in vistos:
                    continue
                vistos.add(id(c._tc))
                partes.append(c.text)
    return d, '\n'.join(partes)


def leer(*p):
    return open(os.path.join(*p), encoding='utf8').read()


def git(*args):
    return subprocess.run(['git'] + list(args), cwd=ADS, capture_output=True, text=True).stdout.strip()


dE, tE = texto_docx(EST + '.docx')
dG, tG = texto_docx(GUI + '.docx')
dC, tC = texto_docx(CUA + '.docx')
pE, pG, pC = pymupdf.open(EST + '.pdf'), pymupdf.open(GUI + '.pdf'), pymupdf.open(CUA + '.pdf')
res_build = json.load(open(os.path.join(L6, 'resultado_construccion.json'), encoding='utf8'))
r6 = cap17.resumen()
todo = tE + '\n' + tG + '\n' + tC
j = lambda x: json.dumps(x, ensure_ascii=False)
heads = [p.text for d in (dE, dG, dC) for p in d.paragraphs if p.style.name.startswith('Heading')]
norm = lambda s: re.sub(r'\s+', ' ', s)

# ---------------------------------------------------------------- A. índice de Work
WORK = {
    17: ('Preparar una entrega reproducible', (33, 34),
         ['17.1. Guía de instalación, configuración y ejecución.', '17.2. Manual breve basado en tareas.',
          '17.3. Respaldo, restauración, licencias y registro de procedencia.',
          '17.4. Prueba de instalación por otro equipo y acta de entrega simulada.']),
    18: ('Demostrar, defender y aprender del proyecto', (35, 36),
         ['18.1. Demostración organizada por requisitos y evidencias.', '18.2. Defensa individual de decisiones y resolución de una variante.',
          '18.3. Recuperación de aprendizajes pendientes.', '18.4. Retrospectiva y mejoras futuras justificadas.']),
}
work_txt = leer('/tmp/claude-0/fuentes/work.txt') if os.path.exists('/tmp/claude-0/fuentes/work.txt') else ''
w = norm(work_txt)
for Cm in CAPS:
    t, enc, secs = WORK[Cm.NUM]
    chk('A', 'Capítulo %d: título y encuentros %s de Work §6' % (Cm.NUM, enc), Cm.TITULO == t and Cm.ENCUENTROS == enc
        and ('%s — %s' % (Cm.CAPITULO, t)) in tE)
    chk('A', 'Capítulo %d: las cuatro secciones con el texto de Work §6, como títulos del libro' % Cm.NUM,
        [s.replace('. ', ' ', 1).rstrip('.') for s in secs] == Cm.SECCIONES and all(s in heads for s in Cm.SECCIONES))
chk('A', 'Capítulo 17: la práctica de Work («entregar código, datos, pruebas y documentación coherentes; H6»)',
    cap17.PRACTICA_WORK == 'entregar código, datos, pruebas y documentación coherentes; H6')
chk('A', 'Capítulo 18: cierre de Work «Evaluación integradora de la segunda etapa y cierre del proyecto»',
    cap18.REVISION_WORK == 'Evaluación integradora de la segunda etapa y cierre del proyecto' and cap18.INTEGRADORA_TITULO in heads)
if work_txt:
    chk('A', 'Secciones, práctica y cierre textuales en el informe de Work',
        all(s in w for v in WORK.values() for s in v[2]) and 'Práctica: entregar código, datos, pruebas y documentación coherentes; H6.' in w
        and 'Evaluación integradora de la segunda etapa y cierre del proyecto.' in w)
    chk('A', 'Resultado de la unidad y evidencias de §5.2 de los Encuentros 33 a 36 están en el informe de Work',
        all(x in w for x in ('Paquete reproducible, defensa y recuperación', 'Documentación probada por el propio',
                             'H6: paquete reproducible; ajustes finales.', 'Rúbrica de producto y preguntas de',
                             'Evidencia faltante, balance y propuesta de')))
chk('A', 'Apertura: resultado de §5.1 y evidencias de §5.2 de los cuatro encuentros',
    all(x in tE for x in ('paquete reproducible, defensa y recuperación', 'Documentación probada por el propio equipo',
                          'H6: paquete reproducible y ajustes finales', 'Rúbrica de producto y preguntas de transferencia',
                          'evidencia faltante, balance y propuesta de continuidad', 'cuatro encuentros y 20 HC')))

# ---------------------------------------------------------------- B. capacidades
CAPW = {'C7': 'Organiza y registra el trabajo colaborativo.', 'C9': 'Verifica y mejora la solución a partir de evidencia.',
        'C10': 'Documenta, entrega y defiende el prototipo.'}
for Cm, esp in ((cap17, ('C7', 'C10')), (cap18, ('C7', 'C9', 'C10'))):
    r = Cm.ROTULO_CAPACIDAD
    chk('B', 'Capítulo %d: rótulo «propuestas para el programa institucional» y «no son capacidades oficiales del MEC»' % Cm.NUM,
        'propuestas para el programa institucional' in r and 'no son capacidades oficiales del MEC' in r and r in tE and r in tG)
    chk('B', 'Capítulo %d: capacidades %s con su texto de Work §4.2' % (Cm.NUM, ', '.join(esp)),
        all(('%s. %s' % (c, CAPW[c])) in Cm.CAPACIDAD for c in esp))
chk('B', 'Encuentros 33 (C10), 34 (C7, C10), 35 (C1 a C10) y 36 (C7, C9, C10) en los rótulos (Work §5.2)',
    'C10 en el Encuentro 33; C7 y C10 en el Encuentro 34' in cap17.ROTULO_CAPACIDAD
    and 'C1 a C10 en el Encuentro 35; C7, C9 y C10 en el Encuentro 36' in cap18.ROTULO_CAPACIDAD)
chk('B', 'Indicador de C10 (Work §4.2): otro equipo instala con la guía, recupera un respaldo, cada integrante explica una decisión y una prueba',
    'otro equipo sin ayuda' in j(cap17.INDICADORES) and 'restauración' in j(cap17.INDICADORES)
    and 'una decisión del proyecto y una prueba' in j(cap18.INDICADORES))
malas_of = [tx[max(0, m.start() - 40):m.end()] for tx in (tE, tG, tC) for m in re.finditer(r'(?i)capacidad(es)? oficial', tx)
            if not re.search(r'\b[Nn]o (es|son)( una)?\b', tx[max(0, m.start() - 40):m.start()])]
chk('B', 'Ninguna capacidad presentada como oficial del MEC', not malas_of, malas_of)

# ---------------------------------------------------------------- C. trazabilidad de fuentes
locales = dict(MB.FUENTES)
locales.update(cap17.EVAL_FUENTES)
locales.update({m[0]: m[4] for m in X.MENSAJES}); locales.update(C4.MENSAJES_PHP)
locales.update({e[0]: e[1] for e in C4.FICHA_ENTORNO})
locales.update({d[0]: d[1] for d in MB.DEFECTOS}); locales.update({t[0]: t[1] for t in E.TAREAS})
locales.update({f[0]: f[1] for f in I.FICHAS_RECUPERACION}); locales.update({v['id']: v['situacion'] for v in I.V6_EJEMPLOS})
locales['ME-14'] = 'ME-14 (SC-02, lote 5)'
sys.path.insert(0, os.path.join(ADS, 'lote_5', 'caso'))
import verificacion as V5
locales.update({c[0]: c[4] for c in V5.CASOS}); locales.update({i[0]: i[2] for i in V5.INCIDENCIAS_PARTIDA})


def fuente(c):
    return (R.TEXTO.get(c) or D.NECESIDAD.get(c) or X.REGLA.get(c) or locales.get(c) or (c if C4.existe(c) else None)
            or (c if c in ('E-6', 'P-06a', 'P-06b') else None))


PAT = (r'(?:M-F\d\d?|MS-\d\d?|MR-\d\d|IN-\d\d|C-\d\d|DM-\d|T-0\d|FR-\d|V6-E\d|LU-\d|E-\d+[ab]?|OA-\d\d|PA-\d\d|MA-\d\d|MV-\d\d'
       r'|RI-\d\d|PT-\d\d|ME-\d\d|OP-\d\d|P-\d\d[ab]?|RF-\d\d|RNF-\d\d|RS-\d\d|RN-\d\d)')
for nombre, tx in (('libro', tE), ('guía', tG), ('cuaderno', tC)):
    usados = sorted(set(re.findall(r'(?<![\w-])' + PAT + r'(?![\w-])', tx)))
    falt = [c for c in usados if fuente(c) is None and c != 'ME-15']
    chk('C', '%s: los %d códigos citados existen en las fuentes' % (nombre, len(usados)), not falt, falt)
chk('C', 'ME-15 solo aparece en la guía como el mensaje inexistente del defecto DM-5', 'ME-15' not in tE + tC and 'ME-15' in tG)
nuevos = set(MB.FUENTES) - set(MB.MB5.FUENTES)
chk('C', 'Mbokaja U6: %d códigos nuevos (M-F13, MS-9, MS-10) que no repiten ninguno anterior' % len(nuevos),
    nuevos == {'M-F13', 'MS-9', 'MS-10'})

# ---------------------------------------------------------------- P. paquete modelo (D-1)
pm = r6['paquete_modelo']
chk('P', 'El paquete modelo es una copia de la referencia v2 con SC-02: código idéntico, solo cambian LEEME.txt y ENTREGA/',
    pm['codigo_identico'] and pm['distintos_del_origen'] == ['LEEME.txt'] and not pm['faltan_del_origen'])
chk('P', 'Las referencias de los lotes 4 y 5 no cambiaron (git)', not git('status', '--porcelain', '--', 'lote_4', 'lote_5')
    and not git('diff', '--stat', 'b2aed1b', '--', 'lote_4', 'lote_5'))
docs_ok = all(open(os.path.join(A.MODELO, ruta), encoding='utf8').read() == f() for ruta, f in E.DOCUMENTOS.items())
chk('P', 'Los documentos del paquete salen de caso/entrega_u6.py (libro y paquete dicen lo mismo)', docs_ok)
chk('P', 'v1.1: ningún paquete trae un comprobador propio; las herramientas del docente viven fuera de los paquetes',
    not [f_ for d_ in (A.MODELO, A.MBOKAJA) for r_, _, fs in os.walk(d_) for f_ in fs if 'comprobar_entrega' in f_
         or f_ in ('inspeccionar_entrega.php', 'probar_en_aislamiento.php')]
    and 'comprobar_entrega' not in j(E.PASOS) + j(E.ARCHIVOS_LISTA) + j(E.PROCEDENCIA) + E.leeme_txt()
    and not os.path.exists(os.path.join(L6, 'proyecto', 'comprobar_entrega.php'))
    and all(os.path.exists(os.path.join(L6, 'proyecto', 'herramientas', f_))
            for f_ in ('inspeccionar_entrega.php', 'probar_en_aislamiento.php', 'LEEME.md')))
insp, ejec = leer(L6, 'proyecto', 'herramientas', 'inspeccionar_entrega.php'), leer(L6, 'proyecto', 'herramientas', 'probar_en_aislamiento.php')
# código PHP sin cadenas entre comillas simples ni comentarios (los ` de las expresiones regulares son texto, no el operador de consola)
sin_cadenas = lambda s: re.sub(r'(?m)//.*$', '', re.sub(r"'(?:[^'\\]|\\.)*'", "''", s))
chk('P', 'La inspección estática no ejecuta ni incluye nada (sin include, require, eval, exec, shell, proc_open, popen)',
    not re.search(r'\b(include|require|eval|exec|shell_exec|system|passthru|proc_open|popen|pcntl_exec)\s*\(|`|\binclude\b|\brequire\b', sin_cadenas(insp)))
chk('P', 'Las pruebas en aislamiento: lista cerrada propia, comando como lista de argumentos, sin consola, con tiempo límite y '
    'negativa con cuenta de administrador',
    "const PRUEBAS_PERMITIDAS = array('prueba_contratos', 'prueba_garantias', 'prueba_respaldo');" in ejec
    and 'proc_open($comando' in ejec and '$comando = array(PHP_BINARY' in ejec
    and not re.search(r'\b(shell_exec|system|passthru|exec|popen|eval)\s*\(|`', sin_cadenas(ejec))
    and 'posix_geteuid() === 0' in ejec and 'proc_terminate' in ejec and 'isLink()' in ejec)
leeme_h = leer(L6, 'proyecto', 'herramientas', 'LEEME.md')
chk('P', 'Procedimiento de aislamiento documentado: equipo aislado sin red, cuenta estándar, copia descartable, registro y limpieza; Windows y laboratorios pendientes',
    all(x in leeme_h for x in ('Inspección estática', 'sin red', 'Cuenta estándar', 'Copia descartable', 'Registro', 'Limpieza', 'pendientes')))
chk('P', 'Ningún paquete trae una base de datos', not [f_ for d_ in (A.MODELO, A.MBOKAJA) for r_, _, fs in os.walk(d_)
                                                     for f_ in fs if f_.endswith(('.sqlite', '.db'))])

# ---------------------------------------------------------------- R. reproducibilidad
c = r6['herramientas']
limpio = lambda x: all(x[k]['codigo_salida'] == 0 and not x[k]['incoherencias'] for k in ('inspeccion', 'aislamiento'))
chk('R', 'Herramientas del docente: 0 incoherencias en el paquete modelo (inspección y pruebas en aislamiento)', limpio(c['modelo']))
chk('R', 'Herramientas: en el paquete de Mbokaja, exactamente las %d incoherencias de la clave (5 defectos): 5 de la inspección y 1 de las pruebas en aislamiento' % len(MB.incoherencias_esperadas()),
    c['mbokaja']['exactas'] and len(MB.DEFECTOS) == 5 and len(c['mbokaja']['inspeccion']['incoherencias']) == 5
    and len(c['mbokaja']['aislamiento']['incoherencias']) == 1)
chk('R', 'Herramientas: con la clave aplicada, 0 incoherencias', limpio(c['mbokaja_corregido']))
pn = r6['pruebas_negativas']
chk('R', 'Pruebas negativas de las herramientas: %d de %d correctas (inyección, guía maliciosa, efectos al incluir, rutas, enlaces, administrador, tiempo límite, paquete intacto, lista cerrada, red)'
    % (sum(x['correcto'] for x in pn), len(pn)), len(pn) == 10 and all(x['correcto'] for x in pn))
chk('R', 'La guía como prueba corrió aislada (sin red y con la cuenta nobody)',
    all('sin red' in (r6['guia'][k].get('aislamiento') or '') and 'nobody' in r6['guia'][k]['aislamiento'] for k in r6['guia']))
g = r6['guia']
chk('R', 'Guía como prueba: el paquete modelo y el de Mbokaja corregido se instalan siguiendo solo su guía (7 de 7 pasos)',
    all(g[k]['correctos'] == g[k]['total'] == len(E.PASOS) for k in ('modelo', 'mbokaja_corregido')))
chk('R', 'Guía como prueba: el paquete con defectos NO se puede instalar siguiendo su guía', g['mbokaja_con_defectos']['correctos'] < len(E.PASOS))
chk('R', 'Las únicas adaptaciones de la prueba de la guía son el puerto y el nombre del respaldo, y se informan',
    {x['nota'] for x in g['modelo']['pasos'] if x['nota']} <= {'nombre del respaldo del paso anterior'} | {x['nota'] for x in g['modelo']['pasos'] if x['nota'].startswith('puerto ')})
chk('R', 'Regresión de la versión empaquetada: las nueve pruebas sin fallos', len(r6['regresion']) == 9
    and all(x['codigo_salida'] == 0 for x in r6['regresion'].values()))
chk('R', 'La prueba automática se rotula como comprobación técnica y no reemplaza la instalación por otro equipo',
    'NO reemplaza la instalación por otro equipo' in r6['entorno']['tipo'] and 'no reemplaza' in j(cap17.S174)
    and 'Pendiente hasta el acta' in j(cap17.tabla_176()) and 'no reemplaza la instalación por otro equipo' in tG)
chk('R', 'La instalación por otro equipo exige un equipo distinto, sin ayuda, y registra condiciones reales',
    'equipo distinto del que lo hizo' in tG and 'sin ayuda' in j(cap17.S174) and 'Computadora, sistema y versiones' in j(E.ACTA_CAMPOS))
chk('R', 'Libro y cuaderno no revelan los defectos del paquete de Mbokaja (la clave está solo en la guía)',
    not any(d_[1] in tE + tC or 'instalar_base' in tE + tC or 'mbokaja_prueba.json' in tE + tC for d_ in MB.DEFECTOS)
    and all(d_[0] in tG for d_ in MB.DEFECTOS))

# ---------------------------------------------------------------- LI. licencias y procedencia (D-2)
lic = leer(A.MODELO, 'ENTREGA', 'LICENCIA_BORRADOR.md')
chk('LI', 'Borrador de licencia: uso educativo institucional; instalar, ejecutar, estudiar y modificar dentro de la institución; sin redistribución comercial',
    'BORRADOR' in lic and 'instalar, ejecutar, estudiar y modificar' in lic and 'redistribución comercial' in lic
    and 'institución que adquirió' in lic)
chk('LI', 'El aspecto jurídico no se declara cerrado (libro, guía y paquete)', 'no está cerrado' in lic and 'no está cerrado' in tE
    and 'no está cerrado' in tG and not re.search(r'(?i)jur[ií]dic[oa][^.]{0,30}\bcerrad[oa]\b(?<!no está cerrado)', todo.replace('no está cerrado', '')))
chk('LI', 'No se atribuye a la editorial código ajeno ni el código de los estudiantes',
    'código original de los estudiantes, que es de sus autores' in lic and 'no lo cubre esta ficha ni se atribuye a la editorial' in leer(A.MODELO, 'ENTREGA', 'PROCEDENCIA.md')
    and all('Titulares del material' not in f_[2] for f_ in E.PROCEDENCIA[3:]))
ext = [f_ for f_ in E.PROCEDENCIA if f_[5].startswith('No se entrega')]
chk('LI', 'Cada recurso externo (%d) con su versión exacta, su licencia y una fuente primaria con fecha de consulta' % len(ext),
    len(ext) >= 4 and all(re.search(r'\d+\.\d+', f_[0]) and f_[3] != '—' and ('consultad' in f_[4] or 'bash --version' in f_[4]) for f_ in ext))
e6 = r6['entorno']
chk('LI', 'Las versiones de la ficha son las del entorno donde corrieron las pruebas',
    ('PHP %s' % e6['php']) in j(E.PROCEDENCIA) and ('SQLite %s' % e6['sqlite']) in j(E.PROCEDENCIA)
    and e6['curl'].replace('curl ', 'curl ') in j(E.PROCEDENCIA) and '5.2.21' in e6['bash'])
chk('LI', 'Licencias verificadas: PHP 8.3.6 «The PHP License, version 3.01»; SQLite 3.45.1 dominio público; curl 8.5.0 aviso propio; bash GPLv3+',
    all(x in j(E.PROCEDENCIA) for x in ('The PHP License, version 3.01', 'dominio público', 'COPYRIGHT AND PERMISSION NOTICE', 'GPLv3+')))
chk('LI', 'El paquete no trae recursos sin procedencia (la inspección estática lo controla con el bloque D)', not c['modelo']['inspeccion']['incoherencias'])

# ---------------------------------------------------------------- H6
rh = I.RUBRICA_H6
chk('H6', 'Rúbrica del H6: seis criterios (C7 y C10), cuatro niveles, evidencia del equipo e individual',
    len(rh) == 6 and all(len(x[2]) == 4 and x[3] and x[4] for x in rh) and sorted(set(x[1] for x in rh)) == ['C10', 'C7']
    and all(('%s (%s)' % (x[0], x[1])) in tG for x in rh))
chk('H6', 'H6: 1 HC de evaluación en el Encuentro 34; 60/40 como propuesta editorial; AE aparte; pendientes del entorno separados',
    '1 de evaluación (H6)' in cap17.TIEMPO and '60 %' in j(cap17.H6_EVALUACION) and 'propuesta editorial' in j(cap17.H6_EVALUACION)
    and 'no se confunde con un nivel Inicial' in j(cap17.H6_EVALUACION) and 'pendiente del entorno' in j(cap17.H6_EVALUACION))
chk('H6', 'La condición mínima «instalación y restauración reproducibles» se completa solo con la evidencia de otro equipo',
    'condición mínima que quedó pendiente en el H5' in j(cap17.S174) and 'Pendiente hasta el acta' in j(cap17.tabla_176()))
chk('H6', 'Guía: planilla de registro del H6 con columna AE', 'Planilla de registro del H6' in tG)

# ---------------------------------------------------------------- D18. Capítulo 18 (D-3 y D-4)
chk('D18', 'Capítulo 18 sin práctica independiente (D-4)', cap18.PRACTICA_TITULO is None and not cap18.NIVELES and 'Práctica 18' not in todo)
chk('D18', 'Capítulo 18 sin prueba escrita adicional (D-3); el Capítulo 17 con su evaluación opcional de 15 puntos',
    cap18.EVAL_ID is None and 'Evaluación del Capítulo 18' not in todo and cap17.EVAL_TITULO in tE)
chk('D18', 'Instrumentos del Capítulo 18: demostración (guion), defensa individual (rúbrica), evaluación integradora',
    cap18.TABLA_181_ROTULO in tE and cap18.TABLA_182_ROTULO in tE and cap18.INTEGRADORA_TITULO in tE)
chk('D18', 'Evaluación integradora: H6, producto final con la ponderación de Work §4.3 y defensa; propuesta editorial',
    [p for _, p, _, _ in I.PRODUCTO] == [15, 20, 20, 20, 15, 10] and 'propuesta editorial' in j(I.INTEGRADORA)
    and all(x in j(I.INTEGRADORA) for x in ('H6', 'rúbrica de producto final', 'defensa individual')))
chk('D18', 'Planilla anual H1 a H6 y defensa, con columna AE', [h for h, *_ in I.PLANILLA_ANUAL] == ['H1', 'H2', 'H3', 'H4', 'H5', 'H6', 'Defensa']
    and cap18.TABLA_185_ROTULO in tE)
chk('D18', 'Preguntas de transferencia con los cuatro proyectos alternativos de Work §7.4',
    [t[0] for t in I.TRANSFERENCIA] == ['Préstamo de equipos audiovisuales', 'Inscripción a talleres escolares',
                                        'Gestión de solicitudes de soporte', 'Reserva de espacios escolares'])

# ---------------------------------------------------------------- TI. planilla de tiempos (D-5)
xl = os.path.join(L6, 'instrumentos', 'planilla_tiempos_encuentro_35.xlsx')
casos_pt = PT.comprobar_casos()
PT.escribir()
chk('TI', 'La planilla XLSX existe y sus fórmulas, recalculadas por un evaluador independiente, coinciden con el cálculo de la guía en los %d casos (límites A/B y B/B+, demostraciones excedidas, datos inválidos)' % len(casos_pt),
    os.path.exists(xl) and len(casos_pt) == 11 and all(b for _, b, _ in casos_pt), [n for n, b, _ in casos_pt if not b])
ej = I.tiempos(**I.EJEMPLO_TIEMPOS)
chk('TI', 'Ejemplo de control (v1.1): pasan 2 defensas enteras = 12 minutos; con 5 recuperaciones de 15 minutos quedan 48 (no 51)',
    ej['defensas_que_pasan_al_36'] == 2 and ej['minutos_trasladados_al_36'] == 12 and ej['diferencia_36'] == 48
    and 'quedan 48 minutos' in tG and '51' not in j(GU.tabla_tiempos_ejemplo()))
casos = [I.EJEMPLO_TIEMPOS, dict(I.EJEMPLO_TIEMPOS, estudiantes=18), dict(I.EJEMPLO_TIEMPOS, estudiantes=40, recuperaciones=8)]
alts = [I.tiempos(**cc)['alternativa'] for cc in casos]
chk('TI', 'Tres casos ejercitan las tres alternativas (A, B y B + calendario institucional)', alts == ['B', 'A', 'B + calendario institucional autorizado'], alts)
invalidos = []
for mal in (dict(I.EJEMPLO_TIEMPOS, min_defensa=0), dict(I.EJEMPLO_TIEMPOS, estudiantes=-1), dict(I.EJEMPLO_TIEMPOS, equipos=True)):
    try:
        I.tiempos(**mal); invalidos.append(False)
    except ValueError:
        invalidos.append(True)
chk('TI', 'Datos inválidos (0 minutos por defensa, negativos, no enteros) se rechazan: no hay división por cero', all(invalidos))
chk('TI', 'La planilla pide cantidad de equipos, estudiantes, minutos de la HC y por defensa; no supone la duración de la HC',
    [k for k, *_ in I.TIEMPOS_CAMPOS][:6] == ['equipos', 'estudiantes', 'min_hc', 'min_demo', 'min_cambio', 'min_defensa']
    and 'no se supone' in j(I.TIEMPOS_CAMPOS) and 'no son la duración de la HC de ninguna institución' in tG)
chk('TI', 'Un solo evaluador: no se supone un segundo evaluador disponible',
    'no se supone un segundo evaluador disponible' in tG and 'Un solo evaluador' in j(I.ALTERNATIVAS)
    and 'segundo evaluador' not in tE.replace('no se supone un segundo evaluador', ''))
chk('TI', 'Si no alcanza: defensas en el Encuentro 36 y tiempo institucional autorizado, sin eliminar evidencias obligatorias',
    'Ninguna evidencia obligatoria se elimina' in j(I.ALTERNATIVAS) and 'calendario' in j(I.ALTERNATIVAS[2]))
chk('TI', 'Preparación anticipada del Encuentro 35 (variantes, planillas, paquetes instalados, orden)', len(I.PREPARACION_35) >= 5
    and all(p in tG for p in I.PREPARACION_35))

# ---------------------------------------------------------------- V6. variantes reservadas
chk('V6', 'Ningún banco de variantes en el repositorio', not [os.path.join(d_, f_) for d_, _, fs in os.walk(REPO) if '.git' not in d_
                                                          for f_ in fs if f_ in ('variantes_v6.json', 'comprobacion_v6.txt')])
q = subprocess.run([sys.executable, '-B', os.path.join(L6, 'caso', 'v6_plantilla.py'), '--banco=' + os.path.join(L6, 'x.json')],
                   capture_output=True, text=True)
chk('V6', 'El validador se niega a usar un banco dentro del repositorio', q.returncode != 0 and 'FUERA del repositorio' in q.stderr)
chk('V6', 'Libro y guía solo muestran ejemplos públicos (V6-E1 a V6-E3), rotulados como que no se usan en la defensa',
    not re.search(r'V6-\d\d', todo) and 'no se usa en la defensa' in tE and 'no se usan en la defensa' in tG)
for var, etiqueta in (('ADS_V6', 'variantes'), ('ADS_DR', 'DR')):
    ruta = os.environ.get(var)
    if ruta and os.path.exists(ruta):
        datos = json.load(open(ruta, encoding='utf8'))
        archivos = [f_ for f_ in git('ls-files').split('\n') + git('ls-files', '--others', '--exclude-standard').split('\n')
                    if f_ and os.path.isfile(os.path.join(ADS, f_))]
        versionado = ''.join(open(os.path.join(ADS, f_), encoding='utf8', errors='ignore').read() for f_ in archivos) + todo
        if etiqueta == 'variantes':
            claves_ = [v['situacion'][:40] for v in datos]
        else:
            claves_ = [x[0] for k in ('obras', 'ediciones', 'ejemplares', 'lectores') for x in datos[k]]
        hallados = [k for k in claves_ if k in versionado]
        chk('V6', '%s: ninguno de los %d textos o códigos reservados está en el repositorio ni en los documentos' % (var, len(claves_)),
            not hallados, '%d hallados' % len(hallados))
        chk('V6', '%s: la carpeta privada está fuera del repositorio' % var,
            not (os.path.realpath(ruta) + os.sep).startswith(os.path.realpath(REPO) + os.sep))

# ---------------------------------------------------------------- W. recuperación y calendario
extra = [s for s in re.split(r'(?<=[.])\s+', todo) if 'horas extraordinarias' in s]
chk('W', 'No se suponen horas extraordinarias (%d menciones, todas negadas)' % len(extra),
    extra and all(re.search(r'[Nn]o se suponen|sin suponer|Suponer horas extraordinarias\nAlternativa B', s) for s in extra) and 'fuera del horario' not in todo)
chk('W', 'Recuperación con prioridad (condiciones mínimas, evidencias ausentes, nivel Inicial); lo que no se recupera no se declara aprobado',
    'no se declara aprobado' in j(I.REGLA_RECUPERACION) and 'condiciones mínimas' in j(I.REGLA_RECUPERACION)
    and [p for *_, p in I.FICHAS_RECUPERACION] == [1, 2, 3, 2])
chk('W', 'Fichas de recuperación FR-1 a FR-4 en el libro y en la guía', all(f_[0] in tE and f_[0] in tG for f_ in I.FICHAS_RECUPERACION))
chk('W', 'Calendario de 28: Encuentros 25 a 28 sin cambios; el 28 recibe las recuperaciones técnicas de la Unidad V',
    'pasan a ser el 25 y el 26, sin cambios' in j(cap17.CALENDARIO) and 'pasan a ser el 27 y el 28, sin cambios' in j(cap18.CALENDARIO)
    and 'recuperaciones técnicas de la Unidad V' in j(cap18.CALENDARIO) and 'Unidad VI en 36 y en 28 encuentros' in tG)
chk('W', 'Las 20 HC de la unidad: 33 Doc 5; 34 V 2, Doc 2, E 1; 35 E 5; 36 E 3, G 2',
    '33: 5 HC de documentación' in cap17.TIEMPO and '2 HC de verificación (instalación por otro equipo), 2 de documentación (ajustes) y 1 de evaluación' in cap17.TIEMPO
    and 'Encuentro 35: 5 HC de evaluación' in cap18.TIEMPO and 'Encuentro 36: 3 HC de evaluación' in cap18.TIEMPO and '2 de gestión' in cap18.TIEMPO)

# ---------------------------------------------------------------- T. fragmentos de código
for fid in F.TODOS:
    ls = F.extraer(fid)
    chk('T', '%s: leído tal cual del paquete modelo (%d líneas), ninguna de más de %d caracteres' % (fid, len(ls), F.ANCHO),
        '\n'.join(ls) in leer(F.PROY, F.TODOS[fid][0]) and max(len(l) for l in ls) <= F.ANCHO and len(ls) <= 40)
    dest = tE if fid in F.FRAGMENTOS else tG
    chk('T', '%s: con su ficha en %s' % (fid, 'el libro' if fid in F.FRAGMENTOS else 'la guía'),
        all(v in dest for v in F.ficha(fid).values()) and all(l.strip() in dest for l in ls if l.strip()))
for fid in F.REFERENCIA:
    del_libro = {l.strip() for f_ in F.FRAGMENTOS for l in F.extraer(f_)}
    propias = [l.strip() for l in F.extraer(fid) if len(l.strip()) > 25 and l.strip() not in del_libro]
    chk('T', '%s no se publica en el libro ni en el cuaderno' % fid, not [l for l in propias if l in tE or l in tC])
orden = [m.group(1) for m in re.finditer(r'^Código (\d+\.\d+) —', tE, re.M)]
chk('T', 'Códigos 17.1 y 17.2 correlativos, en orden de aparición y con explicación', orden == ['17.1', '17.2']
    and all(f_ in cap17.EXPLICA for f_ in F.FRAGMENTOS))
txtpdf = re.sub(r'[ \t]+', ' ', '\n'.join(pg.get_text() for pg in pE) + '\n'.join(pg.get_text() for pg in pG))
faltan_cod = [l for f_ in F.TODOS for l in F.extraer(f_) if l.strip() and re.sub(r'\s+', ' ', l.strip()) not in txtpdf]
chk('T', 'Cada línea de código del libro y de la guía está impresa entera en el PDF', not faltan_cod, faltan_cod[:3])

# ---------------------------------------------------------------- V. pruebas y entorno
chk('V', 'Última ejecución de build/reproducir.py sin fallos y con el entorno registrado', r6['fallos'] == 0 and r6['entorno']['php'])
chk('V', 'Cada resultado guardado tiene su línea de entorno y su condición',
    all(('Entorno' in leer(L6, 'resultados_pruebas', f_) and 'Condición' in leer(L6, 'resultados_pruebas', f_)) or f_ == 'resumen.json'
        for f_ in os.listdir(os.path.join(L6, 'resultados_pruebas'))))
chk('V', 'Cada capítulo y la apertura advierten que el entorno no se probó en la institución', tE.count(C4.AVISO_ENTORNO[0]) >= 3)
chk('V', 'Windows y el laboratorio siguen pendientes (libro, guía y paquete)', 'Pendiente (ficha de entorno)' in tE
    and 'Pendientes' not in '' and 'Windows' in leer(A.MODELO, 'ENTREGA', 'LIMITES.md') and 'en Windows' in tG)
malas = re.findall(r'(?i)(?:entorno|base|sistema|prototipo|computadora|material|libro|paquete)[^.\n]{0,40}\bvalidad[oa]s?\b(?! en aulas)'
                   r'|probad[oa]s? en (?:los )?colegios|\bapto para todos', todo)
chk('V', 'Nada se declara validado ni probado en colegios', not malas, malas[:5])
chk('V', 'Datos ficticios en el paquete y en los casos', 'Datos ficticios' in leer(A.MODELO, 'LEEME.txt') and 'ficticio' in MB.DATOS_TITULO and 'ficticio' in MB.TITULO_LEEME)

# ---------------------------------------------------------------- F. voseo y residuos
TUTEO = (r'\b(tienes|puedes|debes|sabes|quieres|eres|propón|tú)\b|(?:^\d+\.\s+|^\([a-h]\)\s+)'
         r'(Escribe|Lee|Completa|Indica|Explica|Clasifica|Responde|Anota|Marca|Elige|Justifica|Redacta|Compara|Describe|Identifica|Resume|Haz|Ejecuta|Registra|Señala|Decide|Diseña|Aplica|Revisa|Agrega|Calcula|Ordena|Instala|Corrige|Sigue)\b')
sin_codigo = lambda tx: '\n'.join(l for l in tx.split('\n') if l.rstrip() not in {x.rstrip() for f_ in F.TODOS for x in F.extraer(f_)})
for nombre, tx in (('libro', tE), ('guía', tG), ('cuaderno', tC)):
    m = [x.group(0) for x in re.finditer(TUTEO, tx, re.M)]
    chk('F', '%s: sin formas de tuteo frecuentes' % nombre, not m, m[:5])
    chk('F', '%s: sin identificadores de modelos de IA' % nombre, not re.search(r'(?i)\b(claude|gpt|gemini|opus|sonnet|codex|chatgpt)\b', tx))
    chk('F', '%s: sin marcadores sin reemplazar' % nombre, '§§' not in tx and 'TODO' not in tx and 'XXX' not in tx
        and '%s' not in sin_codigo(tx) and '%d' not in sin_codigo(tx))
    chk('F', '%s: no se declara apto ni aprobado para la venta' % nombre,
        not re.search(r'(?i)\bapto\b|aprobado para (la )?(venta|comercializ)', tx.replace('No está aprobado para la venta', '')))
for d_ in (A.MODELO, A.MBOKAJA):
    txt = ''.join(leer(d_, 'ENTREGA', f_) for f_ in os.listdir(os.path.join(d_, 'ENTREGA'))) + leer(d_, 'LEEME.txt')
    chk('F', '%s: documentos de entrega en voseo y sin modelos de IA' % os.path.basename(d_),
        not re.search(TUTEO, txt, re.M) and not re.search(r'(?i)\b(claude|gpt|chatgpt)\b', txt))

# ---------------------------------------------------------------- G. maqueta (DOCX)
tw = int(O.Cm(O.ANCHO).twips)
for nombre, d in (('libro', dE), ('guía', dG), ('cuaderno', dC)):
    s = d.sections[0]
    chk('G', '%s: oficio 21,59 × 33,02 cm, márgenes O3 y espejados' % nombre,
        abs(s.page_width.cm - 21.59) < 0.01 and abs(s.page_height.cm - 33.02) < 0.01 and abs(s.left_margin.cm - 2.5) < 0.01
        and abs(s.right_margin.cm - 1.8) < 0.01 and d.settings.element.find(qn('w:mirrorMargins')) is not None)
    filas = [r_ for t in d.tables for r_ in t.rows]
    chk('G', '%s: %d filas de tabla, todas «no dividir»' % (nombre, len(filas)),
        all(r_._tr.trPr is not None and r_._tr.trPr.find(qn('w:cantSplit')) is not None for r_ in filas))
    tablas = [t for t in d.tables if len(t.columns) > 1 or t.rows[0].cells[0].paragraphs[0].runs and
              t.rows[0].cells[0].paragraphs[0].runs[0].font.name == 'Courier New']
    malas_t = [i for i, t in enumerate(tablas) if abs(sum(int(g_.get(qn('w:w'))) for g_ in t._tbl.tblGrid) - tw) > 15]
    chk('G', '%s: %d tablas y bloques de código al ancho útil' % (nombre, len(tablas)), not malas_t, malas_t)
chk('G', 'Se usan la maqueta del lote 1, el perfil O3 del lote 2 y las piezas del lote 4 (sin copias)',
    os.path.dirname(O.H.__file__).endswith(os.path.join('lote_1', 'build')) and os.path.dirname(B.Z.Z4.__file__).endswith(os.path.join('lote_4', 'build')))
chk('G', 'Portada: versión del lote 6 en los tres documentos', all('v1.1 (lote 6, cierre de auditoría)' in t for t in (tE, tG, tC)))


# ---------------------------------------------------------------- H. maqueta (PDF)
def lineas(pg):
    return [l.strip() for l in pg.get_text().split('\n') if l.strip() and not re.match(r'^Página \d+ de \d+$', l.strip())]


for nombre, pdf in (('libro', pE), ('guía', pG), ('cuaderno', pC)):
    chk('H', '%s: ninguna página en blanco' % nombre, not [i for i, pg in enumerate(pdf) if not lineas(pg)])
    chk('H', '%s: ninguna página empieza con renglones de respuesta sueltos' % nombre,
        not [i for i, pg in enumerate(pdf) if lineas(pg) and lineas(pg)[0].startswith('_' * 20)])
pags = [i + 1 for i, pg in enumerate(pE) if cap17.EVAL_TITULO in lineas(pg)[:2]]
chk('H', 'Evaluación del Capítulo 17 en página propia (pág. %s)' % pags, len(pags) == 1 and lineas(pE[pags[0]])[0].startswith('Capítulo 18'))
for t_ in (cap17.PRACTICA_TITULO, cap17.E34_TITULO, cap18.E36_TITULO):
    chk('H', '«%s» empieza en página nueva' % t_.split(' —')[0], len([i for i, pg in enumerate(pE) if lineas(pg)[:1] == [t_]]) == 1)
chk('H', 'Cuaderno: la Práctica 17 empieza en página nueva', len([i for i, pg in enumerate(pC) if lineas(pg)[:1] == [cap17.PRACTICA_TITULO]]) == 1)
paginas = res_build['paginas_libro']
recalc, _ = K.paginas_de_titulos(EST + '.pdf', list(paginas), desde=2)
chk('H', 'Índice estable: las %d entradas apuntan a su página; la apertura sigue al índice' % len(paginas),
    recalc == paginas and len(paginas) == len(B.entradas_indice()) and lineas(pE[2])[0] == B.UNIDAD_VI)
chk('H', 'Extensión: libro %d, guía %d, cuaderno %d págs. (estimación del plan: 26-34, 18-24, 3-6; no se rellena)' % (len(pE), len(pG), len(pC)),
    16 <= len(pE) <= 34 and 12 <= len(pG) <= 24 and 3 <= len(pC) <= 6)

# ---------------------------------------------------------------- I. evaluaciones e instrumentos
claves = res_build['claves']
ks = [c_ for c_ in claves if c_[0] == cap17.EVAL_ID]
mc = [c_[2] for c_ in ks if len(c_[2]) == 1 and c_[2] in 'abcd']
chk('I', 'Evaluación del Capítulo 17: %d ítems con clave; opción múltiple con claves %s' % (len(ks), ''.join(mc)),
    len(ks) == len(cap17.EVAL_ITEMS) and len(set(mc)) >= 3)
pts = sum(1 for it in cap17.EVAL_ITEMS if it[0] in ('mc', 'vf')) + sum(it[3] for it in cap17.EVAL_ITEMS if it[0] == 'ab')
chk('I', 'Evaluación del Capítulo 17: 15 puntos y texto base nuevo (ludoteca, LU-1 a LU-4)', pts == 15 and 'Ludoteca' in tE)
prev = ['ADS3-L1-C1', 'ADS3-L1-C2-i', 'ADS3-L1-C3', 'ADS3-C4v2-2', 'ADS3-L2-C5-c', 'ADS3-L2-C6', 'ADS3-L3-C7', 'ADS3-L3-C8',
        'ADS3-L3-C9', 'ADS3-L4-C10', 'ADS3-L4-C11', 'ADS3-L4-C12', 'ADS3-L4-C13-f', 'ADS3-L5-C14-c', 'ADS3-L5-C15', 'ADS3-L5-C16']
chk('I', 'Semilla de evaluación distinta de las de los lotes 1 a 5', cap17.EVAL_ID not in prev)
chk('I', 'Guía: clave y criterios de la evaluación, claves de todas las actividades de los Encuentros 33 a 36',
    'Clave de la evaluación del capítulo 17' in tG and all(len(k) == sum(len(it) for _, it in acts) for k, acts in (
        (GU.CLAVES_E33, cap17.ACT_E33), (GU.CLAVES_E34, cap17.ACT_E34), (GU.CLAVES_E35, cap18.ACT_E35), (GU.CLAVES_E36, cap18.ACT_E36))))
chk('I', 'Práctica 17: tres niveles con Mbokaja, claves por nivel y lista de cotejo', [n.split(' —')[0] for n, _ in cap17.NIVELES]
    == ['Nivel 1', 'Nivel 2', 'Nivel 3'] and 'Mbokaja' in j(cap17.PRACTICA_INTRO) and 'Lista de cotejo de la Práctica 17' in tG)
chk('I', 'Guía: secuencia en HC de los Encuentros 33 a 36 con las categorías de la v2.1',
    '(5 HC de documentación)' in tG and '(5 HC: 2 de verificación, 2 de documentación y 1 de evaluación)' in tG
    and '(5 HC de evaluación)' in tG and '(5 HC: 3 de evaluación y 2 de gestión)' in tG)
chk('I', 'Cuaderno: la Práctica 17 con el mismo texto que el libro y sin evaluaciones',
    all(t_ in tC for t_ in [cap17.PRACTICA_TITULO, cap17.PRACTICA_CONTROL] + cap17.PRACTICA_INTRO) and cap17.EVAL_TITULO not in tC)
chk('I', 'Rúbrica de defensa individual con cuatro criterios y cuatro niveles; planilla de defensa en la guía',
    len(I.RUBRICA_DEFENSA) == 4 and all(len(x[2]) == 4 for x in I.RUBRICA_DEFENSA) and 'Planilla de defensa individual' in tG)

# ---------------------------------------------------------------- J. numeración
for Cm in CAPS:
    n = Cm.NUM
    for tipo in ('Tabla', 'Figura', 'Ejemplo resuelto'):
        sec = [int(m.group(1)) for m in re.finditer(r'^%s %d\.(\d+)(?: —|:)' % (tipo, n), tE, re.M)]
        chk('J', 'Capítulo %d: %s %d.1 a %d.%d correlativas' % (n, tipo.lower(), n, n, len(sec)), sec == list(range(1, len(sec) + 1)) and sec, sec)
    for tipo, pat in (('tabla', r'^Tabla %d\.(\d+) —'), ('figura', r'^Figura %d\.(\d+) —'), ('ejemplo resuelto', r'^Ejemplo resuelto %d\.(\d+):'),
                      ('código', r'^Código %d\.(\d+) —')):
        refs = set(int(x) for x in re.findall(r'%s %d\.(\d+)' % (tipo.capitalize(), n), tE + tG + tC))
        defs = set(int(x) for x in re.findall(pat % n, tE, re.M))
        chk('J', 'Capítulo %d: toda %s mencionada existe' % (n, tipo), refs <= defs, refs - defs)
lotes = '\n'.join(texto_docx(os.path.join(ADS, r_))[1] for r_ in (
    'lote_1_oficio_candidata/ADS_3BTI_LOTE1_ESTUDIANTE_OFICIO_O3_CANDIDATA.docx', 'lote_2/ADS_3BTI_LOTE2_ESTUDIANTE_OFICIO_O3.docx',
    'lote_3/ADS_3BTI_LOTE3_ESTUDIANTE_OFICIO_O3.docx', 'lote_4/ADS_3BTI_LOTE4_ESTUDIANTE_OFICIO_O3.docx',
    'lote_5/ADS_3BTI_LOTE5_ESTUDIANTE_OFICIO_O3.docx'))
otras = sorted(set(re.findall(r'\b(?:Tabla|Figura) (?:[1-9]|1[0-6])\.\d+', tE + tG + tC)))
chk('J', 'Tablas y figuras de los lotes 1 a 5 citadas existen allí (%s)' % ', '.join(otras),
    all(re.search(r'^%s —' % re.escape(o), lotes, re.M) for o in otras), [o for o in otras if not re.search(r'^%s —' % re.escape(o), lotes, re.M)])
chk('J', 'Ejemplos resueltos de otros capítulos citados existen (15.1)', re.search(r'^Ejemplo resuelto 15\.1:', lotes, re.M))

# ---------------------------------------------------------------- L. notaciones
reg = json.load(open(os.path.join(AQUI, 'figs', 'simbolos.json'), encoding='utf8'))
for f_ in G.FIGURAS:
    chk('L', '%s (esquema didáctico) no usa símbolos de ninguna notación' % f_, reg[f_]['notacion'] == 'esquema' and not reg[f_]['simbolos'])

# ---------------------------------------------------------------- Q. lotes 1 a 5 intactos
st = git('status', '--porcelain', '--', 'lote_1', 'lote_2', 'lote_3', 'lote_4', 'lote_5', 'lote_1_oficio_candidata', 'lote_5_plan', 'lote_6_plan')
chk('Q', 'git: ningún archivo de los lotes 1 a 5 ni de los planes cambió ni se agregó', not st, st[:300])
for lote, n in (('lote_3', 3), ('lote_4', 4), ('lote_5', 5)):
    m_ = json.load(open(os.path.join(ADS, lote, 'MANIFIESTO_LOTE_%d.json' % n), encoding='utf8'))
    docs = [a for a in m_['archivos'] if a['ruta'].startswith('ADS_3BTI_LOTE%d_' % n)]
    chk('Q', 'Lote %d: sus seis DOCX y PDF coinciden con su manifiesto SHA-256' % n,
        len(docs) == 6 and all(hashlib.sha256(open(os.path.join(ADS, lote, a['ruta']), 'rb').read()).hexdigest() == a['sha256'] for a in docs))
m10 = json.load(open(os.path.join(L6, 'version_1_0', 'MANIFIESTO_LOTE_6.json'), encoding='utf8'))
v10 = [a_ for a_ in m10['archivos'] if os.path.exists(os.path.join(L6, 'version_1_0', a_['ruta']))]
chk('Q', 'Lote 6 v1.0 intacto en version_1_0/: %d archivos coinciden con su manifiesto SHA-256' % len(v10),
    len(v10) >= 40 and all(hashlib.sha256(open(os.path.join(L6, 'version_1_0', a_['ruta']), 'rb').read()).hexdigest() == a_['sha256'] for a_ in v10))
chk('Q', 'No se construyeron el tomo anual ni los planes de clase', not os.path.exists(os.path.join(ADS, 'coleccion'))
    and not [f_ for f_ in os.listdir(os.path.join(ADS, 'planes_clase')) if f_.endswith('.docx')])

print('\nResultado por bloque:')
for b, (o_, f_) in RES.items():
    print('  %s: %d correctos, %d fallos' % (b, o_, f_))
print('Resultado: %d correctos, %d fallos' % (ok, fallos))
json.dump({'correctos': ok, 'fallos': fallos, 'bloques': RES, 'controles': DETALLE},
          open(os.path.join(L6, 'verificacion_resultado.json'), 'w', encoding='utf8'), ensure_ascii=False, indent=1)
sys.exit(1 if fallos else 0)
