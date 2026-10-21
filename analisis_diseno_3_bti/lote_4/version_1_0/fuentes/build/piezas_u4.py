# -*- coding: utf-8 -*-
"""Piezas de armado del lote 4 que no están en la maqueta ni en las piezas de los lotes 2 y 3: portada del lote,
bloques de código con su ficha (decisión 4 de Fernando) y recuadros de contrato de las operaciones auxiliares.
Usa sin modificarlas la maqueta O3 (lote_2/build/oficio.py) y las piezas de los lotes 2 y 3 (portada, leyendas,
contratos de OP-01 a OP-06 y encadenado de tablas largas)."""
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import oficio as O
H, M, K = O.H, O.M, O.K
import piezas_u3 as Z3
import construccion as C
import fragmentos as F
Z2 = Z3.Z2
X = C.X

VERSION = 'v1.0 (lote 4, para auditoría) · oficio O3'
LINEA = 'LOTE EDITORIAL 4 — UNIDAD IV (CAPÍTULOS 10 A 13)'
MONO = 'Courier New'          # en LibreOffice se imprime con Liberation Mono, de la misma métrica
CUERPO_CODIGO = 7.5           # 96 caracteres entran en el ancho útil del perfil O3
FONDO_CODIGO = 'F3F3F3'

codigos = Z2.codigos
contrato = Z3.contrato
encadenar_inicio = Z3.encadenar_inicio
tabla_leyenda = Z3.tabla_leyenda


def portada(doc, tipo_doc, estado):
    Z2.VERSION = VERSION
    Z2.portada(doc, tipo_doc, estado, linea=LINEA)


def contrato_aux(doc, cod):
    """Recuadro con las partes del contrato de una operación auxiliar de la Unidad IV (caso/construccion.py)."""
    c, nombre, pant, ent, pre, res, err, req, prueba, cap = C.AUX[cod]
    H.caja(doc, 'Contrato %s — %s (operación auxiliar de %s)' % (c, nombre, pant),
           ['Entradas: ' + ent, 'Precondiciones: ' + pre, 'Resultado: ' + res,
            'Errores: ' + ', '.join('%s (%s)' % (e, C.USO_MENSAJE.get(e, 'sin permiso' if e == 'ME-13' else e))
                                    for e in err),
            'Requisitos: %s. Prueba: %s.' % (', '.join(req), prueba)], tipo='neutra')


def _p_codigo(celda, linea, primera):
    p = celda.paragraphs[0] if primera else celda.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = Pt(CUERPO_CODIGO * 1.22)
    pf.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(linea if linea else ' ')
    r.font.name = MONO
    r.font.size = Pt(CUERPO_CODIGO)
    rpr = r._element.get_or_add_rPr()
    fuentes = rpr.find(qn('w:rFonts'))
    for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fuentes.set(qn(k), MONO)
    return p


def codigo(doc, fid, explicacion=None):
    """Fragmento de código leído del proyecto, con su ficha (archivo, propósito, requisito o contrato, entradas,
    resultado y prueba) y, si se da, una explicación breve. El rótulo, la ficha y el código se mantienen juntos."""
    f = F.ficha(fid)
    rot = H.p_texto(doc, '%s — %s' % (fid, f['proposito']), bold=True, size=10, color=H.ID['MED'], despues=2)
    H.keep_next(rot)
    t = doc.add_table(rows=5, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    H._bordes(t, color=H.ID['SOFT'])
    datos = [('Archivo', f['archivo']), ('Requisito o contrato', f['requisito']), ('Entradas', f['entradas']),
             ('Resultado', f['resultado']), ('Prueba', f['prueba'])]
    for fila, (k, v) in zip(t.rows, datos):
        rk = fila.cells[0].paragraphs[0].add_run(k)
        rk.bold = True; rk.font.size = Pt(8.5)
        H.sombrear_celda(fila.cells[0], H.ID['LIGHT'])
        rv = fila.cells[1].paragraphs[0].add_run(v); rv.font.size = Pt(8.5)
        for c in fila.cells:
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            c.paragraphs[0].paragraph_format.keep_with_next = True
    M.fijar_anchos(t, [3.4, 13.2])
    H.cant_split(t)
    sep = doc.add_paragraph(); sep.paragraph_format.space_after = Pt(0); H.keep_next(sep)
    b = doc.add_table(rows=1, cols=1)
    b.alignment = WD_TABLE_ALIGNMENT.CENTER
    H._bordes(b, color='BFBFBF')
    celda = b.rows[0].cells[0]
    H.sombrear_celda(celda, FONDO_CODIGO)
    for i, linea in enumerate(F.extraer(fid)):
        _p_codigo(celda, linea, i == 0)
    celda.paragraphs[-1].paragraph_format.keep_with_next = bool(explicacion)
    M.fijar_anchos(b, [16.6])
    H.cant_split(b)
    if explicacion:
        H.p_texto(doc, explicacion, size=10, antes=3)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)


def aviso_entorno(doc):
    H.caja(doc, *C.AVISO_ENTORNO, tipo='neutra')
