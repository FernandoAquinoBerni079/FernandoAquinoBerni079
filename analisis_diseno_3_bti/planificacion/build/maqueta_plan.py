# -*- coding: utf-8 -*-
"""Maqueta del plan anual y de los planes de clase: el formato comercial aprobado el 20-oct-2026
(formatos_comerciales/formatos.py, que se importa sin modificarlo): oficio vertical O3 con márgenes espejados, Calibri 10,
campos editables de institución, docente, sección y año lectivo, anexos extensos en horizontal y letra no menor de 8 pt.
Agrega lo que el formato en blanco no necesitaba: texto con negritas, listas, celdas combinadas y división selectiva de
filas (las filas largas de los momentos pueden pasar de página; las cortas no se parten)."""
import importlib.util, os, re, subprocess
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
ADS = os.path.dirname(os.path.dirname(AQUI))
_spec = importlib.util.spec_from_file_location('formatos', os.path.join(ADS, 'formatos_comerciales', 'formatos.py'))
F = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(F)

VERDE, CLARO, PAPEL, MARG = F.VERDE, F.CLARO, F.PAPEL, F.M
UTIL = PAPEL[0] - MARG['interior'] - MARG['exterior']          # 17,29 cm en vertical
UTIL_H = PAPEL[1] - MARG['superior'] - MARG['inferior']        # 30,02 cm en el anexo horizontal
campo, sombrear, bordes, anchos, pie = F.campo, F.sombrear, F.bordes, F.anchos, F.pie
MIN_PT = 8


def documento():
    d = F.documento()
    st = d.styles['Normal']
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.line_spacing = 1.05
    return d


def rico(p, texto, tam=None, color=None):
    """Agrega a un párrafo un texto con **negritas**."""
    for i, trozo in enumerate(re.split(r'\*\*', texto)):
        if not trozo:
            continue
        r = p.add_run(trozo)
        r.bold = i % 2 == 1
        if tam: r.font.size = Pt(tam)
        if color: r.font.color.rgb = RGBColor.from_string(color)
    return p


def parrafo(d, texto, tam=10, italica=False, antes=0, despues=2, centro=False, color=None, seguir=False):
    p = d.add_paragraph()
    rico(p, texto, tam, color)
    if italica:
        for r in p.runs: r.italic = True
    pf = p.paragraph_format; pf.space_before = Pt(antes); pf.space_after = Pt(despues)
    if seguir: pf.keep_with_next = True
    if centro: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def titulo(d, texto, nivel=1, nueva_pagina=False):
    tam = {1: 14, 2: 11.5, 3: 10.5}[nivel]
    p = d.add_paragraph(style='Heading %d' % nivel)
    if nueva_pagina:
        p.paragraph_format.page_break_before = True
    r = p.add_run(texto); r.font.size = Pt(tam); r.font.color.rgb = RGBColor.from_string(VERDE); r.bold = True
    r.font.name = 'Calibri'
    p.paragraph_format.space_before = Pt(10 if nivel == 1 else 7); p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    return p


def lista(d_o_celda, items, tam=10, primera=False):
    """Lista con viñetas; en una celda reutiliza el primer párrafo vacío."""
    ps = []
    for i, it in enumerate(items):
        if primera and i == 0 and hasattr(d_o_celda, 'paragraphs') and not d_o_celda.paragraphs[0].text:
            p = d_o_celda.paragraphs[0]
        else:
            p = d_o_celda.add_paragraph()
        rico(p, '• ' + it, tam)
        p.paragraph_format.left_indent = Cm(0.35); p.paragraph_format.first_line_indent = Cm(-0.3)
        p.paragraph_format.space_after = Pt(1)
        ps.append(p)
    return ps


def celda(c, contenido, tam=8.5, fondo=None):
    """contenido: texto, o lista de bloques (texto o ('lista', [items]))."""
    if fondo:
        sombrear(c, fondo)
    bloques = contenido if isinstance(contenido, list) else [contenido]
    primero = True
    for b in bloques:
        if isinstance(b, tuple) and b[0] == 'lista':
            lista(c, b[1], tam, primera=primero)
        else:
            p = c.paragraphs[0] if primero else c.add_paragraph()
            rico(p, str(b), tam)
            p.paragraph_format.space_after = Pt(1)
        primero = False


def tabla(d, enc, filas, cm, tam=8.5, dividir_filas=False, fondo_primera_col=False, encabezado=True):
    t = d.add_table(rows=(1 if encabezado else 0) + len(filas), cols=len(cm)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    bordes(t)
    k = 0
    if encabezado:
        for j, h in enumerate(enc):
            c = t.rows[0].cells[j]; sombrear(c, VERDE)
            r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(tam); r.font.color.rgb = RGBColor.from_string('FFFFFF')
        k = 1
    for i, f in enumerate(filas):
        for j, v in enumerate(f):
            celda(t.rows[i + k].cells[j], v, tam, CLARO if (fondo_primera_col and j == 0) else None)
    anchos(t, cm)
    for i, r in enumerate(t.rows):
        tr = r._tr.get_or_add_trPr()
        if not dividir_filas or i == 0:
            tr.append(OxmlElement('w:cantSplit'))
        if i == 0 and encabezado:
            tr.append(OxmlElement('w:tblHeader'))
    return t


def combinar_fila(t, i, texto, tam=8.5, fondo=CLARO, negrita=True):
    fila = t.rows[i]
    c = fila.cells[0].merge(fila.cells[-1])
    for p in c.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    c.paragraphs[0].text = ''
    r = c.paragraphs[0].add_run(texto); r.bold = negrita; r.font.size = Pt(tam)
    sombrear(c, fondo)
    return c


def seccion_horizontal(d):
    sec = d.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(PAPEL[1]), Cm(PAPEL[0])
    sec.left_margin, sec.right_margin = Cm(MARG['superior']), Cm(MARG['inferior'])
    sec.top_margin, sec.bottom_margin = Cm(MARG['exterior']), Cm(MARG['interior'])
    return sec


def seccion_vertical(d):
    sec = d.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(PAPEL[0]), Cm(PAPEL[1])
    sec.left_margin, sec.right_margin = Cm(MARG['interior']), Cm(MARG['exterior'])
    sec.top_margin, sec.bottom_margin = Cm(MARG['superior']), Cm(MARG['inferior'])
    return sec


def salto(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def recuadro(d, titulo_, lineas, tam=9, fondo='FFF8E1', borde='8D6E00'):
    t = d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; sombrear(c, fondo)
    pr = t._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for lado in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + lado); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '8'); e.set(qn('w:color'), borde); b.append(e)
    pr.append(b)
    p = c.paragraphs[0]; r = p.add_run(titulo_); r.bold = True; r.font.size = Pt(tam + 0.5)
    for l in lineas:
        rico(c.add_paragraph(), l, tam)
    anchos(t, [UTIL])
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def a_pdf(ruta, destino):
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', destino, ruta], check=True,
                   capture_output=True, env=dict(os.environ, HOME='/tmp/lohome_plan'))
