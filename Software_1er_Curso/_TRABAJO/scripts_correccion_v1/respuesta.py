import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
d=docx.Document()
s=d.sections[0]; s.page_width,s.page_height=Cm(21.59),Cm(33.02)
for m in ('left_margin','right_margin','top_margin','bottom_margin'): setattr(s,m,Cm(1.27))
st=d.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10)
VERDE=RGBColor(0x1B,0x5E,0x20)
def h(t,l=1):
    p=d.add_heading(t,l)
    for r in p.runs: r.font.color.rgb=VERDE
def par(t,b=False,i=False):
    p=d.add_paragraph(); r=p.add_run(t); r.bold=b; r.italic=i; return p
def shade(c,hexc):
    tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),hexc); tcPr.append(sh)
def tabla(cab,filas,anchos):
    t=d.add_table(rows=1,cols=len(cab)); t.style='Table Grid'
    for i,c in enumerate(cab):
        cell=t.rows[0].cells[i]; cell.text=''; r=cell.paragraphs[0].add_run(c); r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9); shade(cell,'1B5E20')
    for f in filas:
        row=t.add_row()
        for i,v in enumerate(f):
            row.cells[i].text=''; r=row.cells[i].paragraphs[0].add_run(v); r.font.size=Pt(9)
            if i==len(f)-1 or cab[i]=='Estado':
                pass
    for row in t.rows:
        trPr=row._tr.get_or_add_trPr(); cs=OxmlElement('w:cantSplit'); trPr.append(cs)
        for i,w in enumerate(anchos): row.cells[i].width=Cm(w)
    t.autofit=False
    tot=sum(anchos); anchos=[a*19.05/tot for a in anchos]
    for i,gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int(anchos[i]*567)))
    for row in t.rows:
        for i,w in enumerate(anchos): row.cells[i].width=Cm(w)
    tblPr=t._tbl.tblPr; tw=tblPr.find(qn('w:tblW'))
    if tw is None: tw=OxmlElement('w:tblW'); tblPr.append(tw)
    tw.set(qn('w:type'),'dxa'); tw.set(qn('w:w'),str(int(19.05*567)))
    lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
    d.add_paragraph()
    return t

p=d.add_paragraph(); r=p.add_run('05_RESPUESTA_CLAUDE_SOFTWARE_1_ESENCIAL_v1'); r.bold=True; r.font.size=Pt(16); r.font.color.rgb=VERDE
par('Respuesta del constructor a 04_AUDITORIA_CHATGPT_SOFTWARE_1_ESENCIAL_v1 · PROTOCOLO_CLAUDE_CHATGPT_2026 v0.4, §3 paso 5')
par('Gabinete de Informática – Software · 1.er Curso BTI · Edición Esencial · 02/10/2026')
tabla(['Campo','Valor'],[
 ['Paquete','Gabinete_Informatica_Software_1er_Curso (Edición Esencial)'],
 ['Versión auditada','PDF de Drive de 02/10/2026: Libro de 113 págs. (oficio), Solucionario de 22, Planes de 61 y Plan anual'],
 ['Versión corregida (v1)','Libro: 107 págs. y Solucionario: 17 págs., los dos en oficio con márgenes estrechos · Planes de clase: 61 · Plan anual: 12 · Muestra: 14'],
 ['Resultado','14 de 14 hallazgos atendidos: 12 aplicados tal cual o con datos propios y 2 aplicados con otra forma, explicada abajo (H2 y H14). Además, 7 correcciones que no estaban en la auditoría (sección 3).'],
 ['Control técnico','Auditoría automática auditoria_v1.py: 117 controles, 0 fallos. Índices del Libro, del Solucionario y de los Planes verificados contra el PDF (0 diferencias).'],
 ['Pedido','Como hubo hallazgos críticos, el paquete vuelve a ChatGPT para una nueva auditoría (paso 4) sobre esta versión.'],
],[4.5,14.5])

h('1. Hallazgos de la auditoría: qué se aplicó')
F=[
['1','CRÍTICO','Aplicado','La capacidad «Reconoce la importancia de la capacidad de abstracción, análisis y síntesis en los trabajos realizados.» se incorpora textual en la Ficha K del Libro, con un párrafo que la explica y un paso 5 nuevo (síntesis escrita del diagnóstico); en el Plan de Clase 30 (capacidad, indicador nuevo y cierre con la síntesis) y en la secuencia 30 del Plan anual (capacidad, indicador y procedimiento). Respuesta del paso 5 agregada al Solucionario. Con esto quedan cubiertas las 20 capacidades oficiales del programa (pp. 59–63).'],
['2','MEDIO','Aplicado con ajuste de forma','La trazabilidad se resuelve en el Plan anual, pero sin agregar una octava columna: la capacidad MEC va dentro de la celda de cada secuencia, debajo del tema, con el rótulo «Capacidad MEC:» y la transcripción textual entre comillas. El encabezado pasa a decir «TEMAS Y CAPACIDAD MEC». Motivo: el formato canónico del Plan anual tiene 7 columnas en 26,7 cm de ancho útil, y una octava columna estrecharía la de indicadores. En las evaluaciones de unidad se repiten todas las capacidades que consolidan. En las dos integradoras de etapa (secuencias 18 y 36) se remite a las secuencias que las contienen, para que la fila no ocupe más de una hoja; los Planes de clase 18 y 36 sí las transcriben completas.'],
['3','MEDIO','Aplicado','Las capacidades fusionadas o recortadas se reemplazan por las oficiales, separadas y textuales, en las fichas de las Clases 9, 11, 12, 16, 17 y 18 del Libro, en los 36 Planes de clase y en las 36 secuencias del Plan anual. El control automático comprueba que las 20 capacidades aparezcan literalmente en los tres documentos y que no quede ninguna fusión.'],
['4','CRÍTICO','Aplicado','En los Criterios generales del Plan anual se usa el texto propuesto. Verificado: la Ley 7593/2025 se promulgó el 27/11/2025 y tiene una vacatio legis de 24 meses (art. 57), así que entra en vigencia en noviembre de 2027. El Libro ya la presentaba bien.'],
['5','CRÍTICO','Aplicado con datos propios','Los datos propuestos por el auditor no se usaron porque chocaban con otros del Libro: «HDD 2 TB» ya está en la Ficha A y «1 TB» en la Clase 3 y en el ejemplo de la Clase 7. Datos nuevos, que no aparecen en ninguna otra parte del paquete: Actividad 11 de la Clase 7, «Otro proveedor ofrece un SSD de 512 GB a ₲ 320.000 y un HDD de 3 TB a ₲ 560.000»; Integradora de la 1.ª etapa, ítem 3 b), «lista de precios de un mayorista (SSD 256 GB ₲ 210.000; HDD 6 TB ₲ 1.020.000)». Las respuestas del Solucionario se rehicieron e incluyen el costo por GB (≈ ₲ 625 frente a ≈ ₲ 187).'],
['6','CRÍTICO','Aplicado y ampliado','Se armó un control de operandos para toda la Unidad 3 (ejemplos, actividades, Práctica 20, Evaluación de la Unidad 3 e Integradora de la 2.ª etapa): dos operaciones no pueden compartir dos o más números. Además de lo que señaló el auditor, aparecieron 4 casos más (ver sección 3). Cambios en la Clase 20: ejemplo «El control del stock», 1010 + 1001 = 10011 (10 + 9 = 19); ítem 3, 1010 × 11 (opciones nuevas, la correcta sigue siendo la a); ítem 4, 10100 ÷ 100 (opciones nuevas, la correcta sigue siendo la d); ítem 8, 11011 − 1010 = 10001 (VERDADERO). La propuesta del auditor (11001 − 1010 = 1111) compartía 1010 y 1111 con el producto parcial 101 + 1010 = 1111 del ejemplo. Ítem 12: estantes de 1110₂ y 1010₂ (suma 11000; diferencia 100). La Figura 20.1 y la suma 1011 + 110 de la explicación no se tocaron: son el ejemplo resuelto.'],
['7','MEDIO','Aplicado','Clase 9 (ejemplo «La licencia de la notebook»), Clase 12 (párrafo sobre activación) y Solucionario de la Clase 12, ítem 10, reescritos con la propuesta del auditor: la copia sin licencia «puede implicar infracción» de la Ley 1328/98, y las actualizaciones críticas de seguridad siguen disponibles en Windows no genuino (confirmado por Microsoft). Se agrega el riesgo de los activadores no oficiales.'],
['8','MEDIO','Aplicado','El ítem 6 de la Clase 8 se reemplaza por la definición propuesta. La respuesta del Solucionario («grooming») no cambia.'],
['9','MENOR','Aplicado','Plan 6: «Figuras 5.1 y 5.2».'],
['10','MENOR','Aplicado','Plan 11: «Figuras 8.1 y 8.2».'],
['11','MENOR','Aplicado','«del gabinete» en el Plan 6 (indicador y criterio) y en la secuencia 6 del Plan anual.'],
['12','MEDIO','Aplicado','Ver sección 2. El Libro queda sin páginas en blanco ni casi vacías: se midió la parte ocupada de cada página y ninguna interior baja del 25 %. La p. 61 de los Planes ya no está casi vacía. Índice regenerado y verificado contra el PDF.'],
['13','MEDIO','Aplicado','Nuevos nombres: Gabinete_Informatica_Software_1er_Curso_PLAN_ANUAL_COMERCIAL_2026 y …_PLANES_DE_CLASE_COMERCIAL_2026 (.docx y .pdf). Tomo y Solucionario conservan _ESENCIAL_.'],
['14','MENOR','Aplicado con otro diagnóstico','Las 137 páginas de las notas no eran un error de las notas: el .docx maestro estaba en A4 (137 págs.) y el PDF de Drive en oficio (113 págs.), o sea que el PDF y el DOCX no estaban sincronizados. Ahora los dos están en oficio y tienen 107 páginas. En 03_NOTAS hay que poner «(.docx y .pdf, 107 páginas)». Ese documento es de Google Docs y el conector no permite editar su texto, así que queda para Fer.'],
]
tabla(['N.º','Gravedad','Estado','Qué se hizo'],F,[1,1.8,2.6,13.6])

h('2. Formato del tomo (indicación de Fer, 02/10/2026)')
for t in ['El estándar para los tomos es oficio (21,59 × 33,02 cm), con márgenes estrechos (1,27 cm) y tablas ajustadas a la ventana. El Libro se rehízo con ese formato: todas las secciones en oficio, las 181 tablas al 100 % del ancho y con las columnas reescaladas en proporción.',
 'La portada y la contraportada son imágenes compuestas en proporción A4. Para no deformarlas ni recortar su texto, se agregaron filas del mismo verde en las franjas lisas y se llevaron a la proporción de oficio. La foto de Gemini no se tocó.',
 'Páginas: los saltos de página estaban en párrafos vacíos propios, y cuando la hoja anterior quedaba llena generaban una página en blanco. Se pasaron a «salto de página antes» del título siguiente. Los subtítulos de las actividades llevan «mantener con el siguiente». En 7 bloques que dejaban una hoja casi vacía se redujo el espacio entre párrafos, solo en ese bloque. La Figura 15.2 pasó de 16 a 13 cm de ancho: la imagen no se modificó, solo su tamaño en la página.',
 'Se respetaron las opciones en 2 y 4 columnas que Fer armó en Word (Clases 5, 7 y 8).',
 'El Solucionario también pasó a oficio (decisión de Fer, 02/10/2026): 17 páginas, ninguna casi vacía, índice verificado. En dos bloques (Evaluaciones y Prácticas) se bajó el interlineado a 0,95 para que el final de la parte no se pasara a otra hoja. Los Planes de clase y el Plan anual siguen en A4.',
 'El PDF se generó con LibreOffice 24.2, con fuentes métricamente compatibles con Calibri y Cambria. Si al abrirlo en Word cambia algún número, se actualiza con F9 sobre el índice.']:
    par('• '+t)

h('3. Correcciones que no estaban en la auditoría')
tabla(['Documento','Qué se encontró','Corrección'],[
 ['Libro + Solucionario','Clase 20, ítem 4: 1111 ÷ 101 = 11 es la inversa exacta del ejemplo resuelto 101 × 11 = 1111 (los mismos tres números).','10100 ÷ 100 = 101 (20 ÷ 4 = 5).'],
 ['Libro + Solucionario','Clase 20, ítem 3: 110 × 11 = 10010 es la inversa exacta de la Eval. U3, 3 d) 10010 ÷ 11 = 110.','Ítem 3: 1010 × 11 = 11110. Eval. U3, 3 d): 11000 ÷ 11 = 1000 (24 ÷ 3 = 8).'],
 ['Libro + Solucionario','Práctica 20: la división 11110 ÷ 110 = 101 es la inversa exacta del ejemplo 110 × 101 = 11110.','11100 ÷ 100 = 111 (28 ÷ 4 = 7).'],
 ['Libro + Solucionario','La suma de la Práctica 20, 10110 + 1101, comparte dos números con la Integradora 2, 4 b) 10110 − 1001 = 1101 y con el ejemplo de complemento a dos.','Práctica 20: 10111 + 1110 = 100101 (23 + 14 = 37). Integradora 2, 4 b): 11010 − 1001 = 10001 (26 − 9 = 17).'],
 ['Planes de clase','Plan 10 (Ficha C, almacenamiento) tenía la capacidad «Identifica los elementos básicos de la computadora»; Plan 21 (Ficha H, comparación de sistemas actuales) tenía «origen y evolución»; Planes 15 y 17 tenían solo la capacidad de tipos de software y aplicaciones actuales.','Capacidades oficiales que corresponden: Plan 10 → almacenamiento; Plan 21 → ventajas y desventajas de los sistemas operativos actuales; Plan 15 → las que integra la Ficha F; Plan 17 → las de la Unidad 1.'],
 ['Plan anual','Errata «aplicando el convención indicada» (secuencia 3).','«aplicando la convención indicada».'],
 ['Libro','Los saltos de página en párrafos vacíos podían generar páginas en blanco.','Pasados a «salto de página antes» (ver sección 2).'],
],[3,8,8])

h('4. Lo que no se cambió y por qué')
for t in ['Clase 9, ítem 9 («Usar una copia de un programa sin licencia viola la Ley 1328/98») y el desarrollo de software propietario («usar copias sin licencia es ilegal») no se tocaron. El auditor no los marcó, y para el uso no autorizado de un programa protegido las dos afirmaciones son correctas. El matiz de «puede implicar» se aplicó solo donde el auditor lo pidió, porque ahí el texto mezclaba lo legal con las actualizaciones.',
 'La explicación de la Clase 20 conserva 1011 + 110 = 10001, porque es el ejemplo resuelto que muestra la Figura 20.1. Lo que se cambió es todo lo que lo repetía.',
 'La duración de la hora cátedra (40 minutos) declarada en los Planes no se modificó: el documento ya avisa que la institución reajusta si usa otra duración.']:
    par('• '+t)

h('5. Verificación (auditoria_v1.py)')
tabla(['Control','Resultado'],[
 ['20 capacidades oficiales textuales en Libro, Planes y Plan anual','60/60'],
 ['Capacidades fusionadas residuales (6 fórmulas × 3 documentos)','0'],
 ['Datos de almacenamiento: el ejemplo de la Clase 7 aparece una sola vez; datos nuevos, uno por ítem y presentes en el Solucionario','OK'],
 ['Operaciones de la Unidad 3 que comparten 2 o más números (fuera del mismo ejemplo)','0'],
 ['Aritmética de las 9 operaciones nuevas y su transcripción en el Solucionario','OK'],
 ['Ley 7593/2025, Windows no genuino, grooming, erratas y figuras','OK'],
 ['Índice del Libro, del Solucionario y de los Planes contra el PDF','0 diferencias'],
 ['Páginas interiores ocupadas en menos del 25 % (Libro y Solucionario)','0'],
 ['Tamaño del DOCX del Libro (límite 25 MB)','23,3 MB'],
],[13,6])

h('6. Pendiente de Fer')
for t in ['Subir a la carpeta del paquete los archivos de esta versión y pasar a SUPERSEDIDOS los anteriores, incluidos los dos .docx del Libro que hay hoy en Drive (…_2026.docx y …_2026_1.docx). El conector no permite subir el .docx del Libro (23 MB).',
 'Actualizar en 03_NOTAS el número de páginas a 107 (hallazgo 14).',
 'Volver a pasar el paquete a ChatGPT para la auditoría v2.']:
    par('• '+t)
d.core_properties.author='Equipo editorial'; d.core_properties.title='05_RESPUESTA_CLAUDE_SOFTWARE_1_ESENCIAL_v1'
d.save('out/05_RESPUESTA_CLAUDE_SOFTWARE_1_ESENCIAL_v1.docx')
print('ok')
