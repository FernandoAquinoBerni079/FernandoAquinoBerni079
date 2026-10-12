# -*- coding: utf-8 -*-
"""Corrección v3 (conciliación de auditorías, autorizada por Fer el 12/10/2026): H02 y H15.
H02: filas 18 y 36 del Plan Anual con las capacidades MEC literales (Anexo B del informe exhaustivo = caps.py).
H15: se quitan del texto del Libro y del Solucionario los pictogramas ✓ ✗ ✅ ✔ (las imágenes no se tocan).
Uso: python3 fix_v3.py DIR_ENTRADA DIR_SALIDA"""
import sys, copy, re, docx
from docx.oxml.ns import qn
from caps import C, U1, U2, U3
ENT, SAL = sys.argv[1].rstrip('/') + '/', sys.argv[2].rstrip('/') + '/'
N = 'Gabinete_Informatica_Software_1er_Curso_'
T = lambda el: ''.join(t.text or '' for t in el.iter(qn('w:t')))
LOG = []
# ---- H02 ----
d = docx.Document(ENT + N + 'PLAN_ANUAL_COMERCIAL_2026.docx')
FILAS = {'18': (U1, 'Todas las capacidades de la Unidad 1, transcriptas en las secuencias 1 a 15.'),
         '36': (U2 + U3, 'Todas las capacidades de las Unidades 2 y 3, transcriptas en las secuencias 19 a 34.')}
hechas = set()
for t in d.tables:
    for r in t.rows:
        n = r.cells[0].text.strip()
        if n in FILAS and n not in hechas:
            caps, viejo = FILAS[n]
            ps = [p for p in r.cells[1]._tc.iter(qn('w:p')) if T(p).strip() == viejo]
            assert len(ps) == 1, (n, len(ps))
            p0 = ps[0]; ts = list(p0.iter(qn('w:t'))); assert len(ts) == 1
            prev = p0
            for k in caps:
                q = copy.deepcopy(p0); list(q.iter(qn('w:t')))[0].text = '«%s»' % C[k]
                prev.addnext(q); prev = q
            p0.getparent().remove(p0)
            # la fila crece: se le permite partirse entre hojas (el encabezado de la tabla se repite), para no dejar
            # media hoja en blanco antes de ella
            for cs in r._tr.iter(qn('w:cantSplit')): cs.getparent().remove(cs)
            hechas.add(n); LOG.append('H02 fila %s: %d capacidades literales' % (n, len(caps)))
assert hechas == {'18', '36'}
d.save(SAL + N + 'PLAN_ANUAL_COMERCIAL_2026.docx')
# ---- H15 ----
SY = '✓✗✅✔'
for doc in ('LIBRO_ESENCIAL_COMERCIAL_2026', 'SOLUCIONARIO_DOCENTE_ESENCIAL_COMERCIAL_2026'):
    d = docx.Document(ENT + N + doc + '.docx'); quitados = 0; parrafos = 0
    for p in d.element.body.iter(qn('w:p')):
        ts = list(p.iter(qn('w:t')))
        if not any(c in (t.text or '') for t in ts for c in SY): continue
        parrafos += 1
        antes = T(p)
        for i, t in enumerate(ts):
            s = t.text or ''
            for c in SY:
                quitados += s.count(c)
                s = s.replace(' ' + c, '').replace(c + ' ', '').replace(c, '')
            t.text = s
        # si el pictograma iba en su propio run, puede quedar un espacio inicial o doble
        for t in ts:
            if t.text:
                t.text = t.text.lstrip() if t is next(x for x in ts if x.text) and not antes[:1].isspace() else t.text
                break
        despues = T(p)
        # control: el texto final es el original sin pictogramas ni sus espacios de separación
        ref = antes
        for c in SY: ref = ref.replace(' ' + c, '').replace(c + ' ', '').replace(c, '')
        assert re.sub(r'\s+', ' ', despues).strip() == re.sub(r'\s+', ' ', ref).strip(), (antes, despues)
        assert '  ' not in despues.strip() or '  ' in antes
    d.save(SAL + N + doc + '.docx')
    LOG.append('H15 %s: %d pictogramas quitados en %d párrafos' % (doc.split('_')[0], quitados, parrafos))
for l in LOG: print(l)
