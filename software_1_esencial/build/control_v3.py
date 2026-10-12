# -*- coding: utf-8 -*-
"""Controles de la v3 (H02 y H15) sobre ../salida, comparando con el commit 558a335 (v2). Uso: python3 control_v3.py"""
import re, io, zipfile, hashlib, subprocess, docx, pymupdf
from docx.oxml.ns import qn
from caps import C, SEQ, U1, U2, U3
BASE = '558a335'
R = 'software_1_esencial/salida/'
N = 'Gabinete_Informatica_Software_1er_Curso_'
DOCS = {'L': 'LIBRO_ESENCIAL_COMERCIAL_2026', 'S': 'SOLUCIONARIO_DOCENTE_ESENCIAL_COMERCIAL_2026',
        'A': 'PLAN_ANUAL_COMERCIAL_2026', 'P': 'PLANES_DE_CLASE_COMERCIAL_2026'}
SY = '✓✗✅✔'
ok = 0; fallos = []; INFO = []
def chk(c, m):
    global ok
    if c: ok += 1
    else: fallos.append(m)
def git(path): return subprocess.run(['git', 'show', BASE + ':' + path], capture_output=True, cwd='../..').stdout
def cur(path): return open('../../' + path, 'rb').read()
T = lambda el: ''.join(t.text or '' for t in el.iter(qn('w:t')))
def paras(blob): return [T(p) for p in docx.Document(io.BytesIO(blob)).element.body.iter(qn('w:p'))]
def sinpic(s):
    for c in SY: s = s.replace(' ' + c, '').replace(c + ' ', '').replace(c, '')
    return re.sub(r'\s+', ' ', s).strip()
cm = lambda v: round(v / 72 * 2.54, 1)
# 1. Identidad de contenido con la v2, salvo filas 18/36 (Plan Anual) y pictogramas (Libro y Solucionario)
for k in 'LSP':
    a = paras(git(R + N + DOCS[k] + '.docx')); b = paras(cur(R + N + DOCS[k] + '.docx'))
    chk(len(a) == len(b) and all(sinpic(x) == re.sub(r'\s+', ' ', y).strip() for x, y in zip(a, b)), DOCS[k] + ': cambió algo más que los pictogramas')
a = paras(git(R + N + DOCS['A'] + '.docx')); b = paras(cur(R + N + DOCS['A'] + '.docx'))
viejos = ['Todas las capacidades de la Unidad 1, transcriptas en las secuencias 1 a 15.', 'Todas las capacidades de las Unidades 2 y 3, transcriptas en las secuencias 19 a 34.']
esperado = []
for x in a:
    if x == viejos[0]: esperado += ['«%s»' % C[k] for k in U1]
    elif x == viejos[1]: esperado += ['«%s»' % C[k] for k in U2 + U3]
    else: esperado.append(x)
chk(b == esperado, 'Plan Anual: el cambio no se limita a las filas 18 y 36')
INFO.append('Plan Anual: %d → %d párrafos (+%d capacidades literales, −2 remisiones)' % (len(a), len(b), len(U1) + len(U2 + U3)))
# 2. H02: filas 18 y 36 con sus capacidades literales; 20 capacidades en Libro, Plan Anual y Planes
d = docx.Document(io.BytesIO(cur(R + N + DOCS['A'] + '.docx'))); filas = {}
for t in d.tables:
    for r in t.rows:
        n = r.cells[0].text.strip()
        if n.isdigit() and n not in filas: filas[n] = re.sub(r'\s+', ' ', r.cells[1].text)
chk(len(filas) == 36, 'Plan Anual: no tiene 36 filas')
for n in range(1, 37):
    chk(all(C[k] in filas[str(n)] for k in SEQ[n]), 'Plan Anual fila %d: faltan capacidades' % n)
    chk('transcriptas en las secuencias' not in filas[str(n)], 'Plan Anual fila %d: remite en vez de transcribir' % n)
for k in 'LAP':
    txt = re.sub(r'\s+', ' ', ' '.join(paras(cur(R + N + DOCS[k] + '.docx'))))
    chk(all(c in txt for c in C.values()), DOCS[k] + ': no están las 20 capacidades')
# 3. H15: cero pictogramas en .docx y PDF de los cinco documentos
for k in 'LSAP':
    chk(not any(c in p for p in paras(cur(R + N + DOCS[k] + '.docx')) for c in SY), DOCS[k] + '.docx: quedan pictogramas')
for f in list(DOCS.values()) + ['MUESTRA_COMERCIAL_2026']:
    pdf = pymupdf.open(R.replace('software_1_esencial/', '../') + N + f + '.pdf')
    n_ = sum(pg.get_text().count(c) for pg in pdf for c in SY)
    chk(n_ == 0, f + '.pdf: %d pictogramas' % n_)
# 4. Imágenes: idénticas byte a byte a la v2 en los cuatro .docx
for k in 'LSAP':
    za = zipfile.ZipFile(io.BytesIO(git(R + N + DOCS[k] + '.docx'))); zb = zipfile.ZipFile(io.BytesIO(cur(R + N + DOCS[k] + '.docx')))
    ma = {n: hashlib.sha1(za.read(n)).hexdigest() for n in za.namelist() if n.startswith('word/media/')}
    mb = {n: hashlib.sha1(zb.read(n)).hexdigest() for n in zb.namelist() if n.startswith('word/media/')}
    chk(ma == mb, DOCS[k] + ': las imágenes cambiaron')
    if k == 'L': INFO.append('Libro: %d imágenes idénticas a la v2' % len(mb))
# 5. Paginación, tamaño, hojas vacías o casi vacías, texto fuera de margen
ESP = {'L': 107, 'S': 17, 'P': 38}
for k, f in list(DOCS.items()) + [('M', 'MUESTRA_COMERCIAL_2026')]:
    pdf = pymupdf.open('../salida/' + N + f + '.pdf'); INFO.append('%s: %d páginas' % (f, pdf.page_count))
    if k in ESP: chk(pdf.page_count == ESP[k], '%s: %d páginas (se esperaban %d)' % (f, pdf.page_count, ESP[k]))
    for i, pg in enumerate(pdf):
        w, h = cm(pg.rect.width), cm(pg.rect.height)
        chk((w, h) in ((21.6, 33.0), (33.0, 21.6)), '%s p%d: no es oficio' % (f, i + 1))
        bl = [b for b in pg.get_text('blocks') if b[4].strip() and not b[4].strip().startswith('Página')]
        im = [r for x in pg.get_images(full=True) for r in pg.get_image_rects(x[0])]
        es_tapa = (k in 'LSAP' and i == 0) or (k == 'L' and i == pdf.page_count - 1) or (k == 'M' and i in (0, 9, 10, 11))
        # página de presentación del Solucionario (preliminar de dos párrafos, igual en v1 y v2): se informa, no se cuenta como falla
        if k == 'S' and pg.get_text().lstrip().startswith('Presentación'):
            es_tapa = True; INFO.append('Solucionario p%d: «Presentación» (preliminar breve, sin cambios desde la v1)' % (i + 1))
        if not es_tapa:
            ys = [b[3] for b in bl] + [r.y1 for r in im]
            chk(ys and max(ys) / pg.rect.height >= 0.2, '%s p%d: hoja vacía o casi vacía' % (f, i + 1))
            chk(all(b[0] >= 28 and b[2] <= pg.rect.width - 28 for b in bl), '%s p%d: texto a menos de 1 cm del borde' % (f, i + 1))
# 6. Plan Anual: oficio horizontal y sin hoja de firma suelta; Planes: un plan por hoja con firmas
A = pymupdf.open('../salida/' + N + DOCS['A'] + '.pdf')
chk(all(cm(pg.rect.width) == 33.0 for pg in list(A)[1:]), 'Plan Anual: hojas no horizontales')
ult = A[-1].get_text(); chk('Firma del docente' in ult and 'Estrategias de inclusión' in ult, 'Plan Anual: la firma quedó sola')
P = pymupdf.open('../salida/' + N + DOCS['P'] + '.pdf')
nums = [int(m.group(1)) for pg in P for m in [re.search(r'Plan de Clase N\.º (\d+) de 36', pg.get_text())] if m]
chk(nums == list(range(1, 37)), 'Planes: no hay un plan por hoja en orden')
chk(all('Dirección — Visto bueno' in pg.get_text() for pg in list(P)[2:]), 'Planes: falta una firma')
# 7. Márgenes de 1,27 cm en los cuatro .docx
for k in 'LSAP':
    dd = docx.Document(io.BytesIO(cur(R + N + DOCS[k] + '.docx')))
    chk(all(abs(x.cm - 1.27) < 0.01 for s in dd.sections for x in (s.top_margin, s.bottom_margin, s.left_margin, s.right_margin)), DOCS[k] + ': márgenes ≠ 1,27 cm')
# 8. Muestra: páginas idénticas a sus fuentes, sin claves ni continuaciones
M = pymupdf.open('../salida/' + N + 'MUESTRA_COMERCIAL_2026.pdf')
L = pymupdf.open('../salida/' + N + DOCS['L'] + '.pdf'); S = pymupdf.open('../salida/' + N + DOCS['S'] + '.pdf')
mapa = [(L, 1), (L, 2), (L, 3), (L, 4), (L, 6), (L, 7), (L, 8), (L, 9), (L, 10), (L, 107), (S, 1), (P, 1), (P, 3)]
chk(M.page_count == len(mapa), 'Muestra: cantidad de páginas')
for i, (src, p) in enumerate(mapa):
    chk(M[i].get_pixmap(dpi=20).samples == src[p - 1].get_pixmap(dpi=20).samples, 'Muestra p%d distinta de su fuente' % (i + 1))
tm = ''.join(pg.get_text() for pg in M)
chk(not re.search(r'Opción correcta|Resultado esperado|Orientación:|continuación', tm), 'Muestra: claves o página de continuación')
print('Controles v3 (contra %s): %d OK · %d fallos' % (BASE, ok, len(fallos)))
for f in fallos: print(' FALLO', f)
for i in INFO: print(' ·', i)
