"""Utilidades comunes: abre la plantilla v2 de los planes, la vacía y ofrece
constructores de párrafos y tablas con la identidad visual original
(Calibri, morado 4A148C / 6A1B9A / F3E5F5, bordes BFBFBF)."""
import copy
import datetime as dt
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

import os
S = os.environ.get("QUIM3_SCRATCH", os.path.dirname(os.path.abspath(__file__)))  # carpeta con v2/ (plantilla)
PLANTILLA = f"{S}/v2/Quimica_3_CBT_Planes_de_Clase_2026_v2_COMPACTO.docx"

OSCURO = "4A148C"
MEDIO = "6A1B9A"
SUAVE = "F3E5F5"
GRIS_BORDE = "BFBFBF"
BANDA_UNIDAD = "EDE3F3"
TXT = 22  # medios puntos: 11 pt


def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def nuevo_documento(orientacion="vertical", margenes_cm=(1.6, 2.0, 1.6, 2.0), pie=""):
    doc = Document(PLANTILLA)
    body = doc.element.body
    sect = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not sect:
            body.remove(child)
    # Encabezados/pies: solo se conservan los referenciados por la sección final;
    # se quitan las referencias de primera página (no hay titlePg).
    for ref in sect.findall(qn("w:headerReference")) + sect.findall(qn("w:footerReference")):
        if ref.get(qn("w:type")) == "first":
            sect.remove(ref)
    usados = {r.get(qn("r:id")) for r in sect if r.tag in (qn("w:headerReference"), qn("w:footerReference"))}
    rels = doc.part.rels
    for rid in list(rels):
        rt = rels[rid].reltype
        if ((rt.endswith("/header") or rt.endswith("/footer")) and rid not in usados) or rt.endswith("/stylesWithEffects"):
            del rels[rid]
    sec = doc.sections[0]
    if orientacion == "horizontal":
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
        pgsz = sect.find(qn("w:pgSz"))
        pgsz.set(qn("w:orient"), "landscape")
    else:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    t, r, b, l = margenes_cm
    sec.top_margin, sec.right_margin, sec.bottom_margin, sec.left_margin = Cm(t), Cm(r), Cm(b), Cm(l)
    sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.7)
    # Pie con campos reales PAGE / NUMPAGES
    ftr = sec.footer
    ftr.is_linked_to_previous = False
    for p in ftr.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    p = ftr.paragraphs[0]
    for r_ in list(p._p):
        if r_.tag != qn("w:pPr"):
            p._p.remove(r_)
    p.alignment = 1
    def run(texto=None, campo=None):
        partes = [("text", texto)] if campo is None else [
            ("begin", None), ("instr", campo), ("separate", None), ("text", "1"), ("end", None)]
        for tipo, val in partes:
            r_ = el("w:r")
            rPr = el("w:rPr")
            rPr.append(el("w:color", **{"w:val": "595959"}))
            rPr.append(el("w:sz", **{"w:val": 17})); rPr.append(el("w:szCs", **{"w:val": 17}))
            r_.append(rPr)
            if tipo == "instr":
                it = el("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {val} "; r_.append(it)
            elif tipo == "text":
                t_ = el("w:t"); t_.set(qn("xml:space"), "preserve"); t_.text = val; r_.append(t_)
            else:
                r_.append(el("w:fldChar", **{"w:fldCharType": tipo}))
            p._p.append(r_)
    if pie:
        run(pie + "  ·  ")
    run("Página ")
    run(campo="PAGE")
    run(" de ")
    run(campo="NUMPAGES")
    # Encabezado vacío
    hdr = sec.header
    hdr.is_linked_to_previous = False
    # Idioma español para la revisión ortográfica
    rpr = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    lang = rpr.find(qn("w:lang"))
    lang.set(qn("w:val"), "es-PY")
    # Propiedades internas anónimas
    cp = doc.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    cp.comments = ""
    cp.keywords = ""
    cp.subject = ""
    cp.category = ""
    cp.revision = 1
    ahora = dt.datetime(2026, 10, 21, 12, 0, 0)
    cp.created = ahora
    cp.modified = ahora
    return doc


def _ppr(p, keep_next=False, keep_lines=False, page_break=False, before=0, after=0, line=240, jc=None):
    pPr = p._p.get_or_add_pPr()
    for tag in ("w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:widowControl", "w:spacing", "w:jc"):
        for x in pPr.findall(qn(tag)):
            pPr.remove(x)
    # orden del esquema: pStyle, keepNext, keepLines, pageBreakBefore, ..., widowControl, ..., spacing, ..., jc
    anchor = pPr.find(qn("w:pStyle"))
    seq = []
    if keep_next:
        seq.append(el("w:keepNext"))
    if keep_lines:
        seq.append(el("w:keepLines"))
    if page_break:
        seq.append(el("w:pageBreakBefore"))
    seq.append(el("w:widowControl"))
    seq.append(el("w:spacing", **{"w:before": before, "w:after": after, "w:line": line, "w:lineRule": "auto"}))
    if jc:
        seq.append(el("w:jc", **{"w:val": jc}))
    pos = 0 if anchor is None else list(pPr).index(anchor) + 1
    for i, x in enumerate(seq):
        pPr.insert(pos + i, x)
    # mover rPr de párrafo (si existiera) al final
    rpr = pPr.find(qn("w:rPr"))
    if rpr is not None:
        pPr.remove(rpr); pPr.append(rpr)


def add_runs(p, partes, size=TXT, color=None):
    """partes: str o lista de (texto, dict(b=, i=, color=))."""
    if isinstance(partes, str):
        partes = [(partes, {})]
    import re as _re
    expand = []
    for texto, fmt in partes:
        for i, trozo in enumerate(_re.split(r"\^\(([^)]*)\)", texto)):
            if trozo:
                expand.append((trozo, dict(fmt, sup=True) if i % 2 else fmt))
    for texto, fmt in expand:
        r = p.add_run(texto)
        rPr = r._r.get_or_add_rPr()
        if fmt.get("b"):
            rPr.append(el("w:b"))
        if fmt.get("i"):
            rPr.append(el("w:i"))
        c = fmt.get("color", color)
        if c:
            rPr.append(el("w:color", **{"w:val": c}))
        sz = fmt.get("size", size)
        rPr.append(el("w:sz", **{"w:val": sz}))
        rPr.append(el("w:szCs", **{"w:val": sz}))
        if fmt.get("sup"):
            rPr.append(el("w:vertAlign", **{"w:val": "superscript"}))
    return p


def parrafo(contenedor, partes, estilo=None, size=TXT, color=None, **ppr):
    if hasattr(contenedor, "add_paragraph"):
        p = contenedor.add_paragraph(style=estilo)
    else:
        raise TypeError
    _ppr(p, **ppr)
    add_runs(p, partes, size=size, color=color)
    return p


def separador(doc, alto=70, keep_next=True):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    if keep_next:
        pPr.append(el("w:keepNext"))
    pPr.append(el("w:widowControl"))
    pPr.append(el("w:spacing", **{"w:before": 0, "w:after": 0, "w:line": alto, "w:lineRule": "exact"}))
    r = p.add_run("")
    rPr = r._r.get_or_add_rPr()
    rPr.append(el("w:sz", **{"w:val": 4})); rPr.append(el("w:szCs", **{"w:val": 4}))
    return p


def tabla(doc, anchos_cm, header_repeat=False):
    """Crea una tabla de layout fijo con tblGrid real. Devuelve el objeto tabla de python-docx."""
    total = sum(anchos_cm)
    t = doc.add_table(rows=0, cols=len(anchos_cm))
    tbl = t._tbl
    tblPr = tbl.tblPr
    for x in list(tblPr):
        tblPr.remove(x)
    tw = int(round(total / 2.54 * 1440))
    tblPr.append(el("w:tblW", **{"w:w": tw, "w:type": "dxa"}))
    tblPr.append(el("w:jc", **{"w:val": "center"}))
    b = el("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b.append(el(f"w:{side}", **{"w:val": "single", "w:sz": 4, "w:space": 0, "w:color": GRIS_BORDE}))
    tblPr.append(b)
    tblPr.append(el("w:tblLayout", **{"w:type": "fixed"}))
    mar = el("w:tblCellMar")
    mar.append(el("w:top", **{"w:w": 28, "w:type": "dxa"}))
    mar.append(el("w:left", **{"w:w": 85, "w:type": "dxa"}))
    mar.append(el("w:bottom", **{"w:w": 28, "w:type": "dxa"}))
    mar.append(el("w:right", **{"w:w": 85, "w:type": "dxa"}))
    tblPr.append(mar)
    tblPr.append(el("w:tblLook", **{"w:val": "04A0", "w:firstRow": 1, "w:lastRow": 0, "w:firstColumn": 1,
                                   "w:lastColumn": 0, "w:noHBand": 0, "w:noVBand": 1}))
    grid = tbl.tblGrid
    for gc in list(grid):
        grid.remove(gc)
    t._anchos = [int(round(a / 2.54 * 1440)) for a in anchos_cm]
    for w in t._anchos:
        grid.append(el("w:gridCol", **{"w:w": w}))
    return t


def fila(t, celdas, cant_split=True, header=False, keep_next=False, alto_min=None):
    """celdas: lista de dict(txt=partes|lista de párrafos, fill=, span=, bold=, color=, jc=, valign=)."""
    tr = OxmlElement("w:tr")
    t._tbl.append(tr)
    trPr = el("w:trPr")
    if cant_split:
        trPr.append(el("w:cantSplit"))
    if alto_min:
        trPr.append(el("w:trHeight", **{"w:val": alto_min, "w:hRule": "atLeast"}))
    if header:
        trPr.append(el("w:tblHeader"))
    tr.append(trPr)
    col = 0
    from docx.table import _Cell
    for c in celdas:
        span = c.get("span", 1)
        w = sum(t._anchos[col:col + span])
        tc = OxmlElement("w:tc")
        tcPr = el("w:tcPr")
        tcPr.append(el("w:tcW", **{"w:w": w, "w:type": "dxa"}))
        if span > 1:
            tcPr.append(el("w:gridSpan", **{"w:val": span}))
        if c.get("fill"):
            tcPr.append(el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": c["fill"]}))
        if c.get("valign"):
            tcPr.append(el("w:vAlign", **{"w:val": c["valign"]}))
        tc.append(tcPr)
        tr.append(tc)
        cell = _Cell(tc, t)
        parrafos = c.get("pars")
        if parrafos is None:
            parrafos = [c.get("txt", "")]
        for i, partes in enumerate(parrafos):
            p = cell.add_paragraph()
            _ppr(p, keep_next=keep_next, before=c.get("before", 0), after=c.get("after", 0) if i < len(parrafos) - 1 else 0,
                 jc=c.get("jc"))
            if isinstance(partes, str):
                partes = [(partes, {"b": c.get("bold", False)})]
            add_runs(p, partes, size=c.get("size", TXT), color=c.get("color"))
        col += span
    return tr


def guardar(doc, ruta):
    # settings: updateFields antes de compat (ya está en la plantilla); limpia rsids
    st = doc.settings.element
    for x in st.findall(qn("w:rsids")):
        st.remove(x)
    for x in st.findall(qn("w:proofState")):
        st.remove(x)
    for z in st.findall(qn("w:zoom")):
        z.set(qn("w:percent"), "100")
    doc.save(ruta)
