# -*- coding: utf-8 -*-
"""Planes de Clase — Edición Esencial Comercial 2026 (formato V3, modelo Software 1.º).
Una sección por plan: título con estilo «Plan de Clase», banda, datos, capacidad/indicadores, momentos
(cantSplit selectivo: Desarrollo divisible), recursos, evaluación, observaciones y firmas; encabezado de
continuación por STYLEREF; portada a página completa; índice poblado y verificado contra el PDF."""
import os, re, json
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import maqueta as M
import planes_data
from build_libro import a_pdf

B = '/home/claude/alg3/build/'; OUT = '/home/claude/alg3/salida_v2_2/'
NOMBRE = 'Algoritmica_3er_Curso_PLANES_DE_CLASE_ESENCIAL_COMERCIAL_2026_v2_2'
DARK, MED, LIGHT, GRIS = '1B5E20', '2E7D32', 'E8F5E9', 'BFBFBF'
W = M.ANCHO
MIN_HC, HC = 40, 4
PL = planes_data.planes()


def sec_setup(sec, primera=True):
    M.formato_oficio(sec)
    sec.different_first_page_header_footer = primera


def fld(par, instr, size=8.5, italic=False, color=None):
    for el in ('begin', 'instr', 'end'):
        r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
        s = OxmlElement('w:sz'); s.set(qn('w:val'), str(int(size * 2))); rPr.append(s)
        if italic: rPr.append(OxmlElement('w:i'))
        if color:
            c = OxmlElement('w:color'); c.set(qn('w:val'), color); rPr.append(c)
        r.append(rPr)
        if el == 'instr':
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = instr
        else:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), el)
        r.append(e); par._p.append(r)


def limpiar(par):
    for el in list(par._p):
        if el.tag != qn('w:pPr'):
            par._p.remove(el)


def pie(sec):
    for foot in (sec.first_page_footer, sec.footer):
        foot.is_linked_to_previous = False
        par = foot.paragraphs[0]; par.alignment = WD_ALIGN_PARAGRAPH.CENTER; limpiar(par)
        r = par.add_run('Página '); r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string('777777')
        fld(par, ' PAGE ')
        r = par.add_run(' de '); r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string('777777')
        fld(par, ' NUMPAGES ')


def encabezado_cont(sec):
    sec.first_page_header.is_linked_to_previous = False
    for p in sec.first_page_header.paragraphs:
        limpiar(p)
    hdr = sec.header; hdr.is_linked_to_previous = False
    par = hdr.paragraphs[0]; par.alignment = WD_ALIGN_PARAGRAPH.RIGHT; limpiar(par)
    fld(par, ' STYLEREF "Plan de Clase" \\* MERGEFORMAT ', size=8.5, italic=True, color=MED)
    r = par.add_run(' · continuación'); r.font.size = Pt(8.5); r.italic = True; r.font.color.rgb = RGBColor.from_string(DARK)


def tabla(d, filas, cols, anchos, proteger=True):
    t = d.add_table(rows=filas, cols=cols); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    M.bordes(t, GRIS, '4'); M.margenes_celda(t, 40, 90)
    M.ancho_fijo(t, anchos)
    if proteger:
        M.cant_split(t)
    return t


def celda(cell, val, bold=False, blanco=False, size=9.5, center=False, vinetas=False):
    par = cell.paragraphs[0]
    if center: par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    items = val if isinstance(val, list) else [val]
    for i, linea in enumerate(items):
        if i:
            par = cell.add_paragraph()
        par.paragraph_format.space_after = Pt(2)
        r = par.add_run(('• ' if vinetas else '') + linea)
        r.bold = bold; r.font.size = Pt(size)
        if blanco: r.font.color.rgb = RGBColor.from_string('FFFFFF')


def desproteger(row):
    trPr = row._tr.get_or_add_trPr()
    cs = trPr.find(qn('w:cantSplit'))
    if cs is not None: trPr.remove(cs)


def esp(d, pts=2):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(pts); p.paragraph_format.space_before = Pt(0)
    r = p.add_run(); r.font.size = Pt(3)


def plan(d, p, total):
    idx = p['idx']
    tp = d.add_paragraph(style='Plan de Clase'); tp.add_run('Plan de Clase N.º %d — %s' % (idx, p['titulo']))
    t = tabla(d, 1, 1, [W]); M.shd(t.rows[0].cells[0], DARK)
    celda(t.rows[0].cells[0], 'Institución: ______________________________     ·     Plan de Clase N.º %d de %d · 2026' % (idx, total), bold=True, blanco=True, center=True)
    esp(d)
    datos = [('Docente', '______________________________', 'Disciplina', 'Algorítmica'),
             ('Curso', '3.er Curso — Bachillerato Técnico en Servicios, Especialidad Informática', 'Tiempo', '%d horas cátedra (%d min)' % (HC, HC * MIN_HC)),
             ('Unidad', p['unidad'], 'Sección', '____________'),
             ('Material', p['material'], 'Fecha', '____ / ____ / 2026')]
    t = tabla(d, len(datos) + 1, 4, [2.4, 6.4, 2.1, 6.1])
    for i, (l1, v1, l2, v2) in enumerate(datos):
        c = t.rows[i].cells
        M.shd(c[0], MED); celda(c[0], l1, bold=True, blanco=True); celda(c[1], v1)
        M.shd(c[2], MED); celda(c[2], l2, bold=True, blanco=True); celda(c[3], v2)
    fr = t.rows[len(datos)]
    M.shd(fr.cells[0], MED); celda(fr.cells[0], 'Tema', bold=True, blanco=True)
    celda(fr.cells[1].merge(fr.cells[3]), p['tema'])
    esp(d)
    t = tabla(d, 2, 2, [3.4, 13.6])
    M.shd(t.rows[0].cells[0], MED); capv = p['capacidad'] if isinstance(p['capacidad'], list) else [p['capacidad']]; celda(t.rows[0].cells[0], 'Capacidad' if len(capv) == 1 else 'Capacidades', bold=True, blanco=True); celda(t.rows[0].cells[1], capv, vinetas=len(capv) > 1)
    M.shd(t.rows[1].cells[0], MED); celda(t.rows[1].cells[0], ['Indicadores', 'de logro'], bold=True, blanco=True); celda(t.rows[1].cells[1], p['indicadores'], vinetas=True)
    esp(d)
    t = tabla(d, 4, 3, [2.2, 12.7, 2.1])
    for j, hh in enumerate(('MOMENTO', 'ACTIVIDADES', 'TIEMPO')):
        M.shd(t.rows[0].cells[j], DARK); celda(t.rows[0].cells[j], hh, bold=True, blanco=True, center=(j != 1))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for c in t.rows[0].cells:
        for par in c.paragraphs: par.paragraph_format.keep_with_next = True
    for i, (nombre, mins) in enumerate(zip(('Inicio', 'Desarrollo', 'Cierre'), p['tiempos']), 1):
        c = t.rows[i].cells
        M.shd(c[0], LIGHT); celda(c[0], nombre, bold=True)
        celda(c[1], p['momentos'][nombre], vinetas=True); celda(c[2], '%d min' % mins, center=True)
    desproteger(t.rows[2])
    esp(d)
    t = tabla(d, 1, 2, [3.4, 13.6])
    M.shd(t.rows[0].cells[0], MED); celda(t.rows[0].cells[0], ['Recursos y medios', 'auxiliares'], bold=True, blanco=True)
    celda(t.rows[0].cells[1], p['recursos'])
    esp(d)
    t = tabla(d, 4, 2, [3.4, 13.6])
    fb = t.rows[0].cells[0].merge(t.rows[0].cells[1]); M.shd(fb, DARK); celda(fb, 'EVALUACIÓN', bold=True, blanco=True, center=True, size=10)
    for i, (l, v, vin) in enumerate((('Procedimientos', p['proc'], False), ('Instrumentos', p['inst'], False), ('Criterios', p['criterios'], True)), 1):
        M.shd(t.rows[i].cells[0], MED); celda(t.rows[i].cells[0], l, bold=True, blanco=True); celda(t.rows[i].cells[1], v, vinetas=vin)
    esp(d)
    po = d.add_paragraph(); r = po.add_run('Observaciones: '); r.bold = True; r.font.size = Pt(9.5)
    r = po.add_run('_' * 74); r.font.size = Pt(9.5)
    t = tabla(d, 1, 2, [8.5, 8.5])
    for c in t.rows[0].cells:
        tcPr = c._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
        for e in ('top', 'left', 'bottom', 'right'):
            x = OxmlElement('w:' + e); x.set(qn('w:val'), 'nil'); b.append(x)
        tcPr.append(b)
    celda(t.rows[0].cells[0], ['', '____________________________', 'Docente'], center=True)
    celda(t.rows[0].cells[1], ['', '____________________________', 'Dirección — Visto bueno'], center=True)


COMO_SE_USA = [
 'Los planes de clase siguen la secuencia del libro del estudiante: cada clase remite a su Clase y a la Práctica del mismo número. Diez clases continúan en un encuentro propio, dedicado a las Actividades 2 y 3 de su práctica; se suman tres talleres del Proyecto Final Integrador y las dos evaluaciones integradoras de etapa. Los 36 planes completan una programación anual de 4 horas cátedra semanales; la institución puede ajustar fechas sin alterar la progresión.',
 'El tema, los indicadores, los procedimientos y los instrumentos se corresponden con el Plan Anual; los indicadores se toman literalmente de las fichas del libro. Los momentos didácticos, los recursos y los criterios de evaluación son propios de este documento.',
 'Los tiempos se calcularon con una hora cátedra de 40 minutos: 4 horas cátedra equivalen a 160 minutos, repartidos en 40 de inicio, 90 de desarrollo y 30 de cierre en los encuentros de clase, y en 20, 120 (o 100) y 20 (o 40) en las continuaciones prácticas, según incluyan la evaluación de unidad. La institución que trabaje con otra duración de hora cátedra debe reajustar el reparto en la misma proporción.',
 'Los planes que dependen del equipamiento del laboratorio (PSeInt o Microsoft Access) declaran una alternativa de trabajo en papel. La institución, el docente, la sección y la fecha quedan en blanco para su llenado.',
]


def construir(pags=None):
    d = M.nuevo_doc('Algorítmica · 3.er Curso · Planes de clase', 'Material del docente: 36 planes de clase', 'Algorítmica; BTI; 3.er Curso; planes de clase')
    st = d.styles
    st['Normal'].font.size = Pt(9.5)
    try:
        ps = st['Plan de Clase']
    except KeyError:
        ps = st.add_style('Plan de Clase', WD_STYLE_TYPE.PARAGRAPH)
    ps.base_style = st['Normal']; ps.font.name = 'Arial'; ps.font.size = Pt(13.5); ps.font.bold = True
    ps.font.color.rgb = RGBColor.from_string(DARK)
    pf = ps.paragraph_format; pf.space_before = Pt(0); pf.space_after = Pt(6); pf.keep_with_next = True
    ppr = ps.element.get_or_add_pPr()
    if ppr.find(qn('w:outlineLvl')) is None:
        ol = OxmlElement('w:outlineLvl'); ol.set(qn('w:val'), '0'); ppr.append(ol)
    sec_setup(d.sections[0], primera=False)
    M.imagen_pagina_completa(d, B + 'portadas/portada_planes.jpg')
    sec = d.add_section(WD_SECTION.NEW_PAGE); sec_setup(sec, primera=False)
    s0 = d.sections[0]
    for hf in (s0.footer, s0.header):
        hf.is_linked_to_previous = False
        for p in hf.paragraphs: limpiar(p)
    M.heading(d, 'Cómo se usa este documento', 1, estilo='Título preliminar')
    for t in COMO_SE_USA:
        M.par(d, t, 10)
    M.heading(d, 'Índice', 1, estilo='Título preliminar')
    E = [(1, 'Plan de Clase N.º %d — %s' % (p['idx'], p['titulo'])) for p in PL]
    ps_ = M.indice(d, E, pags)
    for q in ps_:
        q.style = d.styles['toc 2']; q.paragraph_format.left_indent = Cm(0)
        for r in q.runs: r.bold = False
        q.paragraph_format.tab_stops.add_tab_stop(Cm(M.ANCHO - 0.05), 2, 1)
    pie(sec)
    for p in PL:
        s = d.add_section(WD_SECTION.NEW_PAGE); sec_setup(s, primera=True); pie(s); encabezado_cont(s)
        plan(d, p, len(PL))
    M.no_actualizar_campos(d)
    M.purgar_relaciones(d)
    return d, E


if __name__ == '__main__':
    tmp = B + 'tmp/'
    d, E = construir(None)
    p1 = tmp + NOMBRE + '.docx'; d.save(p1); pdf1 = a_pdf(p1, tmp)
    import pymupdf
    doc = pymupdf.open(pdf1)
    idx_last = max(i for i, pg in enumerate(doc) if re.search(r'\.{5,}\s*000', pg.get_text())) + 1
    pags, total = M.paginas_de(pdf1, [t for _, t in E], desde=idx_last + 1)
    for it in range(3):
        d, E = construir(pags)
        out = OUT + NOMBRE + '.docx'; d.save(out); pdf = a_pdf(out, OUT)
        pags2, total = M.paginas_de(pdf, [t for _, t in E], desde=idx_last + 1)
        if pags2 == pags: break
        pags = pags2
    json.dump({'entradas': E, 'paginas': pags, 'total': total}, open(B + 'indice_planes.json', 'w'), ensure_ascii=False, indent=1)
    print('Planes: páginas', total, 'índice verificado =', pags2 == pags)
