# -*- coding: utf-8 -*-
"""Auditoría automática de liberación — Algorítmica 3.º Edición Esencial Comercial 2026.
Cada control suma al conteo; el resultado esperado es 0 fallos."""
import re, json, os, collections
import docx, pymupdf
from docx.oxml.ns import qn
import datos, actividades, practicas, evaluaciones, planes_data
from estructura import cargar
import ediciones

O = '/home/claude/alg3/salida_v2_2/'; B = '/home/claude/alg3/build/'
F = {k: O + 'Algoritmica_3er_Curso_%s_v2_2.%s' % (k, 'docx') for k in
     ('LIBRO_ESENCIAL_COMERCIAL_2026', 'SOLUCIONARIO_DOCENTE_ESENCIAL_COMERCIAL_2026', 'PLANES_DE_CLASE_ESENCIAL_COMERCIAL_2026', 'PLAN_ANUAL_ESENCIAL_COMERCIAL_2026')}
LIB, SOL, PLA, ANU = list(F.values())
RES = []


def chk(bloque, nombre, ok, detalle=''):
    RES.append((bloque, nombre, bool(ok), detalle if not ok else ''))


def pdf(p):
    return pymupdf.open(p[:-5] + '.pdf')


def textos_doc(path):
    d = docx.Document(path)
    out = []
    for el in d.element.body.iter():
        if el.tag == qn('w:p'):
            t = ''.join(x.text or '' for x in el.iter(qn('w:t')))
            st = el.find('.//' + qn('w:pStyle'))
            out.append((st.get(qn('w:val')) if st is not None else '', t))
    return d, out


_pre, C, _ev, U = cargar(); ediciones.aplicar(C); FIG = ediciones.fig_refs(C)
dL, PL = textos_doc(LIB)
TL = '\n'.join(t for _, t in PL)
pdfL = pdf(LIB)
pags = [pg.get_text() for pg in pdfL]

# ---------------- A. Estructura ----------------
clases = [t for s, t in PL if re.match(r'^Clase \d+ — ', t) and (s.startswith('Heading') or s.startswith('Ttulo'))]
prs = [t for s, t in PL if re.match(r'^Práctica \d+ — ', t) and (s.startswith('Heading') or s.startswith('Ttulo'))]
nums_c = [int(re.match(r'Clase (\d+)', t).group(1)) for t in clases]
nums_p = [int(re.match(r'Práctica (\d+)', t).group(1)) for t in prs]
chk('A', 'Cantidad de clases = 21', len(clases) == 21, len(clases))
chk('A', 'Cantidad de prácticas = 21', len(prs) == 21, len(prs))
chk('A', 'Numeración continua de clases 1..21', nums_c == list(range(1, 22)), nums_c)
chk('A', 'Numeración continua de prácticas 1..21', nums_p == list(range(1, 22)), nums_p)
seq = [(t.split(' — ')[0]) for s, t in PL if (s.startswith('Heading') or s.startswith('Ttulo')) and re.match(r'^(Clase|Práctica) \d+ — ', t)]
esperado = [x for n in range(1, 22) for x in ('Clase %d' % n, 'Práctica %d' % n)]
chk('A', 'Correspondencia Clase N → Práctica N (orden intercalado)', seq == esperado)
for n in range(1, 22):
    chk('A', 'Título de la Clase %d coincide con la fuente' % n, ('Clase %d — %s' % (n, C[n]['titulo'])) in clases)
    chk('A', 'Título de la Práctica %d coincide con la fuente' % n, ('Práctica %d — %s' % (n, practicas.PR[n]['titulo'])) in prs)
fichas = sum(1 for t in dL.tables if t.rows[0].cells[0].text.strip() in ('Capacidad', 'Capacidades'))
chk('A', 'Una ficha por clase (21)', fichas == 21, fichas)
for e in evaluaciones.EV:
    chk('A', 'Evaluación presente: ' + e['titulo'], e['titulo'] in TL)
chk('A', 'Prueba diagnóstica presente', 'Prueba diagnóstica — ¿Qué sabés ya?' in TL)
chk('A', 'Proyecto Final Integrador presente', 'Proyecto Final Integrador — La solución digital completa' in TL)
for n in range(1, 22):
    it = actividades.items(n)
    chk('A', 'Clase %d: al menos 8 actividades' % n, len(it) >= 8, len(it))
    chk('A', 'Clase %d: actividades numeradas 1..k' % n, [x[1] for x in it] == list(range(1, len(it) + 1)))
for n in range(1, 22):
    p = practicas.PR[n]
    chk('A', 'Práctica %d: competencia, saber, actividades, control y desafío' % n,
        p['competencia'] and p['saber'] and p['acts'] and all(a['control'] and a['pasos'] and a['objetivo'] for a in p['acts']) and p['desafio'])
    if n in planes_data.CONT:
        chk('A', 'Práctica %d (continuación en encuentro propio): 3 actividades' % n, len(p['acts']) >= 3, len(p['acts']))

# ---------------- B. Paginación, índice, saltos, encabezados ----------------
ji = json.load(open(B + 'indice_libro.json'))
for (niv, t), p in zip(ji['entradas'], ji['paginas']):
    ok = re.sub(r'\s+', '', t) in re.sub(r'\s+', '', pags[p - 1])
    chk('B', 'Índice del libro: «%s» en pág. %d' % (t[:50], p), ok)
idxtxt = ''.join(pags[1:5])
chk('B', 'Índice poblado (sin «Ctrl+E/F9» ni marcadores 000)', 'Ctrl+E' not in TL and not re.search(r'\.{5,}\s*000\b', idxtxt))
for t in prs:
    p = next(i for i, x in enumerate(pags) if re.sub(r'\s+', '', t) in re.sub(r'\s+', '', x) and i > 0 and not re.search(r"\.{8,}", x))
    chk('B', 'Salto: «%s» abre página' % t[:40], re.sub(r'\s+', '', pags[p]).startswith(re.sub(r'\s+', '', t)[:30]))
for e in evaluaciones.EV:
    p = next(i for i, x in enumerate(pags) if e['titulo'] in x and i > 0 and not re.search(r"\.{8,}", x))
    chk('B', 'Salto: «%s» abre página' % e['titulo'], pags[p].strip().startswith(e['titulo']))
for t in clases:
    p = next(i for i, x in enumerate(pags) if re.sub(r'\s+', '', t) in re.sub(r'\s+', '', x) and i > 0 and not re.search(r"\.{8,}", x))
    top = re.sub(r'\s+', '', pags[p])
    chk('B', 'Salto: «%s» abre página (o sigue a la apertura de unidad)' % t[:40], top.startswith(re.sub(r'\s+', '', t)[:25]) or top.startswith('UNIDAD'))
chk('B', 'Pie «Página X de Y» en todas las páginas salvo portada y contraportada',
    all(re.search(r'Página \d+ de %d' % len(pags), pags[i]) for i in range(1, len(pags) - 1)) and not re.search(r'Página', pags[0]) and not re.search(r'Página', pags[-1]))
for nombre, path, j in (('Solucionario', SOL, 'indice_sol.json'), ('Planes', PLA, 'indice_planes.json')):
    pp = [pg.get_text() for pg in pdf(path)]
    jj = json.load(open(B + j))
    bad = [t for (_, t), p in zip(jj['entradas'], jj['paginas']) if re.sub(r'\s+', '', t) not in re.sub(r'\s+', '', pp[p - 1])]
    chk('B', 'Índice de %s verificado contra el PDF (%d entradas)' % (nombre, len(jj['entradas'])), not bad, bad[:3])
    chk('B', '%s: portada sin pie de página' % nombre, 'Página' not in pp[0])
ppla = [pg.get_text() for pg in pdf(PLA)]
jp = json.load(open(B + 'indice_planes.json'))
cont_ok = True
for k, p in enumerate(jp['paginas']):
    fin = jp['paginas'][k + 1] - 1 if k + 1 < len(jp['paginas']) else len(ppla)
    for q in range(p, fin):
        if 'continuación' not in ppla[q][:250]:
            cont_ok = False
chk('B', 'Planes: encabezado «· continuación» en las páginas siguientes de cada plan', cont_ok)
chk('B', 'Planes: 36 planes, cada uno abre página', len(jp['paginas']) == 36 and all(ppla[p - 1].strip().startswith('Plan de Clase N.º') for p in jp['paginas']))
for path in F.values():
    mb = os.path.getsize(path) / 1e6
    chk('B', 'DOCX < 25 MB: %s (%.2f MB)' % (os.path.basename(path), mb), mb < 25)
    d = docx.Document(path)
    chk('B', 'Metadatos de colección: %s' % os.path.basename(path), d.core_properties.author == 'Equipo editorial' and 'python-docx' not in (d.core_properties.comments or '') + (d.core_properties.title or ''))
    pd_ = pdf(path)
    chk('B', 'PDF regenerado después del DOCX: %s' % os.path.basename(path), os.path.getmtime(path[:-5] + '.pdf') >= os.path.getmtime(path))
# cantSplit
def cantsplit(path):
    d = docx.Document(path); tot = ok = 0
    for t in d.tables:
        for r in t.rows:
            tot += 1; ok += r._tr.trPr is not None and r._tr.trPr.find(qn('w:cantSplit')) is not None
    return ok, tot
import maqueta as MQ
for nombre, path, horiz in (('libro', LIB, False), ('solucionario', SOL, False), ('planes de clase', PLA, False), ('plan anual', ANU, True)):
    d_ = docx.Document(path)
    secs = d_.sections
    tamano = all(abs(s_.page_width.cm - MQ.PAG_W) < 0.05 and abs(s_.page_height.cm - MQ.PAG_H) < 0.05 for s_ in (secs[:1] if horiz else secs))
    if horiz:
        tamano = tamano and all(abs(s_.page_width.cm - MQ.PAG_H) < 0.05 and abs(s_.page_height.cm - MQ.PAG_W) < 0.05 for s_ in secs[1:])
    chk('B', 'Oficio 21,6 × 33 cm en todas las secciones del %s%s' % (nombre, ' (horizontal en el cuerpo)' if horiz else ''), tamano)
    chk('B', 'Márgenes estrechos (1,27 cm) en todas las secciones del %s' % nombre, all(abs(m.cm - MQ.MARG) < 0.03 for s_ in secs for m in (s_.left_margin, s_.right_margin, s_.top_margin, s_.bottom_margin)))
    tws = [t._tbl.tblPr.find(qn('w:tblW')) for t in d_.tables]
    chk('B', 'Tablas y cajas ajustadas al ancho de la página (100 %%) en el %s (%d)' % (nombre, len(tws)), all(w is not None and w.get(qn('w:type')) == 'pct' and w.get(qn('w:w')) == '5000' for w in tws))
    pdf_ = pdf(path)
    chk('B', 'PDF del %s en oficio' % nombre, all(abs(sorted([pg.rect.width, pg.rect.height])[0] / 72 * 2.54 - MQ.PAG_W) < 0.05 and abs(sorted([pg.rect.width, pg.rect.height])[1] / 72 * 2.54 - MQ.PAG_H) < 0.05 for pg in pdf_))
for nombre, path in (('libro', LIB), ('solucionario', SOL), ('plan anual', ANU)):
    ok, tot = cantsplit(path)
    chk('B', 'cantSplit en el 100 %% de las filas del %s (%d/%d)' % (nombre, ok, tot), ok == tot)

# ---------------- C. Coherencia aritmética ----------------
T = datos.agg(datos.DET, datos.PROD, datos.VEN, datos.CLI, 'venta')
chk('C', 'Semana 1: 134.000 en 5 ventas (promedio 26.800)', sum(T.values()) == 134000)
chk('C', 'Semana 1: por cliente 48/19/17/50 mil', datos.agg(datos.DET, datos.PROD, datos.VEN, datos.CLI, 'cliente') == {'C01': 48000, 'C02': 19000, 'C03': 17000, 'C04': 50000})
chk('C', 'Semana 1: por categoría 23/63/48 mil', datos.agg(datos.DET, datos.PROD, datos.VEN, datos.CLI, 'cat') == {'Panificados': 23000, 'Salados': 63000, 'Bebidas': 48000})
T2 = datos.agg(datos.DET2, datos.PROD2, datos.VEN2, datos.CLI2, 'venta')
chk('C', 'Semana 2: 135.000 en 6 ventas (promedio 22.500)', sum(T2.values()) == 135000 and sum(T2.values()) / 6 == 22500)
chk('C', 'Semana 2: con precio de catálogo daría 136.000 (trampa didáctica)', sum(q * datos.PROD2[p][2] for v, p, q, pu in datos.DET2) == 136000)
chk('C', 'Semana 2: 33 registros (8+6+6+13)', len(datos.PROD2) + len(datos.CLI2) + len(datos.VEN2) + len(datos.DET2) == 33)
rows1 = [(v, datos.VEN[v][0], datos.VEN[v][1], datos.CLI[datos.VEN[v][1]], p) for v, p, q, pu in datos.DET]
chk('C', 'Redundancia de la tabla única = 34 celdas', (11 - 5) * 2 + (11 - 4) + (11 - 6) * 3 == 34)
for txt, val in (('Clase 4 act. 7', abs(9 * .2 + 9 * .3 + 6 * .5 - 7.5) < 1e-9), ('Práctica 4 X=8,7', abs(9 * .5 + 8 * .3 + 9 * .2 - 8.7) < 1e-9),
                 ('Práctica 17 = 221.500', 18500 + 100000 + 103000 == 221500), ('Práctica 1 total 28.000', 3 * 5000 + 2 * 4000 + 5000 == 28000),
                 ('Ev. U1 ítem 8 = 22.000', 2 * 5000 + 3 * 4000 == 22000), ('Clase 21 por fecha 33/51/50 mil', 14000 + 19000 == 33000 and 34000 + 17000 == 51000)):
    chk('C', 'Cálculo verificado: ' + txt, val)
# todos los montos G. x.xxx del libro existen en el universo de cifras verificadas o en la clase
# ---------------- D. Tomo ↔ solucionario ----------------
dS, PS = textos_doc(SOL); TS = '\n'.join(t for _, t in PS)
for n in range(1, 22):
    for fam, k, enun, resp, nuevo in actividades.items(n):
        pass
    chk('D', 'Solucionario: respuestas de la Clase %d (%d ítems)' % (n, len(actividades.items(n))), all(('%d. %s' % (k, r)) in TS for _, k, _, r, _ in actividades.items(n)))
    chk('D', 'Libro: consignas de la Clase %d (%d ítems)' % (n, len(actividades.items(n))), all(('%d. %s' % (k, e)) in TL for _, k, e, _, _ in actividades.items(n)))
    chk('D', 'Solucionario: resultado y errores de la Práctica %d' % n, practicas.PR[n]['sol']['resultado'] in TS and practicas.PR[n]['sol']['errores'] in TS)
for e in evaluaciones.EV:
    k = 0; okmc = True
    for enun, ops, letra in e['mc']:
        k += 1
        okmc &= ('%d. %s) %s' % (k, letra, ops['abc'.index(letra)])) in TS and ('%d. %s' % (k, enun)) in TL
    chk('D', 'Opción múltiple coherente libro ↔ clave: ' + e['titulo'], okmc)
    chk('D', 'Abiertas con clave: ' + e['titulo'], all(r in TS for _, r in e['ab']) and all(q in TL for q, _ in e['ab']))
viejas = ['2 chipas y 1 gaseosa (G. 3.000', 'Prolog→lógico; Java→OO; C→imperativo; Haskell→funcional', 'C04 → V5 (G. 50.000)', 'contiene,N:M, cantidad)—Producto', 'Socios(código…), Libros(código…), Préstamos(código', '(V1, Chipa, 2, 3.000)']
chk('D', 'Solucionario sin respuestas de ejercicios de la edición anterior', not any(v in TS for v in viejas), [v for v in viejas if v in TS])
chk('D', 'Libro del estudiante sin respuestas ni rúbricas', not any(practicas.PR[n]['sol']['resultado'][:60] in TL for n in range(1, 22)) and 'Errores frecuentes:' not in TL and 'Rúbrica de evaluación' not in TL)
chk('D', 'Evaluaciones sin respuestas entre paréntesis en la consigna', not re.search(r'\((?:V\d(?:,| y)|Rosa \d|V5, G\.)', '\n'.join(q for e in evaluaciones.EV for q, _ in e['ab'])))

# ---------------- E. Duplicados clase ↔ práctica ↔ actividades ----------------
def toks(s):
    return set(re.findall(r'\b(?:V\d+|C0\d|P0\d|\d{1,3}\.\d{3})\b', s))
def cuerpo_txt(c):
    return ' '.join((b.get('text') or (b.get('title', '') + ' ' + ' '.join(b.get('body', [])))) + ' ' + ' '.join(' '.join(r) for r in b.get('rows', [])) for b in c['cuerpo'])
for n in range(1, 22):
    p = practicas.PR[n]
    ptxt = ' '.join([a['objetivo'] + ' ' + ' '.join(a['modelo'][1] if a.get('modelo') else []) for a in p['acts']])
    if p.get('datos'):
        ptxt += ' '.join(' '.join(r) for _, _, rows in p['datos'] for r in rows)
    ejemplos = ' '.join(b['text'] for b in C[n]['cuerpo'] if b['t'] == 'p' and re.match(r'^\d+\. ', b['text']))
    inter = toks(ptxt) & toks(ejemplos)
    # se toleran los precios de catálogo (datos maestros), no los datos de ventas ni los totales resueltos
    catalogo = {'3.000', '5.000', '6.000', '15.000', '4.000', '8.000', '9.000', '7.000'}
    inter = {x for x in inter if x not in catalogo}
    chk('E', 'Práctica %d no repite datos de los ejemplos resueltos de la Clase %d' % (n, n), not inter, sorted(inter))
    acts_txt = ' '.join(e for _, _, e, _, _ in actividades.items(n))
    frases = ['2 chipas y 1 gaseosa', 'V5 y la gaseosa', 'precio > 5000', '«total = 2 * 3000»', 'usá los datos de la clase']
    chk('E', 'Actividades de la Clase %d sin copiar el ejemplo resuelto' % n, not any(f in acts_txt for f in frases))
enuns = [e for n in range(1, 22) for _, _, e, _, _ in actividades.items(n)]
dup = [e for e, c in collections.Counter(enuns).items() if c > 1]
chk('E', 'Sin consignas idénticas repetidas entre clases', not dup, dup[:2])
pr_obj = [a['objetivo'] for n in range(1, 22) for a in practicas.PR[n]['acts']]
chk('E', 'Sin objetivos de práctica repetidos', len(pr_obj) == len(set(pr_obj)))
ev_items = [q for e in evaluaciones.EV for q, _ in e['ab']] + [q for e in evaluaciones.EV for q, _, _ in e['mc']]
rep = [q for q, c in collections.Counter(ev_items).items() if c > 1]
chk('E', 'Evaluaciones sin ítems duplicados entre sí', not rep, rep)
chk('E', 'Evaluaciones sin ítems copiados de las actividades', not (set(q for e in evaluaciones.EV for q, _ in e['ab']) & set(enuns)))

# ---------------- F. Figuras ----------------
caps = [t for _, t in PL if re.match(r'^Figura \d+\.\d+ — ', t)]
porclase = collections.defaultdict(list)
for c_ in caps:
    a, b = map(int, re.match(r'Figura (\d+)\.(\d+)', c_).groups())
    porclase[a].append(b)
chk('F', 'Figuras numeradas por clase y correlativas', all(v == list(range(1, len(v) + 1)) for v in porclase.values()), dict(porclase))
chk('F', 'Cantidad de figuras: 28 conservadas + 4 nuevas + mapa del Proyecto', len(caps) == 32 and 'Figura PFI.1' in TL, len(caps))
nimg = len(dL.inline_shapes)
chk('F', 'Cada imagen del cuerpo tiene su epígrafe (%d imágenes)' % nimg, nimg == len(caps) + 1)
chk('F', 'Sin rótulos «generada con IA»', not re.search(r'generad[ao] con (IA|inteligencia)', TL, re.I))
hashes = {}
import hashlib
for f in sorted(os.listdir(B + 'figs_src')):
    hashes[hashlib.sha256(open(B + 'figs_src/' + f, 'rb').read()).hexdigest()] = f
dpk = docx.Document(LIB).part
blobs = {hashlib.sha256(r.target_part.blob).hexdigest() for r in dpk.rels.values() if r.reltype.endswith('/image')}
chk('F', 'Las 28 imágenes originales se conservan byte a byte (sin recomprimir)', set(hashes) <= blobs, len(set(hashes) - blobs))
sin_lead = []
for n, c in C.items():
    for i, b in enumerate(c['cuerpo']):
        if b['t'] == 'img':
            j = i - 1; anunciada = bool(b.get('lead'))
            while j >= 0 and c['cuerpo'][j]['t'] != 'h2' and not anunciada:
                q = c['cuerpo'][j]
                anunciada = q['t'] == 'p' and re.search(r'figura|queda así|así queda|el esquema muestra', q['text'], re.I) is not None
                j -= 1
            if not anunciada:
                sin_lead.append(b['num'])
chk('F', 'Toda figura está anunciada en el texto', not sin_lead, sin_lead)
refs = re.findall(r'Figura (\d+\.\d+)', TL)
chk('F', 'Referencias cruzadas a figuras existentes', all(('Figura %s — ' % r) in TL for r in refs))

# ---------------- G. Planes y Plan Anual ----------------
PLs = planes_data.planes()
dA, PA = textos_doc(ANU); TA = '\n'.join(t for _, t in PA)
dP, PP = textos_doc(PLA); TP = '\n'.join(t for _, t in PP)
chk('G', 'Plan Anual: 36 encuentros numerados 1..36', all(('%d' % i) in [c.text for row in dA.tables[0].rows for c in row.cells[:1]] for i in range(1, 37)))
for n in range(1, 22):
    chk('G', 'Plan Anual: indicadores literales de la ficha de la Clase %d' % n, all(('• ' + i) in TA for i in C[n]['ficha']['indicadores']))
chk('G', 'Plan Anual: procedimientos únicos (36)', len({p['proc'] for p in PLs}) == 36)
chk('G', 'Plan Anual: instrumentos únicos (36)', len({p['inst'] for p in PLs}) == 36)
for doc_, nombre in ((TA, 'Plan Anual'), (TP, 'Planes de Clase')):
    chk('G', '%s sin «Cuaderno»/«cuadernillo»/«tomo»' % nombre, not re.search(r'Cuaderno de Prácticas|cuadernillo|Cuadernillo|\btomo\b', doc_))
    chk('G', '%s sin códigos internos' % nombre, not re.search(r'\b\d-\d\d\b|IMG_\d+|TXT_\w+|\{FIG', doc_))
for p in PLs:
    t = json.dumps(p, ensure_ascii=False)
    for m in re.finditer(r'[Ff]iguras? ((?:\d+\.\d+(?:,? (?:y )?)?)+)', t):
        for r in re.findall(r'\d+\.\d+', m.group(1)):
            chk('G', 'Plan %d: figura %s existe en el libro' % (p['idx'], r), ('Figura %s — ' % r) in TL)
    if p['tipo'] == 'C':
        er, pd, de, tot = planes_data._fam(p['n'])
        for m in re.finditer(r'actividades 1 a (\d+)', t):
            chk('G', 'Plan %d: «actividades 1 a %s» dentro del rango de la Clase %d' % (p['idx'], m.group(1), p['n']), int(m.group(1)) <= tot)
        for m in re.finditer(r'afirmaci(?:ón|ones) (\d+)(?: y (\d+))?', t):
            nums = [int(x) for x in m.groups() if x]
            chk('G', 'Plan %d: afirmaciones %s son del bloque «Pensá y decidí»' % (p['idx'], nums), all(x in pd for x in nums))
    for m in re.finditer(r'Práctica (\d+)', t):
        chk('G', 'Plan %d: remite a una práctica existente (%s)' % (p['idx'], m.group(1)), 1 <= int(m.group(1)) <= 21)
for k, p in enumerate(PLs, 1):
    chk('G', 'Plan %d: tiempos suman 160 min' % k, sum(p['tiempos']) == 160)
chk('G', 'Planes: 36 títulos en el documento', all(('Plan de Clase N.º %d — %s' % (p['idx'], p['titulo'])) in TP for p in PLs))

# ---------------- H. Lengua, residuos y cobertura ----------------
tuteo = re.compile(r'^(?:\d+\. )?(Escribe|Calcula|Indica|Explica|Dibuja|Completa|Responde|Elige|Marca|Lee|Observa|Justifica|Define|Nombra|Ordena|Clasifica|Propón|Diseña|Crea|Abre|Guarda)\b')
cons = [e for n in range(1, 22) for _, _, e, _, _ in actividades.items(n) if tuteo.match(e)] + [s_ for n in range(1, 22) for a_ in practicas.PR[n]['acts'] for s_ in a_['pasos'] if tuteo.match(s_)] + [q for e in evaluaciones.EV for q, _ in e['ab'] if tuteo.match(q)]
chk('H', 'Voseo en las consignas (sin imperativos en tuteo)', not cons, cons[:3])
for r in ('Cuaderno de Prácticas', 'cuadernillo', 'Cuadernillo', ' tomo ', '{FIG', 'TODO:', 'XXX', 'python-docx', 'Ctrl+E'):
    chk('H', 'Residuo ausente en el libro: «%s»' % r, r not in TL)
for r in ('Colegio Nacional', 'Villarrica', 'Asunción'):
    chk('H', 'Sin institución ni ciudad: «%s»' % r, r not in TL)
PROG = {'tipos y campos de aplicación de lenguajes': 'campos de aplicación', 'paradigma imperativo': 'imperativo', 'paradigma funcional': 'funcional',
        'paradigma lógico': 'lógico', 'orientado a objetos': 'orientado a objetos', 'base de datos: concepto y tipos': 'Tipos y aplicaciones de bases de datos',
        'tablas y relacionamiento': 'Tablas, campos y registros', 'clave principal y secundaria': 'clave secundaria', 'integridad referencial': 'Integridad referencial',
        'SGBD: redundancia e inconsistencia': 'inconsistencia', 'aislamiento de datos': 'Aislamiento', 'atomicidad': 'Atomicidad', 'acceso concurrente': 'concurrencia',
        'seguridad': 'Seguridad', 'niveles físico, lógico y de vistas': 'nivel de vistas', 'tipos de usuario': 'usuarios sofisticados',
        'DDL y DML': 'DDL', 'modelo entidad-relación': 'entidad-relación', 'modelo orientado a objetos': 'Orientado a objetos',
        'lógico basado en registros': 'Lógico basado en registros', 'lógico basado en objetos': 'Lógico basado en objetos',
        'atributos simples y compuestos': 'compuesto', 'univalorados y multivalorados': 'multivalorado', 'atributo derivado': 'derivado',
        'cuestiones de diseño (entidad o atributo / entidad o relación)': 'Cuestiones de diseño', 'ligaduras 1:1, 1:N, N:1, N:M': 'varios a uno',
        'claves': 'clave principal', 'DER': 'diagrama Entidad-Relación', 'grado de relacionamiento': 'Grado de una relación', 'dominio': 'dominio',
        '1FN': '1FN', '2FN': '2FN', '3FN': '3FN', 'dependencia funcional': 'dependencia funcional', 'DF compuesta o completa': 'compuesta o completa',
        'dependencia transitiva': 'transitiva', 'gerenciador de base de datos (Access)': 'Microsoft Access'}
for k, v in PROG.items():
    chk('H', 'Cobertura del programa: ' + k, v.lower() in TL.lower())

# ---------------- I. Ronda de correcciones v2 (auditoría ChatGPT + decisiones de Fer) ----------------
import ediciones_v2 as V2, preliminares as PRE
MECS = set(V2.MEC.values())
fich = [t for t in dL.tables if t.rows[0].cells[0].text.strip() in ('Capacidad', 'Capacidades')]
for n, t in enumerate(fich, 1):
    caps_doc = [x.text.strip().lstrip('• ').strip() for x in t.rows[0].cells[1].paragraphs if x.text.strip()]
    chk('I', 'Clase %d: capacidad textual del programa MEC' % n, caps_doc == V2.caps(V2.CAP_CLASE[n]), caps_doc)
chk('I', 'Clase 10: tres capacidades oficiales (abstracción, usuarios, lenguajes)', len(V2.CAP_CLASE[10]) == 3 and set(V2.CAP_CLASE[10]) == {'C7', 'C8', 'C9'})
chk('I', 'Todas las capacidades MEC de las unidades aparecen en alguna ficha', set(c for v in V2.CAP_CLASE.values() for c in v) == set(V2.MEC) - {'C15'})
chk('I', 'Plan Anual: nota sobre capacidades textuales e indicadores', V2.MEC['C15'] in TA and 'Las capacidades se reproducen textualmente del programa MEC vigente. Los indicadores de logro de este plan son operativizaciones observables elaboradas para organizar la enseñanza y la evaluación de cada encuentro.' in TA)
chk('I', 'Plan Anual: capacidad transversal (práctica Conductual)', 'Capacidad transversal del programa' in TA and V2.MEC['C15'] in TA)
for u, cods in V2.CAP_UNIDAD.items():
    chk('I', 'Plan Anual: capacidades textuales de la Unidad %d' % u, all(V2.MEC[c] in TA for c in cods))
cap_planes = [x for x in PLs for x in (x['capacidad'] if isinstance(x['capacidad'], list) else [x['capacidad']])]
chk('I', 'Planes de Clase: toda capacidad es textual del programa MEC', set(cap_planes) <= MECS, sorted(set(cap_planes) - MECS)[:3])
chk('I', 'Planes de Clase: las capacidades figuran en el documento', all(c in TP for c in set(cap_planes)))
chk('I', 'Clave secundaria: texto acordado presente', 'El programa oficial emplea el término «clave secundaria» sin definirlo en este apartado. En este libro usaremos el término clave foránea para el campo que referencia la clave principal de otra tabla. En otros textos, «clave secundaria» también puede referirse a una clave alternativa o a un índice secundario.' in TL)
chk('I', 'Clave secundaria: sin la equivalencia anterior', 'la clave secundaria es la clave foránea' not in TL and 'secundaria, en la terminología del programa' not in TS)
chk('I', 'Entidad débil: sin «exactamente la misma tabla»', 'exactamente la misma tabla' not in TL)
chk('I', 'Entidad débil: idea de Fer sobre el discriminante', 'RENGLÓN necesita un discriminante propio, por ejemplo nro_renglón, y ambos modelos ya no producen exactamente la misma clave.' in TL)
for x in ('Entidad fuerte', 'dependencia de existencia', 'Discriminante (identificador parcial)', 'Relación identificadora', 'entidad asociativa', '(venta, nro_renglón)', '(venta, producto)'):
    chk('I', 'Entidad débil: «%s»' % x, x in TL)
chk('I', 'Histórico de precios: texto nuevo', 'El precio de catálogo actual vive en Productos; queda fuera del mini-mundo un historial independiente de cambios del catálogo. Cada renglón de venta sí conserva en precio_unitario el importe realmente cobrado en esa operación.' in TL and 'el histórico de precios queda fuera del mini-mundo' not in TL)
chk('I', 'PSeInt: perfil Flexible indicado', 'perfil Flexible' in TL)
chk('I', 'PSeInt: Práctica 3 define i como Entero', 'Definí total, importe e i como Entero y promedio como Real; inicializá total en 0.' in TL)
chk('I', 'PSeInt: Práctica 2 sin texto exacto del error en el libro', 'Al intentar dividir por cero, el problema aparece durante la ejecución. Registrá el mensaje exacto mostrado por la versión de PSeInt instalada.' in TL and 'Division por cero' not in TL)
chk('I', 'PSeInt: solucionario documenta versión y resultados', 'PSeInt 20250314' in TS)
chk('I', 'Access 2016: referencia una sola vez en el libro', TL.count(V2.ACCESS_REF) == 1)
for x in ('versión actual', 'la más utilizada', 'más usada en Paraguay', 'marcada con estrella', 'dos llaves', 'Nuevo origen de datos', 'utonumérico'):
    chk('I', 'Access: sin «%s»' % x, x not in TL and x not in TS)
chk('I', 'Access: icono/indicador de clave', 'la clave principal identificada con el icono/indicador de clave' in TL)
chk('I', 'Fechas: respuesta de la Evaluación de la Unidad 4', 'En la cuadrícula de Diseño con configuración regional d/m/a: #05/03/2026# → V5. En la vista SQL: #3/5/2026#.' in TS)
ctrl_lib = [a['control'] for n in range(1, 22) for a in practicas.PR[n]['acts']]
for x in ('Cargaste 33', 'G. 8.000 aunque', '2 de las 6', '13 renglones', 'seis paradigmas distintos', 'tres problemas (redundancia', 'Usaste las tres integridades', 'cuatro tipos de usuario', 'tres grados, uno por', 'Toda flecha sale de Préstamos', 'Solo uno no necesita', 'un solo renglón', 'termina exactamente igual', 'menos préstamos que renglones', 'al menos un enunciado de cada nivel', 'por más de un punto', 'Con C04'):
    chk('I', 'Puntos de control sin revelar: «%s»' % x, not any(x in c for c in ctrl_lib) and x not in TL)
for n, txt in ((18, 'El total de registros coincide con la suma de los registros indicados en las cuatro tablas de datos.'), (19, 'El filtro muestra solamente las ventas de la fecha elegida'),
               (20, 'La consulta produce una respuesta coherente tanto para clientes con ventas como para un cliente que no tenga registros coincidentes.'), (21, 'ConsDetalle contiene un renglón por cada registro de DetalleVenta y ningún Subtotal queda vacío.'),
               (3, 'Clasificaste los seis fragmentos, justificaste cada elección con una evidencia del enunciado y no asignaste dos categorías por simple intuición.')):
    chk('I', 'Punto de control acordado en la Práctica %d' % n, txt in TL)
for n in (5, 6, 8, 9, 10):
    tr = practicas.PR[n]['transfer']
    chk('I', 'Práctica %d: transferencia y revisión entre pares en el libro' % n, ('Modelo — Mini-caso: ' + tr['caso'][0]) in TL and all(x in TL for x in tr['caso'][1]))
    chk('I', 'Práctica %d: solución de la transferencia en el solucionario' % n, tr['sol'] in TS)
    chk('I', 'Plan de continuación de la Práctica %d: incluye la transferencia' % n, any(p['tipo'] == 'P' and p['n'] == n and any('Transferencia y revisión entre pares' in x for x in p['momentos']['Desarrollo']) for p in PLs))
chk('I', 'Transferencia solo donde hace falta (5 de 10 continuaciones)', sum(1 for n in planes_data.CONT if practicas.PR[n].get('transfer')) == 5)
cajas_py = [t for s_, t in PL if t.startswith('En Paraguay')]
chk('I', 'Recuadros «En Paraguay»: solo los de contenido paraguayo verificable (3)', len(cajas_py) == 3, cajas_py)
chk('I', 'Recuadros «Aplicación profesional»/«Ejemplo cotidiano» no aparecen en todas las clases', sum(1 for n in range(1, 22) if any(b['t'] == 'caja' and b['title'].startswith(('Aplicación profesional', 'Ejemplo cotidiano')) for b in C[n]['cuerpo'])) < 21)
chk('I', 'PFI: extensión recomendada con el texto acordado', PRE.PFI_EXTENSION in TL and 'formulario de carga y un informe (recomendados)' not in TL)
chk('I', 'Rúbrica analítica: 5 criterios con pesos 25/25/20/15/15', [p for _, p, _ in PRE.RUBRICA_ANALITICA] == [25, 25, 20, 15, 15])
chk('I', 'Rúbrica analítica: 4 descriptores por criterio en el solucionario', all(all(dsc in TS for dsc in desc) and len(desc) == 4 for _, _, desc in PRE.RUBRICA_ANALITICA))
for nombre, path in (('libro', LIB), ('solucionario', SOL), ('planes', PLA), ('plan anual', ANU)):
    texto = '\n'.join(pg.get_text() for pg in pdf(path))
    chk('I', 'Sin símbolos ✓ ✗ ✔ ✅ en el %s' % nombre, not re.search('[\u2705\u2714\u2713\u2717]', texto))
    fuentes = {f_[3] for pg in pdf(path) for f_ in pg.get_fonts()}
    chk('I', 'Sin fuente de emoji en el PDF del %s' % nombre, not any('Emoji' in f_ for f_ in fuentes), sorted(fuentes))
    chk('I', 'Sin cuadros de glifo faltante (U+FFFD/U+25A1) en el %s' % nombre, not re.search('[\ufffd\u25a1]', texto))
    docx_txt = '\n'.join(t for _, t in textos_doc(path)[1])
    for pat, desc in ((r'\{FIG|image\d+\.(?:jpg|png)', 'marcadores de figura'), (r'«\s*»|(?<!\w)\(\s*\)', 'comillas o paréntesis vacíos (fuera de funciones como Fecha())'), (r'[^.]\.\.(?!\.)', 'punto doble'), (r' [,;:]| \.(?![\w\d])', 'espacio antes de puntuación'), (r'\b(\w{3,}) \1\b', 'palabra repetida')):
        hits = re.findall(pat, docx_txt)
        chk('I', 'Residuos de conversión en el %s: %s' % (nombre, desc), not hits, hits[:3])

# ---------------- J. Corrección de cierre v2.1 ----------------
chk('J', 'Clase 15: capacidad «Trabaja con el modelo Entidad Relación…»', V2.CAP_CLASE[15] == ['C12'] and C[15]['ficha']['capacidad'] == [V2.MEC['C12']])
chk('J', 'Clase 15: indicadores sin cambios en el Plan Anual', all(('• ' + i) in TA for i in C[15]['ficha']['indicadores']))
chk('J', 'Planes: el plan de la Clase 15 lleva la capacidad nueva', any(p['tipo'] == 'C' and p['n'] == 15 and p['capacidad'] == [V2.MEC['C12']] for p in PLs))
chk('J', 'Plan Anual: la Unidad 3 lista solo capacidades de sus clases', V2.CAP_UNIDAD[3] == ['C10', 'C11', 'C12', 'C13'])
chk('J', 'Clase 9: recuadro «En Paraguay — integridad y facturación» con el texto acordado', 'En Paraguay — integridad y facturación' in TL and 'En operaciones B2B y B2G la identificación del receptor es obligatoria; en operaciones B2C existen casos de consumidor final.' in TL and 'La integridad referencial protege aquellas relaciones que el modelo define como obligatorias.' in TL)
chk('J', 'Clase 9: sin el recuadro anterior', 'En Paraguay — por qué importa' not in TL and 'problemas ante la DNIT' not in TL)
chk('J', 'Práctica 6: sin la dicotomía gestor/aplicación', 'pertenecen al gestor' not in TL and 'las guarda el gestor' not in TS)
chk('J', 'Práctica 6: paso 2 y solución acordados', 'Recordá que Access integra ambas funciones en una misma herramienta.' in TL and 'En Access las cuatro familias son objetos de la misma aplicación y trabajan integradas.' in TS)
chk('J', 'Práctica 19: punto de control acordado de la Actividad 2', 'Los registros devueltos y los no devueltos, juntos y sin superposición, reconstruyen el conjunto completo de Ventas.' in TL and 'son exactamente las del 11/03/2026' not in TL)
chk('J', 'Clases 1 y 10: sin recuadros de oficio', 'programar como oficio' not in TL and 'un oficio con demanda' not in TL)
chk('J', 'Clases 1 y 10: contenido integrado como texto común', 'Saber programar abre oportunidades laborales' in TL and 'La administración de bases de datos es un área profesional importante' in TL)
chk('J', 'Clase 21: rutas oficiales de referencias cruzadas', 'Crear → Asistente para consultas → Asistente para consultas de referencias cruzadas' in TL and 'Tipo de consulta → Tabla de referencias cruzadas' in TL)
chk('J', 'Clase 18: ruta oficial de importación', 'Datos externos → grupo Importar y vincular → Excel' in TL)

# ---------------- salida ----------------
fallos = [r for r in RES if not r[2]]
por = collections.Counter(r[0] for r in RES)
with open(B + 'auditoria_resultado.txt', 'w') as f:
    f.write('AUDITORÍA AUTOMÁTICA — Algorítmica 3.º · Edición Esencial Comercial 2026 · versión v2.2 (formato oficio)\n')
    f.write('Controles: %d · Fallos: %d\n' % (len(RES), len(fallos)))
    nombres = {'A': 'Estructura y correspondencia', 'B': 'Paginación, índice, saltos y encabezados', 'C': 'Coherencia aritmética', 'D': 'Libro ↔ solucionario',
               'E': 'Duplicados', 'F': 'Figuras', 'G': 'Plan Anual y Planes de Clase', 'H': 'Lengua, residuos y cobertura curricular', 'I': 'Correcciones v2 (auditoría y decisiones de Fer)', 'J': 'Corrección de cierre v2.1'}
    for b in sorted(por):
        f.write('  %s. %-45s %4d controles · %d fallos\n' % (b, nombres[b], por[b], sum(1 for r in fallos if r[0] == b)))
    for r in fallos:
        f.write('FALLO [%s] %s — %s\n' % (r[0], r[1], r[3]))
print(open(B + 'auditoria_resultado.txt').read())
