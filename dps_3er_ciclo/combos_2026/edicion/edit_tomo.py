# -*- coding: utf-8 -*-
"""Edición diferencial del tomo candidato (ver el informe del grado). Uso: python3 edit_tomo.py 7|8|9"""
import sys, os, re, json, copy
import docx, pymupdf
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_ops as O
import maqueta as M
import render as R
import config_grados as G

g = sys.argv[1]
C = G.GRADOS[g]
B = O.E
OUT = B + '../%s/' % C['carpeta']
os.makedirs(OUT, exist_ok=True)
NOMBRE = C['nombre_tomo']
M.DARK, M.MED, M.LIGHT, M.FICHA = '6A3E7A', '8E5BA6', 'F3ECF7', 'F8F3FA'
M.ANCHO = round(M.PAG_W - 2 * M.MARG, 2)


def es_h(el, nivel):
    st = el.find('.//' + qn('w:pStyle'))
    return el.tag == qn('w:p') and st is not None and st.get(qn('w:val')) in ('Heading%d' % nivel, 'Ttulo%d' % nivel)   # 8.º: ids «Ttulo»


def primer_el(d, pred):
    for el in d.element.body:
        if pred(el):
            return el
    raise ValueError('elemento no encontrado')


def pbb(el):
    ppr = el.find(qn('w:pPr'))
    if ppr is None:
        return False
    x = ppr.find(qn('w:pageBreakBefore'))
    return x is not None and x.get(qn('w:val')) not in ('0', 'false', 'off')


def marcadores(body):
    """Puntos donde empieza una página nueva: (elemento antes del cual insertar, elemento cuyo texto abre la página)."""
    els = [e for e in body if e.tag != qn('w:sectPr')]
    out = []
    for i, el in enumerate(els):
        if pbb(el):   # si el que abre página es un párrafo vacío, el texto que la abre es el del elemento siguiente
            out.append((el, el if O.texto(el).strip() or i + 1 >= len(els) else els[i + 1]))
        elif el.tag == qn('w:p') and any(b.get(qn('w:type')) == 'page' for b in el.iter(qn('w:br'))) and i + 1 < len(els):
            if not O.texto(el).strip():   # párrafo vacío con salto: el texto siguiente abre la página
                out.append((el, els[i + 1]))
    return out


def entradas_indice(d):
    """(nivel, texto) en orden de aparición: unidades, clases, evaluaciones, módulo PEC y talleres."""
    E = []
    for el in d.element.body:
        tx = O.texto(el).strip()
        if el.tag == qn('w:p'):
            sub = re.match(r'^(Clase \d+|Ficha de profundización|Actividad complementaria)', tx)
            if (es_h(el, 1) or es_h(el, 2)) and sub:      # 7.º: clases en Heading 2; 8.º: en Heading 1
                E.append((2, tx))
            elif es_h(el, 1) and tx and tx not in ('Presentación', 'Índice') and not tx.startswith('FICHA'):
                E.append((1, tx))
        elif el.tag == qn('w:tbl'):
            c0 = el.find('.//' + qn('w:tc'))
            t0 = O.texto(c0.find(qn('w:p'))).strip() if c0 is not None and c0.find(qn('w:p')) is not None else ''
            if t0.startswith(('EVALUACIÓN', 'Prueba diagnóstica')):
                E.append((2 if 'UNIDAD' in t0 else 1, t0))
    return E


def escribir_indice(d, E, pags):
    h = primer_el(d, lambda el: es_h(el, 1) and O.texto(el).strip() == 'Índice')
    body = d.element.body
    els = list(body)
    i = els.index(h) + 1
    # quitar el índice estático anterior (hasta el próximo elemento que abre página)
    while i < len(els) and not pbb(els[i]) and els[i].tag == qn('w:p'):
        body.remove(els[i]); i += 1
    ancla = h
    nota = OxmlElement('w:p')
    ancla.addnext(nota)
    p = docx.text.paragraph.Paragraph(nota, d._body)
    M.run(p, 'Los números de página corresponden a esta edición.', italic=True, size=8.5, color='777777')
    p.paragraph_format.space_after = Pt(4)
    ancla = nota
    for (niv, t), pg in zip(E, pags):
        q = OxmlElement('w:p'); ancla.addnext(q); ancla = q
        p = docx.text.paragraph.Paragraph(q, d._body)
        pf = p.paragraph_format
        pf.space_after = Pt(0.6); pf.space_before = Pt(3 if niv == 1 else 0)
        pf.left_indent = Cm(0 if niv == 1 else 0.5)
        pf.tab_stops.add_tab_stop(Cm(M.ANCHO - 0.05), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        M.run(p, t.split(' / ')[0], bold=(niv == 1), size=9.5 if niv == 1 else 9)
        p.add_run('\t')
        M.run(p, str(pg), bold=(niv == 1), size=9.5 if niv == 1 else 9)


def insertar_notas(d, antes_de, n):
    """Inserta «Espacio para anotaciones y respuestas» con n renglones antes del elemento dado."""
    tmp = docx.Document(B + 'plantilla.docx')
    pr = M.par(tmp, 'Espacio para anotaciones y respuestas', 10, True, color=M.DARK, keep=True, antes=8, despues=2)
    t = M.renglones(tmp, n)
    for el in (pr._p, t._tbl):
        antes_de.addprevious(copy.deepcopy(el))


def reemplazar_pec(d):
    body = d.element.body
    pec = primer_el(d, lambda el: es_h(el, 1) and O.texto(el).strip().startswith('PROYECTO EDUCATIVO COMUNITARIO'))
    els = list(body); i = els.index(pec)
    for el in els[i:]:
        if el.tag != qn('w:sectPr'):
            body.remove(el)
    R.NIVELES.update({'#': 1, '##': 2, '###': 3})
    R.render(d, R.leer(C['pec_tomo']))


def insertar_pec(d):
    """9.º: conserva el módulo PEC de la base y agrega, en cada taller, la ficha (después del título) y las consignas
    individuales con el recuadro de cuidados (después de «Producto evaluable»)."""
    R.NIVELES.update({'#': 1, '##': 2, '###': 3})
    B_ = R.bloques(C['pec_tomo'])
    els = list(d.element.body)
    for letra in 'ABCDEF':
        h = next(el for el in els if es_h(el, 1) and O.texto(el).strip().startswith('TALLER PEC %s ·' % letra))
        i = els.index(h)
        prod = next(el for el in els[i:] if O.texto(el).strip().startswith('Producto evaluable'))
        R.render_despues(d, h, B_['FICHA ' + letra])
        R.render_despues(d, prod, B_['CONSIGNAS ' + letra])
        els = list(d.element.body)


def construir(paso_indice=None, notas=None):
    d = docx.Document(C['base_tomo'])
    body = d.element.body
    # 1. portada profesional
    pres = primer_el(d, lambda el: es_h(el, 1) and O.texto(el).strip() == 'Presentación')
    O.portada_imagen(d, B + 'portadas/%s/portada_libro.png' % g, pres)
    # 2. márgenes de 1,27 cm y tablas al nuevo ancho
    O.margenes(d)
    # 3. correcciones de texto puntuales (informe diferencial)
    for viejo, nuevo in C['tomo_reemplazos']:
        O.reemplazar_en_doc(d, viejo, nuevo)
    # 4. capacidades literales en las fichas de clase
    fichas = [t for t in body.iter(qn('w:tbl')) if O.texto(t.find('.//' + qn('w:tc'))).strip() in ('Capacidad', 'Capacidades')]
    assert len(fichas) == len(C['caps_clase']), 'fichas: %d frente a %d' % (len(fichas), len(C['caps_clase']))
    for t, caps in zip(fichas, C['caps_clase']):
        tr = t.findall(qn('w:tr'))[0]; tcs = tr.findall(qn('w:tc'))
        O.reemplazar_parrafo(tcs[0].find(qn('w:p')), 'Capacidad' if len(caps) == 1 else 'Capacidades')
        O.fijar_celda(tcs[1], [C['CAP'][k] for k in caps] if len(caps) == 1 else ['• ' + C['CAP'][k] for k in caps])
    # 4 bis. bandas de unidad (8.º): «Capacidades de la unidad: …» parafraseadas → literales del programa
    u = None
    for el in body:
        tx = O.texto(el).strip()
        m = re.match(r'^UNIDAD (\d+) — ', tx)
        if m and es_h(el, 1):
            u = int(m.group(1))
        elif u and el.tag == qn('w:p') and tx.startswith('Capacidades de la unidad: '):
            O.reemplazar_parrafo(el, 'Capacidades de la unidad%s: ' % C.get('rotulo_fuente', ' (programa MEC)') + ' '.join('«%s»' % C['CAP'][k] for k in C['unidad_caps'][u]))
            u = None
    # 5. módulo PEC: se reemplaza por la versión desarrollada (7.º y 8.º) o se completa en su lugar (9.º)
    if C.get('pec_modo') == 'insertar':
        insertar_pec(d)
    else:
        reemplazar_pec(d)
    # 5 bis. saltos redundantes: párrafo vacío con salto de página seguido de un elemento que ya abre página
    els = [e for e in body if e.tag != qn('w:sectPr')]
    for a, b in zip(els, els[1:]):
        if a.tag == qn('w:p') and not O.texto(a).strip() and pbb(b):
            for br in [x for x in a.iter(qn('w:br')) if x.get(qn('w:type')) == 'page']:
                br.getparent().remove(br)
            # un párrafo vacío justo antes de un salto puede desbordar solo a una página en blanco: se quita
            if a.find('.//' + qn('w:drawing')) is None and a.find('.//' + qn('w:sectPr')) is None:
                body.remove(a)
    # 7. índice
    E = entradas_indice(d)
    escribir_indice(d, E, paso_indice or [0] * len(E))
    # 6. espacio para anotaciones antes de los elementos que abren página (regla N-02)
    if notas:
        abre = marcadores(body)
        for k, n in notas.items():
            insertar_notas(d, abre[k][0], n)
    # ilustraciones corregidas puntualmente: se reemplazan solo sus bytes (mismas dimensiones, mismo lugar)
    for nombre, ruta in C.get('imagenes_corregidas', {}).items():
        parte = next(p for p in d.part.package.iter_parts() if str(p.partname) == '/word/media/' + nombre)
        parte._blob = open(ruta, 'rb').read()
    for upd in d.settings.element.findall(qn('w:updateFields')):   # índice estático: sin aviso de actualizar campos
        d.settings.element.remove(upd)
    return d, E


def paginas(pdf, E):
    textos = [t.split(' / ')[0] for _, t in E]
    norm = [re.sub(r'\s+', '', pg.get_text()) for pg in pymupdf.open(pdf)]
    desde = max(j for j, x in enumerate(norm) if 'Losnúmerosdepáginacorresponden' in x or 'Índice' in x[:20]) + 2
    return M.paginas_de(pdf, textos, desde=desde)


if __name__ == '__main__':
    ruta = OUT + NOMBRE + '.docx'
    d, E = construir(); d.save(ruta); pdf = O.a_pdf(ruta, OUT)
    pags, total = paginas(pdf, E)
    # regla N-02: páginas que terminan una pieza con más del 55 % libre. Se mide sobre el PDF del mismo documento
    # (con las anotaciones ya puestas) y solo se agregan anotaciones en los cortes que todavía no tienen.
    notas = {}
    for ronda in range(4):
        doc = pymupdf.open(pdf)
        abre = marcadores(d.element.body)
        textos = [re.sub(r'\s+', '', O.texto(t))[:40] for _, t in abre]
        norm = [re.sub(r'\s+', '', pg.get_text()) for pg in doc]
        cur = max(j for j, x in enumerate(norm) if 'Losnúmerosdepáginacorresponden' in x) + 1
        nuevas_notas = {}
        for k, t in enumerate(textos):
            if not t:
                continue
            pg = next((j for j in range(cur, len(norm)) if norm[j].startswith(t[:25]) or t in norm[j][:200]), None)
            if pg is None:
                continue
            cur = pg + 1
            prev = doc[pg - 1]
            fr = O.llenado(prev)
            if pg - 1 >= 3 and fr < 0.45 and k not in notas:
                nuevas_notas[k] = max(3, int((0.86 - fr) * prev.rect.height / 72 * 2.54 / 0.74) - 3)
            elif k in notas and norm[pg - 1].startswith('Espacioparaanotaciones'):
                nuevas_notas[k] = None          # el espacio quedó solo en una página: se quita
        if not nuevas_notas:
            break
        for k, n in nuevas_notas.items():
            if n is None:
                notas.pop(k)
            else:
                notas[k] = n
        d, E = construir(pags, notas); d.save(ruta); pdf = O.a_pdf(ruta, OUT)
        pags, total = paginas(pdf, E)
    for v in range(3):
        d, E = construir(pags, notas); d.save(ruta); pdf = O.a_pdf(ruta, OUT)
        nuevas, total = paginas(pdf, E)
        if nuevas == pags:
            break
        pags = nuevas
    doc = pymupdf.open(pdf)
    vac = [(i + 1, round(O.llenado(p), 2)) for i, p in enumerate(doc) if 2 < i < len(doc) - 1 and O.llenado(p) < 0.30]
    json.dump({'entradas': E, 'paginas': pags, 'total': total, 'notas': notas}, open(B + '../verificacion/%s/indice_tomo.json' % C['carpeta'], 'w'), ensure_ascii=False, indent=0)
    print('páginas', total, '· índice estable', nuevas == pags, '· anotaciones', len(notas), '· casi vacías', vac)
