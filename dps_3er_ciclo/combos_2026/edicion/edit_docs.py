# -*- coding: utf-8 -*-
"""Edición diferencial del Plan Anual, los Planes de Clase y el Solucionario candidatos. Uso: python3 edit_docs.py 7|8|9
Plan Anual y Planes: solo .docx (el PDF va a verificacion/ para control visual). Solucionario: .docx y .pdf."""
import sys, os, re, copy
import docx, pymupdf
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_ops as O
import maqueta as M
import render as R
import config_grados as G

g = sys.argv[1]
C = G.GRADOS[g]
B = O.E
OUT = B + '../%s/' % C['carpeta']
CTRL = B + '../verificacion/%s/control_pdf/' % C['carpeta']
M.DARK, M.MED, M.LIGHT, M.FICHA = '6A3E7A', '8E5BA6', 'F3ECF7', 'F8F3FA'
CAP = C['CAP']
ROT = C.get('rotulo_fuente', ' (programa MEC)')   # 9.º: vacío hasta cotejar con el programa oficial
TALLER = C['talleres']            # {'A': (encuentro, [caps]), ...}


def lit(ks, sep=' '):
    return sep.join('«%s»' % CAP[k] for k in ks)


def pec_indicadores():
    """Indicadores de las fichas PEC del tomo, por taller (A–F en orden): la misma fuente para tomo, Plan Anual y Planes."""
    L = R.leer(C['pec_tomo']); out = []; i = 0
    while True:
        try:
            i = L.index('[[FICHA_PEC]]', i)
        except ValueError:
            break
        blk, i = R.bloque(L, i, '[[/FICHA_PEC]]')
        out.append([x[2:].strip() for x in blk if x.startswith('- ')])
    assert len(out) == 6, 'fichas PEC en el tomo: %d' % len(out)
    return dict(zip('ABCDEF', out))


IND = pec_indicadores()


def portada_seccion(d, ruta, horizontal_resto=False):
    """Portada a página completa en una sección vertical propia al principio del documento (sin pie)."""
    body = d.element.body
    p = M.imagen_pagina_completa(d, ruta)
    body.remove(p._p); body.insert(0, p._p)
    # sectPr de la sección de portada (oficio vertical, sin pie)
    sp = OxmlElement('w:sectPr')
    pgSz = OxmlElement('w:pgSz'); pgSz.set(qn('w:w'), str(int(21.59 / 2.54 * 1440))); pgSz.set(qn('w:h'), str(int(33.02 / 2.54 * 1440)))
    pgMar = OxmlElement('w:pgMar')
    for k in ('top', 'right', 'bottom', 'left'):
        pgMar.set(qn('w:' + k), '720')
    for k in ('header', 'footer', 'gutter'):
        pgMar.set(qn('w:' + k), '0')
    sp.append(pgSz); sp.append(pgMar)
    ppr = p._p.get_or_add_pPr(); ppr.append(sp)


def pdf_control(ruta):
    return O.a_pdf(ruta, CTRL)


# ---------------- Plan Anual ----------------

def plan_anual():
    d = docx.Document(C['base_pa'])
    O.margenes(d)
    for viejo, nuevo in C['pa_reemplazos']:
        O.reemplazar_en_doc(d, viejo, nuevo)
    # capacidades literales en las bandas de unidad
    n = 0; unidad = None
    for p in d.element.body.iter(qn('w:p')):
        tx = O.texto(p)
        m = re.match(r'^(Unidad (\d+) · [^·]+ · )Capacidades: ', tx)
        if m:
            u = int(m.group(2))
            O.reemplazar_parrafo(p, m.group(1) + 'Capacidades%s: ' % ROT + lit(C['unidad_caps'][u]))
            n += 1
            continue
        # formato 8.º/9.º: «UNIDAD N — …» [salto] «Capacidades: …» [salto] «Período…» en una sola corrida
        m = re.match(r'^UNIDAD (\d+) — ', tx)
        if m and 'Capacidades: ' in tx:
            t = next(t for t in p.iter(qn('w:t')) if (t.text or '').startswith('Capacidades: '))
            t.text = 'Capacidades%s: ' % ROT + lit(C['unidad_caps'][int(m.group(1))])
            n += 1
        elif m:
            unidad = int(m.group(1))
            continue
        # formato 9.º: «Capacidades: …» en el párrafo siguiente al de «UNIDAD N — …», dentro de la misma celda
        if tx.startswith('Capacidades: ') and unidad:
            O.reemplazar_parrafo(p, 'Capacidades%s: ' % ROT + lit(C['unidad_caps'][unidad]))
            n += 1; unidad = None
    assert n >= 7, 'bandas de unidad: %d' % n
    # talleres PEC: capacidades del área vinculadas, en lugar de la vinculación genérica
    for letra, (enc, caps) in TALLER.items():
        O.reemplazar_en_doc(d, C['pa_pec_generico'][letra],
                            C['pa_pec_generico'][letra].replace('Vinculado a la unidad de Orientación correspondiente.',
                                                               'Capacidades del área vinculadas%s: ' % ROT + lit(caps) + ' Indicadores: elaboración editorial.'))
    # indicadores de las filas PEC: los de la ficha del taller en el tomo
    k = 0
    for tr in d.element.body.iter(qn('w:tr')):
        tcs = tr.findall(qn('w:tc'))
        m = re.match(r'^Taller PEC ([A-F]) — ', O.texto(tcs[1]).strip()) if len(tcs) >= 3 else None
        if m:
            O.fijar_celda(tcs[2], ['• ' + x for x in IND[m.group(1)]]); k += 1
    assert k == 6, 'filas PEC del Plan Anual: %d' % k
    # DPS-03: la columna TIEMPO dice «3 HC» en los 36 encuentros, con una nota sobre la equivalencia
    body = d.element.body
    t_tiempo = []
    for t in body.iter(qn('w:tbl')):
        hdr = [O.texto(c).strip() for c in t.findall(qn('w:tr'))[0].findall(qn('w:tc'))]
        if 'TIEMPO' in hdr:
            t_tiempo.append(t); j = hdr.index('TIEMPO')
            for tr in t.findall(qn('w:tr'))[1:]:
                tcs = tr.findall(qn('w:tc'))
                if len(tcs) > j and O.texto(tcs[j]).strip() == '3 h':
                    O.fijar_celda(tcs[j], ['3 HC'])
    nota = OxmlElement('w:p'); t_tiempo[0].addprevious(nota)
    pn = docx.text.paragraph.Paragraph(nota, d._body)
    M.run(pn, 'En este esquema editorial, cada encuentro dura 3 HC (horas cátedra), que equivalen a 120 minutos; '
              'los 36 encuentros suman 108 HC.', italic=True, size=9, color='555555')
    pn.paragraph_format.space_after = Pt(4)
    # DPS-01: ninguna fila se parte entre páginas; el encabezado de columnas se repite; la tabla de talleres PEC va entera
    for t in body.iter(qn('w:tbl')):
        trs = t.findall(qn('w:tr'))
        for k, tr in enumerate(trs):
            trPr = tr.find(qn('w:trPr'))
            if trPr is None:
                trPr = OxmlElement('w:trPr'); tr.insert(0, trPr)
            if trPr.find(qn('w:cantSplit')) is None:
                trPr.append(OxmlElement('w:cantSplit'))
            if k == 0 and t in t_tiempo and trPr.find(qn('w:tblHeader')) is None:
                trPr.append(OxmlElement('w:tblHeader'))
        if O.texto(trs[0]).startswith('Encuentro'):      # tabla de correspondencia de los talleres PEC
            for tr in trs[:-1]:
                for p in tr.iter(qn('w:p')):
                    docx.text.paragraph.Paragraph(p, d._body).paragraph_format.keep_with_next = True
            prev = t.getprevious()                         # el párrafo que la presenta va con ella
            if prev is not None and prev.tag == qn('w:p'):
                docx.text.paragraph.Paragraph(prev, d._body).paragraph_format.keep_with_next = True
    # DPS-01 (7.º): las unidades corren seguidas; se quitan los saltos manuales que dejaban páginas casi vacías
    # y el título de cada unidad (etapa + banda de capacidades) queda unido a su tabla
    if C.get('pa_flujo_continuo'):
        els = list(body)
        for i, t in enumerate(els):
            if t in t_tiempo:
                k = i - 1
                while k >= 0 and els[k].tag == qn('w:p') and k >= i - 4:
                    p = els[k]
                    brs = [x for x in p.iter(qn('w:br')) if x.get(qn('w:type')) == 'page']
                    if brs and not O.texto(p).strip():
                        body.remove(p); break
                    pf = docx.text.paragraph.Paragraph(p, d._body).paragraph_format
                    pf.keep_with_next = True
                    if re.match(r'^\d\.ª ETAPA', O.texto(p).strip()):
                        pf.space_before = Pt(10)
                    k -= 1
        # «Criterios transversales…» sigue a la última tabla sin salto fijo (si la tabla se pasa por una fila, no queda
        # una página casi vacía); la firma queda unida a las líneas que la preceden
        els = list(body)
        crit = next(i for i, e in enumerate(els) if O.texto(e).strip().startswith('CRITERIOS TRANSVERSALES'))
        pf = docx.text.paragraph.Paragraph(els[crit], d._body).paragraph_format
        pf.page_break_before = False; pf.space_before = Pt(14); pf.keep_with_next = True
        docx.text.paragraph.Paragraph(els[crit + 1], d._body).paragraph_format.keep_with_next = True
        firma = next(i for i, e in enumerate(els) if 'Firma del docente' in O.texto(e))
        for e in els[firma - 2:firma]:
            if e.tag == qn('w:p'):
                docx.text.paragraph.Paragraph(e, d._body).paragraph_format.keep_with_next = True
        # cierre robusto: «Articulación curricular…», su introducción, la tabla PEC y el párrafo final forman un bloque
        # unido que nunca se separa, sea cual sea la fuente con que cada programa sustituya a Aptos
        els = list(body)
        art = next(i for i, e in enumerate(els) if O.texto(e).strip().startswith('Articulación curricular y organización'))
        pa_art = docx.text.paragraph.Paragraph(els[art], d._body).paragraph_format
        pa_art.page_break_before = False; pa_art.space_before = Pt(14)   # sin salto fijo: el bloque unido va donde entre entero
        for e in els[art:]:
            if e.tag == qn('w:p'):
                docx.text.paragraph.Paragraph(e, d._body).paragraph_format.keep_with_next = True
            elif e.tag == qn('w:tbl'):
                for p in e.iter(qn('w:p')):
                    docx.text.paragraph.Paragraph(p, d._body).paragraph_format.keep_with_next = True
            else:
                break
        ultimo = [e for e in els[art:] if e.tag == qn('w:p')][-1]
        docx.text.paragraph.Paragraph(ultimo, d._body).paragraph_format.keep_with_next = False
    sin_actualizar_campos(d)
    portada_seccion(d, B + 'portadas/%s/portada_plan_anual.png' % g)
    ruta = OUT + C['nombre_pa'] + '.docx'; d.save(ruta)
    return ruta


# ---------------- Planes de Clase ----------------

def es_plan(el):
    st = el.find('.//' + qn('w:pStyle')) if el.tag == qn('w:p') else None
    return st is not None and st.get(qn('w:val')) == 'PlandeClase'


def planes():
    pags = None
    for vuelta in range(4):
        d, E = _planes(pags)
        ruta = OUT + C['nombre_planes'] + '.docx'; d.save(ruta)
        if not E:
            return ruta
        pdf = pdf_control(ruta)
        norm = [re.sub(r'\s+', '', p.get_text()) for p in pymupdf.open(pdf)]
        idx = next(j for j, x in enumerate(norm) if 'Losnúmerosdepáginacorresponden' in x)
        nuevas, _ = M.paginas_de(pdf, [t for _, t, _ in E], desde=idx + 2)
        if nuevas == pags:
            break
        pags = nuevas
    return ruta


def _planes(pags):
    d = docx.Document(C['base_planes'])
    O.margenes(d)
    for viejo, nuevo in C['planes_reemplazos']:
        O.reemplazar_en_doc(d, viejo, nuevo, min_veces=1)
    body = d.element.body
    fichas = [t for t in body.iter(qn('w:tbl')) if O.texto(t.find('.//' + qn('w:tc'))).strip() in ('Capacidad', 'Capacidades')]
    assert len(fichas) == 36, 'fichas de planes: %d' % len(fichas)
    for n, t in enumerate(fichas, 1):
        tipo, caps, extra = C['plan_caps'][n]
        tr = t.findall(qn('w:tr'))[0]; tcs = tr.findall(qn('w:tc'))
        rot = {'C': 'Capacidad' if len(caps) == 1 else 'Capacidades', 'P': 'Capacidades del área vinculadas',
               'E': 'Capacidades evaluadas', 'I': 'Capacidades evaluadas'}[tipo]
        O.reemplazar_parrafo(tcs[0].find(qn('w:p')), rot)
        viejo = O.texto(tcs[1]).strip()
        lineas = [CAP[k] for k in caps] if len(caps) == 1 else ['• ' + CAP[k] for k in caps]
        if tipo == 'P':
            obj = re.sub(r'\s*\(objetivo didáctico.*$', '', viejo).strip()
            lineas += ['Objetivo del taller (elaboración editorial; el PEC no tiene capacidades propias en el programa): ' + obj]
        O.fijar_celda(tcs[1], lineas)
        if tipo == 'P':
            letra = next(l for l, (enc, _) in TALLER.items() if enc == n)
            O.fijar_celda(O.celda_de(t, 'Indicadores'), ['• ' + x for x in IND[letra]])
    # talleres: los momentos remiten a las consignas agregadas en el tomo
    for letra in TALLER:
        O.reemplazar_en_doc(d, C['planes_cierre_pec'][letra], C['planes_cierre_pec'][letra] + C.get('planes_cierre_texto',
                            ' Resolución de las consignas 1 a 3 («Pensá y respondé») y acuerdo de la tarea «En equipo» (consigna 4).'))
    # DPS-02: cada plan empieza en página nueva (8.º: los planes 26 y 32 quedaban partidos entre dos hojas)
    if C.get('planes_pagina_propia'):
        for el in body:
            if es_plan(el) and O.texto(el).strip().startswith('Plan de Clase N.º '):   # incluido el Plan 1, después del índice
                docx.text.paragraph.Paragraph(el, d._body).paragraph_format.page_break_before = True
    # índice: el de la candidata quedó vacío (solo el pedido de actualizar con F9)
    E = []
    h = next((el for el in body if el.tag == qn('w:p') and O.texto(el).strip() == 'Índice'), None)
    if h is None and C.get('planes_indice_nuevo'):
        # DPS-04 (7.º): índice en página propia, después de la portada interna; los títulos de los planes están en tablas
        tabs = [t for t in body.iter(qn('w:tbl')) if O.texto(t).startswith('PLAN DE CLASE N.º')]
        E = [(2, re.sub(r'\s+', ' ', O.texto(t.find('.//' + qn('w:tc')))).strip(), True) for t in tabs]
        assert len(E) == 36, 'planes en el índice: %d' % len(E)
        primero = tabs[0]
        while primero.getprevious() is not None and not any(b.get(qn('w:type')) == 'page' for b in primero.getprevious().iter(qn('w:br'))):
            primero = primero.getprevious()
        salto = primero.getprevious()                      # párrafo con el salto que abre el primer plan
        h = OxmlElement('w:p'); salto.addnext(h)
        ph = docx.text.paragraph.Paragraph(h, d._body)
        M.run(ph, 'Índice', bold=True, size=14, color=M.DARK); ph.paragraph_format.space_after = Pt(4)
        fin = OxmlElement('w:p'); h.addnext(fin)
        r = OxmlElement('w:r'); b = OxmlElement('w:br'); b.set(qn('w:type'), 'page'); r.append(b); fin.append(r)
        indice_sol(d, E, pags or [0] * 36, h=h, es_fin=lambda el: el is fin)
    elif h is not None:
        E = [(2, O.texto(el).strip(), True) for el in body if es_plan(el)]
        assert len(E) == 36, 'planes en el índice: %d' % len(E)
        indice_sol(d, E, pags or [0] * 36, h=h, es_fin=es_plan)
    sin_actualizar_campos(d)
    portada_seccion(d, B + 'portadas/%s/portada_planes.png' % g)
    return d, E


# ---------------- Solucionario ----------------

def nivel_h(el):
    st = el.find('.//' + qn('w:pStyle')) if el.tag == qn('w:p') else None
    v = st.get(qn('w:val')) if st is not None else ''
    m = re.match(r'^(?:Heading|Ttulo)(\d)$', v)
    return int(m.group(1)) if m else 0


def entradas_sol(d):
    """(nivel, título, ¿va después del índice?) de los títulos de nivel 1 y 2."""
    E, tras = [], False
    for el in d.element.body:
        tx = O.texto(el).strip()
        if nivel_h(el) == 1 and tx == 'Índice':
            tras = True
        elif nivel_h(el) in (1, 2) and tx:
            E.append((nivel_h(el), tx, tras))
    return E


def sin_actualizar_campos(d):
    """Los índices son estáticos: se quita el pedido de actualizar campos al abrir (Word mostraría un aviso)."""
    for upd in d.settings.element.findall(qn('w:updateFields')):
        d.settings.element.remove(upd)


def indice_sol(d, E, pags, h=None, es_fin=None):
    """Reemplaza el índice heredado (campo TOC con números de otra versión, o vacío) por uno con las páginas de esta edición."""
    body = d.element.body
    h = h if h is not None else next(el for el in body if nivel_h(el) == 1 and O.texto(el).strip() == 'Índice')
    es_fin = es_fin or (lambda el: nivel_h(el) == 1)
    els = list(body); i = els.index(h) + 1
    while i < len(els) and not es_fin(els[i]) and els[i].tag == qn('w:p'):
        body.remove(els[i]); i += 1
    ancla = OxmlElement('w:p'); h.addnext(ancla)
    p = docx.text.paragraph.Paragraph(ancla, d._body)
    M.run(p, 'Los números de página corresponden a esta edición.', italic=True, size=8.5, color='777777')
    ancho = round(M.PAG_W - 2 * M.MARG, 2)
    for (niv, t, _), pg in zip(E, pags):
        q = OxmlElement('w:p'); ancla.addnext(q); ancla = q
        p = docx.text.paragraph.Paragraph(q, d._body); pf = p.paragraph_format
        pf.space_after = Pt(0.6); pf.space_before = Pt(3 if niv == 1 else 0); pf.left_indent = Cm(0 if niv == 1 else 0.5)
        pf.tab_stops.add_tab_stop(Cm(ancho - 0.05), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        M.run(p, t, bold=(niv == 1), size=9.5 if niv == 1 else 9); p.add_run('\t')
        M.run(p, str(pg), bold=(niv == 1), size=9.5 if niv == 1 else 9)
    sin_actualizar_campos(d)


def solucionario():
    pags = None
    for vuelta in range(4):
        d, E, tiene_indice = _solucionario(pags)
        ruta = OUT + C['nombre_sol'] + '.docx'; d.save(ruta)
        pdf = O.a_pdf(ruta, OUT)
        if not tiene_indice:
            break
        norm = [re.sub(r'\s+', '', p.get_text()) for p in pymupdf.open(pdf)]
        idx = next(j for j, x in enumerate(norm) if 'Losnúmerosdepáginacorresponden' in x)
        n_antes = E.index(next(e for e in E if e[2])) if any(e[2] for e in E) else len(E)
        antes, _ = M.paginas_de(pdf, [t for _, t, _ in E[:n_antes]], desde=2)
        despues, _ = M.paginas_de(pdf, [t for _, t, _ in E[n_antes:]], desde=idx + 2)
        nuevas = antes + despues
        if nuevas == pags:
            break
        pags = nuevas
    if tiene_indice:
        import json
        json.dump({'entradas': E, 'paginas': pags}, open(B + '../verificacion/%s/indice_sol.json' % C['carpeta'], 'w'), ensure_ascii=False, indent=0)
    return ruta, pdf


def _solucionario(pags):
    d = docx.Document(C['base_sol'])
    O.margenes(d)
    M.ANCHO = round(M.PAG_W - 2 * M.MARG, 2)
    for viejo, nuevo in C['sol_reemplazos']:
        O.reemplazar_en_doc(d, viejo, nuevo)
    body = d.element.body
    R.NIVELES.update({'#': 1, '##': 2, '###': 3})
    if C.get('sol_modo') == 'insertar':
        # 9.º: se conserva el anexo de la base y se agregan, al final de cada taller, las respuestas y el instrumento
        B_ = R.bloques(C['pec_sol'])
        for letra, (enc, _) in TALLER.items():
            els = list(body)
            h = next(el for el in els if nivel_h(el) == 2 and O.texto(el).strip().startswith('Encuentro %d de 36 · Taller PEC %s' % (enc, letra)))
            i = els.index(h) + 1
            while i < len(els) and nivel_h(els[i]) not in (1, 2) and els[i].tag != qn('w:sectPr'):
                i += 1
            R.render_despues(d, els[i - 1], B_['SOL ' + letra])
    else:
        ini = next(el for el in body if el.tag == qn('w:p') and O.texto(el).strip().startswith('ANEXO DOCENTE'))
        els = list(body); i = els.index(ini)
        for el in els[i:]:
            if el.tag != qn('w:sectPr'):
                body.remove(el)
        R.render(d, R.leer(C['pec_sol']))
    # la portada de texto de la candidata queda reemplazada por la portada de imagen
    for el in list(body):
        if nivel_h(el) == 1:
            break
        if el.tag != qn('w:sectPr') and el.find('.//' + qn('w:sectPr')) is None:
            body.remove(el)
    sin_actualizar_campos(d)
    E = entradas_sol(d)
    tiene_indice = any(nivel_h(el) == 1 and O.texto(el).strip() == 'Índice' for el in body)
    if tiene_indice:
        indice_sol(d, E, pags or [0] * len(E))
    portada_seccion(d, B + 'portadas/%s/portada_solucionario.png' % g)
    return d, E, tiene_indice


if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    qué = sys.argv[2:] or ['pa', 'planes', 'sol']
    if 'pa' in qué:
        r = plan_anual(); p = pdf_control(r); print('Plan Anual', pymupdf.open(p).page_count, 'págs. (PDF solo de control)')
    if 'planes' in qué:
        r = planes(); p = pdf_control(r); print('Planes', pymupdf.open(p).page_count, 'págs. (PDF solo de control)')
    if 'sol' in qué:
        r, p = solucionario(); print('Solucionario', pymupdf.open(p).page_count, 'págs.')
