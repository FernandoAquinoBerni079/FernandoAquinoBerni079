# -*- coding: utf-8 -*-
"""Plan Anual — Desarrollo Personal y Social 7.º. 36 encuentros de 3 HC (108 HC), misma numeración que los Planes de Clase.
Temas, capacidades (literales) e indicadores se LEEN de las fichas del libro (fuente/tomo); procedimientos e instrumentos,
de los planes de clase (fuente/planes), únicos por fila. Oficio horizontal. Entregable: solo .docx (sin PDF para el docente)."""
import os, json, re, subprocess
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
import maqueta as M
import curriculo as C

B = os.path.dirname(os.path.abspath(__file__)) + '/'
F = B + '../fuente/'
OUT = B + '../salida/'
VER = B + '../verificacion/'
NOMBRE = 'Desarrollo_Personal_y_Social_7mo_Grado_PLAN_ANUAL_2026'
ANCHOS = [1.0, 6.6, 6.0, 1.2, 5.0, 4.3, 2.3]
COLS = ['N.º', 'TEMA Y CAPACIDAD', 'INDICADORES DE LOGRO', 'TIEMPO', 'PROCEDIMIENTOS', 'INSTRUMENTOS DE EVALUACIÓN', 'OBSERVACIÓN']
DARK, MED, LIGHT, AZUL, AZUL_CLARO, PEC_CLARO, AMB = M.DARK, M.MED, M.LIGHT, '1B4F72', 'EAF0F6', 'FFF6E0', 'B8860B'
NOTA = ('Las capacidades se reproducen textualmente del programa de estudio de Desarrollo Personal y Social del MEC (7.º grado). '
        'El Proyecto Educativo Comunitario no tiene capacidades propias en el programa: sus filas muestran las capacidades del área que '
        'movilizan y su propósito, y sus indicadores son de elaboración editorial, igual que los indicadores de logro de las clases.')
IND_EVAL = {
    'E': ['Responde consignas de comprensión sobre los temas de la unidad.', 'Analiza situaciones ficticias aplicando los conceptos trabajados.',
          'Fundamenta una posición personal con argumentos, sin exponer datos privados.'],
    18: ['Integra los aprendizajes de las Unidades 1 a 4 en consignas de comprensión.', 'Resuelve casos ficticios que combinan pertenencia, autoconocimiento, vida afectiva y derechos.',
         'Reconoce el propósito del Proyecto Educativo Comunitario y la etapa de identificación de necesidades.'],
    36: ['Integra los aprendizajes de las Unidades 5 a 7 en consignas de comprensión.', 'Resuelve casos ficticios sobre participación, organización del estudio y gestión de riesgo.',
         'Ordena y explica las etapas del proyecto comunitario realizado por el grado.'],
}


def ficha_de(n):
    L = open(F + 'tomo/e%02d.md' % n, encoding='utf-8').read().split('\n')
    tag = '[[FICHA]]' if '[[FICHA]]' in L else '[[FICHA_PEC]]'
    i = L.index(tag); j = L.index(tag.replace('[[', '[[/'))
    blk = L[i + 1:j]
    caps = [x[len('Capacidad: '):].strip() for x in blk if x.startswith('Capacidad: ')]
    tema = next((x.split(': ', 1)[1].strip() for x in blk if x.startswith(('Tema: ', 'Etapa del proyecto: '))), '')
    inds = [x[2:].strip() for x in blk if x.startswith('- ')]
    return caps, tema, inds


def plan(n):
    return json.load(open(F + 'planes/e%02d.json' % n, encoding='utf-8'))


def filas():
    rows = []
    for e in C.ENCUENTROS:
        n, t, k, u, caps_n, tit, _ = e
        p = plan(n)
        if t in 'CP':
            caps, tema, inds = ficha_de(n)
            assert caps == [C.CAP[c] for c in caps_n], 'capacidades de la ficha distintas de curriculo.py en el encuentro %d' % n
            if t == 'C':
                txt = [C.rotulo(e) + '. ' + tema] + ['Capacidad: «%s»' % c for c in caps]
            else:
                txt = [C.rotulo(e) + '. Etapa: ' + tema] + ['Capacidad del área vinculada: «%s»' % c for c in caps]
        else:
            inds = IND_EVAL.get(n, IND_EVAL['E'])
            txt = [C.rotulo(e)]
        obs = p.get('observacion', '').strip()
        if t == 'E':
            obs = obs or 'Cierre de la Unidad %d. Prueba en página propia del libro.' % u
        if n == 1 and 'diagnóstica' not in obs:
            obs = ('Prueba diagnóstica (sin nota). ' + obs).strip()
        if t == 'P' and C.PEC_MES[k] not in obs:
            obs = ('%s. ' % C.PEC_MES[k] + obs).strip()
        rows.append(dict(n=n, tipo=t, unidad=u, tema=txt, indicadores=inds, tiempo='3 HC', proc=p['procedimientos'].strip(),
                         inst=p['instrumentos'].strip(), obs=obs))
    assert len(rows) == 36
    rep = [r['n'] for r in rows if sum(1 for x in rows if x['proc'] == r['proc']) > 1]
    assert not rep, 'procedimientos repetidos en %s' % rep
    rep = [r['n'] for r in rows if sum(1 for x in rows if x['inst'] == r['inst']) > 1]
    assert not rep, 'instrumentos repetidos en %s' % rep
    cod = [r['n'] for r in rows for v in [r['proc'], r['inst'], r['obs']] + r['tema'] if re.search(r'\b(e\d\d|image\d+|C-\d\d)\b', v)]
    assert not cod, 'códigos internos en %s' % cod
    return rows


def celda(cell, texto, bold=False, blanco=False, size=7.8, italic_desde=None):
    par = cell.paragraphs[0]
    for i, linea in enumerate(texto if isinstance(texto, list) else [texto]):
        if i:
            par = cell.add_paragraph()
        par.paragraph_format.space_after = Pt(1)
        r = par.add_run(linea); r.bold = bold; r.font.size = Pt(size)
        if italic_desde is not None and i >= italic_desde:
            r.italic = True; r.font.color.rgb = RGBColor.from_string(DARK)
        if blanco:
            r.font.color.rgb = RGBColor.from_string('FFFFFF')


def construir():
    d = M.nuevo_doc('Desarrollo Personal y Social 7.º grado — Plan anual', 'Material del docente: plan anual de 36 encuentros',
                    'Desarrollo Personal y Social; plan anual; Proyecto Educativo Comunitario')
    s0 = d.sections[0]
    M.imagen_pagina_completa(d, B + 'portadas/portada_plan_anual.png')
    sec = d.add_section(WD_SECTION.NEW_PAGE)
    for hf in (s0.footer, s0.header):
        hf.is_linked_to_previous = False
        for p in hf.paragraphs:
            for el in list(p._p):
                if el.tag != qn('w:pPr'):
                    p._p.remove(el)
    M.formato_oficio(sec, horizontal=True)
    M.ANCHO = round(M.PAG_H - 2 * M.MARG, 2)
    M.par(d, 'PLAN ANUAL DETALLADO', 14, True, color=DARK, align='c', despues=2)
    M.par(d, 'Área: Desarrollo Personal y Social (Orientación Educacional y Vocacional · Proyecto Educativo Comunitario)  ·  Curso: 7.º grado  ·  Educación Escolar Básica, Tercer Ciclo  ·  3 horas cátedra semanales', 9.5, align='c', despues=1)
    M.par(d, 'República del Paraguay · 2026', 9, italic=True, align='c', despues=1)
    M.par(d, 'Institución educativa: ______________________________     Docente: ______________________________     Sección: __________', 9, align='c', despues=4)
    pp = M.par(d, '', 8.5, despues=3); M.run(pp, 'Nota sobre capacidades e indicadores. ', True, size=8.5, color=DARK); M.run(pp, NOTA, size=8.5)
    pp = M.par(d, '', 8.5, despues=4); M.run(pp, 'Distribución anual. ', True, size=8.5, color=DARK)
    M.run(pp, '36 encuentros de 3 horas cátedra (108 HC de referencia, sujetas al calendario institucional), 18 en cada etapa: 20 clases del componente '
              'Orientación Educacional y Vocacional, 7 talleres del componente Proyecto Educativo Comunitario (junio a noviembre, filas sombreadas), '
              '7 evaluaciones de unidad y 2 evaluaciones integradoras de etapa. Las 3 HC semanales son las del área: el proyecto no suma horas. '
              'Cada fila tiene su Plan de Clase con el mismo número. Proyecto del grado (perspectiva del programa): «La escuela como una comunidad».', size=8.5)
    rows = filas()
    t = d.add_table(rows=1, cols=7)
    M.bordes(t, '808080', '4'); M.margenes_celda(t, 30, 70); M.ancho_fijo(t, ANCHOS)
    for j, h in enumerate(COLS):
        M.shd(t.rows[0].cells[j], '5D4037'); celda(t.rows[0].cells[j], h, bold=True, blanco=True, size=8)

    def fila():
        r = t.add_row()
        for j, c in enumerate(r.cells):
            c.width = Cm(ANCHOS[j] * M.ANCHO / sum(ANCHOS))
        return r

    def banda(texto, fill, size=9):
        r = fila(); m = r.cells[0].merge(r.cells[6]); M.shd(m, fill); celda(m, texto, bold=True, blanco=True, size=size)

    u_prev = None
    for r_ in rows:
        if r_['n'] == 1:
            banda('PRIMERA ETAPA · febrero – junio · 18 encuentros / 54 HC', AZUL)
        if r_['n'] == 19:
            banda('SEGUNDA ETAPA · julio – noviembre · 18 encuentros / 54 HC', AZUL)
        if r_['tipo'] == 'C' and r_['unidad'] != u_prev:
            u_prev = r_['unidad']
            banda('UNIDAD %d — %s · Período aproximado: %s' % (u_prev, C.UNIDADES[u_prev].upper(), C.PERIODO[u_prev]), MED, 8.5)
        rr = fila()
        ntema = 1
        vals = [str(r_['n']), r_['tema'], ['• ' + i for i in r_['indicadores']], r_['tiempo'], r_['proc'], r_['inst'], r_['obs']]
        for j, v in enumerate(vals):
            celda(rr.cells[j], v, bold=(j == 0 or (j == 1 and r_['tipo'] in 'EIP')) if j != 1 else False, italic_desde=(ntema if j == 1 else None))
        if j == 6 and r_['tipo'] in 'EIP':
            pass
        if r_['tipo'] in 'EI':
            for c in rr.cells:
                M.shd(c, AZUL_CLARO)
        elif r_['tipo'] == 'P':
            for c in rr.cells:
                M.shd(c, PEC_CLARO)
            M.shd(rr.cells[0], 'F2C14E')
        else:
            M.shd(rr.cells[0], LIGHT)
        if r_['tipo'] in 'EIP':
            rr.cells[1].paragraphs[0].runs[0].bold = True
    M.cant_split(t, header=True)
    M.par(d, 'Criterios generales de planificación y evaluación', 10, True, color=DARK, keep=True, antes=6)
    for x in ['El área tiene carácter eminentemente formativo y vivencial: se valoran procesos (participación, producciones individuales y grupales, compromisos) y no solo resultados.',
              'Las secciones «Mirate» del libro son privadas: no se califican ni se exponen. Ningún instrumento exige revelar experiencias o datos personales; las evaluaciones usan casos ficticios.',
              'Las 7 evaluaciones de unidad y las 2 integradoras son encuentros reales de 3 HC, en páginas propias del libro, con tiempo para la devolución y la recuperación.',
              'El Proyecto Educativo Comunitario se evalúa como proceso y como producto con la bitácora del equipo, las listas de cotejo de cada taller y la rúbrica final del solucionario (indicadores Logrado / No logrado / No manifestado del esquema MEC). Su peso en la calificación lo define la institución según su reglamento de evaluación.',
              'Las actividades del proyecto fuera del aula se realizan en horario escolar, con autorización de la dirección y, cuando corresponde, de las familias; la docente registra las evidencias de seguimiento.',
              'La escala de calificaciones y el porcentaje mínimo de logro son los del reglamento de evaluación vigente del MEC y de la institución.']:
        M.par(d, '• ' + x, 8.5, despues=1)
    M.par(d, 'Estrategias de inclusión', 10, True, color=DARK, keep=True, antes=4)
    for x in ['Consignas orales y escritas; tiempo adicional en las evaluaciones para quienes lo requieran.',
              'Equipos heterogéneos y rotativos, con roles diferenciados en las dinámicas y en el proyecto, para que cada estudiante aporte desde sus fortalezas.',
              'Uso del castellano y del guaraní según la necesidad comunicativa del grupo.',
              'Adecuaciones razonables de acceso para estudiantes con discapacidad en las dinámicas corporales, expositivas y en las tareas del proyecto.']:
        M.par(d, '• ' + x, 8.5, despues=1)
    M.par(d, '__________________________________                                        __________________________________', 9.5, align='c', despues=0, antes=14)
    M.par(d, 'Firma del docente                                                                         Dirección — Visto bueno', 9, align='c')
    M.pie_paginas(d, sec)
    M.no_actualizar_campos(d)
    M.purgar_relaciones(d)
    return d, rows


if __name__ == '__main__':
    d, rows = construir()
    out = OUT + NOMBRE + '.docx'; d.save(out)
    # PDF solo para control visual (no es entregable): va a verificacion/, no a salida/
    os.makedirs(VER + 'control_pdf', exist_ok=True)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', VER + 'control_pdf', out], check=True, capture_output=True,
                   env=dict(os.environ, HOME='/tmp/lohome'))
    import pymupdf
    print('Plan anual: páginas', pymupdf.open(VER + 'control_pdf/' + NOMBRE + '.pdf').page_count, '· filas', len(rows))
    json.dump(rows, open(B + 'pa_rows.json', 'w'), ensure_ascii=False, indent=0)
