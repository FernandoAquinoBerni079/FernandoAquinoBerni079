# -*- coding: utf-8 -*-
"""Render del texto marcado de las fuentes (fuente/tomo, fuente/sol) a .docx. Marcado: fuente/GUIA_REDACCION.md §5."""
import re, os
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import maqueta as M
import curriculo as C

B = os.path.dirname(os.path.abspath(__file__)) + '/'
FIGS = [B + '../figs/originales/', B + '../figs/nuevas/']
PREFIJO = {'ejemplo': 'Ejemplo — ', 'concepto': 'Concepto clave — ', 'paraguay': 'En Paraguay — ', 'cuidado': 'Para cuidarte — ',
           'mirate': 'Mirate — para pensar sobre vos', 'taller': 'Actividad práctica — ', 'bitacora': 'Para la bitácora del proyecto — ',
           'docente': 'Orientación docente — ', 'neutra': ''}
ROTULOS = ('Pensá y respondé:', 'En grupo:', 'En equipo:', 'Para la bitácora:', 'Aplicá:', 'Desafío:')
USADAS = []  # figuras usadas (control: cada figura una vez)


# Figuras de la edición anterior con texto desbordado o cortado: se usan sus redibujos fieles (figuras_nuevas.py).
# Los archivos originales quedan sin modificar en figs/originales/.
REDIBUJO = {'image3.png': 'fig_norma.png', 'image5.png': 'fig_escalera.png', 'image10.png': 'fig_derechos.png',
            'image14.png': 'fig_participacion.png', 'image18.png': 'fig_metas.png',
            'image21.png': 'fig_red_proyecto.png'}


def figura(nombre):
    nombre = REDIBUJO.get(nombre, nombre)
    for d in FIGS:
        if os.path.exists(d + nombre):
            return d + nombre
    if os.environ.get('DPS_PARCIAL'):
        return B + 'portadas/portada_libro.png'  # solo para pruebas parciales
    raise FileNotFoundError(nombre)


def leer(ruta):
    return open(ruta, encoding='utf-8').read().split('\n')


def bloque(L, i, fin):
    j = L.index(fin, i)
    return L[i + 1:j], j + 1


def renglones_celda(cell, n):
    for _ in range(n):
        p = cell.add_paragraph(); pf = p.paragraph_format
        pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = Cm(0.72)
        pPr = p._p.get_or_add_pPr(); bd = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
        b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:space'), '1'); b.set(qn('w:color'), 'A0A0A0')
        bd.append(b); pPr.append(bd)
        p.add_run(' ').font.size = Pt(9)


def caja(d, tipo, cuerpo):
    tit = PREFIJO.get(tipo, '')
    lineas = [x for x in cuerpo if x.strip()]
    if tipo in ('ejemplo', 'concepto', 'paraguay', 'cuidado', 'taller', 'bitacora', 'docente') and lineas and not any(
            l.startswith(('[[', '    ')) or re.match(r'^\d+\. ', l) for l in lineas):
        # título en la primera línea, texto a continuación (como el paquete anterior)
        return M.caja(d, tit[:-3] if tit.endswith(' — ') else tit, lineas)
    # caja con consignas o renglones: se arma celda por celda
    t = M.caja(d, tit[:-3] if tit.endswith(' — ') else tit, [])
    c = t.rows[0].cells[0]
    for l in lineas:
        m = re.match(r'\[\[RENGLONES (\d+)\]\]', l.strip())
        if m:
            renglones_celda(c, int(m.group(1))); continue
        p = c.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        ind = 0.9 if l.startswith('    ') else 0.3 if re.match(r'^\d+\. ', l) else 0
        if ind: p.paragraph_format.left_indent = Cm(ind)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        M.run(p, l.strip(), size=9.5, italic=(tipo == 'mirate' and not re.match(r'^\d+\. ', l)))
    return t


def tabla(d, filas):
    filas = [[c.strip() for c in f.split(' | ')] for f in filas if f.strip()]
    hdr, rows = filas[0], filas[1:]
    n = len(hdr)
    rows = [r + [''] * (n - len(r)) if len(r) < n else r[:n] for r in rows]
    vacia = rows and all(not c for r in rows for c in r[1:])
    t = M.tabla(d, hdr, rows)
    if vacia or any(not c for r in rows for c in r):
        for row in t.rows[1:]:
            if all(not cell.text.strip() for cell in row.cells[1:]):
                trPr = row._tr.get_or_add_trPr(); h = OxmlElement('w:trHeight'); h.set(qn('w:val'), '520'); h.set(qn('w:hRule'), 'atLeast'); trPr.append(h)
    return t


def ficha(d, blk, pec=False):
    caps = [x[len('Capacidad: '):].strip() for x in blk if x.startswith('Capacidad: ')]
    inds = [x[2:].strip() for x in blk if x.startswith('- ')]
    if not pec:
        tema = next(x[len('Tema: '):].strip() for x in blk if x.startswith('Tema: '))
        return M.ficha(d, caps if len(caps) > 1 else caps[0], tema, inds), caps, tema, inds
    etapa = next(x[len('Etapa del proyecto: '):].strip() for x in blk if x.startswith('Etapa del proyecto: '))
    filas = [('Componente', ['Proyecto Educativo Comunitario (componente del área Desarrollo Personal y Social)']),
             ('Propósito del componente (programa MEC)', ['«' + C.PEC_PROPOSITO + '»']),
             ('Etapa del proyecto', [etapa]),
             ('Capacidades del área vinculadas', caps if len(caps) == 1 else ['• ' + x for x in caps]),
             ('Indicadores de logro (elaboración editorial)', ['• ' + i for i in inds])]
    M.ficha_filas(d, filas)
    return None, caps, etapa, inds


def render(d, L, solucionario=False):
    i = 0
    prev = None
    while i < len(L):
        l = L[i]; s = l.strip()
        if not s or s.startswith('@@'):
            i += 1; continue
        if s == '[[FICHA]]':
            blk, i = bloque(L, i, '[[/FICHA]]'); ficha(d, blk); continue
        if s == '[[FICHA_PEC]]':
            blk, i = bloque(L, i, '[[/FICHA_PEC]]'); ficha(d, blk, pec=True); continue
        m = re.match(r'\[\[CAJA (\w+)\]\]$', s)
        if m:
            blk, i = bloque(L, i, '[[/CAJA]]'); caja(d, m.group(1), blk); continue
        if s == '[[TABLA]]':
            blk, i = bloque(L, i, '[[/TABLA]]')
            if prev is not None and prev.text.rstrip().endswith(':'):
                prev.paragraph_format.keep_with_next = True
            tabla(d, blk); continue
        m = re.match(r'\[\[IMG (\S+)\|([\d.]+)\]\]', s)
        if m:
            if prev is not None:
                prev.paragraph_format.keep_with_next = True
            USADAS.append(m.group(1))
            M.imagen(d, figura(m.group(1)), L[i + 1].strip(), float(m.group(2)))
            i += 2; continue
        m = re.match(r'\[\[RENGLONES (\d+)\]\]', s)
        if m:
            M.renglones(d, int(m.group(1))); i += 1; continue
        if s == '[[SALTO]]':
            M.salto(d); i += 1; continue
        if l.startswith('### '):
            prev = M.heading(d, s[4:], 4); i += 1; continue
        if l.startswith('## '):
            prev = M.heading(d, s[3:], 3); i += 1; continue
        if s in ROTULOS or (s.endswith(':') and len(s) < 40 and s.split()[0] in ('Pensá', 'En', 'Para', 'Aplicá', 'Desafío')):
            prev = M.par(d, s, 10.5, True, color=M.MED, keep=True, antes=4, despues=2); i += 1; continue
        if re.match(r'^[a-e]\) ', s) and l.startswith('    '):
            sig = L[i + 1] if i + 1 < len(L) else ''
            prev = M.par(d, s, indent=0.9, despues=1, keep=sig.startswith('    ')); i += 1; continue
        if re.match(r'^\d+\. ', s):
            sig = L[i + 1] if i + 1 < len(L) else ''
            prev = M.par(d, s, indent=0.3, despues=2, keep=sig.startswith(('    ', '[[RENGLONES'))); i += 1; continue
        if s.startswith('- '):
            prev = M.par(d, '• ' + s[2:], indent=0.4, despues=2); i += 1; continue
        prev = M.par(d, s); i += 1
    return d
