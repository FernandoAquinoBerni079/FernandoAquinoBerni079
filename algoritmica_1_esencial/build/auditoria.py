# -*- coding: utf-8 -*-
"""Auditoría automática de liberación — Algorítmica 1.º Edición Esencial Comercial 2026 (v1).
Cada control suma al conteo; el resultado esperado es 0 fallos. Incluye la re-ejecución en PSeInt de
los algoritmos de las prácticas 14 a 21 y del Puente a PSeInt, y el control de las lecciones del piloto de 3.º (v2 y v2.1)."""
import re, json, os, collections, subprocess, hashlib
import docx, pymupdf
from docx.oxml.ns import qn
import actividades, practicas, evaluaciones, planes_data, preliminares as PRE
from estructura import cargar
import ediciones

O = '/home/claude/alg1/salida/'; B = '/home/claude/alg1/build/'; PS = '/home/claude/alg1/pseint/'
F = {k: O + 'Algoritmica_1er_Curso_%s.docx' % k for k in
     ('LIBRO_ESENCIAL_COMERCIAL_2026', 'SOLUCIONARIO_DOCENTE_ESENCIAL_COMERCIAL_2026', 'PLANES_DE_CLASE_ESENCIAL_COMERCIAL_2026', 'PLAN_ANUAL_ESENCIAL_COMERCIAL_2026')}
LIB, SOL, PLA, ANU = list(F.values())
RES = []


def chk(bloque, nombre, ok, detalle=''):
    RES.append((bloque, nombre, bool(ok), detalle if not ok else ''))


def pdf(p):
    return pymupdf.open(p[:-5] + '.pdf')


def textos_doc(path):
    d = docx.Document(path)
    out = []
    for el in d.element.body.iter():
        if el.tag == qn('w:p'):
            t = ''.join(x.text or '' for x in el.iter(qn('w:t')))
            st = el.find('.//' + qn('w:pStyle'))
            out.append((st.get(qn('w:val')) if st is not None else '', t))
    return d, out


def H(s):
    return s.startswith('Heading') or s.startswith('Ttulo')


_pre, C, _ev, U = cargar(); ediciones.aplicar(C); FIG = ediciones.fig_refs(C)
U = {1: 'Teoría de Conjuntos', 2: 'Lógica Simbólica', 3: 'Introducción a la Algoritmia'}
PFI_TIT = 'Proyecto Final Integrador — Presentamos nuestro emprendimiento'
dL, PL = textos_doc(LIB)
TL = '\n'.join(t for _, t in PL)
pdfL = pdf(LIB)
pags = [pg.get_text() for pg in pdfL]
dS, PSx = textos_doc(SOL); TS = '\n'.join(t for _, t in PSx)

# ---------------- A. Estructura ----------------
clases = [t for s, t in PL if re.match(r'^Clase \d+ — ', t) and H(s)]
prs = [t for s, t in PL if re.match(r'^Práctica \d+ — ', t) and H(s)]
chk('A', 'Cantidad de clases = 21', len(clases) == 21, len(clases))
chk('A', 'Cantidad de prácticas = 21', len(prs) == 21, len(prs))
seq = [t.split(' — ')[0] for s, t in PL if H(s) and re.match(r'^(Clase|Práctica) \d+ — ', t)]
chk('A', 'Correspondencia Clase N → Práctica N (orden intercalado, 1 a 21)', seq == [x for n in range(1, 22) for x in ('Clase %d' % n, 'Práctica %d' % n)])
for n in range(1, 22):
    chk('A', 'Título de la Clase %d coincide con la fuente' % n, ('Clase %d — %s' % (n, C[n]['titulo'])) in clases)
    chk('A', 'Título de la Práctica %d coincide con la fuente' % n, ('Práctica %d — %s' % (n, practicas.PR[n]['titulo'])) in prs)
fichas = [t for t in dL.tables if t.rows[0].cells[0].text.strip() in ('Capacidad', 'Capacidades')]
chk('A', 'Una ficha por clase (21)', len(fichas) == 21, len(fichas))
unidades = [t for s, t in PL if H(s) and t.startswith('UNIDAD ')]
chk('A', 'Tres unidades con los títulos del tomo vigente', unidades == ['UNIDAD %d — %s' % (u, U[u]) for u in range(1, 4)], unidades)
chk('A', 'Unidades por clases: 1–5, 6–12, 13–21', [C[n]['unidad'] for n in range(1, 22)] == [1] * 5 + [2] * 7 + [3] * 9)
orden_ev = [(e['id'], e['tras']) for e in evaluaciones.EV]
chk('A', 'Evaluaciones en su lugar: U1 tras la Clase 5, U2 y E1 tras la 12, U3 y E2 tras la 21', orden_ev == [('U1', 5), ('U2', 12), ('E1', 12), ('U3', 21), ('E2', 21)])
for e in evaluaciones.EV:
    chk('A', 'Evaluación presente: ' + e['titulo'], e['titulo'] in TL)
chk('A', 'Prueba diagnóstica presente', 'Prueba diagnóstica — ¿Qué sabés ya?' in TL)
chk('A', 'Proyecto Final Integrador presente', PFI_TIT in TL)
for n in range(1, 22):
    it = actividades.items(n)
    chk('A', 'Clase %d: 13 actividades numeradas 1..13' % n, [x[1] for x in it] == list(range(1, 14)))
    p = practicas.PR[n]
    chk('A', 'Práctica %d: competencia, saber, actividades, control y desafío' % n,
        p['competencia'] and p['saber'] and p['acts'] and all(a['control'] and a['pasos'] and a['objetivo'] for a in p['acts']) and p['desafio'])
    chk('A', 'Práctica %d: 3 actividades, como en el cuaderno vigente' % n, len(p['acts']) == 3)

# ---------------- B. Paginación, índice, saltos, encabezados ----------------
ji = json.load(open(B + 'indice_libro.json'))
for (niv, t), p in zip(ji['entradas'], ji['paginas']):
    chk('B', 'Índice del libro: «%s» en pág. %d' % (t[:50], p), re.sub(r'\s+', '', t) in re.sub(r'\s+', '', pags[p - 1]))
chk('B', 'Índice poblado (sin «Ctrl+E/F9» ni marcadores 000)', 'Ctrl+E' not in TL and not re.search(r'\.{5,}\s*000\b', ''.join(pags[1:5])))
for t in prs:
    p = next(i for i, x in enumerate(pags) if re.sub(r'\s+', '', t) in re.sub(r'\s+', '', x) and i > 0 and not re.search(r"\.{8,}", x))
    chk('B', 'Salto: «%s» abre página' % t[:40], re.sub(r'\s+', '', pags[p]).startswith(re.sub(r'\s+', '', t)[:30]))
for e in evaluaciones.EV:
    p = next(i for i, x in enumerate(pags) if e['titulo'] in x and i > 0 and not re.search(r"\.{8,}", x))
    chk('B', 'Salto: «%s» abre página' % e['titulo'], pags[p].strip().startswith(e['titulo']))
for t in clases:
    p = next(i for i, x in enumerate(pags) if re.sub(r'\s+', '', t) in re.sub(r'\s+', '', x) and i > 0 and not re.search(r"\.{8,}", x))
    top = re.sub(r'\s+', '', pags[p])
    chk('B', 'Salto: «%s» abre página (o sigue a la apertura de unidad)' % t[:40], top.startswith(re.sub(r'\s+', '', t)[:25]) or top.startswith('UNIDAD'))
chk('B', 'Pie «Página X de Y» en todas las páginas salvo portada y contraportada',
    all(re.search(r'Página \d+ de %d' % len(pags), pags[i]) for i in range(1, len(pags) - 1)) and 'Página' not in pags[0] and 'Página' not in pags[-1])
for nombre, path, j in (('Solucionario', SOL, 'indice_sol.json'), ('Planes', PLA, 'indice_planes.json')):
    pp = [pg.get_text() for pg in pdf(path)]
    jj = json.load(open(B + j))
    bad = [t for (_, t), p in zip(jj['entradas'], jj['paginas']) if re.sub(r'\s+', '', t) not in re.sub(r'\s+', '', pp[p - 1])]
    chk('B', 'Índice de %s verificado contra el PDF (%d entradas)' % (nombre, len(jj['entradas'])), not bad, bad[:3])
    chk('B', '%s: portada sin pie de página' % nombre, 'Página' not in pp[0])
ppla = [pg.get_text() for pg in pdf(PLA)]
jp = json.load(open(B + 'indice_planes.json'))
cont_ok = all('continuación' in ppla[q][:250] for k, p in enumerate(jp['paginas']) for q in range(p, (jp['paginas'][k + 1] - 1 if k + 1 < len(jp['paginas']) else len(ppla))))
chk('B', 'Planes: encabezado «· continuación» en las páginas siguientes de cada plan', cont_ok)
chk('B', 'Planes: 37 planes, cada uno abre página', len(jp['paginas']) == 37 and all(ppla[p - 1].strip().startswith('Plan de Clase N.º') for p in jp['paginas']))
for path in F.values():
    mb = os.path.getsize(path) / 1e6
    chk('B', 'DOCX < 25 MB: %s (%.2f MB)' % (os.path.basename(path), mb), mb < 25)
    d = docx.Document(path)
    chk('B', 'Metadatos de colección: %s' % os.path.basename(path), d.core_properties.author == 'Equipo editorial' and '1.er Curso' in (d.core_properties.title or ''))
    chk('B', 'PDF regenerado después del DOCX: %s' % os.path.basename(path), os.path.getmtime(path[:-5] + '.pdf') >= os.path.getmtime(path))


def cantsplit(path):
    d = docx.Document(path); tot = ok = 0
    for t in d.tables:
        for r in t.rows:
            tot += 1; ok += r._tr.trPr is not None and r._tr.trPr.find(qn('w:cantSplit')) is not None
    return ok, tot


import maqueta as MQ
for nombre, path, horiz in (('libro', LIB, False), ('solucionario', SOL, False), ('planes de clase', PLA, False), ('plan anual', ANU, True)):
    d_ = docx.Document(path)
    secs = d_.sections
    tamano = all(abs(s_.page_width.cm - MQ.PAG_W) < 0.05 and abs(s_.page_height.cm - MQ.PAG_H) < 0.05 for s_ in (secs[:1] if horiz else secs))
    if horiz:
        tamano = tamano and all(abs(s_.page_width.cm - MQ.PAG_H) < 0.05 and abs(s_.page_height.cm - MQ.PAG_W) < 0.05 for s_ in secs[1:])
    chk('B', 'Oficio 21,6 × 33 cm en todas las secciones del %s%s' % (nombre, ' (horizontal en el cuerpo)' if horiz else ''), tamano)
    chk('B', 'Márgenes estrechos (1,27 cm) en todas las secciones del %s' % nombre, all(abs(m.cm - MQ.MARG) < 0.03 for s_ in secs for m in (s_.left_margin, s_.right_margin, s_.top_margin, s_.bottom_margin)))
    tws = [t._tbl.tblPr.find(qn('w:tblW')) for t in d_.tables]
    chk('B', 'Tablas y cajas ajustadas al ancho de la página (100 %%) en el %s (%d)' % (nombre, len(tws)), all(w is not None and w.get(qn('w:type')) == 'pct' and w.get(qn('w:w')) == '5000' for w in tws))
    pdf_ = pdf(path)
    chk('B', 'PDF del %s en oficio' % nombre, all(abs(sorted([pg.rect.width, pg.rect.height])[0] / 72 * 2.54 - MQ.PAG_W) < 0.05 and abs(sorted([pg.rect.width, pg.rect.height])[1] / 72 * 2.54 - MQ.PAG_H) < 0.05 for pg in pdf_))
for nombre, path in (('libro', LIB), ('solucionario', SOL), ('plan anual', ANU)):
    ok, tot = cantsplit(path)
    chk('B', 'cantSplit en el 100 %% de las filas del %s (%d/%d)' % (nombre, ok, tot), ok == tot)
chk('B', 'Muestra comercial generada', os.path.exists(O + 'Algoritmica_1er_Curso_MUESTRA_COMERCIAL_2026.pdf'))

# ---------------- C. Coherencia aritmética y ejecución real en PSeInt ----------------
def venn(n, a, b, ambos):
    return a - ambos, ambos, b - ambos, n - (a + b - ambos)
chk('C', 'Práctica 5: encuesta 45/28/21/12 → 16, 12, 9, 8', venn(45, 28, 21, 12) == (16, 12, 9, 8))
chk('C', 'Práctica 5: transferencia biblioteca 40/22/17, 9 ninguna → ambas 8', 22 + 17 - (40 - 9) == 8 and venn(40, 22, 17, 8) == (14, 8, 9, 9))
chk('C', 'Actividad 5.6: 36/21/13/7 → 14, 7, 6, 9', venn(36, 21, 13, 7) == (14, 7, 6, 9))
chk('C', 'Evaluación U1: 35/20/14/6 → 14, 6, 8, 7', venn(35, 20, 14, 6) == (14, 6, 8, 7))
chk('C', 'Evaluación E1: 30/17/12/5 → 12, 5, 7, 6', venn(30, 17, 12, 5) == (12, 5, 7, 6))
chk('C', 'Clase 5 (ejemplo): 30/18/15/8 → 10, 8, 7, 5', venn(30, 18, 15, 8) == (10, 8, 7, 5))
from itertools import product
def tabla(f, k=2):
    return ''.join('V' if f(*v) else 'F' for v in product([True, False], repeat=k))
imp = lambda a, b: (not a) or b
chk('C', 'Práctica 8: p ∧ ¬q = FVFF; ¬(p ∨ q) = FFFV; (p → q) ∧ p = VFFF; ¬p ↔ q = FVVF',
    [tabla(lambda p, q: p and not q), tabla(lambda p, q: not (p or q)), tabla(lambda p, q: imp(p, q) and p), tabla(lambda p, q: (not p) == q)] == ['FVFF', 'FFFV', 'VFFF', 'FVVF'])
chk('C', 'Práctica 9: p → (q ∨ r) tiene una sola F (V F F)', tabla(lambda p, q, r: imp(p, q or r), 3) == 'VVVFVVVV')
chk('C', 'Práctica 9: transferencia R1 tautología, R2 contradicción, R3 = VFVV',
    [tabla(lambda p, q: imp(p and q, p)), tabla(lambda p, q: (p or q) and (not p and not q)), tabla(lambda p, q: imp(p, p and q))] == ['VVVV', 'FFFF', 'VFVV'])
chk('C', 'Práctica 10: contrarrecíproco p → q ≡ ¬q → ¬p', tabla(lambda p, q: imp(p, q)) == tabla(lambda p, q: imp(not q, not p)) == 'VFVV')
chk('C', 'Actividad 9.7: (p ∨ q) → r falsa en las filas 2, 4 y 6', [i + 1 for i, c in enumerate(tabla(lambda p, q, r: imp(p or q, r), 3)) if c == 'F'] == [2, 4, 6])
chk('C', 'Actividad 8.9: p ↔ q ≡ (p ∧ q) ∨ (¬p ∧ ¬q)', tabla(lambda p, q: p == q) == tabla(lambda p, q: (p and q) or (not p and not q)) == 'VFFV')
chk('C', 'Evaluación U2: (p → q) ∧ ¬q = FFFV (indeterminación)', tabla(lambda p, q: imp(p, q) and not q) == 'FFFV')
chk('C', 'Evaluación E1: ¬p → q ≡ p ∨ q', tabla(lambda p, q: imp(not p, q)) == tabla(lambda p, q: p or q))
p, q, r, s_ = True, False, True, False
chk('C', 'Práctica 7: con p=V, q=F, r=V, s=F solo a) y g) son F',
    [p and q, p or q, imp(q, r), (not p) != r, imp(p and q, r), (p or s_) == (q or r), not (p and r)] == [False, True, True, True, True, True, False])
chk('C', 'Práctica 15: 13 + 27 + 15 + 14 + 50', [7 + 2 * 3, (7 + 2) * 3, 18 - 6 / 2, 4 * 5 - 3 * 2, 100 - (20 + 5) * 2] == [13, 27, 15, 14, 50])
def caja(v):
    suma = c = cd = 0
    for x in v:
        cob = x - x * 10 / 100 if x >= 50000 else x
        cd += x >= 50000; suma += cob; c += 1
    return suma, c, cd, suma / c
chk('C', 'Práctica 21: jueves 80.000 · 30.000 · 70.000 → 165.000, 3, 2, 55.000', caja([80000, 30000, 70000]) == (165000, 3, 2, 55000))
chk('C', 'Clase 21: día real → 134.000, 4, 2, 33.500', caja([20000, 60000, 15000, 50000]) == (134000, 4, 2, 33500))
chk('C', 'Actividad 21.7: 142.000, 4, 1, 35.500', caja([45000, 80000, 10000, 15000]) == (142000, 4, 1, 35500))
chk('C', 'Actividad 21.9: 160.500, 3, 2, 53.500', caja([30000, 55000, 90000]) == (160500, 3, 2, 53500))
chk('C', 'Evaluación U3: 65.000, 2, 1, 32.500', caja([50000, 20000]) == (65000, 2, 1, 32500))
chk('C', 'Evaluación E2: 138.000, 3, 2, 46.000', caja([30000, 50000, 70000]) == (138000, 3, 2, 46000))
chk('C', 'Actividad 20.13: 51.000, 3, 17.000, mayor 30.000', sum([12000, 30000, 9000]) == 51000 and 51000 / 3 == 17000)
chk('C', 'Actividad 20.9: propinas 12.000, 3, 4.000', sum([2000, 5000, 5000]) == 12000)
PSE = [('pr/P14_doble', ['35'], ['70']), ('pr/P14_area', ['8', '5'], ['40']), ('pr/P14_harina', ['6000'], ['21000']),
       ('pr/P15_expr', [], ['13 27 15 14 50', 'FALSO VERDADERO FALSO VERDADERO VERDADERO']),
       ('pr/P16_vuelto', ['3500', '4', '25000'], ['14000 11000']), ('pr/P16_swap', ['7', '3'], ['3 7']), ('pr/P16_swap', ['10', '4'], ['4 10']),
       ('pr/P16_prom', ['2000', '4500', '3500'], ['10000 3333.3333333333']), ('pr/P16_recargo', ['18000'], ['19800 1800']),
       ('pr/P17_envio', ['70000'], ['¡Envío gratis!', '70000']), ('pr/P17_envio', ['45000'], ['53000']), ('pr/P17_par', ['0'], ['par']), ('pr/P17_par', ['45'], ['impar']),
       ('pr/P17_croq', ['10'], ['20000']), ('pr/P17_croq', ['6'], ['15000']),
       ('pr/P18_jugo', ['500'], ['grande 6000']), ('pr/P18_jugo', ['300'], ['mediano 4500']), ('pr/P18_jugo', ['250'], ['chico 3000']),
       ('pr/P18_promo', ['30000', 'QR'], ['promo']), ('pr/P18_promo', ['30000', 'efectivo'], ['sin promo']),
       ('pr/P18_rota', ['45'], ['familiar']), ('pr/P18_ok', ['45'], ['evento']), ('pr/P18_ok', ['18'], ['familiar']), ('pr/P18_ok', ['4'], ['individual']),
       ('pr/P19_mientras', [], ['2\n5\n8\n11\nfin 14']), ('pr/P19_para', [], ['Bandeja 9 lista']), ('pr/P19_fix', [], ['sale con 6', '¡Cerrado!']),
       ('pr/P20_caja', ['7000', '13000', '4000', '0'], ['24000 8000 3 1']), ('pr/P20_caja', ['0'], ['sin ventas']),
       ('pr/P20_bolsas', ['25', '25', '10', '0'], ['60 3 20 25']), ('pr/P20_rifa', ['35000', '50000', '20000', '45000', '0'], ['150000 4 37500 2']),
       ('pr/P21_caja', ['80000', '30000', '70000', '0'], ['165000 3 2 55000 bruto 180000 desc 15000 mayor 80000']),
       ('T_puente', ['20000', '3'], ['54000']), ('T_puente', ['25000', '2'], ['45000']),
       ('T_miles', ['30000'], ['27000'])]
for f_, ins_, esp in PSE:
    out = subprocess.run([PS + 'run.sh', PS + f_ + '.psc'] + ins_, capture_output=True, text=True).stdout
    chk('C', 'PSeInt 20250314 (Flexible): %s %s' % (f_, ' '.join(ins_)[:30]), all(e in out for e in esp), out[-200:])
for n in range(14, 22):
    chk('C', 'Práctica %d: el solucionario declara la comprobación en PSeInt' % n, 'comprobado en PSeInt 20250314' in practicas.PR[n]['sol']['resultado'])

# ---------------- D. Libro ↔ solucionario ----------------
for n in range(1, 22):
    it = actividades.items(n)
    chk('D', 'Solucionario: respuestas de la Clase %d (13 ítems)' % n, all(('%d. %s' % (k, r)) in TS for _, k, _, r, _ in it))
    chk('D', 'Libro: consignas de la Clase %d (13 ítems)' % n, all(('%d. %s' % (k, e)) in TL for _, k, e, _, _ in it))
    chk('D', 'Solucionario: resultado y errores de la Práctica %d' % n, practicas.PR[n]['sol']['resultado'] in TS and practicas.PR[n]['sol']['errores'] in TS)
for e in evaluaciones.EV:
    k = 0; okmc = True
    for enun, ops, letra in e['mc']:
        k += 1
        okmc &= ('%d. %s) %s' % (k, letra, ops['abcd'.index(letra)])) in TS and ('%d. %s' % (k, enun)) in TL
    chk('D', 'Opción múltiple coherente libro ↔ clave: ' + e['titulo'], okmc)
    chk('D', 'Abiertas con clave: ' + e['titulo'], all(r in TS for _, r in e['ab']) and all(q in TL for q, _ in e['ab']))
chk('D', 'Libro del estudiante sin respuestas ni rúbricas', not any(practicas.PR[n]['sol']['resultado'][:60] in TL for n in range(1, 22)) and 'Errores frecuentes:' not in TL and 'Rúbrica analítica' not in TL)
chk('D', 'Evaluaciones sin la pista ambigua «(contando desde 2)» de E1 A1', 'contando desde' not in TL)
for d_ in PRE.DIAGNOSTICA:
    chk('D', 'Diagnóstica: orientación en el solucionario («%s…»)' % d_[0][:30], d_[1] in TS)

# ---------------- E. Duplicados ----------------
def nums(s):
    return set(re.findall(r'\b\d{1,3}(?:\.\d{3})+\b|\b\d{4,}\b', s))


for n in range(1, 22):
    ej = ' '.join(' '.join([b.get('title', '')] + b.get('body', [])) for b in C[n]['cuerpo'] if b['t'] == 'caja' and b['title'].startswith('Ejemplo'))
    ej_n = {x.replace('.', '') for x in nums(ej)}
    p = practicas.PR[n]
    ptxt = ' '.join(a['objetivo'] + ' ' + ' '.join(a['pasos']) + ' ' + ' '.join(a['modelo'][1] if a.get('modelo') else []) for a in p['acts']) + ' ' + p['antes']
    p_n = {x.replace('.', '') for x in nums(ptxt)}
    umbrales = {'50000', '10000', '2000', '4000', '5000', '6000', '3000', '8000'}
    inter = (p_n & ej_n) - umbrales
    chk('E', 'Práctica %d sin datos de los ejemplos resueltos de la Clase %d' % (n, n), not inter, sorted(inter))
ACTS = {n: ' '.join(e for _, _, e, _, _ in actividades.items(n)) for n in range(1, 22)}
COPIAS = {1: ['Ña Rosa arma el menú del día'], 2: ['Ña Rosa separa los fritos F'], 3: ['C = comestibles'], 4: ['Los pedidos del lunes'],
          5: ['30 clientes, 18 chipa'], 6: ['«Hay chipa y hay cocido»'], 7: ['si hay harina y hay queso', 'comprás dos chipas'],
          8: ['¬p∨q e indicá con qué conectivo'], 9: ['(p∧q)→p', '(p∧q)→r'], 10: ['«hay chipa y hay cocido»'], 11: ['«x cuesta G. 500»', '«x es frito»'],
          12: ['«Si es viernes, hay chipa. Hoy no hay chipa.»'], 13: ['«preparar el mostrador del copetín antes de abrir»'],
          14: ['total = 12.000 y pago = 20.000', 'Probalo con precio = 48.000'], 15: ['precio = 4.000 y cantidad = 3'], 16: ['probalo con 4.000, 3 y 20.000', 'usando una variable auxiliar, y hacé la prueba de escritorio con a = 8'],
          17: ['hacé la prueba de escritorio con 60.000, 50.000 y 30.000'], 18: ['(≥ 80.000 mayorista'], 19: ['escriba los números del 1 al 5'], 20: ['suma = 35.000 y c = 3'], 21: ['Con una venta de G. 60.000']}
for n, frases in COPIAS.items():
    for frase in frases:
        chk('E', 'Clase %d: actividades sin copiar el ejemplo («%s»)' % (n, frase[:30]), frase not in ACTS[n])
enuns = [e for n in range(1, 22) for _, _, e, _, _ in actividades.items(n)]
dup = [e for e, c in collections.Counter(enuns).items() if c > 1]
chk('E', 'Sin consignas idénticas repetidas entre clases', not dup, dup[:2])
pr_obj = [a['objetivo'] for n in range(1, 22) for a in practicas.PR[n]['acts']]
chk('E', 'Sin objetivos de práctica repetidos', len(pr_obj) == len(set(pr_obj)))
ev_items = [q for e in evaluaciones.EV for q, _ in e['ab']]
chk('E', 'Evaluaciones sin ítems de resolución duplicados entre sí', len(ev_items) == len(set(ev_items)))
chk('E', 'Evaluaciones sin ítems copiados de las actividades', not (set(ev_items) & set(enuns)))
chk('E', 'Evaluaciones sin los ejemplos de clase (harina y queso · mayorista/frecuente · para llevar o para la mesa)',
    not re.search(r'harina y hay queso|«mayorista»|para llevar o es para la mesa|precio cargado', '\n'.join(q for e in evaluaciones.EV for q, _ in e['ab'])))
ev_mc = [m[0] for e in evaluaciones.EV for m in e['mc']]
chk('E', 'Opción múltiple sin copiar actividades (n({a,l,g,o,r,i,t,m}), A′ de {2,4,6}, 4 elementos → 16, venta de 50.000)',
    not any(x in ' '.join(ev_mc) for x in ('a, l, g, o', 'A = {2, 4, 6}', 'Un conjunto de 4 elementos', 'una venta de G. 50.000 con la regla')))
chk('E', 'Opción múltiple: la respuesta correcta cambia de posición dentro de cada evaluación', all(len({m[2] for m in e['mc']}) == 3 for e in evaluaciones.EV))

# ---------------- F. Figuras ----------------
caps_ = [t for _, t in PL if re.match(r'^Figura \d+\.\d+ — ', t)]
porclase = collections.defaultdict(list)
for c_ in caps_:
    a, b = map(int, re.match(r'Figura (\d+)\.(\d+)', c_).groups())
    porclase[a].append(b)
chk('F', 'Figuras numeradas por clase y correlativas', all(v_ == list(range(1, len(v_) + 1)) for v_ in porclase.values()), dict(porclase))
chk('F', 'Toda clase tiene al menos una figura (antes 10 no tenían)', sorted(porclase) == list(range(1, 22)), sorted(set(range(1, 22)) - set(porclase)))
chk('F', 'Cantidad de figuras de clase: 11 conservadas + 11 nuevas = 22', len(caps_) == 22, len(caps_))
chk('F', 'Figura del Proyecto: PFI.1', 'Figura PFI.1 — ' in TL)
nimg = len(dL.inline_shapes)
chk('F', 'Cada imagen del cuerpo tiene su epígrafe (%d imágenes)' % nimg, nimg == len(caps_) + 1)
chk('F', 'La figura de De Morgan del tomo (src_04) se reemplazó por una tabla redibujada', 'fig_n_demorgan.png' in [b['file'] for b in C[10]['cuerpo'] if b['t'] == 'img'])
RET = set(os.listdir(B + 'figs_ret'))
hashes = {hashlib.sha256(open(B + ('figs_ret/' if f_ in RET else 'figs_src/') + f_, 'rb').read()).hexdigest(): f_ for f_ in os.listdir(B + 'figs_src') if f_ != 'src_04.png'}
chk('F', 'Retoques mínimos solo en 4 imágenes del tomo (✓, marca de agua, «blend», comillas «» en el código)', RET == {'src_01.png', 'src_02.png', 'src_07.png', 'src_12.png'})
blobs = {hashlib.sha256(r.target_part.blob).hexdigest() for r in dL.part.rels.values() if r.reltype.endswith('/image')}
chk('F', 'Las imágenes del tomo se conservan byte a byte (salvo las 4 retocadas, que se usan en su versión retocada)', set(hashes) <= blobs, [hashes[h] for h in set(hashes) - blobs])
sin_lead = []
for n, c in C.items():
    for i, b in enumerate(c['cuerpo']):
        if b['t'] == 'img':
            j = i - 1; anunciada = bool(b.get('lead'))
            while j >= 0 and c['cuerpo'][j]['t'] != 'h2' and not anunciada:
                q = c['cuerpo'][j]
                anunciada = q['t'] == 'p' and re.search(r'figura', q['text'], re.I) is not None
                j -= 1
            if not anunciada:
                sin_lead.append(b['num'])
chk('F', 'Toda figura está anunciada en el texto', not sin_lead, sin_lead)
refs = re.findall(r'Figura (\d+\.\d+)', TL)
chk('F', 'Referencias cruzadas a figuras existentes', all(('Figura %s — ' % r) in TL for r in refs), [r for r in refs if ('Figura %s — ' % r) not in TL][:3])

# ---------------- G. Plan Anual y Planes de Clase ----------------
PLs = planes_data.planes()
dA, PA = textos_doc(ANU); TA = '\n'.join(t for _, t in PA)
dP, PP = textos_doc(PLA); TP = '\n'.join(t for _, t in PP)
col1 = [r.cells[0].text.strip() for r in dA.tables[0].rows]
chk('G', 'Plan Anual: 37 filas de encuentro numeradas 1..37 (148 HC, como el plan vigente)', [x for x in col1 if x.isdigit()] == [str(i) for i in range(1, 38)])
chk('G', 'Plan Anual: 20 encuentros en la 1.ª etapa (E1 es el 20) y 17 en la 2.ª', PLs[19]['tipo'] == 'E' and PLs[19]['n'] == 'E1' and PLs[36]['n'] == 'E2')
chk('G', 'Continuaciones: las 14 clases de 8 HC del plan vigente (2, 4, 5, 7, 8, 9, 12, 14, 16 a 21)', sorted(p['n'] for p in PLs if p['tipo'] == 'P') == [2, 4, 5, 7, 8, 9, 12, 14, 16, 17, 18, 19, 20, 21])
for n in range(1, 22):
    chk('G', 'Plan Anual: indicadores literales de la ficha de la Clase %d' % n, all(('• ' + i) in TA for i in C[n]['ficha']['indicadores']))
chk('G', 'Plan Anual: procedimientos únicos (37)', len({p['proc'] for p in PLs}) == 37)
chk('G', 'Plan Anual: instrumentos únicos (37)', len({p['inst'] for p in PLs}) == 37)
for doc_, nombre in ((TA, 'Plan Anual'), (TP, 'Planes de Clase')):
    chk('G', '%s sin «Cuaderno»/«cuadernillo»/«tomo»' % nombre, not re.search(r'Cuaderno de Prácticas|cuadernillo|Cuadernillo|\btomo\b', doc_))
    chk('G', '%s sin códigos internos' % nombre, not re.search(r'\b\d-\d\d\b|IMG_\d+|TXT_\w+|\{FIG|\bB\d{1,2}\b', doc_))
    chk('G', '%s remite al libro' % nombre, 'libro' in doc_)
for p in PLs:
    t = json.dumps(p, ensure_ascii=False)
    for m in re.finditer(r'Figuras? ((?:\d+\.\d+(?:,? (?:y )?)?)+)', t):
        for r in re.findall(r'\d+\.\d+', m.group(1)):
            chk('G', 'Plan %d: figura %s existe en el libro' % (p['idx'], r), ('Figura %s — ' % r) in TL)
    if p['tipo'] == 'C':
        er, pd, de = planes_data._fam(p['n'])
        for m in re.finditer(r'afirmaciones ([\d, y]+) \(Pensá', t):
            ns = [int(x) for x in re.findall(r'\d+', m.group(1))]
            chk('G', 'Plan %d: afirmaciones %s son del bloque «Pensá y decidí»' % (p['idx'], ns), ns == pd)
    chk('G', 'Plan %d: tiempos suman 160 min' % p['idx'], sum(p['tiempos']) == 160)
chk('G', 'Planes de continuación: el Tema coincide con la fila del Plan Anual (auditoría independiente, hallazgo 4)', all(p['tema'] in TA and p['tema'] in TP for p in PLs if p['tipo'] == 'P'))
chk('G', 'Planes: 37 títulos en el documento', all(('Plan de Clase N.º %d — %s' % (p['idx'], p['titulo'])) in TP for p in PLs))
chk('G', 'Planes: alternativa en papel en los dos planes que usan PSeInt (Clase 21 y su continuación)', all('Alternativa sin equipamiento' in p['recursos'] for p in PLs if p['n'] == 21))
chk('G', 'Planes: evaluación de unidad al cierre de las continuaciones 5, 12 y 21', all(any('evaluación de la Unidad' in x for x in p['momentos']['Cierre']) for p in PLs if p['tipo'] == 'P' and p['n'] in (5, 12, 21)))
chk('G', 'Planes: el Proyecto Final Integrador figura en las continuaciones 17, 19, 20 y 21', all(any('Proyecto Final Integrador' in x for x in p['momentos']['Desarrollo']) for p in PLs if p['tipo'] == 'P' and p['n'] in (17, 19, 20, 21)))
chk('G', 'Plan Anual: capacidad actitudinal de cada unidad en su banda', all(ediciones.MEC[c] in TA for c in ediciones.ACTITUDINAL.values()))

# ---------------- H. Lengua, residuos y cobertura del programa ----------------
tuteo = re.compile(r'^(?:\d+\. )?(Escribe|Calcula|Indica|Explica|Dibuja|Completa|Responde|Elige|Marca|Lee|Observa|Justifica|Define|Nombra|Ordena|Clasifica|Propón|Diseña|Crea|Abre|Guarda|Ejecuta|Prueba|Copia|Busca|Determina|Simboliza|Construye|Niega)\b')
cons = [e for n in range(1, 22) for _, _, e, _, _ in actividades.items(n) if tuteo.match(e)] + [s_ for n in range(1, 22) for a_ in practicas.PR[n]['acts'] for s_ in a_['pasos'] if tuteo.match(s_)] + [q for e in evaluaciones.EV for q, _ in e['ab'] if tuteo.match(q)]
chk('H', 'Voseo en las consignas (sin imperativos en tuteo)', not cons, cons[:3])
for r in ('Cuaderno de Prácticas', 'cuadernillo', 'Cuadernillo', ' tomo ', '{FIG', 'TODO:', 'python-docx', 'Ctrl+E'):
    chk('H', 'Residuo ausente en el libro: «%s»' % r, r not in TL)
PROG = {'definición y determinación de conjuntos': 'por comprensión', 'clasificación de conjuntos': 'unitario', 'tipos de subconjuntos': 'subconjunto propio',
        'representación (Venn)': 'diagrama de Venn', 'unión e intersección': 'intersección', 'diferencia y complemento': 'complemento',
        'problemas con conjuntos': 'control del universo', 'proposición: definición y clasificación': 'molecular', 'conectivos lógicos': 'bicondicional',
        'signos de agrupación': 'paréntesis', 'tablas de verdad': 'tabla de verdad', 'tautología, contradicción, indeterminación': 'indeterminación',
        'equivalencias y leyes': 'De Morgan', 'adjunción y simplificación': 'adjunción', 'funciones proposicionales': 'función proposicional',
        'cuantificadores': 'cuantificador', 'inferencia: MPP, MTT, MTP': 'Modus Tollendo Ponens', 'silogismo hipotético y disyuntivo': 'silogismo disyuntivo',
        'algoritmo: concepto y características': 'finito', 'cualitativos y cuantitativos': 'cuantitativ', 'pseudocódigo': 'pseudocódigo', 'diagrama de flujo': 'diagrama de flujo',
        'expresiones aritméticas, relacionales y lógicas': 'relacional', 'jerarquía de operadores': 'jerarquía', 'enunciado de lectura, asignación y escritura': 'asignación',
        'enunciado de decisión': 'Si', 'decisiones anidadas': 'anidad', 'ciclos repetitivos': 'Mientras', 'contadores y acumuladores': 'acumulador', 'prueba de escritorio': 'prueba de escritorio'}
for k, v_ in PROG.items():
    chk('H', 'Cobertura del programa: ' + k, v_.lower() in TL.lower())
chk('H', 'Clases con al menos 750 palabras de desarrollo (antes las 21 tenían entre 344 y 549)', all(sum(len(' '.join([b.get('text', ''), b.get('title', '')] + b.get('body', []) + [' '.join(r) for r in b.get('rows', [])]).split()) for b in C[n]['cuerpo']) >= 750 for n in C))

# ---------------- I. Lecciones del piloto de 3.º (v2 y v2.1) ----------------
for n, t in enumerate(fichas, 1):
    caps_doc = [x.text.strip().lstrip('• ').strip() for x in t.rows[0].cells[1].paragraphs if x.text.strip()]
    chk('I', 'Clase %d: capacidad textual del programa MEC' % n, caps_doc == ediciones.caps(ediciones.CAP_CLASE[n]), caps_doc)
cog = set(ediciones.MEC) - set(ediciones.ACTITUDINAL.values())
chk('I', 'Las 12 capacidades cognitivas del programa aparecen en alguna ficha', set(c for v_ in ediciones.CAP_CLASE.values() for c in v_) == cog)
chk('I', 'Clase 13: suma «Analiza las características de los algoritmos cualitativos y cuantitativos», que faltaba', ediciones.CAP_CLASE[13] == ['B9', 'B10'])
chk('I', 'Plan Anual: nota sobre capacidades textuales e indicadores', 'Las capacidades se reproducen textualmente del programa MEC vigente. Los indicadores de logro de este plan son operativizaciones observables elaboradas para organizar la enseñanza y la evaluación de cada encuentro.' in TA)
for u, cods in ediciones.CAP_UNIDAD.items():
    chk('I', 'Plan Anual: capacidades textuales de la Unidad %d' % u, all(ediciones.MEC[c] in TA for c in cods))
cap_planes = {x for p in PLs for x in p['capacidad']}
chk('I', 'Planes de Clase: toda capacidad es textual del programa MEC', cap_planes <= set(ediciones.MEC.values()))
chk('I', 'Planes de Clase: las capacidades figuran en el documento', all(c in TP for c in cap_planes))
chk('I', 'PSeInt: perfil Flexible indicado en el libro', 'Flexible' in TL)
chk('I', 'PSeInt: solucionario documenta versión y perfil', 'PSeInt 20250314' in TS and 'perfil Flexible' in TS)
for x in ('versión actual', 'la más utilizada', 'más usada', 'la más usada', 'el más usado'):
    chk('I', 'Sin «%s» en libro ni solucionario' % x, x not in TL and x not in TS)
ctrl = [a['control'] for n in range(1, 22) for a in practicas.PR[n]['acts']] + [practicas.PR[n]['transfer']['control'] for n in (5, 9, 20)]
reveal = [c for c in ctrl if re.search(r'Deberías ver|Debe darte|Deben darte|\b\d{3,}\b|\d+\.\d{3}|= \d', c)]
chk('I', 'Puntos de control sin revelar resultados (sin «Deben darte», cifras ni igualdades)', not reveal, reveal[:3])
chk('I', 'Sin «✅» del cuaderno anterior', '✅' not in TL)
for n in (5, 9, 20):
    tr = practicas.PR[n]['transfer']
    chk('I', 'Práctica %d: transferencia y revisión entre pares en el libro' % n, ('Modelo — Mini-caso: ' + tr['caso'][0]) in TL and all(x in TL for x in tr['caso'][1]))
    chk('I', 'Práctica %d: solución de la transferencia en el solucionario' % n, tr['sol'] in TS)
    chk('I', 'Plan de continuación de la Práctica %d incluye la transferencia' % n, any(p['tipo'] == 'P' and p['n'] == n and any('Transferencia y revisión entre pares' in x for x in p['momentos']['Desarrollo']) for p in PLs))
chk('I', 'Transferencia solo donde hace falta (3 de 14 continuaciones)', sum(1 for n in planes_data.CONT if practicas.PR[n].get('transfer')) == 3)
cajas_py = [b['title'] for n in C for b in C[n]['cuerpo'] if b['t'] == 'caja' and b['title'].startswith('En Paraguay')]
chk('I', 'Recuadros «En Paraguay» solo con contenido paraguayo concreto (1)', cajas_py == ['En Paraguay — conjuntos del país'], cajas_py)
chk('I', 'Sin tramos de IRP ni afirmaciones sobre factura/boleta (T5)', 'IRP' not in TL and not re.search(r'factura (?:o|y) boleta|boleta de venta', TL))
chk('I', '«Aplicación profesional»/«Ejemplo cotidiano» no aparecen en todas las clases', sum(1 for n in C if any(b['t'] == 'caja' and b['title'].startswith(('Aplicación profesional', 'Ejemplo cotidiano')) for b in C[n]['cuerpo'])) < 21)
chk('I', 'PFI: extensión recomendada con la fórmula acordada', PRE.PFI_EXTENSION in TL and 'No son requisitos mínimos para aprobar el proyecto' in PRE.PFI_EXTENSION)
chk('I', 'Rúbrica analítica: 5 criterios con los pesos del paquete vigente 30/25/20/15/10', [p for _, p, _ in PRE.RUBRICA_ANALITICA] == [30, 25, 20, 15, 10])
chk('I', 'Rúbrica analítica: niveles 100/60/30/0 %', [f_ for _, f_ in PRE.NIVELES] == [1.0, 0.6, 0.3, 0.0])
chk('I', 'Rúbrica analítica: 4 descriptores por criterio en el solucionario', all(len(desc) == 4 and all(dsc in TS for dsc in desc) for _, _, desc in PRE.RUBRICA_ANALITICA))
for nombre, path in (('libro', LIB), ('solucionario', SOL), ('planes', PLA), ('plan anual', ANU)):
    texto = '\n'.join(pg.get_text() for pg in pdf(path))
    chk('I', 'Sin símbolos ✓ ✗ ✔ ✅ en el %s' % nombre, not re.search('[✅✔✓✗]', texto))
    chk('I', 'Sin fuente de emoji en el PDF del %s' % nombre, not any('Emoji' in f_[3] for pg in pdf(path) for f_ in pg.get_fonts()))
    chk('I', 'Sin cuadros de glifo faltante en el %s' % nombre, not re.search('[�□]', texto))
    docx_txt = '\n'.join(t for _, t in textos_doc(path)[1])
    for pat, desc in ((r'\{FIG|image\d+\.(?:jpg|png)', 'marcadores de figura'), (r'«\s*»', 'comillas vacías'), (r'[^.]\.\.(?!\.)', 'punto doble'), (r' ,| ;(?! )| \.(?![\w\d.])', 'espacio antes de puntuación'), (r'\b(?!Fin)(\w{3,}) \1\b', 'palabra repetida (salvo cierres anidados FinSi FinSi)')):
        hits = re.findall(pat, docx_txt)
        chk('I', 'Residuos en el %s: %s' % (nombre, desc), not hits, hits[:3])

# ---------------- J. Correcciones técnicas propias de 1.º (diagnóstico T1–T9) ----------------
cod_libro = [l[1:] for n in C for b in C[n]['cuerpo'] if b['t'] == 'caja' for l in b['body'] if l.startswith('§')]
cod_libro += [l[1:] for n in practicas.PR for a in practicas.PR[n]['acts'] if a.get('modelo') for l in a['modelo'][1] if l.startswith('§')]
chk('J', 'T1: ningún número con punto de miles dentro del código (50.000 se lee 50 en PSeInt)', not [l for l in cod_libro if re.search(r'\d\.\d{3}\b', l)], [l for l in cod_libro if re.search(r'\d\.\d{3}\b', l)][:3])
prosa_cod = [t for n in C for b in C[n]['cuerpo'] for t in [b.get('text', '')] + [x for x in b.get('body', []) if not x.startswith('§')] + [' | '.join(r) for r in b.get('rows', [])]
             if re.search(r'←[^.;:«]*\d\.\d{3}|(?:[<>]=?|<>)\s*\d{1,3}\.\d{3}', t) and not t.startswith('Números sin separador') and 'vale 50' not in t and not t.startswith('vuelto ←')]
chk('J', 'T1: tampoco hay expresiones de código con punto de miles escritas en el texto (salvo el ejemplo del error en la Clase 15)', not prosa_cod, [x[:80] for x in prosa_cod])
chk('J', 'T1: la Clase 15 explica por qué 50.000 no es «cincuenta mil» dentro de una expresión', '50.000 vale 50' in TL)
chk('J', 'T2: comillas rectas en el código (sin «» después de Escribir o en comparaciones)', not [l for l in cod_libro if '«' in l])
chk('J', 'T3: sin «≈» en libro ni solucionario (promedios periódicos escritos con su desarrollo)', '≈' not in TL and '≈' not in TS)
chk('J', 'T4: «Puente a PSeInt» es un apartado de la Clase 21, no una clase suelta', any(b['t'] == 'h2' and b['text'].startswith('Puente a PSeInt') for b in C[21]['cuerpo']) and not any(t.startswith('Puente a PSeInt') for s_, t in PL if s_ == 'Heading 2'))
chk('J', 'T5: «En Paraguay» reducido a un recuadro verificable', len(cajas_py) == 1)
elegi = [e for n in range(1, 22) for _, _, e, _, _ in actividades.items(n) if e.startswith('Elegí y justificá')]
chk('J', 'T6: opciones de «Elegí y justificá» con una sola forma («a) X; b) Y; c) Z.»)', len(elegi) == 21 and all(re.search(r' a\) [^;]+; b\) [^;]+(?:; c\) [^;]+)?\.$', e) for e in elegi), [e for e in elegi if not re.search(r' a\) [^;]+; b\) [^;]+(?:; c\) [^;]+)?\.$', e)][:2])
chk('J', 'T8: el solucionario trae las respuestas de las 21 prácticas', all(practicas.PR[n]['sol']['resultado'] in TS for n in range(1, 22)))
chk('J', 'T9: hay Planes de Clase (37)', len(jp['paginas']) == 37)
chk('J', 'Dilema constructivo (p ∨ q, p → r, q → s ⊢ r ∨ s) sin atribuir la nomenclatura al programa', 'r ∨ s' in TL and 'convención del programa' not in TL and 'Dilema constructivo' in TL)
chk('J', 'Clase 14: aclaración diagrama de flujo del algoritmo / DFD', 'diagrama de flujo del algoritmo' in TL and 'Pseudocódigo y DFD' not in TL)
chk('J', 'Sin «favorita», «el error más común» ni «la pregunta favorita»', not re.search(r'favorit|el error más común', TL))
chk('J', 'Puente a PSeInt: comprobación 54.000 y frontera 45.000 en el libro', '54.000' in ' '.join(b.get('text', '') + ' '.join(b.get('body', [])) for b in C[21]['cuerpo']) and '45.000' in TL)
chk('J', 'Recurso «tomo» y «cuadernillo» reemplazados por «libro» en las clases', not re.search(r'\btomo\b|cuadernillo', TL))

# ---------------- salida ----------------
fallos = [r for r in RES if not r[2]]
por = collections.Counter(r[0] for r in RES)
with open(B + 'auditoria_resultado.txt', 'w') as f:
    f.write('AUDITORÍA AUTOMÁTICA — Algorítmica 1.º · Edición Esencial Comercial 2026 · versión v1.1\n')
    f.write('Controles: %d · Fallos: %d\n' % (len(RES), len(fallos)))
    nombres = {'A': 'Estructura y correspondencia', 'B': 'Paginación, índice, saltos y encabezados', 'C': 'Coherencia aritmética y ejecución en PSeInt', 'D': 'Libro ↔ solucionario',
               'E': 'Duplicados', 'F': 'Figuras', 'G': 'Plan Anual y Planes de Clase', 'H': 'Lengua, residuos y cobertura curricular', 'I': 'Lecciones del piloto de 3.º (v2 y v2.1)', 'J': 'Correcciones técnicas de 1.º (T1–T9)'}
    for b in sorted(por):
        f.write('  %s. %-45s %4d controles · %d fallos\n' % (b, nombres[b], por[b], sum(1 for r in fallos if r[0] == b)))
    for r in fallos:
        f.write('FALLO [%s] %s — %s\n' % (r[0], r[1], r[3]))
print(open(B + 'auditoria_resultado.txt').read())
