# -*- coding: utf-8 -*-
"""Plan Anual — Gabinete de Informática – Software · 2.º Curso. 36 secuencias de 2 HC (72 HC), misma numeración que
los Planes de Clase. Temas, capacidades MEC (textuales) e indicadores leídos de las fichas del libro; procedimientos e
instrumentos de la versión V2_4, sin códigos internos y únicos por fila."""
import re, json, sys
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
import maqueta as M
from build_libro import pdf, leer, titulo_de, FICHA_TRAS, OUT, B
sys.path.insert(0, B)
from mec import MEC

NOMBRE = 'Gabinete_Informatica_Software_2do_Curso_PLAN_ANUAL_COMERCIAL_2026'
ANCHOS = [1.0, 6.2, 6.0, 1.3, 5.2, 4.3, 2.6]
COLS = ['N.º', 'TEMAS Y CAPACIDAD MEC', 'INDICADORES DE LOGRO', 'TIEMPO', 'PROCEDIMIENTOS', 'INSTRUMENTOS DE EVALUACIÓN', 'OBSERVACIÓN']
DARK, MED, LIGHT, AZUL, AZUL_CLARO = '1B5E20', '2E7D32', 'E8F5E9', '1B4F72', 'EAF0F6'
ORIG = json.load(open(B + 'pa_orig.json'))
NOTA_CAP = ('Las capacidades se reproducen textualmente del programa del MEC (Diseño Curricular de Informática, mayo 2026). Los indicadores '
            'de logro son operativizaciones observables elaboradas para organizar la enseñanza y la evaluación de cada secuencia.')


def ficha_de(lineas):
    i = lineas.index('[[FICHA]]'); j = lineas.index('[[/FICHA]]')
    blk = lineas[i + 1:j]
    caps = next(x[len('Capacidad: '):] for x in blk if x.startswith('Capacidad: ')).split(' || ')
    tema = next(x[len('Tema: '):] for x in blk if x.startswith('Tema: '))
    inds = [x[2:] for x in blk if x.startswith('- ')]
    return caps, tema, inds


def limpiar(t):
    t = t.replace('del tomo', 'del libro').replace('el tomo', 'el libro').replace('Cuaderno de Laboratorio', 'libro')
    return re.sub(r'\s+', ' ', t).strip()


def secuencias():
    """Devuelve las 36 secuencias en orden, con tipo C (clase), F (ficha), U (evaluación de unidad), E (integradora)."""
    S = []
    for n in range(1, 21):
        S.append(('C', n))
        if n in FICHA_TRAS:
            S.append(('F', FICHA_TRAS[n]))
        if n == 10:
            S += [('U', 1), ('E', 1)]
        if n == 20:
            S += [('U', 2), ('E', 2)]
    assert len(S) == 36
    return S


def filas():
    rows = []
    for idx, (t, k) in enumerate(secuencias(), 1):
        o = ORIG[idx - 1]
        assert int(o[0]) == idx
        if t == 'C':
            L = leer('tomo_clase%02d' % k); caps, tema, inds = ficha_de(L)
            titulo = titulo_de(L); temas = ['%s: %s' % (titulo, tema)]
            obs = 'Clase %d y Práctica %d del libro.' % (k, k)
            unidad = 1 if k <= 10 else 2
            assert 'Clase %d —' % k in o[1]
        elif t == 'F':
            L = leer('lab_ficha_' + k); caps, tema, inds = ficha_de(L)
            temas = [titulo_de(L)]
            obs = 'Ficha %s del libro (continuidad).' % k
            unidad = 1 if k <= 'F' else 2
            assert 'Continuación %s' % k in o[1]
        elif t == 'U':
            unidad = k
            caps = MEC[0:9] if k == 1 else MEC[9:16]
            temas = ['Evaluación de la Unidad %d — %s' % (k, 'Componentes básicos de una computadora' if k == 1 else 'Redes de computadoras')]
            inds = [x.lstrip('• ').strip() for x in o[2].split('\n') if x.strip()]
            obs = 'Cierre de la Unidad %d.' % k
            o = list(o)
            o[4] = ('Aplicación individual de la evaluación de la Unidad 1 del libro, con interpretación de una especificación de equipo y cálculo de costo por página; devolución focalizada en compatibilidad.' if k == 1 else
                    'Aplicación individual de la evaluación de la Unidad 2 del libro, con dimensionamiento de equipos para un local y regla de los 100 metros; devolución focalizada en el diseño de la red.')
        else:
            unidad = k
            caps = None
            temas = ['Evaluación integradora de la %s etapa (Unidad %d)' % ('1.ª' if k == 1 else '2.ª', k)]
            inds = [x.lstrip('• ').strip() for x in o[2].split('\n') if x.strip()]
            obs = 'Cierre de la %s etapa — %s.' % (('1.ª', 'junio') if k == 1 else ('2.ª', 'noviembre'))
        if caps:
            cap_txt = ['Capacidad MEC: «%s»' % c for c in caps]
        else:
            cap_txt = ['Capacidades MEC: las de la Unidad %d, consolidadas en las secuencias %s.' % (k, '1 a 17' if k == 1 else '19 a 35')]
        rows.append(dict(n=idx, tipo=t, unidad=unidad, tema=temas + cap_txt, indicadores=inds, tiempo='2 h',
                         proc=limpiar(o[4]), inst=limpiar(o[5]), obs=obs))
    assert len({r['proc'] for r in rows}) == 36, 'procedimientos repetidos'
    assert len({r['inst'] for r in rows}) == 36, 'instrumentos repetidos'
    for r in rows:
        for campo in ('proc', 'inst', 'obs'):
            assert not re.search(r'\b\d-\d\d\b|IMG_\d+|TXT_\w+|[Tt]omo|Cuaderno', r[campo]), r
    return rows


def celda(cell, texto, bold=False, blanco=False, size=8, italic_desde=None):
    par = cell.paragraphs[0]
    for i, linea in enumerate(texto if isinstance(texto, list) else [texto]):
        if i:
            par = cell.add_paragraph()
        par.paragraph_format.space_after = Pt(1)
        r = par.add_run(linea); r.bold = bold; r.font.size = Pt(size)
        if italic_desde is not None and i >= italic_desde:
            r.italic = True; r.font.color.rgb = RGBColor.from_string('1B4F72')
        if blanco:
            r.font.color.rgb = RGBColor.from_string('FFFFFF')


def construir():
    d = M.nuevo_doc('Gabinete de Informática – Software · 2.º Curso · Plan anual', 'Material del docente: plan anual de 36 secuencias',
                    'BTI; Gabinete de Informática; Software; plan anual')
    s0 = d.sections[0]
    M.imagen_pagina_completa(d, B + 'portadas/portada_plan_anual.jpg')
    sec = d.add_section(WD_SECTION.NEW_PAGE)
    for hf in (s0.footer, s0.header):
        hf.is_linked_to_previous = False
        for p in hf.paragraphs:
            for el in list(p._p):
                if el.tag != qn('w:pPr'):
                    p._p.remove(el)
    M.formato_oficio(sec, horizontal=True)
    M.ANCHO = round(M.PAG_H - 2 * M.MARG, 2)
    M.par(d, 'PLAN ANUAL DETALLADO', 14, True, color=DARK, align='c', despues=2)
    M.par(d, 'Disciplina: Gabinete de Informática – Software  ·  Curso: 2.º Curso  ·  Bachillerato Técnico en Servicios · Especialidad Informática  ·  2 horas cátedra semanales', 9.5, align='c', despues=1)
    M.par(d, 'República del Paraguay · 2026', 9, italic=True, align='c', despues=1)
    M.par(d, 'Institución educativa: ______________________________     Docente: ______________________________', 9, align='c', despues=4)
    pp = M.par(d, '', 8.5, despues=4); M.run(pp, 'Nota sobre capacidades e indicadores. ', True, size=8.5, color=DARK); M.run(pp, NOTA_CAP, size=8.5)
    M.par(d, 'Distribución anual: 36 secuencias de 2 horas cátedra (72 HC), 18 en cada etapa: 20 clases (cada una con su práctica), 12 fichas de continuidad (A a L), 2 evaluaciones de unidad y 2 evaluaciones integradoras de etapa. Cada fila tiene su Plan de Clase con el mismo número. Material de referencia: el libro del estudiante, que reúne cada clase con su práctica y las fichas de continuidad.', 8.5, despues=4)
    rows = filas()
    t = d.add_table(rows=1, cols=7)
    M.bordes(t, '808080', '4'); M.margenes_celda(t, 30, 70); M.ancho_fijo(t, ANCHOS)
    for j, h in enumerate(COLS):
        M.shd(t.rows[0].cells[j], DARK); celda(t.rows[0].cells[j], h, bold=True, blanco=True, size=8)

    def fila():
        r = t.add_row()
        for j, c in enumerate(r.cells):
            c.width = Cm(ANCHOS[j] * M.ANCHO / sum(ANCHOS))
        return r

    def banda(texto, fill, size=9):
        r = fila(); m = r.cells[0].merge(r.cells[6]); M.shd(m, fill); celda(m, texto, bold=True, blanco=True, size=size)

    u_prev = None
    for r_ in rows:
        if r_['n'] == 1:
            banda('1.ª ETAPA · febrero – junio · 18 secuencias de 2 HC', AZUL)
        if r_['n'] == 19:
            banda('2.ª ETAPA · julio – noviembre · 18 secuencias de 2 HC', AZUL)
        if r_['unidad'] != u_prev:
            u_prev = r_['unidad']
            banda('UNIDAD %d — %s' % (u_prev, 'Componentes básicos de una computadora' if u_prev == 1 else 'Redes de computadoras'), MED, 8.5)
        rr = fila()
        ntema = len(r_['tema']) - sum(1 for x in r_['tema'] if x.startswith('Capacidad'))
        vals = [str(r_['n']), r_['tema'], ['• ' + i for i in r_['indicadores']], r_['tiempo'], r_['proc'], r_['inst'], r_['obs']]
        for j, v in enumerate(vals):
            celda(rr.cells[j], v, bold=(j == 0 or (j == 1 and r_['tipo'] in 'UE')), italic_desde=(ntema if j == 1 else None))
        if r_['tipo'] in 'UE':
            for c in rr.cells:
                M.shd(c, AZUL_CLARO)
        else:
            M.shd(rr.cells[0], LIGHT)
    M.cant_split(t, header=True)
    M.par(d, '', despues=4)
    M.par(d, 'Criterios generales de planificación', 10, True, color=DARK, keep=True)
    for x in ['Carga oficial: 2 horas cátedra semanales. Las 36 secuencias de 2 HC (72 HC) se adaptan al calendario efectivo de la institución sin alterar la progresión.',
              "Enfoque constructivista «Aprender Haciendo»: el caso integrador del Servicio Técnico Panambi (equipo y red del Instituto Ko'ẽju) articula las dos unidades hasta el informe técnico final; las prácticas de la Unidad 2 transfieren lo aprendido a otro cliente (la Farmacia Tajy).",
              'La prueba diagnóstica, al inicio del libro, orienta el punto de partida del grupo y la formación de parejas de apoyo para las prácticas.',
              'Las producciones de las 20 prácticas y de las Fichas A a L se reutilizan como piezas del Proyecto Final Integrador; las fichas no agregan capacidades nuevas: consolidan las oficiales.',
              'Los contenidos regulatorios y de mercado se trabajan con fuentes oficiales vigentes (CONATEL, Ley 642/95 de Telecomunicaciones) y precios en guaraníes, con fecha de consulta.']:
        M.par(d, '• ' + x, 8.5, despues=1)
    M.par(d, 'Estrategias de inclusión', 10, True, color=DARK, keep=True, antes=4)
    for x in ['Prácticas con pasos numerados y puntos de control verificables que permiten avanzar a ritmos distintos sin perder el hilo.',
              'Desafíos opcionales en cada práctica para quienes terminan antes, y parejas de apoyo definidas a partir de la prueba diagnóstica.',
              'Modelos completos a reproducir (tablas y textos listos) que reducen la barrera de entrada para estudiantes con menor experiencia digital.',
              'Alternativas impresas provistas por el docente cuando un equipo del laboratorio no dispone de la herramienta requerida.',
              'Adecuaciones razonables de acceso (ubicación, tamaño de letra, apoyos técnicos) según las necesidades documentadas del estudiante.']:
        M.par(d, '• ' + x, 8.5, despues=1)
    M.par(d, '', despues=14)
    M.par(d, '__________________________________                                        __________________________________', 9.5, align='c', despues=0)
    M.par(d, 'Firma del docente                                                                         Dirección — Visto bueno', 9, align='c')
    M.pie_paginas(d, sec)
    M.no_actualizar_campos(d)
    M.purgar_relaciones(d)
    return d, rows


if __name__ == '__main__':
    d, rows = construir()
    out = OUT + NOMBRE + '.docx'; d.save(out); p = pdf(out)
    import pymupdf
    print('Plan anual: páginas', pymupdf.open(p).page_count, '· filas', len(rows))
    json.dump(rows, open(B + 'pa_rows.json', 'w'), ensure_ascii=False, indent=0)
