"""Muestra del solucionario del piloto (Clase 18, Continuación F y Evaluación
de la Unidad 5). A4 vertical, letra de 10,5 pt (criterios de corrección en 10 pt). La numeración es la del tomo
(1 a 12), cada bloque dice en qué página del tomo está y cada respuesta
abierta trae su criterio de corrección. La continuación va en su lugar de la
secuencia, no al final.

Uso: python3 solucionario_piloto.py SALIDA.docx
"""
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, __import__('os').path.dirname(__file__))
from planes_piloto import AZUL, BANDA, GRIS, encabezado_repetido, fijar_anchos, no_partir, pie_de_pagina, sombrear, texto_celda  # noqa: E402

# Páginas del tomo de octubre (PDF oficio de 83 págs.); se actualizan al maquetar.
SECCIONES = [
    {
        'titulo': 'Clase 18 — La hidrografía y la Hidrovía',
        'ref': 'Tomo, págs. 67 a 69 · Encuentro 28',
        'bloques': [
            ('Marcá la opción correcta', [
                ('1', 'b) hidrografía', ''),
                ('2', 'c) Paraguay', 'Se acepta también si el alumno aclara que en parte es límite con Brasil y con Argentina.'),
                ('3', 'b) Itaipú', ''),
                ('4', 'a) Hidrovía Paraguay-Paraná', ''),
            ]),
            ('Completá los espacios', [
                ('5', 'Paraná', ''),
                ('6', 'Yacyretá', ''),
                ('7', 'renovable', 'Se acepta «escaso» o «vital» si el alumno lo fundamenta; la idea buscada es la del Concepto clave: se renueva, pero no es inagotable.'),
            ]),
            ('Verdadero o falso', [
                ('8', 'Falso: el Paraguay es rico en ríos (Paraguay, Paraná, Pilcomayo y muchos otros).', 'La justificación es obligatoria para obtener el punto.'),
                ('9', 'Verdadero.', ''),
                ('10', 'Verdadero.', ''),
            ]),
            ('Trabajá con la fuente', [
                ('1', 'Porque por ella sale gran parte de la producción del país; sin ella, exportar sería mucho más caro.', ''),
                ('2', 'El Paraguay no tiene mar (mediterraneidad); la Hidrovía lo conecta por agua con los puertos del océano.', 'Debe nombrar la falta de litoral marítimo.'),
                ('3', 'Soja, carne y otros productos.', ''),
            ]),
            ('Respondé y fundamentá', [
                ('11', 'Dan agua para el consumo y la producción, pesca, transporte (la Hidrovía) y energía (Itaipú y Yacyretá, sobre el río Paraná).',
                 'Completo: tres usos con un ejemplo del Paraguay. Parcial: usos sin ejemplos.'),
                ('12', 'Porque el agua se renueva, pero no es inagotable: la contaminación y la basura afectan el consumo, la pesca y la salud. Cuidarla es tarea del Estado, de las empresas y de cada persona.',
                 'Completo: un riesgo concreto y quién debe actuar. Parcial: solo «porque es importante».'),
            ]),
            ('Ubicate', [
                ('—', 'Río Paraguay de norte a sur, separando la Región Oriental de la Occidental; Asunción sobre su margen izquierda (oriental).',
                 'Se valora, sin exigirlo, que marque los tramos en que el río es límite con Brasil (norte) y con Argentina (sur).'),
            ]),
        ],
    },
    {
        'titulo': 'Continuación F — Posición, ríos e Hidrovía: beneficios e impactos',
        'ref': 'Tomo, pág. 70 · Encuentro 29',
        'bloques': [
            ('Caso para decidir', [
                ('1', 'Un beneficio económico: barcazas con más carga, menor costo de transporte, exportaciones más estables en épocas de aguas bajas.', ''),
                ('2', 'Un aspecto social para consultar o monitorear: la pesca y el trabajo de las comunidades ribereñas, la seguridad de la navegación, el uso del río por la población.', ''),
                ('3', 'Un posible impacto ecológico: cambios en el fondo del río y en los hábitats, sedimentos en suspensión, efectos en la calidad del agua. Medida: estudio previo, monitoreo del agua, elegir la época, consultar a quienes usan el río.',
                 'El impacto y la medida tienen que corresponderse.'),
            ]),
            ('Producto: matriz y conclusión', [
                ('—', 'Matriz Económico / Social / Ecológico con al menos un beneficio o un impacto por columna, y conclusión de 6 a 8 líneas que tome posición y proponga una medida.',
                 'Usar la mini-rúbrica de las continuaciones (8 puntos). No se califica la posición tomada, sino que esté fundamentada.'),
            ]),
        ],
    },
    {
        'titulo': 'Evaluación de la Unidad 5',
        'ref': 'Tomo, págs. 70 y 71 · Encuentro 30 · Total 18 puntos (Partes A y B: 12; Parte C complementaria: 6)',
        'bloques': [
            ('Parte A — Opción correcta (1 punto cada una)', [
                ('1', 'c) Argentina', ''),
                ('2', 'b) mediterráneo', ''),
                ('3', 'b) llano', ''),
                ('4', 'b) la Hidrovía Paraguay-Paraná', ''),
            ]),
            ('Parte B — Completá y respondé (2 puntos cada una)', [
                ('5', 'La posición absoluta se da con coordenadas (latitud y longitud) y no cambia. La posición relativa se describe respecto de otros lugares: por ejemplo, el Paraguay está en el centro de América del Sur y limita con Argentina, Bolivia y Brasil.',
                 '1 punto por cada definición con su ejemplo.'),
                ('6', 'Paraguay.', 'Se acepta la aclaración de que en parte es límite (ver Clase 18).'),
                ('7', 'Porque el Paraguay no tiene salida al mar: el transporte por agua es la forma más barata de llegar a los puertos del océano.',
                 '1 punto por nombrar la mediterraneidad; 1 por explicar el costo o la salida al océano.'),
                ('8', 'El relieve llano facilita la agricultura mecanizada, la ganadería y la construcción de caminos; pero no lo explica todo: también influyen el suelo, el clima y los mercados.',
                 '2 puntos si relaciona el relieve con una actividad y nombra otro factor; 1 punto si solo nombra actividades.'),
            ]),
            ('Parte C — Complementaria (la agrega el docente; ver Plan 30)', [
                ('9', 'Río Paraguay, río Paraná, Itaipú (sobre el Paraná, en el límite con Brasil) y Asunción.', '1 punto por elemento bien ubicado.'),
                ('10', 'Beneficio: exportar a menor costo (económico). Impacto: por ejemplo, efectos de una obra sobre la pesca (social) o sobre los hábitats (ecológico).', '1 punto por fila completa.'),
            ]),
        ],
    },
]


def main():
    salida = sys.argv[1]
    doc = docx.Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0)
    s.top_margin = s.bottom_margin = Cm(1.8)
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    pie_de_pagina(s, 'Solucionario docente')
    p = doc.add_paragraph(); r = p.add_run('Solucionario docente — muestra del piloto'); r.bold = True
    r.font.size = Pt(16); r.font.color.rgb = AZUL
    doc.add_paragraph('Numeración igual a la del tomo. Cada respuesta abierta trae su criterio de corrección.')
    for k, sec in enumerate(SECCIONES):
        p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(14 if k else 6)
        r = p.add_run(sec['titulo']); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = AZUL
        p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
        r = p.add_run(sec['ref']); r.italic = True; r.font.size = Pt(9.5)
        for nombre, items in sec['bloques']:
            t = doc.add_table(rows=1, cols=3); t.style = 'Table Grid'
            enc = t.rows[0]; encabezado_repetido(enc); no_partir(enc)
            m = enc.cells[0].merge(enc.cells[2]); sombrear(m, BANDA)
            texto_celda(m, nombre, negrita=True, color=RGBColor(0xFF, 0xFF, 0xFF), tam=10.5)
            for n, resp, crit in items:
                f = t.add_row(); no_partir(f)
                sombrear(f.cells[0], GRIS)
                texto_celda(f.cells[0], n, negrita=True, tam=10.5, alinear=WD_ALIGN_PARAGRAPH.CENTER)
                texto_celda(f.cells[1], resp, tam=10.5)
                texto_celda(f.cells[2], crit, tam=10)
            fijar_anchos(t, [1.0, 9.6, 6.4])
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
    doc.save(salida)


if __name__ == '__main__':
    main()
