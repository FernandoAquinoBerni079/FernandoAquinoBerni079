"""Planes de clase del piloto: A4 vertical, márgenes de 2 cm a los lados,
cuerpo de 10,5 pt y tablas de 10 pt. Portada e índice en hojas propias; el
índice lleva los números de página medidos en una primera pasada. Cada plan
empieza en hoja nueva y puede ocupar más de una.

Uso: python3 planes_piloto.py SALIDA.docx [paginas.json]
"""
import json
import sys

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, __import__('os').path.dirname(__file__))
from planes_contenido import COMUNES, PLANES  # noqa: E402

AZUL = RGBColor(0x1F, 0x3A, 0x5F)
GRIS = 'E8EDF3'
BANDA = '1F3A5F'
ANCHO = 17.0  # 21,0 − 2 × 2,0


def sombrear(celda, color):
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), color)
    tcpr.append(shd)


def fijar_anchos(tabla, anchos_cm):
    """Anchos fijos de verdad: tblLayout fixed + tblGrid + tcW en cada celda."""
    tbl = tabla._tbl
    tblpr = tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblpr.append(lay)
    tw = tblpr.find(qn('w:tblW'))
    if tw is None:
        tw = OxmlElement('w:tblW'); tblpr.append(tw)
    tw.set(qn('w:type'), 'dxa'); tw.set(qn('w:w'), str(int(sum(anchos_cm) / 2.54 * 1440)))
    grid = tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(int(anchos_cm[i] / 2.54 * 1440)))
    for fila in tabla.rows:
        for i, c in enumerate(fila.cells):
            if i < len(anchos_cm):
                c.width = Cm(anchos_cm[i])


def no_partir(fila):
    trpr = fila._tr.get_or_add_trPr()
    trpr.append(OxmlElement('w:cantSplit'))


def encabezado_repetido(fila):
    trpr = fila._tr.get_or_add_trPr()
    trpr.append(OxmlElement('w:tblHeader'))


def texto_celda(celda, texto, negrita=False, tam=10, color=None, alinear=None):
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(texto)
    r.bold = negrita; r.font.size = Pt(tam)
    if color:
        r.font.color.rgb = color
    if alinear:
        p.alignment = alinear
    return p


def vinetas_celda(celda, items, tam=10):
    primero = True
    for it in items:
        p = celda.paragraphs[0] if primero else celda.add_paragraph()
        primero = False
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Cm(0.35)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        r = p.add_run('• ' + it); r.font.size = Pt(tam)


def titulo_seccion(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(texto); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = AZUL
    return p


def tabla_clave_valor(doc, filas, anchos=(4.0, 13.0)):
    t = doc.add_table(rows=0, cols=2); t.style = 'Table Grid'
    for k, v in filas:
        f = t.add_row(); no_partir(f)
        sombrear(f.cells[0], GRIS)
        texto_celda(f.cells[0], k, negrita=True)
        if isinstance(v, list):
            vinetas_celda(f.cells[1], v)
        else:
            texto_celda(f.cells[1], v)
    fijar_anchos(t, list(anchos))
    return t


def portada(doc):
    for _ in range(6):
        doc.add_paragraph()
    for texto, tam, neg in [('PLANES DE CLASE', 26, True), ('Historia y Geografía · 7.º grado', 18, True),
                            ('Educación Escolar Básica · Tercer Ciclo', 13, False),
                            ('Muestra del piloto de diseño · planes 1, 28, 29, 30 y 31', 12, False)]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(texto); r.font.size = Pt(tam); r.bold = neg; r.font.color.rgb = AZUL
    for _ in range(3):
        doc.add_paragraph()
    for texto in ['Carga horaria: 3 horas cátedra por semana · 36 encuentros de 120 minutos',
                  'Institución: ______________________________________',
                  'Docente: ______________________________________',
                  'Año lectivo: ____________']:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12)
        p.add_run(texto).font.size = Pt(12)


def indice(doc, paginas):
    p = doc.add_paragraph(); p.paragraph_format.page_break_before = True
    r = p.add_run('Índice'); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = AZUL
    p.paragraph_format.space_after = Pt(12)
    for pl in PLANES:
        q = doc.add_paragraph()
        q.paragraph_format.tab_stops.add_tab_stop(Cm(ANCHO), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        q.paragraph_format.space_after = Pt(6)
        clave = f"Plan de Clase N.º {pl['n']}"
        q.add_run(f"{clave} — {pl['titulo']}\t{paginas.get(clave, '')}").font.size = Pt(11)


def plan(doc, pl):
    p = doc.add_paragraph(); p.paragraph_format.page_break_before = True
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.keep_with_next = True
    r = p.add_run(f"Plan de Clase N.º {pl['n']}"); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = AZUL
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.keep_with_next = True
    r = p.add_run(pl['titulo']); r.bold = True; r.font.size = Pt(12.5)

    t = doc.add_table(rows=4, cols=4); t.style = 'Table Grid'
    datos = [('Institución', '', 'Docente', ''),
             ('Disciplina', COMUNES['disciplina'], 'Grado y sección', '7.º  ____'),
             ('Unidad', pl['unidad'], 'Encuentro', f"{pl['encuentro']} de 36"),
             ('Material', pl['material'], 'Duración', COMUNES['duracion'] + '\nFecha: ___ / ___ / ______')]
    for i, fila in enumerate(datos):
        no_partir(t.rows[i])
        for j, v in enumerate(fila):
            c = t.cell(i, j)
            if j % 2 == 0:
                sombrear(c, GRIS); texto_celda(c, v, negrita=True)
            else:
                texto_celda(c, v)
    fijar_anchos(t, [2.6, 6.4, 3.0, 5.0])

    titulo_seccion(doc, 'Capacidad, tema e indicadores')
    tabla_clave_valor(doc, [('Capacidad', pl['capacidad']), ('Tema', pl['tema']),
                            ('Indicadores de logro', pl['indicadores']), ('Propósito del encuentro', pl['proposito'])])

    titulo_seccion(doc, 'Secuencia didáctica (120 minutos)')
    t = doc.add_table(rows=1, cols=4); t.style = 'Table Grid'
    enc = t.rows[0]; encabezado_repetido(enc); no_partir(enc)
    for j, h in enumerate(['Minutos', 'Qué hace el docente', 'Qué hacen los estudiantes', 'Producto o evidencia']):
        sombrear(enc.cells[j], BANDA)
        texto_celda(enc.cells[j], h, negrita=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for momento, pasos in pl['secuencia']:
        f = t.add_row(); no_partir(f)
        m = f.cells[0].merge(f.cells[3]); sombrear(m, GRIS)
        ini, fin = pasos[0][0].split('–')[0], pasos[-1][0].split('–')[1]
        texto_celda(m, f"{momento} ({int(fin) - int(ini)} min)", negrita=True)
        for minutos, docente, alumnos, producto in pasos:
            f = t.add_row(); no_partir(f)
            texto_celda(f.cells[0], minutos, alinear=WD_ALIGN_PARAGRAPH.CENTER)
            texto_celda(f.cells[1], docente)
            texto_celda(f.cells[2], alumnos)
            texto_celda(f.cells[3], producto)
    fijar_anchos(t, [1.7, 6.3, 5.3, 3.7])

    titulo_seccion(doc, 'Atención a la diversidad')
    for it in pl['diversidad']:
        q = doc.add_paragraph(style='List Bullet'); q.add_run(it)

    titulo_seccion(doc, 'Evaluación')
    ev = pl['evaluacion']
    tabla_clave_valor(doc, [('Tipo', ev['tipo']), ('Procedimientos', ev['procedimientos']),
                            ('Instrumentos', ev['instrumentos']), ('Criterios', ev['criterios'])])

    titulo_seccion(doc, 'Precisiones para el docente')
    for it in pl['precisiones']:
        q = doc.add_paragraph(style='List Bullet'); q.add_run(it)

    titulo_seccion(doc, 'Tarea')
    doc.add_paragraph(pl['tarea'])

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10)
    p.add_run('Observaciones: ' + '_' * 78)
    t = doc.add_table(rows=1, cols=2)
    for j, txt in enumerate(['____________________________\nDocente', '____________________________\nDirección — Visto bueno']):
        texto_celda(t.cell(0, j), txt, alinear=WD_ALIGN_PARAGRAPH.CENTER)
    no_partir(t.rows[0])
    fijar_anchos(t, [8.5, 8.5])


def pie_de_pagina(seccion, rotulo='Planes de clase'):
    p = seccion.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f'Historia y Geografía 7.º · {rotulo} · Página ')
    r.font.size = Pt(9)
    for instr in ['PAGE']:
        r1 = p.add_run(); f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin'); r1._r.append(f1)
        r2 = p.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr; r2._r.append(it)
        r3 = p.add_run(); f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end'); r3._r.append(f3)
        for rr in (r1, r2, r3):
            rr.font.size = Pt(9)


def main():
    salida = sys.argv[1]
    paginas = json.load(open(sys.argv[2])) if len(sys.argv) > 2 else {}
    doc = docx.Document()
    s = doc.sections[0]
    s.orientation = WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0)
    s.top_margin = s.bottom_margin = Cm(1.8)
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    st.paragraph_format.space_after = Pt(3)
    doc.styles['List Bullet'].font.size = Pt(10.5)
    pie_de_pagina(s)
    portada(doc)
    indice(doc, paginas)
    for pl in PLANES:
        plan(doc, pl)
    doc.save(salida)


if __name__ == '__main__':
    main()
