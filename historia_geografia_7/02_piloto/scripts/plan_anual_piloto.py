"""Plan anual del piloto: misma tabla y mismo texto que el original, con otra
maqueta. Columnas proporcionales a su contenido (en el original eran 7 iguales
de 3,82 cm), letra de la tabla más grande, encabezado repetido en cada hoja y
filas que no se parten.

Uso: python3 plan_anual_piloto.py ORIGINAL.docx {a4|oficio} SALIDA.docx
"""
import sys

import docx
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

# N.º, Temas, Indicadores, Tiempo, Procedimientos, Instrumentos, Observación
PROPORCION = [0.9, 5.0, 6.3, 1.7, 5.2, 4.2, 3.4]   # suma 26,7 = ancho útil A4 horizontal

VARIANTES = {
    'a4':     dict(papel=(29.7, 21.0), margen=1.5, letra=9.5),
    'oficio': dict(papel=(33.02, 21.59), margen=1.5, letra=10.0),
}


def main():
    orig, var, salida = sys.argv[1:4]
    v = VARIANTES[var]
    doc = docx.Document(orig)
    s = doc.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(v['papel'][0]), Cm(v['papel'][1])
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(v['margen'])
    util = v['papel'][0] - 2 * v['margen']
    anchos = [x * util / sum(PROPORCION) for x in PROPORCION]

    t = doc.tables[0]
    tbl = t._tbl
    tw = tbl.tblPr.find(qn('w:tblW'))
    tw.set(qn('w:type'), 'dxa'); tw.set(qn('w:w'), str(int(util / 2.54 * 1440)))
    cols = tbl.tblGrid.findall(qn('w:gridCol'))
    for c, a in zip(cols, anchos):
        c.set(qn('w:w'), str(int(a / 2.54 * 1440)))
    for fila in t.rows:
        trpr = fila._tr.get_or_add_trPr()
        if trpr.find(qn('w:cantSplit')) is None:
            trpr.append(OxmlElement('w:cantSplit'))
        # ancho de cada celda según las columnas que abarca
        col = 0
        for tc in fila._tr.findall(qn('w:tc')):
            tcpr = tc.get_or_add_tcPr()
            span = tcpr.find(qn('w:gridSpan'))
            n = int(span.get(qn('w:val'))) if span is not None else 1
            w = tcpr.find(qn('w:tcW'))
            if w is None:
                w = OxmlElement('w:tcW'); tcpr.insert(0, w)
            w.set(qn('w:type'), 'dxa'); w.set(qn('w:w'), str(int(sum(anchos[col:col + n]) / 2.54 * 1440)))
            col += n
    enc = t.rows[0]._tr.get_or_add_trPr()
    if enc.find(qn('w:tblHeader')) is None:
        enc.append(OxmlElement('w:tblHeader'))
    for fila in t.rows:
        for c in fila.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    if r.font.size is None or r.font.size.pt < v['letra']:
                        r.font.size = Pt(v['letra'])
    doc.save(salida)


if __name__ == '__main__':
    main()
