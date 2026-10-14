"""Plan anual en oficio horizontal, segunda versión del piloto.

Parte de `02_piloto/scripts/plan_anual_piloto.py` (variante oficio: mismo
texto, columnas proporcionales, encabezado repetido, filas que no se parten)
y corrige los saltos que se vieron en el PDF del primer piloto:
- las bandas de etapa y de unidad quedaban solas al pie de una hoja: ahora
  quedan unidas a la fila siguiente;
- el bloque final (criterios, inclusión y firma) se partía: ahora va junto.

Uso: python3 plan_anual_oficio_v2.py PLAN_ANUAL_ORIGINAL.docx SALIDA.docx
"""
import os
import subprocess
import sys
import tempfile

import docx
from docx.oxml.ns import qn

AQUI = os.path.dirname(os.path.abspath(__file__))
PILOTO = os.path.join(AQUI, '..', '..', '02_piloto', 'scripts', 'plan_anual_piloto.py')


def main():
    orig, salida = sys.argv[1:3]
    tmp = os.path.join(tempfile.mkdtemp(), 'base.docx')
    subprocess.run([sys.executable, PILOTO, orig, 'oficio', tmp], check=True)
    doc = docx.Document(tmp)
    t = doc.tables[0]
    for fila in t.rows[1:]:
        celdas = fila._tr.findall(qn('w:tc'))
        unica = len(celdas) == 1 or any(tc.find(qn('w:tcPr') + '/' + qn('w:gridSpan')) is not None and
                                        int(tc.find(qn('w:tcPr') + '/' + qn('w:gridSpan')).get(qn('w:val'))) >= 5
                                        for tc in celdas)
        if unica:   # banda de etapa o de unidad
            for p in fila.cells[0].paragraphs:
                p.paragraph_format.keep_with_next = True
    # bloque final: todos los párrafos después de la tabla quedan juntos
    despues = False
    parrafos = []
    for e in doc.element.body:
        if e is t._tbl:
            despues = True
            continue
        if despues and e.tag == qn('w:p'):
            parrafos.append(docx.text.paragraph.Paragraph(e, doc))
    if parrafos:
        parrafos[0].paragraph_format.page_break_before = False
        for p in parrafos[:-1]:
            p.paragraph_format.keep_with_next = True
    doc.save(salida)


if __name__ == '__main__':
    main()
