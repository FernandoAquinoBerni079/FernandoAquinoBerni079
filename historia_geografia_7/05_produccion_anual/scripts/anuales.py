"""Documentos anuales consolidados que no salen de `construir_bloque.py anual`:
plan anual (oficio horizontal) y matriz curricular (A4 horizontal), más los
controles de consolidación de los 36 encuentros.

Todo se lee de los mismos módulos `contenido/eNN.py` y de `curriculo.py` que
generan el tomo, los planes y el solucionario: no hay texto copiado a mano ni
se usa la planificación del maestro (condición del cierre: el encuentro 22 se
alinea con el contenido efectivamente desarrollado, sin la antigua
planificación centrada en la Guerra Guaranítica).

Uso: python3 anuales.py   (después de `construir_bloque.py anual` y `verificar.py anual`)
"""
import collections
import importlib
import json
import os
import re
import sys

import docx
import pymupdf
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
sys.path[:0] = [AQUI, os.path.join(AQUI, 'contenido'), os.path.join(AQUI, 'motor'),
                os.path.join(AQUI, '..', '..', '02_piloto', 'scripts')]
import curriculo as C  # noqa: E402
import pdf as PDF  # noqa: E402
import planes_piloto as PP  # noqa: E402

ANUAL = os.path.join(RAIZ, 'anual')
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
TIPO = {'clase': 'Clase', 'continuacion': 'Continuación', 'evaluacion': 'Evaluación de unidad', 'integradora': 'Integradora'}
VINCULO_CORTO = {C.LITERAL: 'Literal', C.ASOCIACION: 'Asociación curricular', C.EDITORIAL: 'Formulación editorial'}
ETAPAS = {1: 'PRIMERA ETAPA — Unidades 1 a 3 (encuentros 1 a 18)', 19: 'SEGUNDA ETAPA — Unidades 4 a 6 (encuentros 19 a 36)'}


def modulo(n):
    return importlib.import_module(f'e{n:02d}')


def indicadores(n):
    m = modulo(n)
    return list(getattr(m, 'INDICADORES', None) or m.PLAN['indicadores'])


def tema(n):
    e = C.encuentro(n)
    if e['temas']:
        return ' '.join(C.TEMAS[t] for t in e['temas'])
    return modulo(n).PLAN['tema']


def capacidad(n):
    e = C.encuentro(n)
    if e['caps']:
        return '; '.join(f'{c}: {C.CAPACIDADES[c]}' for c in e['caps'])
    return 'Formulación editorial (clase introductoria): ' + modulo(n).PLAN.get('capacidad', '')


def observacion(n):
    e = C.encuentro(n)
    if n == 22:
        return ('Núcleo del siglo XVII: bandeirantes, pérdida del Guairá, éxodo de 1631 y Mbororé (1641). Selección editorial '
                'dentro de los temas H5c y H5d. Ampliación optativa, no evaluada: Guerra Guaranítica (1754-1756, siglo XVIII).')
    if e['tipo'] == 'continuacion':
        return 'Profundiza ' + ', '.join(e['caps']) + ' (temas del programa).'
    if e['tipo'] in ('evaluacion', 'integradora'):
        sig = n + 1
        dev = (f'Devolución con claves en los primeros 20 minutos del encuentro {sig}.' if n < 36 else
               'Devolución en el cierre que fije la institución.')
        return f'Puntaje: {modulo(n).SOL["puntaje"]}. ' + dev
    if n in (6, 10, 19, 26, 31):
        prev = {6: 'U1', 10: 'U2', 19: 'Integradora 1', 26: 'U4', 31: 'U5'}[n]
        return f'Primeros 20 minutos: devolución de la evaluación {prev}.'
    return ''


# ---------------------------------------------------------------------------- documentos
def documento(ancho_cm, alto_cm, margen=1.5):
    d = docx.Document()
    s = d.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(ancho_cm), Cm(alto_cm)
    for lado in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(s, lado, Cm(margen))
    st = d.styles['Normal']; st.font.name = 'Liberation Sans'; st.font.size = Pt(9)
    return d


def titulo(d, texto, sub):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texto); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = PP.AZUL
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(8)
    r = p.add_run(sub); r.font.size = Pt(9.5)


def banda(t, texto, color, ncol, tam=9.5):
    f = t.add_row(); PP.no_partir(f)
    m = f.cells[0].merge(f.cells[ncol - 1]); PP.sombrear(m, color)
    PP.texto_celda(m, texto, negrita=True, color=BLANCO if color == PP.BANDA else None, tam=tam)
    for q in m.paragraphs:
        q.paragraph_format.keep_with_next = True


def celda_lista(c, items, tam):
    p = c.paragraphs[0]
    for k, it in enumerate(items):
        q = p if k == 0 else c.add_paragraph()
        q.paragraph_format.space_after = Pt(0)
        q.add_run('• ' + it).font.size = Pt(tam)


def plan_anual(salida):
    d = documento(33.0, 21.6, 1.4)   # oficio horizontal
    PP.pie_de_pagina(d.sections[0], 'Plan anual')
    titulo(d, 'PLAN ANUAL — Historia y Geografía · 7.º grado (EEB, Tercer Ciclo)',
           'Programa de Estudio MEC (2014), capacidades literales con su código (págs. 55 a 57) · 3 horas cátedra semanales de '
           '40 minutos · 36 encuentros de 120 minutos · Institución: ____________________ · Docente: ____________________ · '
           'Año lectivo: ________')
    cab = ['N.º', 'Sección', 'Capacidad (código y vínculo)', 'Tema', 'Indicadores de logro', 'Procedimientos e instrumentos',
           'Observaciones']
    anchos = [1.0, 3.3, 5.6, 5.0, 6.2, 4.4, 4.7]
    t = d.add_table(rows=1, cols=len(cab)); t.style = 'Table Grid'
    PP.encabezado_repetido(t.rows[0]); PP.no_partir(t.rows[0])
    for j, h in enumerate(cab):
        PP.sombrear(t.rows[0].cells[j], PP.BANDA); PP.texto_celda(t.rows[0].cells[j], h, negrita=True, color=BLANCO, tam=8.5)
    u_prev = None
    for n in range(1, 37):
        e = C.encuentro(n); m = modulo(n)
        if n in ETAPAS:
            banda(t, ETAPAS[n], PP.BANDA, len(cab), 9.5)
        if e['unidad'] and e['unidad'] != u_prev:
            u = e['unidad']; caps = sorted({c for x in C.ENCUENTROS if x[3] == u for c in x[4]})
            banda(t, f'{C.UNIDADES[u].upper()} · Capacidades: ' + ', '.join(caps), PP.GRIS, len(cab), 9)
            u_prev = u
        f = t.add_row(); PP.no_partir(f)
        cap = (('; '.join(e['caps']) if e['caps'] else 'Editorial') + ' — ' + VINCULO_CORTO.get(e['vinculo'], e['vinculo']))
        if e['caps']:
            cap += '. ' + ' '.join(C.CAPACIDADES[c] for c in e['caps'])
        ev = m.PLAN['evaluacion']
        valores = [str(n), m.TITULO, cap, tema(n), None, f"{ev['procedimientos']} Instrumentos: {ev['instrumentos']}", observacion(n)]
        for j, v in enumerate(valores):
            if j == 4:
                celda_lista(f.cells[j], indicadores(n), 7.8)
            else:
                PP.texto_celda(f.cells[j], v, tam=7.8, negrita=(j in (0, 1)))
    PP.fijar_anchos(t, anchos)
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    p.add_run('Criterios generales de evaluación: los criterios de cada actividad abierta y de cada prueba se anuncian en el tomo '
              'y se repiten sin agregados en el solucionario. Equidad: tiempo adicional supervisado, apoyo excepcional en horario '
              'acordado con la dirección y devolución con claves solo cuando todas las pruebas están entregadas. Atención a la '
              'diversidad: ver cada plan de clase.').font.size = Pt(8.5)
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(16)
    p.add_run('______________________________            ______________________________\n'
              '            Docente                                                     Dirección — Visto bueno').font.size = Pt(9)
    d.save(salida)
    return PDF.render(salida)


def matriz(salida, paginas):
    d = documento(29.7, 21.0, 1.3)   # A4 horizontal
    PP.pie_de_pagina(d.sections[0], 'Matriz curricular')
    titulo(d, 'MATRIZ CURRICULAR — 36 encuentros', 'Historia y Geografía 7.º · Programa MEC (2014) · Tipo de vínculo: Literal '
           '(capacidad y temas del programa), Asociación curricular (evaluaciones que integran capacidades literales), Formulación '
           'editorial (Clase 1). Generada desde los módulos de contenido de la producción anual.')
    cab = ['N.º', 'Tipo', 'Sección', 'Unidad', 'Capacidades', 'Temas MEC', 'Tipo de vínculo', 'Tomo (págs.)', 'Observaciones']
    anchos = [0.9, 2.0, 4.4, 1.5, 2.2, 1.9, 2.4, 1.6, 10.2]
    t = d.add_table(rows=1, cols=len(cab)); t.style = 'Table Grid'
    PP.encabezado_repetido(t.rows[0]); PP.no_partir(t.rows[0])
    for j, h in enumerate(cab):
        PP.sombrear(t.rows[0].cells[j], PP.BANDA); PP.texto_celda(t.rows[0].cells[j], h, negrita=True, color=BLANCO, tam=8.5)
    for n in range(1, 37):
        e = C.encuentro(n); m = modulo(n)
        f = t.add_row(); PP.no_partir(f)
        pg = paginas.get(n)
        valores = [str(n), TIPO[e['tipo']], m.TITULO, str(e['unidad'] or '1-3' if n == 18 else e['unidad'] or '4-6'),
                   ', '.join(e['caps']) or '—', ', '.join(e['temas']) or '—', VINCULO_CORTO.get(e['vinculo'], e['vinculo']),
                   f'{pg[0]}-{pg[1]}' if pg else '—', observacion(n)]
        for j, v in enumerate(valores):
            PP.texto_celda(f.cells[j], v, tam=7.8, negrita=(j == 0))
    PP.fijar_anchos(t, anchos)
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(6)
    r = p.add_run('Capacidades del programa (texto literal): ' +
                  ' '.join(f'{c}: {C.CAPACIDADES[c]} (pág. {C.PAGINA_CAPACIDAD[c]}).' for c in C.CAPACIDADES))
    r.font.size = Pt(7.8)
    d.save(salida)
    return PDF.render(salida)


# ---------------------------------------------------------------------------- controles de consolidación
def textos(bl):
    for b in bl:
        for x in b[1:]:
            if isinstance(x, str):
                yield x
            elif isinstance(x, (list, tuple)):
                for y in x:
                    if isinstance(y, str):
                        yield y
                    elif isinstance(y, (list, tuple)):
                        yield from (z for z in y if isinstance(z, str))


def controles():
    res = collections.OrderedDict()

    def control(nombre, ok, detalle):
        r = res.setdefault(nombre, {'correctos': 0, 'fallos': []})
        if ok:
            r['correctos'] += 1
        else:
            r['fallos'].append(detalle)

    tipos = collections.Counter(C.encuentro(n)['tipo'] for n in range(1, 37))
    control('A. Arquitectura', tipos == {'clase': 21, 'continuacion': 7, 'evaluacion': 6, 'integradora': 2}, f'tipos {dict(tipos)}')
    for n in range(1, 37):
        sec = modulo(n).PLAN['secuencia']
        total = sum(1 for _, pasos in sec for _ in pasos)
        control('A. Arquitectura', total > 0 and sec[-1][1][-1][0].endswith('120'), f'E{n}: la secuencia no termina en 120')

    # B. continuidad de evaluaciones: devolución en 0-20 del encuentro siguiente; equidad
    for n in range(1, 37):
        e = C.encuentro(n)
        if e['tipo'] not in ('evaluacion', 'integradora'):
            continue
        pl = modulo(n).PLAN
        txt = ' '.join(x for _, pasos in pl['secuencia'] for paso in pasos for x in paso) + ' '.join(pl['diversidad'])
        control('B. Continuidad de evaluaciones', 'sin claves' in txt and re.search('tiempo adicional', txt, re.I),
                f'E{n}: falta tiempo adicional o actividad sin claves')
        if n < 36:
            primero = modulo(n + 1).PLAN['secuencia'][0][1][0]
            control('B. Continuidad de evaluaciones', primero[0] == '0–20' and 'Devolución' in primero[1],
                    f'E{n + 1}: no reserva 0-20 para la devolución de E{n}')

    # C. repeticiones de ítems entre secciones (enunciados de opción múltiple, V/F y completar)
    vistos = {}
    for n in range(1, 37):
        for b in modulo(n).TOMO:
            if b[0] == 'item':
                k = re.sub(r'^\d+\.\s*(\(\.\.\.\.\)\s*)?', '', b[1]).strip().lower()
                if len(k) > 25 and not k.startswith('ordená del 1 al 6'):   # consigna genérica de ordenamiento
                    if k in vistos and vistos[k] != n:
                        control('C. Repeticiones', False, f'E{n} repite un ítem de E{vistos[k]}: «{k[:70]}»')
                    else:
                        control('C. Repeticiones', True, '')
                        vistos.setdefault(k, n)

    # D. capacidades: cobertura anual, por etapa e integradoras
    clases_caps = {c for n in range(1, 37) for c in C.encuentro(n)['caps'] if C.encuentro(n)['tipo'] in ('clase', 'continuacion')}
    control('D. Capacidades', clases_caps == set(C.CAPACIDADES), f'capacidades sin clase: {set(C.CAPACIDADES) - clases_caps}')
    for n, etapa in ((18, range(1, 18)), (36, range(19, 36))):
        enseñadas = {c for k in etapa for c in C.encuentro(k)['caps']}
        mz = {c.strip() for _, cap, _, _ in modulo(n).SOL['matriz'] for c in cap.split(',')}
        control('D. Capacidades', mz == set(C.encuentro(n)['caps']) <= enseñadas,
                f'E{n}: la matriz cubre {sorted(mz)}; capacidades de la etapa {sorted(enseñadas)}')

    # E. Guerra Guaranítica: solo como ampliación optativa del encuentro 22, nunca en ítems ni en claves
    for n in range(1, 37):
        m = modulo(n)
        for b in m.TOMO:
            if any('Guaranítica' in x for x in textos([b])):
                ok = n == 22 and b[0] == 'caja' and 'Ampliación optativa' in b[2]
                control('E. Guerra Guaranítica (A05)', ok, f'E{n}: mención fuera del recuadro de ampliación ({b[0]})')
        for _, filas in m.SOL['bloques']:
            for k, r, c in filas:
                control('E. Guerra Guaranítica (A05)', 'Guaranítica' not in r, f'E{n}, ítem {k}: la respuesta la exige')
    control('E. Guerra Guaranítica (A05)', 'Guaranítica' not in tema(22) and all('Guaranítica' not in i for i in indicadores(22)),
            'E22: tema o indicadores mencionan la Guerra Guaranítica')

    # F. figuras: numeración única en el año y archivos presentes
    figdir = os.path.join(ANUAL, 'figuras')
    numeros = collections.Counter()
    for n in range(1, 37):
        m = modulo(n)
        for b in list(m.TOMO) + [x for f in getattr(m, 'FICHAS', []) for x in f[2]]:
            if b[0] == 'fig':
                control('F. Figuras y cartografía', os.path.exists(os.path.join(figdir, b[1])), f'E{n}: falta {b[1]}')
                mm = re.match(r'Figura (\S+) —', b[2])
                if mm:
                    numeros[mm.group(1)] += 1
                if b[1].startswith(('fig_', 'ficha_', 'eval_')) and 'mapa' in b[1]:
                    control('F. Figuras y cartografía', True, '')
    for k, v in numeros.items():
        control('F. Figuras y cartografía', v == 1, f'la Figura {k} está numerada {v} veces en el año')

    # G. N01: el Acuífero Patiño no aparece como uso de un río
    r14 = next(r for nombre, filas in modulo(28).SOL['bloques'] for k, r, c in filas if k == '14')
    control('G. N01 (Clase 18)', 'Patiño' not in r14, 'la respuesta de la actividad 14 usa el Acuífero Patiño')
    rios = next(b[1] for b in modulo(28).TOMO if b[0] == 'p' and b[1].startswith('Los ríos dan agua'))
    control('G. N01 (Clase 18)', 'Patiño' not in rios, 'el apartado de usos de los ríos menciona el Patiño')

    # H. Plan 31: 20 + 100 minutos
    sec = modulo(31).PLAN['secuencia']
    control('H. Plan 31 (A10/N03)', sec[0][1][0][0] == '0–20' and 'Evaluación de la Unidad 5' in sec[0][1][0][1]
            and sec[1][1][0][0].startswith('20'), 'el Plan 31 no reserva 0-20 para la U5')
    return res


def main():
    for s in ('plan_anual', 'matriz', 'datos', 'revision_liviana'):
        os.makedirs(os.path.join(ANUAL, s), exist_ok=True)
    traz = json.load(open(os.path.join(ANUAL, 'datos', 'trazabilidad.json')))
    paginas = {x['encuentro']: x['paginas_tomo'] for x in traz['encuentros']}
    pa = plan_anual(os.path.join(ANUAL, 'plan_anual', 'PLAN_ANUAL_oficio.docx'))
    mz = matriz(os.path.join(ANUAL, 'matriz', 'MATRIZ_CURRICULAR_36_ENCUENTROS_A4.docx'), paginas)
    for p in (pa, mz):
        dst = os.path.join(ANUAL, 'revision_liviana', os.path.basename(p)[:-4] + '_REVISION.pdf')
        PDF.liviano(p, dst)
        print(os.path.basename(p), pymupdf.open(p).page_count, 'págs.')
    res = controles()
    ok = sum(r['correctos'] for r in res.values()); fallos = [f for r in res.values() for f in r['fallos']]
    json.dump({'correctos': ok, 'fallos': fallos, 'por_control': res},
              open(os.path.join(ANUAL, 'datos', 'controles_consolidacion.json'), 'w'), ensure_ascii=False, indent=1)
    print(f'Controles de consolidación: {ok} correctos, {len(fallos)} fallos')
    for k, r in res.items():
        print(f'  {k}: {r["correctos"]} correctos, {len(r["fallos"])} fallos')
    for f in fallos:
        print(' -', f)


if __name__ == '__main__':
    main()
