"""Solucionario docente (A4 vertical), con la maqueta aprobada en la muestra 3.

Generaliza `construir_muestra3.solucionario()`. Flujo continuo, sin salto de
hoja entre secciones. Cada sección puede llevar:
- bloques de respuestas (número, respuesta, criterio);
- mapas o figuras de clave gráfica docente (`mapas`);
- matriz ítem-capacidad-puntaje (`matriz`), obligatoria en las evaluaciones.

Numeración: la del tomo. Fuente: F1, F2, F3. Ubicate: U.
"""
import os
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', '..', '02_piloto', 'scripts'))
import planes_piloto as PP  # noqa: E402

BLANCO = RGBColor(0xFF, 0xFF, 0xFF)


def _tabla_respuestas(doc, nombre, items):
    t = doc.add_table(rows=1, cols=3); t.style = 'Table Grid'
    enc = t.rows[0]; PP.encabezado_repetido(enc); PP.no_partir(enc)
    m = enc.cells[0].merge(enc.cells[2]); PP.sombrear(m, PP.BANDA)
    PP.texto_celda(m, nombre, negrita=True, color=BLANCO, tam=10.5)
    for p in m.paragraphs:
        p.paragraph_format.keep_with_next = True
    for n, resp, crit in items:
        f = t.add_row(); PP.no_partir(f)
        PP.sombrear(f.cells[0], PP.GRIS)
        PP.texto_celda(f.cells[0], n, negrita=True, tam=10.5, alinear=WD_ALIGN_PARAGRAPH.CENTER)
        PP.texto_celda(f.cells[1], resp, tam=10.5)
        PP.texto_celda(f.cells[2], crit, tam=10)
    PP.fijar_anchos(t, [1.0, 9.6, 6.4])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _matriz(doc, filas, total, entera=False):
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(4)
    r = p.add_run('Matriz ítem-capacidad-puntaje'); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = PP.AZUL
    t = doc.add_table(rows=1, cols=4); t.style = 'Table Grid'
    enc = t.rows[0]; PP.encabezado_repetido(enc); PP.no_partir(enc)
    for j, h in enumerate(['Ítem', 'Capacidad', 'Indicador evaluado', 'Puntaje']):
        PP.sombrear(enc.cells[j], PP.BANDA); PP.texto_celda(enc.cells[j], h, negrita=True, color=BLANCO, tam=10)
    for item, cap, ind, pts in filas:
        f = t.add_row(); PP.no_partir(f)
        for j, v in enumerate((item, cap, ind, str(pts))):
            PP.texto_celda(f.cells[j], v, tam=9.5, negrita=(j == 0))
    f = t.add_row(); PP.no_partir(f)
    m = f.cells[0].merge(f.cells[2]); PP.sombrear(m, PP.GRIS); PP.sombrear(f.cells[3], PP.GRIS)
    PP.texto_celda(m, 'Total', negrita=True, tam=10); PP.texto_celda(f.cells[3], str(total), negrita=True, tam=10)
    PP.fijar_anchos(t, [2.2, 2.4, 10.6, 1.8])
    # las últimas filas y el total van juntos, para no dejar una o dos filas sueltas en otra hoja
    for fila in (t.rows[:-1] if entera else t.rows[-4:-1]):   # la última matriz del volumen va entera
        for c in fila.cells:
            for q in c.paragraphs:
                q.paragraph_format.keep_with_next = True
    # COM-01: si hay filas con más de una capacidad, se declara el puntaje compartido (no hay peso exclusivo).
    compartidas = [pts for _, cap, _, pts in filas if ',' in cap or ' y ' in cap]
    if compartidas:
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(2)
        r = p.add_run('Puntaje compartido: %d de los %d puntos corresponden a ítems que evalúan más de una capacidad. '
                      'Cada ítem se cuenta una sola vez en el total; esos puntos no se reparten entre capacidades. La matriz '
                      'muestra qué capacidades cubre cada ítem, no un peso exclusivo por capacidad.' % (sum(compartidas), total))
        r.italic = True; r.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _mapa(doc, figuras, archivo, ancho, epigrafe):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(os.path.join(figuras, archivo), width=Cm(ancho))
    p = doc.add_paragraph(); r = p.add_run(epigrafe); r.italic = True; r.font.size = Pt(9.5)


def construir(secciones, rangos, figuras, salida, titulo, referencias=None):
    doc = docx.Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0)
    s.top_margin = s.bottom_margin = Cm(1.8)
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    PP.pie_de_pagina(s, 'Solucionario docente')
    p = doc.add_paragraph(); r = p.add_run(titulo); r.bold = True
    r.font.size = Pt(16); r.font.color.rgb = PP.AZUL
    doc.add_paragraph('Numeración igual a la del tomo. «F1-F3» son las preguntas de «Trabajá con la fuente» y «U» la consigna '
                      '«Ubicate». Cuando el tomo anuncia criterios, el solucionario los repite y solo agrega niveles parciales y '
                      'respuestas modelo. En las preguntas abiertas se aceptan otras respuestas correctas y fundamentadas.')
    for k, sec in enumerate(secciones):
        p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(14 if k > 0 else 6)
        r = p.add_run(sec['titulo']); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = PP.AZUL
        rango = next((v for t, v in rangos.items() if t.startswith(sec['clave'])), None)
        ref = (f"Tomo, págs. {rango[0]} a {rango[1]} · Encuentro {sec['encuentro']}" if rango and rango[0] != rango[1]
               else f"Tomo, pág. {rango[0]} · Encuentro {sec['encuentro']}" if rango else f"Encuentro {sec['encuentro']}")
        if sec.get('total'):
            ref += ' · ' + sec['total']
        p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
        r = p.add_run(ref); r.italic = True; r.font.size = Pt(10)
        for nombre, items in sec['bloques']:
            _tabla_respuestas(doc, nombre, items)
        for archivo, ancho, epigrafe in sec.get('mapas', []):
            _mapa(doc, figuras, archivo, ancho, epigrafe)
        if sec.get('matriz'):
            _matriz(doc, sec['matriz'], sec['puntaje'], entera=(k == len(secciones) - 1))
    if referencias:
        _referencias(doc, *referencias)
    doc.save(salida)


def _referencias(doc, intro, filas):
    """COM-02: referencias para el docente, al final del volumen."""
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(14)
    r = p.add_run('Referencias para el docente'); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = PP.AZUL
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
    r = p.add_run(intro); r.font.size = Pt(9.5)
    t = doc.add_table(rows=1, cols=4); t.style = 'Table Grid'
    enc = t.rows[0]; PP.encabezado_repetido(enc); PP.no_partir(enc)
    for j, h in enumerate(['Tema', 'Referencia', 'Uso en el libro', 'Estado']):
        PP.sombrear(enc.cells[j], PP.BANDA); PP.texto_celda(enc.cells[j], h, negrita=True, color=BLANCO, tam=9.5)
    for fila in filas:
        f = t.add_row(); PP.no_partir(f)
        for j, v in enumerate(fila):
            PP.texto_celda(f.cells[j], v, tam=8.5, negrita=(j == 0))
    PP.fijar_anchos(t, [3.6, 8.2, 2.6, 2.6])
