"""Constructor del tomo del estudiante (legal vertical, márgenes estrechos; antes oficio, variante B).

Generaliza `04_muestra3/scripts/construir_tomo_m3.py`, la maqueta aprobada
como modelo (commit 841f5bc9). Del maestro del 3 de octubre de 2026
(SHA-256 8dbc4f44…179413) conserva solo portada, presentación, estilos e
índice; todo el contenido sale de `scripts/contenido/`. No usa ninguna
ilustración del maestro.

Bloques disponibles: unidad, h1, h2, h3, p, lista, nota, ficha, caja, tabla,
cuadro, fuente, ubicate, fig, consigna, item, opc, lista_orden, renglones,
banda, datos, ficha_repro.
"""
import copy
import json
import os
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(AQUI, '..', '..', '..', '02_piloto', 'scripts'))
import tomo_piloto as TP  # noqa: E402

W, A, R = TP.W, TP.A, TP.R
TAM_CAJA, TAM_TABLA, TAM_ACT, TAM_EPI = 9.5, 9.0, 10.0, 8.5   # tamaños del maestro
# Edición compacta (pedido de Fernando, 15-oct-2026): se vuelve a los tamaños del maestro (sin el punto extra de la
# variante B), interlineado sencillo, márgenes algo menores y separadores mínimos. Cada clase sigue en hoja nueva.
# Papel legal (21,59 × 35,56 cm) con márgenes estrechos de 1,27 cm, como el «Estrecho» de Word (pedido del 15-oct-2026).
COMPACTO = dict(suma_pt=0.0, papel=(21.59, 35.56), margen=(1.27, 1.27, 1.27, 1.27), encabezado_pie=0.5,
                interlineado=240, despues_pt=4, renglon=400,
                h_antes={'Heading 1': 0, 'Heading 2': 6, 'Heading 3': 4}, separador=80)
COLOR = {'ficha_k': '21374D', 'ficha_v': 'E6EBF1', 'concepto': 'FFF8E1', 'ejemplo': 'E6EBF1', 'paraguay': 'E8F5E9',
         'fuente': 'F3EEE2', 'fuente_t': '8A5A2B', 'ubicate': 'E6EBF1', 'ubicate_t': '21374D', 'cab': '3D5A80',
         'banda': 'B8860B', 'epi': '5D6775', 'repro': 'F7F7F7'}


class Constructor:
    def __init__(self, doc, figuras):
        self.doc, self.fig = doc, figuras
        self.body = doc.element.body
        self.sect = self.body.find(W + 'sectPr')

    # ---------- utilidades ----------
    def _mover(self, el):
        self.sect.addprevious(el)
        return el

    def parrafo(self, texto='', estilo=None, tam=None, negrita=False, cursiva=False, color=None, alinear=None, despues=None):
        p = self.doc.add_paragraph(style=estilo) if estilo else self.doc.add_paragraph()
        if texto:
            r = p.add_run(texto)
            r.bold, r.italic = negrita, cursiva
            if tam:
                r.font.size = Pt(tam)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        if alinear:
            p.alignment = alinear
        if despues is not None:
            p.paragraph_format.space_after = Pt(despues)
        self._mover(p._p)
        return p

    def _tabla(self, filas, cols, anchos_rel):
        t = self.doc.add_table(rows=filas, cols=cols)
        tbl = t._tbl
        tblpr = tbl.tblPr
        tw = tblpr.find(qn('w:tblW'))
        if tw is None:
            tw = OxmlElement('w:tblW'); tblpr.append(tw)
        tw.set(qn('w:type'), 'pct'); tw.set(qn('w:w'), '5000')
        mar = OxmlElement('w:tblCellMar')
        for lado, v in (('top', 60), ('left', 100), ('bottom', 60), ('right', 100)):
            e = OxmlElement('w:' + lado); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
        tblpr.append(mar)
        bordes = OxmlElement('w:tblBorders')
        for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            e = OxmlElement('w:' + lado)
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), 'A0A8B3')
            bordes.append(e)
        tblpr.append(bordes)
        tot = sum(anchos_rel)
        for gc, a in zip(tbl.tblGrid.findall(qn('w:gridCol')), anchos_rel):
            gc.set(qn('w:w'), str(int(10800 * a / tot)))
        for fila in t.rows:
            trpr = fila._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
            for c, a in zip(fila.cells, anchos_rel):
                c.width = Cm(18.0 * a / tot)
        self._mover(tbl)
        return t

    @staticmethod
    def _sombrear(celda, color):
        tcpr = celda._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), color)
        tcpr.append(shd)

    @staticmethod
    def _escribir(celda, texto, tam, negrita=False, color=None, cursiva=False, nuevo=False):
        p = celda.add_paragraph() if nuevo else celda.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(texto); r.font.size = Pt(tam); r.bold = negrita; r.italic = cursiva
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return p

    def separador(self):
        p = self.parrafo(despues=0)
        p.paragraph_format.space_before = Pt(0)
        # separador mínimo entre recuadros: alto exacto de 4 pt
        sp = p._p.get_or_add_pPr().find(qn('w:spacing'))
        if sp is None:
            sp = OxmlElement('w:spacing'); p._p.get_or_add_pPr().append(sp)
        sp.set(qn('w:line'), str(COMPACTO['separador'])); sp.set(qn('w:lineRule'), 'exact')
        return p

    # ---------- bloques ----------
    def h(self, nivel, texto):
        self.parrafo(texto, estilo=f'Heading {nivel}')

    def ficha(self, cap, tema, indicadores):
        t = self._tabla(3, 2, [1, 2.4])
        for i, (k, v) in enumerate([('Capacidad', cap), ('Tema', tema), ('Indicadores de logro', None)]):
            self._sombrear(t.cell(i, 0), COLOR['ficha_k']); self._sombrear(t.cell(i, 1), COLOR['ficha_v'])
            self._escribir(t.cell(i, 0), k, TAM_CAJA, negrita=True, color='FFFFFF')
            if v:
                self._escribir(t.cell(i, 1), v, TAM_CAJA)
            else:
                for j, ind in enumerate(indicadores):
                    self._escribir(t.cell(i, 1), '• ' + ind, TAM_CAJA, nuevo=j > 0)
        self.separador()

    def caja(self, tipo, titulo, texto):
        t = self._tabla(1, 1, [1])
        c = t.cell(0, 0); self._sombrear(c, COLOR[tipo])
        self._escribir(c, titulo, TAM_CAJA, negrita=True)
        self._escribir(c, texto, TAM_CAJA, nuevo=True)
        self.separador()

    def tabla(self, cab, filas, anchos):
        t = self._tabla(1 + len(filas), len(cab), anchos)
        enc = t.rows[0]._tr.get_or_add_trPr(); enc.append(OxmlElement('w:tblHeader'))
        for j, h in enumerate(cab):
            self._sombrear(t.cell(0, j), COLOR['cab']); self._escribir(t.cell(0, j), h, TAM_TABLA, negrita=True, color='FFFFFF')
        for i, fila in enumerate(filas, 1):
            for j, v in enumerate(fila):
                self._escribir(t.cell(i, j), v, TAM_TABLA, negrita=(j == 0))
        self.separador()

    def cuadro(self, cab, etiquetas, anchos, alto=None):
        t = self._tabla(1 + len(etiquetas), len(cab), anchos)
        for j, h in enumerate(cab):
            self._sombrear(t.cell(0, j), COLOR['cab']); self._escribir(t.cell(0, j), h, TAM_TABLA, negrita=True, color='FFFFFF')
        for i, et in enumerate(etiquetas, 1):
            tr = t.rows[i]._tr.get_or_add_trPr()
            hgt = OxmlElement('w:trHeight'); hgt.set(qn('w:val'), str(int((alto or (1.6 if et else 2.6)) / 2.54 * 1440))); hgt.set(qn('w:hRule'), 'atLeast')
            tr.append(hgt)
            # cada celda lleva texto (al menos un espacio) para que la fila no colapse
            self._escribir(t.cell(i, 0), et or ' ', TAM_TABLA, negrita=True)
            for j in range(1, len(cab)):
                self._escribir(t.cell(i, j), ' ', TAM_TABLA)
        self.separador()

    def fuente(self, subtitulo, citas, vocab, preguntas, renglones=0):
        """Dos filas sin línea entre ellas, que se ven como un solo recuadro: la
        fuente (puede seguir en la hoja siguiente) y las preguntas con sus
        renglones (nunca se separan)."""
        t = self._tabla(2, 1, [1])
        c = t.cell(0, 0); self._sombrear(c, COLOR['fuente'])
        self._escribir(c, 'Trabajá con la fuente', 10.5, negrita=True, color=COLOR['fuente_t'])
        self._escribir(c, subtitulo, TAM_TABLA, cursiva=True, nuevo=True)
        for ci in citas:
            self._escribir(c, ci, TAM_CAJA, nuevo=True)
        if vocab:
            self._escribir(c, vocab, TAM_TABLA, cursiva=True, nuevo=True)
        c2 = t.cell(1, 0); self._sombrear(c2, COLOR['fuente'])
        self._escribir(c2, 'Analizá la fuente:', TAM_CAJA, negrita=True)
        for k, q in enumerate(preguntas, 1):
            self._escribir(c2, f'{k}. {q}', TAM_CAJA, nuevo=True)
        self._renglones_en(c2, min(renglones, 3))   # las preguntas van con sus tres primeros renglones
        if renglones > 3:   # el resto de los renglones puede seguir en la hoja siguiente
            tr = OxmlElement('w:tr'); tr.append(OxmlElement('w:trPr'))
            tc = copy.deepcopy(t.rows[1]._tr.find(qn('w:tc')))
            for p in tc.findall(qn('w:p')):
                tc.remove(p)
            tr.append(tc); t._tbl.append(tr)
            for _ in range(renglones - 3):
                tc.append(TP.renglon(None))
        fila1 = t.rows[0]._tr.get_or_add_trPr()
        for x in fila1.findall(qn('w:cantSplit')):
            fila1.remove(x)
        bordes = t._tbl.tblPr.find(qn('w:tblBorders'))
        bordes.find(qn('w:insideH')).set(qn('w:val'), 'nil')
        self.separador()

    def ubicate(self, titulo, texto, renglones=0):
        t = self._tabla(1, 1, [1])
        c = t.cell(0, 0); self._sombrear(c, COLOR['ubicate'])
        self._escribir(c, titulo, TAM_CAJA, negrita=True, color=COLOR['ubicate_t'])
        self._escribir(c, texto, TAM_CAJA, nuevo=True)
        self._renglones_en(c, renglones)
        self.separador()

    def recuadro(self, titulo, lineas, renglones=0):
        """Recuadro con título y consignas numeradas (por ejemplo, la prueba
        diagnóstica); los renglones van dentro, después de cada consigna si
        `renglones` es una lista, o todos al final si es un número."""
        t = self._tabla(1, 1, [1])
        c = t.cell(0, 0); self._sombrear(c, COLOR['ubicate'])
        self._escribir(c, titulo, TAM_CAJA, negrita=True, color=COLOR['ubicate_t'])
        for k, q in enumerate(lineas, 1):
            self._escribir(c, f'{k}. {q}', TAM_CAJA, nuevo=True)
            if isinstance(renglones, (list, tuple)):
                self._renglones_en(c, renglones[k - 1])
        if not isinstance(renglones, (list, tuple)):
            self._renglones_en(c, renglones)
        self.separador()

    @staticmethod
    def _renglones_en(celda, n):
        """Renglones para responder dentro del recuadro: la fila no se parte,
        así que consigna y renglones quedan siempre en la misma página."""
        for _ in range(n):
            celda._tc.append(TP.renglon(None))

    def figura(self, archivo, epigrafe, ancho):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(os.path.join(self.fig, archivo), width=Cm(ancho))
        self._mover(p._p)
        self.parrafo(epigrafe, tam=TAM_EPI, cursiva=True, color=COLOR['epi'], alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=6)

    def banda(self, titulo, sub):
        t = self._tabla(1, 1, [1])
        c = t.cell(0, 0); self._sombrear(c, COLOR['banda'])
        p = self._escribir(c, titulo, 13, negrita=True, color='FFFFFF'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = self._escribir(c, sub, 10, color='FFFFFF', nuevo=True); p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def ficha_repro(self, titulo, bloques):
        self.h(1, titulo)
        self.parrafo('Se permite fotocopiar esta ficha para uso en el aula.', tam=TAM_EPI, cursiva=True, color=COLOR['epi'])
        self.parrafo('Nombre y Apellido: ______________________________   Curso: ______   Fecha: ___ / ___ / ______', tam=TAM_CAJA, despues=8)
        for b in bloques:
            self.bloque(b)

    def bloque(self, b):
        tipo = b[0]
        if tipo == 'h1':
            self.h(1, b[1])
        elif tipo == 'h2':
            self.h(2, b[1])
        elif tipo == 'h3':
            self.h(3, b[1])
        elif tipo == 'p':
            self.parrafo(b[1])
        elif tipo == 'ficha':
            self.ficha(*b[1:])
        elif tipo == 'caja':
            self.caja(*b[1:])
        elif tipo == 'tabla':
            self.tabla(*b[1:])
        elif tipo == 'cuadro':
            self.cuadro(*b[1:])
        elif tipo == 'fuente':
            self.fuente(*b[1:])
        elif tipo == 'ubicate':
            self.ubicate(*b[1:])
        elif tipo == 'fig':
            self.figura(*b[1:])
        elif tipo == 'recuadro':
            self.recuadro(*b[1:])
        elif tipo == 'lista':
            for it in b[1]:
                p = self.parrafo('•  ' + it, despues=2)
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.first_line_indent = Cm(-0.4)
        elif tipo == 'nota':
            self.parrafo(b[1], tam=TAM_EPI, cursiva=True, color=COLOR['epi'], despues=4)
        elif tipo == 'unidad':
            self.h(1, b[1])
        elif tipo == 'consigna':
            p = self.parrafo(b[1], tam=TAM_ACT, negrita=True, despues=2)
            p.paragraph_format.keep_with_next = True
        elif tipo == 'item':
            p = self.parrafo(b[1], tam=TAM_ACT, despues=2)
            p.paragraph_format.keep_with_next = True   # la consigna no queda sola al pie
            n = b[2] if len(b) > 2 else (1 if '(....)' in b[1] else 0)
            if n:
                p.paragraph_format.keep_with_next = True
                self.bloque(('renglones', n))
        elif tipo == 'opc':
            for letra, o in zip('abcd', b[1]):
                self.parrafo(f'    {letra}) {o}', tam=TAM_ACT, despues=0)
        elif tipo == 'lista_orden':
            for k, o in enumerate(b[1]):
                p = self.parrafo(f'(     )  {o}', tam=TAM_ACT, despues=2)
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.keep_with_next = k < len(b[1]) - 1
        elif tipo == 'renglones':
            for k in range(b[1]):
                n = TP.renglon(self.doc)
                if k < min(b[1] - 1, 3):   # la consigna va con sus 4 primeros renglones; los demás pueden seguir en la hoja siguiente
                    n.find(W + 'pPr').insert(0, OxmlElement('w:keepNext'))
                self._mover(n)
            self.parrafo(despues=2)
        elif tipo == 'banda':
            self.banda(b[1], b[2])
        elif tipo == 'datos':
            self.parrafo('Nombre y Apellido: _______________________________  Curso: ______  Fecha: ____ / ____ / ____', tam=TAM_CAJA, despues=6)
        elif tipo == 'ficha_repro':
            self.ficha_repro(b[1], b[2])
        else:
            raise ValueError(tipo)




MOVIBLES = ('p', 'lista', 'h2', 'h3', 'caja', 'tabla', 'nota', 'fuente', 'ubicate')   # bloques que una figura puede dejar pasar


def secuencia(c, bloques, diferir=None):
    """Recorre los bloques; un recuadro de fuente o Ubicate seguido de renglones
    lleva los renglones dentro, para que consigna y espacio no se separen.

    `diferir` ({archivo de figura: k}): la figura se imprime después de los k
    bloques de texto que la siguen (párrafos, listas, subtítulos, recuadros o
    tablas), para que ese texto llene el hueco que quedaría al pie si la figura
    no entra. Nunca pasa una actividad, otra figura ni el fin de la sección."""
    diferir = diferir or {}
    pend = []

    def vaciar(hasta_cero=False):
        for f in [x for x in pend if not hasta_cero or x[1] <= 0]:
            c.bloque(f[0]); pend.remove(f)

    k = 0
    while k < len(bloques):
        b = bloques[k]
        sig = bloques[k + 1] if k + 1 < len(bloques) else None
        if b[0] == 'fig' and diferir.get(b[1], 0) > 0:
            vaciar()
            pend.append([b, diferir[b[1]]])
            k += 1
            continue
        if b[0] in MOVIBLES and pend:
            if b[0] in ('fuente', 'ubicate') and sig and sig[0] == 'renglones':
                getattr(c, b[0])(*b[1:], renglones=sig[1])
                k += 1
            else:
                c.bloque(b)
            for x in pend:
                x[1] -= 1
            vaciar(hasta_cero=True)
            k += 1
            continue
        vaciar()
        if b[0] in ('fuente', 'ubicate', 'recuadro') and sig and sig[0] == 'renglones':
            getattr(c, b[0])(*b[1:], renglones=sig[1])
            k += 2
            continue
        c.bloque(b)
        k += 1
    vaciar()


def primer_texto(b):
    """Texto con que empieza el bloque en la página impresa (para ubicarlo en el PDF)."""
    t = {'p': lambda: b[1], 'nota': lambda: b[1], 'h2': lambda: b[1], 'h3': lambda: b[1], 'lista': lambda: '•  ' + b[1][0],
         'caja': lambda: b[2], 'tabla': lambda: next((h for h in b[1] if h), b[2][0][0]),
         'fuente': lambda: 'Trabajá con la fuente', 'ubicate': lambda: b[1]}[b[0]]()
    return t


def tablas_a_la_ventana(doc):
    """Toda tabla del cuerpo ocupa el 100 % del ancho útil («Autoajustar a la
    ventana»): ancho en porcentaje y sin diseño fijo. Las columnas conservan sus
    proporciones."""
    for t in doc.element.body.iter(W + 'tbl'):
        if t.getparent().tag != W + 'body':
            continue
        pr = t.find(W + 'tblPr')
        tw = pr.find(W + 'tblW')
        if tw is None:
            tw = OxmlElement('w:tblW'); pr.append(tw)
        tw.set(qn('w:type'), 'pct'); tw.set(qn('w:w'), '5000')
        for lay in pr.findall(W + 'tblLayout'):
            pr.remove(lay)
        ind = pr.find(W + 'tblInd')
        if ind is not None:
            ind.set(qn('w:w'), '0'); ind.set(qn('w:type'), 'dxa')


def compactar(doc):
    """Interlineado sencillo, menos espacio después de cada párrafo y antes de
    los títulos, y renglones de respuesta de 7 mm. No toca el texto."""
    from docx.shared import Twips
    cfg = COMPACTO
    for st in doc.styles:
        try:
            pf = st.paragraph_format
        except AttributeError:
            continue
        if st.name in cfg['h_antes']:
            pf.space_before = Pt(cfg['h_antes'][st.name])
    normal = doc.styles['Normal'].paragraph_format
    normal.space_after = Pt(cfg['despues_pt'])
    normal.line_spacing = 1.0
    dflt = doc.styles.element.find(qn('w:docDefaults'))
    if dflt is not None:
        for sp in dflt.iter(qn('w:spacing')):
            sp.set(qn('w:line'), str(cfg['interlineado'])); sp.set(qn('w:after'), str(cfg['despues_pt'] * 20))
    for sp in doc.element.body.iter(qn('w:spacing')):
        if sp.get(qn('w:lineRule')) == 'exact' and sp.get(qn('w:line')) == '440':   # renglón de respuesta
            sp.set(qn('w:line'), str(cfg['renglon']))
        elif sp.get(qn('w:after')) == '120':   # 6 pt explícitos de los párrafos de texto
            sp.set(qn('w:after'), str(cfg['despues_pt'] * 20))


def ajustar(elementos, modo):
    """Ajuste fino de una sección cuya última hoja quedaría casi vacía. «apretar»:
    renglones de 6,3 mm y 2 pt después de cada párrafo, para que la sección entre
    en una hoja menos. «airear» (si apretar no alcanza): renglones de 8,5 mm y
    6 pt después de cada párrafo, para que la última hoja no quede casi vacía.
    No toca el texto ni la letra."""
    linea, despues = ('360', 40) if modo == 'apretar' else ('480', 120)
    for el in elementos:
        for p in ([el] if el.tag == W + 'p' else list(el.iter(W + 'p'))):
            ppr = p.find(W + 'pPr')
            if ppr is None:
                ppr = OxmlElement('w:pPr'); p.insert(0, ppr)
            sp = ppr.find(qn('w:spacing'))
            if sp is None:
                sp = OxmlElement('w:spacing'); ppr.append(sp)
            if sp.get(qn('w:lineRule')) == 'exact' and sp.get(qn('w:line')) == str(COMPACTO['renglon']):
                sp.set(qn('w:line'), linea)
            elif sp.get(qn('w:lineRule')) != 'exact' and int(sp.get(qn('w:after')) or 80) > 20:
                sp.set(qn('w:after'), str(despues))


def construir(maestro, figuras, secciones, salida, paginas=None, diferir=None, ajustes=None):
    """secciones: lista de listas de bloques, en orden de impresión."""
    doc = docx.Document(maestro)
    body = doc.element.body
    els = [e for e in body if e.tag != W + 'sectPr']
    fin_frente = next(i for i, e in enumerate(els) if e.tag == W + 'tbl' and TP.texto(e).startswith('Prueba diagnóstica'))
    for e in els[fin_frente:]:
        body.remove(e)
    for e in [e for e in body if TP.es_p(e)]:
        if not TP.texto(e).strip() and e.find('.//' + W + 'pageBreakBefore') is not None and not e.findall('.//' + A + 'blip'):
            body.remove(e)
    c = Constructor(doc, figuras)
    rangos_el = []
    for bloques in secciones:
        antes = len(body)
        secuencia(c, bloques, diferir)
        rangos_el.append(list(body)[antes - 1:len(body) - 1])   # el último hijo es sectPr
    TP.limpiar_relaciones(doc)
    TP.corregir_saltos(doc)
    ancho = TP.ajustar_pagina(doc, COMPACTO['papel'], COMPACTO['margen'])
    for s in doc.sections:
        s.header_distance = s.footer_distance = Cm(COMPACTO['encabezado_pie'])
    tablas_a_la_ventana(doc)
    TP.agrandar_letra(doc, COMPACTO['suma_pt'])
    compactar(doc)
    for i, modo in (ajustes or {}).items():
        if modo in ('apretar', 'airear'):
            ajustar(rangos_el[i], modo)
        elif modo == 'airear final':   # solo después de la última figura (o la segunda mitad, si no hay figuras)
            els = rangos_el[i]
            ult = max((k for k, e in enumerate(els) if e.findall('.//' + A + 'blip')), default=len(els) // 2)
            ajustar(els[ult + 1:], 'airear')
    TP.agrupar_opciones(doc, ancho)
    TP.indice_estatico(doc, ancho, paginas or {})
    doc.save(salida)
