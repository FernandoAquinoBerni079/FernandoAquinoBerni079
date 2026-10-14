"""Planes de clase de la tercera muestra.

Reutiliza las funciones de formato de `02_piloto/scripts/planes_piloto.py`
(A4 vertical, tablas de anchos fijos, filas cortas que no se parten) y agrega
dos elementos pedidos por la auditoría de la segunda muestra:
- «Tipo de vínculo curricular»: distingue capacidad literal del programa MEC,
  asociación curricular, formulación editorial y ampliación editorial;
- «Actividades del tomo según su función»: esenciales, de consolidación y
  complementarias o domiciliarias.

La duración de cada momento se calcula de los minutos de la secuencia.
"""
import os
import sys

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', '02_piloto', 'scripts'))
import planes_piloto as PP  # noqa: E402


def junto(tabla, filas=None):
    """Mantiene unidas las filas indicadas con la siguiente (keep_with_next en sus párrafos)."""
    for fila in (filas if filas is not None else tabla.rows[:-1]):
        for c in fila.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True


def plan(doc, pl, comunes):
    p = doc.add_paragraph(); p.paragraph_format.page_break_before = True
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.keep_with_next = True
    r = p.add_run(f"Plan de Clase N.º {pl['n']}"); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = PP.AZUL
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.keep_with_next = True
    r = p.add_run(pl['titulo']); r.bold = True; r.font.size = Pt(12.5)

    t = doc.add_table(rows=4, cols=4); t.style = 'Table Grid'
    datos = [('Institución', '', 'Docente', ''),
             ('Disciplina', comunes['disciplina'], 'Grado y sección', '7.º  ____'),
             ('Unidad', pl['unidad'], 'Encuentro', f"{pl['encuentro']} de 36"),
             ('Material', pl['material'], 'Duración', comunes['duracion'] + '\nFecha: ___ / ___ / ______')]
    for i, fila in enumerate(datos):
        PP.no_partir(t.rows[i])
        for j, v in enumerate(fila):
            c = t.cell(i, j)
            if j % 2 == 0:
                PP.sombrear(c, PP.GRIS); PP.texto_celda(c, v, negrita=True)
            else:
                PP.texto_celda(c, v)
    PP.fijar_anchos(t, [2.6, 6.4, 3.0, 5.0])

    PP.titulo_seccion(doc, 'Capacidad, tema e indicadores')
    PP.tabla_clave_valor(doc, [('Capacidad', pl['capacidad']), ('Tipo de vínculo curricular', pl['vinculo']),
                               ('Tema', pl['tema']), ('Indicadores de logro', pl['indicadores']),
                               ('Propósito del encuentro', pl['proposito'])])

    PP.titulo_seccion(doc, 'Actividades del tomo según su función')
    junto(PP.tabla_clave_valor(doc, list(pl['funciones'].items()), anchos=(5.0, 12.0)))

    PP.titulo_seccion(doc, 'Secuencia didáctica (120 minutos)')
    t = doc.add_table(rows=1, cols=4); t.style = 'Table Grid'
    enc = t.rows[0]; PP.encabezado_repetido(enc); PP.no_partir(enc)
    for j, h in enumerate(['Minutos', 'Qué hace el docente', 'Qué hacen los estudiantes', 'Producto o evidencia']):
        PP.sombrear(enc.cells[j], PP.BANDA)
        PP.texto_celda(enc.cells[j], h, negrita=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for momento, pasos in pl['secuencia']:
        f = t.add_row(); PP.no_partir(f)
        m = f.cells[0].merge(f.cells[3]); PP.sombrear(m, PP.GRIS)
        ini, fin = pasos[0][0].split('–')[0], pasos[-1][0].split('–')[1]
        PP.texto_celda(m, f"{momento} ({int(fin) - int(ini)} min)", negrita=True)
        junto(t, [f])   # la banda de cada momento no queda sola al pie de la hoja
        for minutos, docente, alumnos, producto in pasos:
            f = t.add_row(); PP.no_partir(f)
            PP.texto_celda(f.cells[0], minutos, alinear=WD_ALIGN_PARAGRAPH.CENTER)
            PP.texto_celda(f.cells[1], docente)
            PP.texto_celda(f.cells[2], alumnos)
            PP.texto_celda(f.cells[3], producto)
    PP.fijar_anchos(t, [1.7, 6.3, 5.3, 3.7])

    PP.titulo_seccion(doc, 'Atención a la diversidad')
    for it in pl['diversidad']:
        q = doc.add_paragraph(style='List Bullet'); q.add_run(it)

    PP.titulo_seccion(doc, 'Evaluación')
    ev = pl['evaluacion']
    junto(PP.tabla_clave_valor(doc, [('Tipo', ev['tipo']), ('Procedimientos', ev['procedimientos']),
                                     ('Instrumentos', ev['instrumentos']), ('Criterios', ev['criterios'])]))

    PP.titulo_seccion(doc, 'Precisiones para el docente')
    for it in pl['precisiones']:
        q = doc.add_paragraph(style='List Bullet'); q.add_run(it)

    PP.titulo_seccion(doc, 'Tarea')
    doc.add_paragraph(pl['tarea'])

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.keep_with_next = True
    p.add_run('Observaciones: ' + '_' * 78)
    t = doc.add_table(rows=1, cols=2)
    for j, txt in enumerate(['____________________________\nDocente', '____________________________\nDirección — Visto bueno']):
        PP.texto_celda(t.cell(0, j), txt, alinear=WD_ALIGN_PARAGRAPH.CENTER)
    PP.no_partir(t.rows[0])
    PP.fijar_anchos(t, [8.5, 8.5])
