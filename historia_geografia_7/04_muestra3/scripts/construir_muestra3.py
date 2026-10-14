"""Construye la tercera muestra (corrección de la segunda).

Uso: python3 construir_muestra3.py MAESTRO.docx CARPETA_GEOJSON SALIDA [PLAN_ANUAL_ORIGINAL.docx]

SALIDA debe contener cotejo_capacidades.json y estructura_maestro.json (de la
muestra 2, que no se modifican); la matriz escribe trazabilidad_curricular.json.

Genera en SALIDA: figuras, tomo (DOCX y PDF), planes 15, 28, 29 y 30 (DOCX y
PDF), solucionario (DOCX y PDF), matriz curricular (DOCX y PDF), versiones
livianas de los PDF para revisión y paginas_tomo.json.
"""
import json
import os
import re
import subprocess
import sys

import pymupdf

AQUI = os.path.dirname(os.path.abspath(__file__))
PILOTO = os.path.join(AQUI, '..', '..', '02_piloto', 'scripts')
sys.path.insert(0, AQUI)
sys.path.insert(0, PILOTO)


def render(docx_path):
    out = os.path.dirname(docx_path)
    env = dict(os.environ, HOME='/tmp/lohome')
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', out, docx_path], check=True, env=env,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=900)
    return docx_path[:-5] + '.pdf'


def titulos_pdf(pdf, patron, desde=0):
    """Página de cada título que cumple el patrón, después del índice."""
    d = pymupdf.open(pdf)
    res = {}
    for n in range(desde, d.page_count):
        for linea in d[n].get_text().split('\n'):
            l = linea.strip()
            if re.match(patron, l) and l not in res:
                res[l] = n + 1
    return res


def pagina_indice(pdf):
    d = pymupdf.open(pdf)
    return next(n for n in range(d.page_count) if 'Índice' in d[n].get_text().split('\n')[0:3])


def liviano(pdf, salida, lado_max=1100, calidad=72):
    """Copia del PDF para revisión visual: cada imagen se reemplaza en el mismo
    objeto por un JPEG de lado mayor 1100 píxeles; el texto no cambia."""
    import io
    from PIL import Image
    d = pymupdf.open(pdf)
    for x in range(1, d.xref_length()):
        if d.xref_get_key(x, 'Subtype')[1] != '/Image' or d.xref_get_key(x, 'ImageMask')[1] == 'true':
            continue
        if any(d.xref_get_key(y, 'SMask')[1] == f'{x} 0 R' for y in ()):
            continue
        try:
            pix = pymupdf.Pixmap(d, x)
        except Exception:
            continue
        if pix.n - pix.alpha < 3:      # máscaras y grises: se dejan
            continue
        if pix.alpha or pix.n > 3:
            pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
        img = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
        img.thumbnail((lado_max, lado_max))
        buf = io.BytesIO(); img.save(buf, 'JPEG', quality=calidad, optimize=True)
        if buf.tell() >= len(d.xref_stream_raw(x) or b''):
            continue
        d.update_stream(x, buf.getvalue(), compress=False)
        for k, v in (('Filter', '/DCTDecode'), ('Width', str(img.width)), ('Height', str(img.height)),
                     ('ColorSpace', '/DeviceRGB'), ('BitsPerComponent', '8'), ('DecodeParms', 'null'), ('SMask', 'null'),
                     ('Length', str(buf.tell()))):
            d.xref_set_key(x, k, v)
    d.save(salida, garbage=4, deflate=True)
    return os.path.getsize(salida)


def tomo(maestro, fig, out):
    s = os.path.join(AQUI, 'construir_tomo_m3.py')
    dx = os.path.join(out, 'TOMO_MUESTRA3_oficio.docx')
    subprocess.run([sys.executable, s, maestro, fig, dx], check=True)
    pdf = render(dx)
    patron = r'^(Clase \d+ —|Continuación [A-G] —|Ficha reproducible|EVALUACIÓN DE LA UNIDAD \d)'
    ind = pagina_indice(pdf)
    pags = titulos_pdf(pdf, patron, ind + 1)
    jp = os.path.join(out, 'paginas_tomo.json')
    # el índice se escribe con las claves completas de los títulos del documento
    import docx
    import tomo_piloto as TP
    d = docx.Document(dx)
    claves = {}
    for e in d.element.body:
        t = TP.texto(e) if TP.es_p(e) else (TP.texto(e.find('.//' + TP.W + 'p')) if e.tag == TP.W + 'tbl' else '')
        for k, v in pags.items():
            if t and t.startswith(k[:40]):
                claves[t] = v
    json.dump(claves, open(jp, 'w'), ensure_ascii=False, indent=1)
    subprocess.run([sys.executable, s, maestro, fig, dx, jp], check=True)
    pdf = render(dx)
    pags2 = titulos_pdf(pdf, patron, ind + 1)
    assert pags2 == pags, (pags, pags2)
    rangos = {}
    orden = sorted(pags.items(), key=lambda x: x[1])
    total = pymupdf.open(pdf).page_count
    for k, (t, p) in enumerate(orden):
        fin = (orden[k + 1][1] - 1) if k + 1 < len(orden) else total
        rangos[t] = (p, fin)
    json.dump(rangos, open(jp, 'w'), ensure_ascii=False, indent=1)
    return dx, pdf, rangos


def planes(out):
    import docx
    import planes_m3 as PM
    import planes_piloto as PP
    from contenido_planes import COMUNES, PLANES
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
    from docx.oxml.ns import qn
    PP.PLANES = PLANES
    PP.COMUNES.update(COMUNES)

    def construir(salida, paginas):
        doc = docx.Document()
        s = doc.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(1.8)
        st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
        st.paragraph_format.space_after = Pt(3)
        doc.styles['List Bullet'].font.size = Pt(10.5)
        PP.pie_de_pagina(s)
        for _ in range(6):
            doc.add_paragraph()
        for texto, tam, neg in [('PLANES DE CLASE', 26, True), ('Historia y Geografía · 7.º grado', 18, True),
                                ('Educación Escolar Básica · Tercer Ciclo', 13, False),
                                ('Tercera muestra · planes 15, 28, 29 y 30', 12, False)]:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(10)
            r = p.add_run(texto); r.font.size = Pt(tam); r.bold = neg; r.font.color.rgb = PP.AZUL
        for _ in range(3):
            doc.add_paragraph()
        for texto in ['Programa de Estudio: Historia y Geografía, 7.º grado (MEC, 2014)',
                      'Carga horaria: 3 horas cátedra semanales de 40 minutos · 36 encuentros de 120 minutos',
                      'Institución: ______________________________________',
                      'Docente: ______________________________________', 'Año lectivo: ____________']:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
            p.add_run(texto).font.size = Pt(12)
        PP.indice(doc, paginas)
        for pl in PLANES:
            PM.plan(doc, pl, PP.COMUNES)
        doc.save(salida)

    dx = os.path.join(out, 'PLANES_MUESTRA3_A4.docx')
    construir(dx, {})
    pdf = render(dx)
    d = pymupdf.open(pdf)
    pags = {}
    for n in range(2, d.page_count):
        for m in re.finditer(r'Plan de Clase N\.º (\d+)\n', d[n].get_text()):
            pags.setdefault(f'Plan de Clase N.º {m.group(1)}', n + 1)
    construir(dx, pags)
    pdf = render(dx)
    return dx, pdf


def solucionario(out, rangos):
    import docx
    import planes_piloto as PP
    from contenido_solucionario import SECCIONES
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    doc = docx.Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0)
    s.top_margin = s.bottom_margin = Cm(1.8)
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    PP.pie_de_pagina(s, 'Solucionario docente')
    p = doc.add_paragraph(); r = p.add_run('Solucionario docente — tercera muestra'); r.bold = True
    r.font.size = Pt(16); r.font.color.rgb = PP.AZUL
    doc.add_paragraph('Numeración igual a la del tomo de la muestra. Cada respuesta abierta trae su criterio de corrección o su rúbrica.')
    for k, sec in enumerate(SECCIONES):
        p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
        # flujo continuo: las secciones se siguen sin forzar hoja nueva (evita hojas casi vacías)
        p.paragraph_format.space_before = Pt(14 if k > 0 else 6)
        r = p.add_run(sec['titulo']); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = PP.AZUL
        rango = next((v for t, v in rangos.items() if t.startswith(sec['clave'])), None)
        ref = f"Tomo de la muestra, págs. {rango[0]} a {rango[1]} · Encuentro {sec['encuentro']}" if rango else f"Encuentro {sec['encuentro']}"
        if sec.get('total'):
            ref += ' · ' + sec['total']
        p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
        r = p.add_run(ref); r.italic = True; r.font.size = Pt(10)
        for nombre, items in sec['bloques']:
            t = doc.add_table(rows=1, cols=3); t.style = 'Table Grid'
            enc = t.rows[0]; PP.encabezado_repetido(enc); PP.no_partir(enc)
            m = enc.cells[0].merge(enc.cells[2]); PP.sombrear(m, PP.BANDA)
            PP.texto_celda(m, nombre, negrita=True, color=RGBColor(0xFF, 0xFF, 0xFF), tam=10.5)
            for n, resp, crit in items:
                f = t.add_row(); PP.no_partir(f)
                PP.sombrear(f.cells[0], PP.GRIS)
                PP.texto_celda(f.cells[0], n, negrita=True, tam=10.5, alinear=WD_ALIGN_PARAGRAPH.CENTER)
                PP.texto_celda(f.cells[1], resp, tam=10.5)
                PP.texto_celda(f.cells[2], crit, tam=10)
            PP.fijar_anchos(t, [1.0, 9.6, 6.4])
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
    from contenido_solucionario import MAPA_REFERENCIA
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(12)
    r = p.add_run('Mapa de referencia para la corrección cartográfica'); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = PP.AZUL
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(os.path.join(out, 'figuras', MAPA_REFERENCIA[0]), width=Cm(11.5))
    p = doc.add_paragraph(); r = p.add_run(MAPA_REFERENCIA[1]); r.italic = True; r.font.size = Pt(9.5)
    dx = os.path.join(out, 'SOLUCIONARIO_MUESTRA3_A4.docx')
    doc.save(dx)
    return dx, render(dx)


def matriz(out):
    import collections
    import docx
    import planes_piloto as PP
    from contenido_matriz import (CAPACIDADES, ENCUENTROS, ESTADO_FICHA, MUESTRA3, RECURSOS_M3, TEMAS, VINCULOS,
                                  estado_ficha, vinculo)
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    cot = {c['seccion']: c for c in json.load(open(os.path.join(out, 'cotejo_capacidades.json')))}
    estr = {s['t']: s for s in json.load(open(os.path.join(out, 'estructura_maestro.json')))}

    def buscar(dic, nombre):
        for k, v in dic.items():
            if k.upper().startswith(nombre.upper()):
                return v
        return None

    def clave_de(sec):
        if sec.startswith('Evaluación de la Unidad'):
            return 'EVALUACIÓN DE LA UNIDAD ' + sec.split()[-1]
        if sec.startswith('Integradora'):
            return 'EVALUACIÓN INTEGRADORA · 1.ª' if '1.ª' in sec else 'EVALUACIÓN INTEGRADORA · 2.ª'
        return sec + ' —'

    # trazabilidad curricular (reemplaza, sin modificarlo, el cotejo de la muestra 2)
    traz = []
    for n, sec, cap, temas, ampl in ENCUENTROS:
        c = buscar(cot, clave_de(sec))
        ef = estado_ficha(sec, c)
        traz.append({'encuentro': n, 'seccion': sec, 'tipo_vinculo': vinculo(n, sec),
                     'capacidades_mec': [x.strip() for x in cap.split(',')] if cap[0] in 'HG' else [],
                     'temas_mec': temas, 'ficha_maestro_estado': ef,
                     'ficha_maestro_texto': c['cap_tomo'] if c else '',
                     'en_muestra3': n in MUESTRA3, 'observacion': ampl})
    conteo = collections.Counter(t['ficha_maestro_estado'] for t in traz)
    json.dump({'programa': 'MEC (2014), Programa de Estudio: Historia y Geografía, 7.º grado, EEB, reproducción en PDF '
                           '(SHA-256 389190d302b668f62bf1bd05ac4ca01a78cfba83953cd18cd8ce282aefc8d27e), págs. 55 a 57',
               'tipos_de_vinculo': list(VINCULOS), 'conteo_fichas_maestro': dict(conteo),
               'capacidades': CAPACIDADES, 'encuentros': traz},
              open(os.path.join(out, 'trazabilidad_curricular.json'), 'w'), ensure_ascii=False, indent=1)

    doc = docx.Document()
    s = doc.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    s.left_margin = s.right_margin = Cm(1.5)
    s.top_margin = s.bottom_margin = Cm(1.5)
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    PP.pie_de_pagina(s, 'Matriz curricular')
    p = doc.add_paragraph(); r = p.add_run('Matriz curricular de los 36 encuentros — Historia y Geografía 7.º'); r.bold = True
    r.font.size = Pt(15); r.font.color.rgb = PP.AZUL
    doc.add_paragraph('Programa: MEC (2014), Programa de Estudio: Historia y Geografía, 7.º grado, EEB, págs. 55 a 57 '
                      '(reproducción en PDF de 92 págs.; ficha catalográfica con autoridades Cartes y Lafuente). La carga horaria '
                      'es de 3 horas semanales, «horas cátedras de 40 minutos» (pág. 46). Estado del cotejo: literal contra esa '
                      'reproducción; falta el ejemplar impreso o el PDF oficial para declararlo cotejado oficialmente.')
    p = doc.add_paragraph(); r = p.add_run('1. Capacidades del programa'); r.bold = True; r.font.color.rgb = PP.AZUL
    t = doc.add_table(rows=1, cols=2); t.style = 'Table Grid'
    PP.encabezado_repetido(t.rows[0])
    for j, h in enumerate(['Código', 'Capacidad (texto literal)']):
        PP.sombrear(t.cell(0, j), PP.BANDA); PP.texto_celda(t.cell(0, j), h, negrita=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for k, v in CAPACIDADES.items():
        f = t.add_row(); PP.no_partir(f); PP.texto_celda(f.cells[0], k, negrita=True); PP.texto_celda(f.cells[1], v)
    PP.fijar_anchos(t, [2.0, 24.7])
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10)
    r = p.add_run('Tipos de vínculo curricular'); r.bold = True; r.font.color.rgb = PP.AZUL
    for nombre, texto in [
            ('Capacidad literal MEC', 'el encuentro desarrolla temas del programa bajo una capacidad citada al pie de la letra.'),
            ('Asociación curricular', 'el encuentro no tiene capacidad propia: integra o evalúa capacidades literales (evaluaciones).'),
            ('Formulación editorial', 'propósito redactado por la editorial, sin capacidad del programa (Clase 1).'),
            ('Ampliación editorial', 'contenido que el programa no exige; se ofrece como optativo y no se atribuye al MEC '
                                     '(Continuación E, siglo XVIII).')]:
        q = doc.add_paragraph(style='List Bullet'); rr = q.add_run(nombre + ': '); rr.bold = True; q.add_run(texto)
    q = doc.add_paragraph()
    q.add_run('Fichas de capacidad del maestro del 3 de octubre: ' +
              '; '.join(f'{ESTADO_FICHA[k].lower()}, {conteo.get(k, 0)}' for k in
                        ('literal', 'truncada', 'abreviada', 'variante', 'editorial', 'sin_ficha')) +
              ' (36 en total).')

    p = doc.add_paragraph(); p.paragraph_format.page_break_before = True
    r = p.add_run('2. Encuentro por encuentro'); r.bold = True; r.font.color.rgb = PP.AZUL
    doc.add_paragraph('«Ficha en el maestro» compara la ficha del maestro del 3 de octubre con el texto literal. '
                      'Las filas sombreadas se desarrollan en esta tercera muestra.')
    cab = ['N.º', 'Sección', 'Cap.', 'Tipo de vínculo', 'Temas del programa', 'Ficha en el maestro', 'Contenido desarrollado',
           'Recursos', 'Ampliación prevista / estado']
    anchos = [0.8, 2.1, 1.3, 2.0, 4.2, 2.9, 4.6, 3.9, 4.9]
    t = doc.add_table(rows=1, cols=len(cab)); t.style = 'Table Grid'
    PP.encabezado_repetido(t.rows[0]); PP.no_partir(t.rows[0])
    for j, h in enumerate(cab):
        PP.sombrear(t.cell(0, j), PP.BANDA); PP.texto_celda(t.cell(0, j), h, negrita=True, color=RGBColor(0xFF, 0xFF, 0xFF), tam=9)
    for (n, sec, cap, temas, ampl), tr in zip(ENCUENTROS, traz):
        e = buscar(estr, clave_de(sec))
        ef = tr['ficha_maestro_estado']
        est = ESTADO_FICHA[ef] + ('' if ef in ('literal', 'sin_ficha') else ': «' + tr['ficha_maestro_texto'] + '»')
        if n in MUESTRA3:
            est += (' → en la muestra 3: evaluación asociada a G1 y G2; las capacidades se explicitan en el Plan 30; '
                    'sin ficha literal en el tomo.' if n == 30 else ' → en la muestra 3: literal.')
            from contenido_tomo import SECCIONES as ST
            clave_m3 = {15: 'clase10', 28: 'clase18', 29: 'contF', 30: 'evalU5'}[n]
            cont = '; '.join(b[1] for b in ST[clave_m3] if b[0] == 'h2') or 'Partes A a F: opción múltiple, V/F, mapa, fuente cartográfica, matriz, desarrollo.'
            rec = RECURSOS_M3[n]
        else:
            cont = '; '.join(e['h2']) if e and e['h2'] else '—'
            rec = '; '.join([x.split(' — ')[0] for x in (e['fig'] if e else [])] + (e['cajas'] if e else [])) or '—'
        if temas:
            tem = '; '.join(f'{x} {TEMAS[x]}' for x in temas)
        elif n == 22:
            tem = '(ninguno: siglo XVIII, fuera del alcance del programa)'
        else:
            tem = '(evaluación integradora)'
        fila = [str(n), sec, cap, tr['tipo_vinculo'], tem, est, cont, rec, ampl]
        f = t.add_row(); PP.no_partir(f)
        for j, v in enumerate(fila):
            PP.texto_celda(f.cells[j], v, tam=8.5, negrita=(j == 0))
            if n in MUESTRA3:
                PP.sombrear(f.cells[j], 'FFF8E1')
    PP.fijar_anchos(t, anchos)
    dx = os.path.join(out, 'MATRIZ_CURRICULAR_36_ENCUENTROS_M3.docx')
    doc.save(dx)
    return dx, render(dx)


def main():
    maestro, geo, out = sys.argv[1:4]
    os.makedirs(out, exist_ok=True)
    fig = os.path.join(out, 'figuras')
    subprocess.run([sys.executable, os.path.join(AQUI, 'figuras.py'), geo, fig], check=True)
    _, tpdf, rangos = tomo(maestro, fig, out)
    _, ppdf = planes(out)
    _, spdf = solucionario(out, rangos)
    _, mpdf = matriz(out)
    pdfs = [tpdf, ppdf, spdf, mpdf]
    if len(sys.argv) > 4:   # plan anual original
        pa = os.path.join(out, 'PLAN_ANUAL_OFICIO_M3.docx')
        subprocess.run([sys.executable, os.path.join(AQUI, 'plan_anual_oficio_v3.py'), sys.argv[4], pa], check=True)
        pdfs.append(render(pa))
    os.makedirs(os.path.join(out, 'revision_liviana'), exist_ok=True)
    for pdf in pdfs:
        dst = os.path.join(out, 'revision_liviana', os.path.basename(pdf)[:-4] + '_REVISION.pdf')
        print(os.path.basename(pdf), pymupdf.open(pdf).page_count, 'págs.;', os.path.getsize(pdf), 'bytes; liviano', liviano(pdf, dst))
    print(json.dumps(rangos, ensure_ascii=False))


if __name__ == '__main__':
    main()
