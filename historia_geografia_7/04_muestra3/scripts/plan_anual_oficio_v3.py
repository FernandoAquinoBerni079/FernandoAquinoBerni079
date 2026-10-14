"""Plan anual en oficio horizontal, tercera muestra.

Parte de la maqueta de `03_muestra2/scripts/plan_anual_oficio_v2.py` (que a su
vez parte del piloto) y corrige la trazabilidad curricular señalada por la
auditoría de la segunda muestra:
- las bandas de unidad citan las capacidades literales del programa MEC 2014
  (págs. 55 a 57) con su código; la Clase 1 queda marcada como formulación
  editorial;
- la Guerra Guaranítica deja de presentarse como «requerimiento curricular» o
  «priorización MEC»: es una ampliación editorial optativa del siglo XVIII;
- la observación de cada continuación dice qué capacidad profundiza;
- las filas 15, 28, 29 y 30 toman temas e indicadores de las fichas del tomo
  de la muestra (y del Plan 30 para la evaluación, que no tiene ficha), y
  procedimientos e instrumentos de los planes de clase;
- las integradoras declaran su cobertura: H1 a H4 y H5, H6, G1 a G4.

Uso: python3 plan_anual_oficio_v3.py PLAN_ANUAL_ORIGINAL.docx SALIDA.docx
"""
import copy
import os
import subprocess
import sys
import tempfile

import docx
from docx.oxml.ns import qn

AQUI = os.path.dirname(os.path.abspath(__file__))
V2 = os.path.join(AQUI, '..', '..', '03_muestra2', 'scripts', 'plan_anual_oficio_v2.py')
sys.path.insert(0, AQUI)
from contenido_matriz import CAPACIDADES  # noqa: E402
from contenido_planes import PLANES  # noqa: E402
from contenido_tomo import SECCIONES as TOMO  # noqa: E402

BANDAS = {
    'UNIDAD 1': ('UNIDAD 1 — Cómo mirar el tiempo y el espacio', ['H1'],
                 'La Clase 1 es una formulación editorial de inicio (sin capacidad del programa).'),
    'UNIDAD 2': ('UNIDAD 2 — La prehistoria americana', ['H2'], ''),
    'UNIDAD 3': ('UNIDAD 3 — El Paraguay y América en el siglo XVI', ['H4', 'H3'], ''),
    'UNIDAD 4': ('UNIDAD 4 — Paraguay y América colonial: siglo XVII', ['H5', 'H6'],
                 'Ampliación editorial optativa, no exigida por el programa MEC 2014: Continuación E, Guerra Guaranítica '
                 '(1754-1756, siglo XVIII).'),
    'UNIDAD 5': ('UNIDAD 5 — El espacio geográfico: posición, relieve e hidrografía', ['G1', 'G2'], ''),
    'UNIDAD 6': ('UNIDAD 6 — Clima, regiones naturales y ambiente', ['G3', 'G4'], ''),
}

OBS_CONT = {
    4: 'Profundización de H1 (temas del programa). Ficha del maestro editorial: usar la capacidad literal.',
    8: 'Profundización de H2 (temas del programa).',
    13: 'Profundización de H3 (temas del programa).',
    16: 'Profundización de H3 (temas del programa). Ficha del maestro con variante: usar la capacidad literal.',
    22: 'Ampliación editorial (siglo XVIII); no exigida por el programa MEC 2014, cuyo alcance son los siglos XVI y XVII. '
        'Optativa. Núcleo recomendado del encuentro: bandeirantes, pérdida del Guairá y Mbororé (H5).',
    29: 'Profundización de G2 (tema «Hidrovía Paraná-Paraguay. Impacto ecológico, social y económico»). Desarrollada en la '
        'tercera muestra.',
    34: 'Profundización de G3 y G4 (temas del programa). Ficha del maestro editorial: usar las capacidades literales.',
}

OBS_OTRAS = {
    15: 'Capacidad literal H3. Desarrollada en la tercera muestra.',
    18: 'Cierre de la 1.ª etapa — junio. Cubre H1, H2, H3 y H4 (Unidades 1 a 3).',
    28: 'Capacidad literal G2. Desarrollada en la tercera muestra.',
    30: 'Asociación curricular: integra G1 y G2. Desarrollada en la tercera muestra.',
    36: 'Cierre de la 2.ª etapa — noviembre. Cubre H5, H6, G1, G2, G3 y G4 (Unidades 4 a 6). El ítem del maestro sobre la '
        'Guerra Guaranítica evalúa la ampliación editorial: optativo.',
}

FICHA = {15: 'clase10', 28: 'clase18', 29: 'contF'}


def texto_celda(celda, lineas, negrita=None):
    """Reemplaza el texto de una celda conservando el formato de su primer run de texto."""
    p = celda.paragraphs[0]
    for q in celda.paragraphs[1:]:
        q._p.getparent().remove(q._p)
    runs = [r for r in p.runs if r.text.strip()]
    modelo = copy.deepcopy((runs[0] if runs else p.runs[-1])._r)
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for i, (txt, neg) in enumerate(lineas if isinstance(lineas[0], tuple) else [(x, negrita) for x in lineas]):
        if i:
            br = copy.deepcopy(modelo)
            for t in br.findall(qn('w:t')):
                br.remove(t)
            br.append(br.makeelement(qn('w:br'), {}))
            p._p.append(br)
        r = copy.deepcopy(modelo)
        for t in r.findall(qn('w:t')):
            r.remove(t)
        t = r.makeelement(qn('w:t'), {}); t.text = txt; t.set(qn('xml:space'), 'preserve')
        r.append(t)
        p._p.append(r)
        if neg is not None:
            docx.text.run.Run(r, p).bold = neg


def celdas(fila):
    vistas, out = set(), []
    for c in fila.cells:
        if id(c._tc) not in vistas:
            vistas.add(id(c._tc)); out.append(c)
    return out


def main():
    orig, salida = sys.argv[1:3]
    tmp = os.path.join(tempfile.mkdtemp(), 'v2.docx')
    subprocess.run([sys.executable, V2, orig, tmp], check=True)
    doc = docx.Document(tmp)

    # encabezado: criterio histórico y fuente de las capacidades
    for p in doc.paragraphs:
        for r in p.runs:
            if 'Criterio histórico' in r.text:
                r.text = ('\nCapacidades: texto literal del Programa de Estudio MEC 2014, págs. 55 a 57 (códigos H1-H6 y G1-G4 '
                          'de la matriz curricular).\nCriterio histórico: el programa abarca los siglos XVI y XVII; la Guerra '
                          'Guaranítica (1754-1756, siglo XVIII) se ofrece como ampliación editorial optativa en la Continuación E '
                          'y no responde a un requerimiento del programa.')

    planes = {pl['n']: pl for pl in PLANES}
    t = doc.tables[0]
    for fila in t.rows[1:]:
        cs = celdas(fila)
        txt = cs[0].text
        if len(cs) == 1 and txt.startswith('UNIDAD'):
            titulo, caps, nota = BANDAS[txt[:8]]
            periodo = [x for x in txt.split('\n') if x.startswith('Período')][0]
            lineas = [(titulo, True)] + [(f'Capacidad {c} (MEC 2014, literal): {CAPACIDADES[c]}', True) for c in caps]
            if nota:
                lineas.append((nota, False))
            lineas.append((periodo, True))
            texto_celda(cs[0], lineas)
            continue
        if len(cs) != 7 or not txt.strip().isdigit():
            continue
        n = int(txt.strip())
        if n in OBS_CONT:
            texto_celda(cs[6], [OBS_CONT[n]])
        if n in OBS_OTRAS:
            texto_celda(cs[6], [OBS_OTRAS[n]])
        if n == 22:
            texto_celda(cs[4], ['Lectura de mapa y fuente; esquema causa→decisión imperial→resistencia→campañas→consecuencias; '
                                'nota explícita: contenido del siglo XVIII, fuera del alcance del programa MEC 2014.'])
        if n == 25:
            texto_celda(cs[5], ['Prueba escrita de la Unidad 4; lista de cotejo de mediterraneidad, economía colonial, misiones '
                                'y colonizaciones (la Guerra Guaranítica, solo si se desarrolló la ampliación optativa).'])
        if n in planes:
            pl = planes[n]
            if n in FICHA:
                ficha = next(b for b in TOMO[FICHA[n]] if b[0] == 'ficha')
                tema, indic = ficha[2], ficha[3]
            else:
                tema, indic = pl['tema'], pl['indicadores']
            if n == 29:
                tema = next(b[1] for b in TOMO['contF'] if b[0] == 'h1') + '. ' + tema
            if n == 30:
                banda = next(b for b in TOMO['evalU5'] if b[0] == 'banda')
                tema = f'Evaluación de la Unidad 5 — {banda[2]}. ' + tema
            texto_celda(cs[1], [tema])
            texto_celda(cs[2], ['• ' + x for x in indic])
            texto_celda(cs[4], [pl['evaluacion']['procedimientos']])
            texto_celda(cs[5], [pl['evaluacion']['instrumentos']])
    doc.save(salida)


if __name__ == '__main__':
    main()
