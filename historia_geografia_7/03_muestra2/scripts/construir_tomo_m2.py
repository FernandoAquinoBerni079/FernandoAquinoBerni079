"""Tomo de la segunda muestra, sobre el maestro del 3 de octubre de 2026
(SHA-256 8dbc4f44…179413), con la maqueta de la variante B (oficio mejorado).

Conserva del maestro: portada, presentación, estilos y las ilustraciones
originales de las secciones (Figuras 10.1 y 18.1, con la misma imagen y el
mismo ancho). Reemplaza el texto de la Clase 10, la Clase 18, la
Continuación F y la Evaluación de la Unidad 5 por el contenido de
`contenido_tomo.py` y agrega las fichas reproducibles.

Uso: python3 construir_tomo_m2.py MAESTRO.docx CARPETA_FIGURAS SALIDA.docx [paginas.json]
"""
import io
import json
import os
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
sys.path.insert(0, os.path.join(AQUI, '..', '..', '02_piloto', 'scripts'))
import tomo_piloto as TP  # noqa: E402
from contenido_tomo import SECCIONES  # noqa: E402

W, A, R = TP.W, TP.A, TP.R
ORDEN = ['clase10', 'clase18', 'ficha18A', 'contF', 'fichasF', 'evalU5']
TAM_CAJA, TAM_TABLA, TAM_ACT, TAM_EPI = 9.5, 9.0, 10.0, 8.5   # tamaños del maestro; la variante B suma 1 pt
COLOR = {'ficha_k': '21374D', 'ficha_v': 'E6EBF1', 'concepto': 'FFF8E1', 'ejemplo': 'E6EBF1', 'paraguay': 'E8F5E9',
         'fuente': 'F3EEE2', 'fuente_t': '8A5A2B', 'ubicate': 'E6EBF1', 'ubicate_t': '21374D', 'cab': '3D5A80',
         'banda': 'B8860B', 'epi': '5D6775', 'repro': 'F7F7F7'}


class Constructor:
    def __init__(self, doc, figuras, originales):
        self.doc, self.fig, self.orig = doc, figuras, originales
        self.body = doc.element.body
        self.sect = self.body.find(W + 'sectPr')

    # ---------- utilidades ----------
    def _mover(self, el):
        self.sect.addprevious(el)
        return el

    def parrafo(self, texto='', estilo=None, tam=None, negrita=False, cursiva=False, color=None, alinear=None, despues=None):
        p = self.doc.add_paragraph(style=estilo) if estilo else self.doc.add_paragraph()
        if texto:
            r = p.add_run(texto)
            r.bold, r.italic = negrita, cursiva
            if tam:
                r.font.size = Pt(tam)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        if alinear:
            p.alignment = alinear
        if despues is not None:
            p.paragraph_format.space_after = Pt(despues)
        self._mover(p._p)
        return p

    def _tabla(self, filas, cols, anchos_rel):
        t = self.doc.add_table(rows=filas, cols=cols)
        tbl = t._tbl
        tblpr = tbl.tblPr
        tw = tblpr.find(qn('w:tblW'))
        if tw is None:
            tw = OxmlElement('w:tblW'); tblpr.append(tw)
        tw.set(qn('w:type'), 'pct'); tw.set(qn('w:w'), '5000')
        mar = OxmlElement('w:tblCellMar')
        for lado, v in (('top', 60), ('left', 100), ('bottom', 60), ('right', 100)):
            e = OxmlElement('w:' + lado); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
        tblpr.append(mar)
        bordes = OxmlElement('w:tblBorders')
        for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            e = OxmlElement('w:' + lado)
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), 'A0A8B3')
            bordes.append(e)
        tblpr.append(bordes)
        tot = sum(anchos_rel)
        for gc, a in zip(tbl.tblGrid.findall(qn('w:gridCol')), anchos_rel):
            gc.set(qn('w:w'), str(int(10800 * a / tot)))
        for fila in t.rows:
            trpr = fila._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
            for c, a in zip(fila.cells, anchos_rel):
                c.width = Cm(18.0 * a / tot)
        self._mover(tbl)
        return t

    @staticmethod
    def _sombrear(celda, color):
        tcpr = celda._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), color)
        tcpr.append(shd)

    @staticmethod
    def _escribir(celda, texto, tam, negrita=False, color=None, cursiva=False, nuevo=False):
        p = celda.add_paragraph() if nuevo else celda.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(texto); r.font.size = Pt(tam); r.bold = negrita; r.italic = cursiva
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return p

    def separador(self):
        p = self.parrafo(despues=0)
        p.paragraph_format.space_before = Pt(0)
        return p

    # ---------- bloques ----------
    def h(self, nivel, texto):
        self.parrafo(texto, estilo=f'Heading {nivel}')

    def ficha(self, cap, tema, indicadores):
        t = self._tabla(3, 2, [1, 2.4])
        for i, (k, v) in enumerate([('Capacidad', cap), ('Tema', tema), ('Indicadores de logro', None)]):
            self._sombrear(t.cell(i, 0), COLOR['ficha_k']); self._sombrear(t.cell(i, 1), COLOR['ficha_v'])
            self._escribir(t.cell(i, 0), k, TAM_CAJA, negrita=True, color='FFFFFF')
            if v:
                self._escribir(t.cell(i, 1), v, TAM_CAJA)
            else:
                for j, ind in enumerate(indicadores):
                    self._escribir(t.cell(i, 1), '• ' + ind, TAM_CAJA, nuevo=j > 0)
        self.separador()

    def caja(self, tipo, titulo, texto):
        t = self._tabla(1, 1, [1])
        c = t.cell(0, 0); self._sombrear(c, COLOR[tipo])
        self._escribir(c, titulo, TAM_CAJA, negrita=True)
        self._escribir(c, texto, TAM_CAJA, nuevo=True)
        self.separador()

    def tabla(self, cab, filas, anchos):
        t = self._tabla(1 + len(filas), len(cab), anchos)
        enc = t.rows[0]._tr.get_or_add_trPr(); enc.append(OxmlElement('w:tblHeader'))
        for j, h in enumerate(cab):
            self._sombrear(t.cell(0, j), COLOR['cab']); self._escribir(t.cell(0, j), h, TAM_TABLA, negrita=True, color='FFFFFF')
        for i, fila in enumerate(filas, 1):
            for j, v in enumerate(fila):
                self._escribir(t.cell(i, j), v, TAM_TABLA, negrita=(j == 0))
        self.separador()

    def cuadro(self, cab, etiquetas, anchos):
        t = self._tabla(1 + len(etiquetas), len(cab), anchos)
        for j, h in enumerate(cab):
            self._sombrear(t.cell(0, j), COLOR['cab']); self._escribir(t.cell(0, j), h, TAM_TABLA, negrita=True, color='FFFFFF')
        for i, et in enumerate(etiquetas, 1):
            tr = t.rows[i]._tr.get_or_add_trPr()
            hgt = OxmlElement('w:trHeight'); hgt.set(qn('w:val'), str(int((1.6 if et else 2.6) / 2.54 * 1440))); hgt.set(qn('w:hRule'), 'atLeast')
            tr.append(hgt)
            # cada celda lleva texto (al menos un espacio) para que la fila no colapse
            self._escribir(t.cell(i, 0), et or ' ', TAM_TABLA, negrita=True)
            for j in range(1, len(cab)):
                self._escribir(t.cell(i, j), ' ', TAM_TABLA)
        self.separador()

    def fuente(self, subtitulo, citas, vocab, preguntas, renglones=0):
        t = self._tabla(1, 1, [1])
        c = t.cell(0, 0); self._sombrear(c, COLOR['fuente'])
        self._escribir(c, 'Trabajá con la fuente', 10.5, negrita=True, color=COLOR['fuente_t'])
        self._escribir(c, subtitulo, TAM_TABLA, cursiva=True, nuevo=True)
        for ci in citas:
            self._escribir(c, ci, TAM_CAJA, nuevo=True)
        if vocab:
            self._escribir(c, vocab, TAM_TABLA, cursiva=True, nuevo=True)
        self._escribir(c, 'Analizá la fuente:', TAM_CAJA, negrita=True, nuevo=True)
        for k, q in enumerate(preguntas, 1):
            self._escribir(c, f'{k}. {q}', TAM_CAJA, nuevo=True)
        self._renglones_en(c, renglones)
        self.separador()

    def ubicate(self, titulo, texto, renglones=0):
        t = self._tabla(1, 1, [1])
        c = t.cell(0, 0); self._sombrear(c, COLOR['ubicate'])
        self._escribir(c, titulo, TAM_CAJA, negrita=True, color=COLOR['ubicate_t'])
        self._escribir(c, texto, TAM_CAJA, nuevo=True)
        self._renglones_en(c, renglones)
        self.separador()

    @staticmethod
    def _renglones_en(celda, n):
        """Renglones para responder dentro del recuadro: la fila no se parte,
        así que consigna y renglones quedan siempre en la misma página."""
        for _ in range(n):
            celda._tc.append(TP.renglon(None))

    def figura(self, archivo, epigrafe, ancho):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(os.path.join(self.fig, archivo), width=Cm(ancho))
        self._mover(p._p)
        self.parrafo(epigrafe, tam=TAM_EPI, cursiva=True, color=COLOR['epi'], alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=6)

    def figura_original(self, clave, epigrafe):
        blob, ancho = self.orig[clave]
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(io.BytesIO(blob), width=Cm(ancho))
        self._mover(p._p)
        self.parrafo(epigrafe, tam=TAM_EPI, cursiva=True, color=COLOR['epi'], alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=6)

    def banda(self, titulo, sub):
        t = self._tabla(1, 1, [1])
        c = t.cell(0, 0); self._sombrear(c, COLOR['banda'])
        p = self._escribir(c, titulo, 13, negrita=True, color='FFFFFF'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = self._escribir(c, sub, 10, color='FFFFFF', nuevo=True); p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def ficha_repro(self, titulo, bloques):
        self.h(1, titulo)
        self.parrafo('Se permite fotocopiar esta ficha para uso en el aula.', tam=TAM_EPI, cursiva=True, color=COLOR['epi'])
        self.parrafo('Nombre y Apellido: ______________________________   Curso: ______   Fecha: ___ / ___ / ______', tam=TAM_CAJA, despues=8)
        for b in bloques:
            self.bloque(b)

    def bloque(self, b):
        tipo = b[0]
        if tipo == 'h1':
            self.h(1, b[1])
        elif tipo == 'h2':
            self.h(2, b[1])
        elif tipo == 'h3':
            self.h(3, b[1])
        elif tipo == 'p':
            self.parrafo(b[1])
        elif tipo == 'ficha':
            self.ficha(*b[1:])
        elif tipo == 'caja':
            self.caja(*b[1:])
        elif tipo == 'tabla':
            self.tabla(*b[1:])
        elif tipo == 'cuadro':
            self.cuadro(*b[1:])
        elif tipo == 'fuente':
            self.fuente(*b[1:])
        elif tipo == 'ubicate':
            self.ubicate(*b[1:])
        elif tipo == 'fig':
            self.figura(*b[1:])
        elif tipo == 'fig_original':
            self.figura_original(*b[1:])
        elif tipo == 'consigna':
            p = self.parrafo(b[1], tam=TAM_ACT, negrita=True, despues=2)
            p.paragraph_format.keep_with_next = True
        elif tipo == 'item':
            p = self.parrafo(b[1], tam=TAM_ACT, despues=2)
            p.paragraph_format.keep_with_next = True   # la consigna no queda sola al pie
            n = b[2] if len(b) > 2 else (1 if '(....)' in b[1] else 0)
            if n:
                p.paragraph_format.keep_with_next = True
                self.bloque(('renglones', n))
        elif tipo == 'opc':
            for letra, o in zip('abcd', b[1]):
                self.parrafo(f'    {letra}) {o}', tam=TAM_ACT, despues=0)
        elif tipo == 'lista_orden':
            for k, o in enumerate(b[1]):
                p = self.parrafo(f'(     )  {o}', tam=TAM_ACT, despues=2)
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.keep_with_next = k < len(b[1]) - 1
        elif tipo == 'renglones':
            for k in range(b[1]):
                n = TP.renglon(self.doc)
                if k < b[1] - 1:
                    n.find(W + 'pPr').insert(0, OxmlElement('w:keepNext'))
                self._mover(n)
            self.parrafo(despues=2)
        elif tipo == 'banda':
            self.banda(b[1], b[2])
        elif tipo == 'datos':
            self.parrafo('Nombre y Apellido: _______________________________  Curso: ______  Fecha: ____ / ____ / ____', tam=TAM_CAJA, despues=6)
        elif tipo == 'ficha_repro':
            self.ficha_repro(b[1], b[2])
        else:
            raise ValueError(tipo)


def originales(doc):
    """Imagen y ancho de las Figuras 10.1 y 18.1 del maestro."""
    res = {}
    els = list(doc.element.body)
    for i, e in enumerate(els):
        t = TP.texto(e)
        for clave, rotulo in (('10.1', 'Figura 10.1'), ('18.1', 'Figura 18.1')):
            if t.strip().startswith(rotulo) and clave not in res:
                for k in range(i - 1, max(i - 7, 0), -1):
                    blips = els[k].findall('.//' + A + 'blip')
                    if blips:
                        part = doc.part.related_parts[blips[0].get(R + 'embed')]
                        ext = els[k].find('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent')
                        res[clave] = (part.blob, int(ext.get('cx')) / 360000)
                        break
    assert set(res) == {'10.1', '18.1'}, res.keys()
    return res


def main():
    maestro, figuras, salida = sys.argv[1:4]
    paginas = json.load(open(sys.argv[4])) if len(sys.argv) > 4 else {}
    doc = docx.Document(maestro)
    orig = originales(doc)
    body = doc.element.body
    els = [e for e in body if e.tag != W + 'sectPr']
    fin_frente = next(i for i, e in enumerate(els) if e.tag == W + 'tbl' and TP.texto(e).startswith('Prueba diagnóstica'))
    for e in els[fin_frente:]:
        body.remove(e)
    for e in [e for e in body if TP.es_p(e)]:
        if not TP.texto(e).strip() and e.find('.//' + W + 'pageBreakBefore') is not None and not e.findall('.//' + A + 'blip'):
            body.remove(e)
    c = Constructor(doc, figuras, orig)
    for clave in ORDEN:
        bloques = SECCIONES[clave]
        k = 0
        while k < len(bloques):
            b = bloques[k]
            sig = bloques[k + 1] if k + 1 < len(bloques) else None
            if b[0] in ('fuente', 'ubicate') and sig and sig[0] == 'renglones':
                getattr(c, b[0])(*b[1:], renglones=sig[1])
                k += 2
                continue
            c.bloque(b)
            k += 1
    TP.limpiar_relaciones(doc)
    TP.corregir_saltos(doc)
    v = TP.VARIANTES['oficio_mejorado']
    ancho = TP.ajustar_pagina(doc, v['papel'], v['margen'])
    TP.agrandar_letra(doc, v['suma_pt'])
    TP.agrupar_opciones(doc, ancho)
    TP.indice_estatico(doc, ancho, paginas)
    doc.save(salida)


if __name__ == '__main__':
    main()
